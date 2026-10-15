"""Offline all-strategy equity graph, metrics workbook and complete final appendix.

Reads only finished replay tables. The builder writes its final manifest after
this function, so manifest.json and independent_audit.json are never inputs.
"""
from __future__ import annotations

import argparse
import hashlib
import html
import json
from datetime import datetime, timezone
from pathlib import Path

import matplotlib
matplotlib.use("Agg")
import matplotlib.dates as mdates
import matplotlib.pyplot as plt
from matplotlib.colors import to_hex
import numpy as np
import pandas as pd
import plotly.graph_objects as go
from docx import Document
from docx.enum.section import WD_ORIENT, WD_SECTION_START
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt
from PIL import Image as PILImage
from pypdf import PdfReader
from reportlab.lib import colors
from reportlab.lib.pagesizes import landscape, letter
from reportlab.lib.styles import ParagraphStyle
from reportlab.pdfbase import pdfmetrics
from reportlab.pdfbase.ttfonts import TTFont
from reportlab.platypus import BaseDocTemplate, Frame, Image, KeepTogether, NextPageTemplate, PageBreak, PageTemplate, Paragraph, Spacer, Table, TableStyle

from epl_strategy.report import number

ROOT=Path(__file__).resolve().parents[1]
DRAFT_DATE="2026-10-15"
TITLE="EPL strategy laboratory: every account starts with 100,000"
INITIAL=100000.0
ENSEMBLES={
    "ensemble_logit4": ["logit_team_only","logit_team_xg","logit_team_players","logit_combined"],
    "ensemble_learned8": [f"{kind}_{group}" for kind in ("logit","residual") for group in ("team_only","team_xg","team_players","combined")],
    "ensemble_market_blend": ["logit_combined","baseline_market_only"],
}


def _sha(path):
    digest=hashlib.sha256()
    with path.open("rb") as stream:
        for chunk in iter(lambda:stream.read(1024*1024),b""):digest.update(chunk)
    return digest.hexdigest()


def _relative(path):
    try:return path.relative_to(ROOT).as_posix()
    except ValueError:return str(path)


def _json(value):
    return json.dumps(value,ensure_ascii=False,sort_keys=True,separators=(",",": "),allow_nan=False)


def _model_text(key):
    if key in ENSEMBLES:
        parts=ENSEMBLES[key]
        return "Equal-weight probability mean of "+", ".join(parts)+f"; each weight is 1/{len(parts)}."
    labels={"first_accuracy":"Frozen first-study accuracy-selected SVM.","first_calibration":"Frozen first-study ECE-selected MLP.",
            "baseline_class_frequency":"Rolling-study training-class-frequency probabilities.",
            "baseline_market_only":"Raw normalized inverse Bet365 reference-price probabilities.",
            "market_temperature":"Rolling-study held-out-temperature-adjusted reference-market probabilities."}
    if key in labels:return labels[key]
    if key.startswith("logit_"):return "Frozen rolling-study logistic probabilities; feature group "+key.removeprefix("logit_")+"."
    if key.startswith("residual_"):return "Frozen rolling-study market-residual probabilities; feature group "+key.removeprefix("residual_")+"."
    raise ValueError("Unknown catalog probability series: "+str(key))


def _human(config):
    selection={
        "max_edge":"Choose the eligible H/D/A outcome with largest p×odds−1.",
        "model_favorite":"Choose the eligible outcome with greatest model probability.",
        "market_favorite":"Choose the eligible outcome with greatest raw reference-market probability.",
        "market_underdog":"Choose the eligible outcome with smallest raw reference-market probability.",
        "fixed_outcome":"Choose the configured eligible outcome class.",
    }
    rule=config.get("selection_rule")
    if rule not in selection:raise ValueError("Unknown catalog selection rule: "+str(rule))
    quotes={"best":"best eligible individual-book prices","reference":"Bet365 reference prices",
            "best_individual_book":"best eligible individual-book prices","bet365_reference":"Bet365 reference prices"}
    text=[_model_text(config["model_key"])]
    if config.get("cash_only"):
        text.append("Hold the entire balance in cash; place no wager. The saved no-bet threshold is "+number(config["min_edge"],0)+".")
    else:
        text.append(selection[rule])
        text.append("Use "+quotes[config["odds_mode"]]+".")
        allowed=", ".join(config["allowed_outcomes"])
        bounds="no decimal-odds bounds" if config.get("min_odds") is None and config.get("max_odds") is None else "inclusive decimal-odds range "+number(config.get("min_odds"))+" to "+number(config.get("max_odds"))
        text.append("Apply outcomes "+allowed+" and "+bounds+" before outcome choice.")
        if config.get("min_edge") is None:text.append("No edge qualification is required.")
        else:text.append("Require selected p×odds−1 strictly above "+number(config["min_edge"],percent=True)+".")
        if config.get("consensus_min") is not None:
            keys=config["consensus_model_keys"]
            text.append(f"After choosing the mean-eight maximum-edge outcome, require at least {config['consensus_min']} of {len(keys)} constituent models to have strictly positive edge at that same quote; otherwise abstain without searching another outcome.")
        staking=config["staking"]
        if staking=="fixed":text.append("Request a fixed stake of "+number(config["fixed_stake"],0)+" monetary units.")
        elif staking=="dynamic_fraction":text.append("Request "+number(config["bankroll_fraction"],percent=True)+" of the date-start balance B.")
        elif staking=="kelly":
            text.append("Request B×min(cap, f×max((p×odds−1)/(odds−1),0)), where B is date-start balance, f="+number(config["kelly_fraction"],4)+" and cap="+number(config["max_bet_fraction"],percent=True)+".")
        else:raise ValueError("Unknown sizing rule: "+str(staking))
        if config.get("date_cap_fraction") is not None:text.append("Scale all same-date requests proportionally to at most "+number(config["date_cap_fraction"],percent=True)+" of date-start cash.")
        else:text.append("Scale same-date requests proportionally to available cash; no extra date cap.")
    if config.get("diagnostic"):text.append("Historical threshold-diagnostic recipe replayed on this common period; not a newly selected primary rule.")
    return " ".join(part for part in text if part)


def _appendix(catalog):
    entries=[{"strategy_id":c["strategy_id"],"title":c["strategy_id"]+" — "+str(c.get("name",c.get("strategy_name",c["strategy_id"]))),
              "origin":str(c.get("origin","")),"explanation":_human(c),"exact_config":_json(c)} for c in catalog]
    if len({entry["strategy_id"] for entry in entries})!=len(catalog):raise ValueError("Duplicate appendix ID")
    if any(json.loads(entry["exact_config"])!=config for entry,config in zip(entries,catalog)):raise ValueError("Appendix configuration lost a parameter")
    return entries


def _validate(catalog,metrics,graph,equity,bets):
    ids=[str(c["strategy_id"]) for c in catalog]
    if len(ids)!=84 or len(set(ids))!=84:raise ValueError("Expected exactly84 unique catalog strategies")
    expected=set(ids)
    for label,frame in (("metrics",metrics),("graph",graph),("equity",equity)):
        if set(frame.strategy_id.astype(str))!=expected:raise ValueError(f"{label} IDs differ from catalog")
    if len(metrics)!=len(ids) or metrics.strategy_id.duplicated().any():raise ValueError("Metrics are not one-to-one by catalog ID")
    if not set(bets.strategy_id.astype(str)).issubset(expected):raise ValueError("A bet has an unknown strategy ID")
    if not metrics.prediction_count.eq(374).all() or not metrics.matched_valid_odds_count.eq(374).all():raise ValueError("Strategies do not share all374 priced opportunities")
    graph=graph.copy();equity=equity.copy()
    graph["match_date"]=pd.to_datetime(graph.match_date);equity["match_date"]=pd.to_datetime(equity.match_date)
    if not pd.api.types.is_bool_dtype(graph.is_initial):
        values=graph.is_initial.astype(str).str.lower()
        if not values.isin(["true","false"]).all():raise ValueError("Invalid initial-point flags")
        graph["is_initial"]=values.eq("true")
    if graph.duplicated(["strategy_id","match_date"]).any() or equity.duplicated(["strategy_id","match_date"]).any():raise ValueError("Duplicate equity date")
    if not np.isfinite(graph.end_bankroll.to_numpy(float)).all():raise ValueError("Nonfinite plotted balance")
    if (graph.end_bankroll.astype(float)<-1e-6).any():raise ValueError("A plotted account borrows money")
    native_dates=sorted(equity.match_date.unique())
    if len(native_dates)!=113 or pd.Timestamp(native_dates[0])!=pd.Timestamp("2025-08-15") or pd.Timestamp(native_dates[-1])!=pd.Timestamp("2026-05-24"):
        raise ValueError("Native equity dates differ from the declared common season")
    anchor=pd.Timestamp(native_dates[0])-pd.Timedelta(days=1)
    for sid,part in graph.groupby("strategy_id",sort=False):
        initial=part.loc[part.is_initial]
        if len(initial)!=1 or initial.match_date.iloc[0]!=anchor or not np.isclose(initial.end_bankroll.iloc[0],INITIAL,rtol=0,atol=1e-7):
            raise ValueError("Missing exact100000 pre-match anchor: "+sid)
        native=part.loc[~part.is_initial]
        if len(native)!=113 or set(native.match_date)!=set(pd.to_datetime(native_dates)):raise ValueError("Missing date-close point: "+sid)
    joined=graph.loc[~graph.is_initial].merge(equity[["strategy_id","match_date","end_bankroll"]],on=["strategy_id","match_date"],validate="one_to_one",suffixes=("_graph","_ledger"))
    if len(joined)!=len(equity) or not np.allclose(joined.end_bankroll_graph,joined.end_bankroll_ledger,atol=1e-6,rtol=0):raise ValueError("Graph and saved cash ledger differ")
    final=graph.sort_values("match_date").groupby("strategy_id").tail(1).set_index("strategy_id").end_bankroll
    aligned=metrics.set_index("strategy_id").reindex(ids)
    if not np.allclose(final.reindex(ids),INITIAL+aligned.profit,atol=1e-5,rtol=0):raise ValueError("Ending graph balances do not equal initial capital plus profit")
    if not np.allclose(aligned.initial_capital_sum_units,INITIAL,atol=1e-8,rtol=0):raise ValueError("A strategy uses different initial capital")
    counts=bets.groupby("strategy_id").size().reindex(ids,fill_value=0)
    if not np.array_equal(counts.to_numpy(),aligned.bet_count.to_numpy()):raise ValueError("Metrics disagree with recorded trade counts")
    return graph,equity,ids,anchor


def _styles(count):
    palette=plt.get_cmap("tab20").colors
    return [(to_hex(palette[i%20]),["solid","dash","dashdot","dot","longdash"][(i//20)%5]) for i in range(count)]


def _graphs(catalog,graph,figures,entries):
    figures.mkdir(exist_ok=True);styles=_styles(len(catalog));plotly=go.Figure()
    fig,ax=plt.subplots(figsize=(11.8,7.4));fig.subplots_adjust(left=.085,right=.985,top=.91,bottom=.25)
    dash={"solid":"-","dash":"--","dashdot":"-.","dot":":","longdash":(0,(8,3))}
    for i,config in enumerate(catalog):
        sid=config["strategy_id"];part=graph.loc[graph.strategy_id.eq(sid)].sort_values("match_date")
        color,line=styles[i];name=str(config.get("name",config.get("strategy_name",sid)))
        ax.step(part.match_date,part.end_bankroll,where="post",color=color,linestyle=dash[line],linewidth=.85,alpha=.88,label=sid)
        plotly.add_trace(go.Scatter(x=part.match_date.dt.strftime("%Y-%m-%d").tolist(),y=part.end_bankroll.astype(float).tolist(),
            mode="lines",name=sid,uid=sid,legendgroup=sid,visible=True,line={"color":color,"dash":line,"width":1.5,"shape":"hv"},
            customdata=[[sid,name,bool(flag)] for flag in part.is_initial],
            hovertemplate="%{customdata[0]} · %{customdata[1]}<br>Date: %{x|%Y-%m-%d}<br>Balance: %{y:,.2f} monetary units<extra></extra>"))
    ax.set_title(f"All {len(catalog)} strategies · each starts with 100,000 monetary units",fontsize=12)
    ax.set_xlabel("Date (balance after date-level settlement)");ax.set_ylabel("Balance (monetary units)")
    ax.yaxis.set_major_formatter(matplotlib.ticker.StrMethodFormatter("{x:,.0f}"));ax.xaxis.set_major_locator(mdates.MonthLocator(interval=2));ax.xaxis.set_major_formatter(mdates.DateFormatter("%b %Y"))
    ax.grid(alpha=.2);ax.spines[["top","right"]].set_visible(False);ax.set_xlim(graph.match_date.min(),graph.match_date.max())
    handles,labels=ax.get_legend_handles_labels();fig.legend(handles,labels,loc="lower center",ncol=12,fontsize=7.5,frameon=False,bbox_to_anchor=(.5,.012),handlelength=2,columnspacing=1.3,labelspacing=.6)
    with plt.rc_context({"svg.fonttype":"none"}):
        for suffix in (".png",".svg"):fig.savefig(figures/("all_strategies_100k"+suffix),dpi=240)
    plt.close(fig)
    plotly.update_layout(title={"text":f"All {len(catalog)} independently funded strategy accounts · initial 100,000 each","x":.01,"y":.99,"yanchor":"top"},template="plotly_white",
        height=1050,margin={"l":85,"r":35,"t":140,"b":260},xaxis={"title":"Date (date-close settlement)","type":"date"},
        yaxis={"title":"Balance (monetary units)","type":"linear","tickformat":",.0f"},hovermode="closest",
        legend={"title":{"text":"Strategy IDs","side":"top"},"orientation":"h","x":0,"xanchor":"left","y":-.16,"yanchor":"top","entrywidth":70,"entrywidthmode":"pixels","font":{"size":10},"groupclick":"toggleitem"},
        updatemenus=[{"type":"buttons","direction":"right","x":0,"xanchor":"left","y":1.025,"yanchor":"bottom","buttons":[
            {"label":"Show all","method":"restyle","args":[{"visible":True}]},
            {"label":"Hide all (keep legend)","method":"restyle","args":[{"visible":"legendonly"}]}]}])
    if len(plotly.data)!=len(catalog) or [t.uid for t in plotly.data]!=[c["strategy_id"] for c in catalog] or any(t.visible is not True for t in plotly.data):raise AssertionError("Not every strategy is initially visible")
    graph_html=plotly.to_html(full_html=False,include_plotlyjs=True,config={"responsive":True,"displaylogo":False,"toImageButtonOptions":{"format":"png","filename":"all_strategies_100k"}})
    appendix_html="".join("<article id='"+html.escape(e["strategy_id"])+"'><h3>"+html.escape(e["title"])+"</h3><p>Origin: "+html.escape(e["origin"])+". "+html.escape(e["explanation"])+"</p><pre>"+html.escape(e["exact_config"])+"</pre></article>" for e in entries)
    page="<!doctype html><html lang='en'><head><meta charset='utf-8'><meta name='viewport' content='width=device-width,initial-scale=1'><title>All strategy accounts</title><style>body{font:15px system-ui;margin:24px;color:#152e3c}p{max-width:1100px}pre{white-space:pre-wrap;overflow-wrap:anywhere;font-size:12px;background:#f0f4f6;padding:12px}article{border-top:1px solid #ccd6dc;padding-top:12px}h1{font-size:26px}</style></head><body><h1>Every strategy starts with 100,000</h1><p>Complete retrospective 2025/26 comparison. Every line is a separate account in monetary units; overlapping cash curves are retained. All strategies are visible initially. Use the legend or buttons to explore. This file contains its plotting library and works offline.</p>"+graph_html+"<h2>Scope and interpretation</h2><p>All 374 common historical fixtures and 113 settlement dates; no model refitting or new untouched test. Recipes were defined after inspecting this history. Archived prices have unverified execution clocks, fees and fills. Expected gains/losses are model-implied and conditional on recorded stakes.</p><h2>Appendix — complete strategy definitions (last)</h2>"+appendix_html+"</body></html>"
    (figures/"all_strategies_100k.html").write_text(page,encoding="utf-8")
    return {"trace_ids":[t.uid for t in plotly.data],"initially_visible_traces":len(plotly.data),"plotly_js_bundled":True,"linear_money_y_axis":True}


def _metric_table(metrics,ids):
    frame=metrics.set_index("strategy_id").reindex(ids)
    return pd.DataFrame([{"ID":sid,"Trades":number(row.bet_count,0),"Hit %":number(row.hit_ratio,percent=True),
        "Trade\nfreq. %":number(row.trading_frequency,percent=True),"Avg. winning\ngain":number(row.avg_winning_gain_units),
        "Avg. losing\nloss":number(row.avg_losing_loss_units),"Model expected\ngain/trade":number(row.expected_gain_per_trade_units),
        "Model expected\nloss/trade":number(row.expected_loss_per_trade_units),"ROI on\nstakes":number(row.roi,percent=True),"Return on\n100,000":number(row.return_on_initial_capital,percent=True),
        "Max.\ndrawdown":number(row.max_drawdown,percent=True)} for sid,row in frame.iterrows()])


def _body(catalog,metrics,ids,figure):
    table=_metric_table(metrics,ids);no_trades=int(metrics.bet_count.eq(0).sum())
    origin=pd.Series([c.get("origin","") for c in catalog]).value_counts().to_dict()
    pages=[]
    def page(title,orientation="portrait"):
        item={"title":title,"orientation":orientation,"blocks":[]};pages.append(item);return item["blocks"]
    def p(blocks,text):blocks.append({"type":"paragraph","text":text})
    b=page(TITLE)
    p(b,"DRAFT FOR REVIEW | October 15, 2026 | Historical EPL research; each strategy has its own 100,000 monetary-unit account.")
    p(b,f"This comparison includes every one of the {len(catalog)} catalog strategies on the same 374 eligible 2025/26 fixtures, with 113 settlement dates from August 15, 2025 through May 24, 2026. The initial point is August 14, 2025 at 100,000 for every account. The single graph shows all accounts, including {no_trades} policies with no recorded trade; curves are not summed into a portfolio.")
    p(b,"The catalog preserves 38 distinct prior first/rolling-study rules, adds 12 historical threshold-diagnostic recipes, and includes 34 new exploratory benchmarks, ensembles, allocation, side and eligibility rules. Repeated historical season/scope entries are consolidated into distinct rules and replayed in this common season. Threshold diagnostics are not newly selected primary policies. Pending future forecasts have no settled account performance and are kept outside this historical graph. Earlier cost stresses remain a separate sensitivity study.")
    p(b,"Probability inputs comprise 13 frozen series: the first-study accuracy-selected SVM and calibration-selected MLP, and 11 rolling-study model/baseline series. Their training/calibration histories differ, so differences combine model, information-window and rule effects. Three fixed equal-weight ensembles add probability combinations without fitting new estimators or selecting weights by return. No source model is refitted.")
    p(b,"All accounts start with 100,000; fixed bets request 1,000. Dynamic and fractional-Kelly sizing use each date's starting cash. Same-date requests are allocated before settlement and uniformly scaled to available cash or the configured date cap; same-date winnings cannot finance another request. There are no account resets, borrowing or invented returns for missing future observations. The appendix gives every exact rule, including model sources, edge tests, price selection, filters and caps.")
    p(b,"H, D and A mean home win, draw and away win. Outcomes and inclusive odds bounds filter candidates before choice; edge qualification is strict. Ties use H, then D, then A. At most one outcome is selected per match.")
    p(b,"This history was already examined before the expanded catalog was defined. All new recipes and the enlarged comparison are exploratory. A high ending balance, a profitable historical curve or an unadjusted conditional interval does not establish a reliable future edge. The current-club cohort omits six fixtures between excluded clubs, so it is 374 rather than all 380 league matches and retains survivor/selection bias.")
    p(b,"Open figures/all_strategies_100k.html for the offline interactive graph: every curve is visible initially; click a legend entry to toggle it or double-click to isolate it. Hover shows ID, date and monetary balance. PNG/SVG and the printed graph use the same complete data. Full-precision metrics and exact configurations are in the workbook. The final part of this report explains every catalog ID individually.")
    b=page("One graph: all strategy balances over time","landscape")
    p(b,f"All {len(catalog)} independent accounts start at 100,000 on August 14, 2025. Step lines join date-close balances; coincident cash/rule paths remain separate traces. Legend IDs map to the final complete appendix.")
    b.append({"type":"figure","path":figure,"caption":"Figure 1. Every catalog strategy, in catalog order, on linear monetary and calendar-date axes. No losing curve is removed. Archived quotes, fills and simultaneity remain unverified; this is a historical replay."})
    for start in range(0,len(table),28):
        b=page(f"Complete strategy metrics: {ids[start]}–{ids[min(start+27,len(ids)-1)]}","landscape")
        p(b,"All money columns are units per recorded trade. Expected positive/negative components use model probabilities; average winning/losing payoffs are realised conditional means. Cash-policy per-trade means, hit ratio and ROI are undefined, shown as dashes.")
        b.append({"type":"table","frame":table.iloc[start:start+28].copy(),"caption":"Table 1 (continued across all catalog IDs). Frequency=bets/374 priced fixtures; ROI=profit/turnover. Monetary totals, bankroll return, expected net, uncertainty, stakes and exposure remain in the full workbook."})
    b=page("Definitions, uncertainty and source trail")
    p(b,"For stake s, decimal odds O and selected win probability p, model-implied expected gain is p·s·(O−1) and expected loss magnitude is (1−p)·s. Their difference is expected net profit. Totals sum those components across recorded trades; per-trade quantities divide totals by trade count. They condition on actual allocated stakes, including the realised bankroll-dependent Kelly path, and are not unconditional predictions of terminal wealth.")
    p(b,"Hit ratio is winning selected outcomes / recorded trades, distinct from full-match three-class prediction accuracy. Average winning gain excludes returned principal; losing loss is a positive magnitude. Realised expectancy is net profit / trades. ROI uses turnover as denominator, while return on initial capital uses 100,000. Profit factor compares total positive and negative net P&L, and drawdown includes initial capital and subsequent date-close peaks. No-trade cash totals/activity are zero; undefined ratios remain unavailable.")
    p(b,"Bootstrap intervals, where supplied in the workbook, resample recorded whole dates conditional on the model, selections and stake path. They do not refit models, regenerate a counterfactual Kelly path or adjust for trying many recipes. This common-season replay does not turn previously inspected outcomes into a new test. Probability-weighted expectations can diverge substantially from realised returns, especially when odds amplify probability errors.")
    p(b,"The replay uses saved preclosing reference and best-individual-book prices. It adds no actual bookmaker fees, deposit interest, fill restrictions or price-expiry mechanism. Compatible simultaneous quotes and accepted execution are unverified. The previous robustness companion separately studies explicit hypothetical costs. Neither the chart nor these calculations establish deployable arbitrage or a current player-availability signal.")
    p(b,"Method/source references before the final appendix: epl_modeling/METHODOLOGY.md; epl_modeling/ABLATION_PROTOCOL.md; epl_strategy/METHODOLOGY.md; epl_strategy_lab/PROTOCOL.md. These preserve chronological design, payoff definitions and exact replay conventions. The reporter manifest records consumed files, code and output hashes. The builder's final manifest is written afterward and is not a report input, avoiding a circular hash dependency.")
    p(b,"Rebuild the report from a finished output directory with python -m epl_strategy_lab.report --results PATH. All original acquisition data, fitted studies, prior reports and their outputs remain separate. The appendix follows and is the final report section; it includes one standalone explanation and exact configuration for every ID.")
    return pages,table,origin


def _markdown(pages,entries,out):
    chunks=[]
    for i,page in enumerate(pages):
        chunks.extend([("# " if i==0 else "## ")+page["title"],""])
        for block in page["blocks"]:
            if block["type"]=="paragraph":chunks.extend([block["text"],""])
            elif block["type"]=="figure":chunks.extend(["![Every strategy equity curve](../figures/all_strategies_100k.png)","",block["caption"],""])
            else:
                chunks.extend([block["caption"],""]);frame=block["frame"]
                chunks.append("| "+" | ".join(str(c).replace("\n","<br>") for c in frame.columns)+" |")
                chunks.append("| "+" | ".join("---" for _ in frame.columns)+" |")
                for row in frame.itertuples(index=False,name=None):chunks.append("| "+" | ".join(str(v) for v in row)+" |")
                chunks.append("")
    chunks.extend(["## Appendix — complete strategy definitions (final section)",""])
    for entry in entries:chunks.extend(["### "+entry["title"],"","Origin: "+entry["origin"]+".","",entry["explanation"],"","Exact configuration:","","```json",entry["exact_config"],"```",""])
    return "\n".join(chunks)


def _fonts():
    base=Path("C:/Windows/Fonts");names={"body":"Times-Roman","sans":"Helvetica","bold":"Helvetica-Bold","mono":"Courier"}
    for role,file in (("body","times.ttf"),("sans","arial.ttf"),("bold","arialbd.ttf"),("mono","consola.ttf")):
        if (base/file).exists():name="Lab"+role.title();pdfmetrics.registerFont(TTFont(name,str(base/file)));names[role]=name
    return names


def _pdf(pages,entries,path):
    fonts=_fonts()
    styles={"title":ParagraphStyle("LabTitle",fontName=fonts["bold"],fontSize=14,leading=16,spaceAfter=7,textColor=colors.HexColor("#12364D")),
        "body":ParagraphStyle("LabBody",fontName=fonts["body"],fontSize=10.2,leading=12.4,spaceAfter=7),
        "caption":ParagraphStyle("LabCaption",fontName=fonts["body"],fontSize=8.5,leading=10,spaceAfter=4),
        "cell":ParagraphStyle("LabCell",fontName=fonts["sans"],fontSize=8.5,leading=10),
        "head":ParagraphStyle("LabHead",fontName=fonts["bold"],fontSize=8.5,leading=10,textColor=colors.white),
        "entrytitle":ParagraphStyle("LabEntryTitle",fontName=fonts["bold"],fontSize=9.7,leading=11.4,spaceBefore=5,spaceAfter=4),
        "entry":ParagraphStyle("LabEntry",fontName=fonts["body"],fontSize=9.2,leading=10.6,spaceAfter=4),
        "config":ParagraphStyle("LabConfig",fontName=fonts["mono"],fontSize=6.8,leading=8.1,spaceAfter=7,splitLongWords=True,wordWrap="CJK")}
    def header(canvas,doc):
        w,h=canvas._pagesize;canvas.setFont(fonts["sans"],8);canvas.setFillColor(colors.HexColor("#52616B"));canvas.drawString(36,h-21,"EPL STRATEGY LAB — EXPLORATORY DRAFT | October 15, 2026");canvas.drawRightString(w-36,21,str(doc.page))
    doc=BaseDocTemplate(str(path),pagesize=letter,title=TITLE,author="Aman Dhillon and Panagiotis Housos")
    for orientation,size in (("portrait",letter),("landscape",landscape(letter))):
        w,h=size;doc.addPageTemplates(PageTemplate(id=orientation,pagesize=size,onPage=header,frames=[Frame(36,36,w-72,h-72,leftPadding=0,rightPadding=0,topPadding=0,bottomPadding=0)]))
    story=[]
    for i,page in enumerate(pages):
        if i:story.extend([NextPageTemplate(page["orientation"]),PageBreak()])
        width=(letter if page["orientation"]=="portrait" else landscape(letter))[0]-72
        story.append(Paragraph(html.escape(page["title"]),styles["title"]))
        for block in page["blocks"]:
            if block["type"]=="paragraph":story.append(Paragraph(html.escape(block["text"]),styles["body"]))
            elif block["type"]=="figure":
                with PILImage.open(block["path"]) as im:ratio=im.height/im.width
                story.extend([Image(str(block["path"]),width=width-10,height=(width-10)*ratio),Paragraph(html.escape(block["caption"]),styles["caption"])])
            else:
                story.append(Paragraph(html.escape(block["caption"]),styles["caption"]));frame=block["frame"]
                raw=[list(frame.columns)]+frame.values.tolist();cells=[[Paragraph(html.escape(str(v)).replace("\n","<br/>"),styles["head" if j==0 else "cell"]) for v in row] for j,row in enumerate(raw)]
                table=Table(cells,colWidths=[width*x for x in (.07,.055,.065,.075,.10,.10,.12,.12,.095,.10,.10)],repeatRows=1)
                table.setStyle(TableStyle([("BACKGROUND",(0,0),(-1,0),colors.HexColor("#12364D")),("ROWBACKGROUNDS",(0,1),(-1,-1),[colors.white,colors.HexColor("#F0F4F6")]),("VALIGN",(0,0),(-1,-1),"MIDDLE"),("LEFTPADDING",(0,0),(-1,-1),3),("RIGHTPADDING",(0,0),(-1,-1),3),("TOPPADDING",(0,0),(-1,-1),.8),("BOTTOMPADDING",(0,0),(-1,-1),.8)]));story.append(table)
    story.extend([NextPageTemplate("portrait"),PageBreak(),Paragraph("Appendix — complete strategy definitions (final section)",styles["title"])])
    for entry in entries:
        block=[Paragraph(html.escape(entry["title"]),styles["entrytitle"]),Paragraph(html.escape("Origin: "+entry["origin"]+". "+entry["explanation"]),styles["entry"]),
               Paragraph(html.escape("Exact configuration: "+entry["exact_config"]),styles["config"])]
        story.append(KeepTogether(block))
    doc.build(story)


def _docx(pages,entries,path):
    doc=Document();normal=doc.styles["Normal"];normal.font.name="Times New Roman";normal.font.size=Pt(10.2);normal.paragraph_format.line_spacing=Pt(12.4);normal.paragraph_format.space_after=Pt(7)
    h=doc.styles["Heading 1"];h.font.name="Arial";h.font.size=Pt(14);h.paragraph_format.space_before=Pt(0);h.paragraph_format.space_after=Pt(7)
    doc.core_properties.title=TITLE;doc.core_properties.author="Aman Dhillon and Panagiotis Housos"
    def section(index,land):
        s=doc.sections[0] if index==0 else doc.add_section(WD_SECTION_START.NEW_PAGE);s.orientation=WD_ORIENT.LANDSCAPE if land else WD_ORIENT.PORTRAIT
        s.page_width=Inches(11 if land else 8.5);s.page_height=Inches(8.5 if land else 11);s.top_margin=s.bottom_margin=s.left_margin=s.right_margin=Inches(.5);s.header_distance=s.footer_distance=Inches(.2)
        if index==0:
            r=s.header.paragraphs[0].add_run("EPL STRATEGY LAB — EXPLORATORY DRAFT | October 15, 2026");r.font.name="Arial";r.font.size=Pt(8)
            p=s.footer.paragraphs[0];p.alignment=2;f=OxmlElement("w:fldSimple");f.set(qn("w:instr"),"PAGE");p._p.append(f)
    for i,page in enumerate(pages):
        land=page["orientation"]=="landscape";section(i,land);doc.add_heading(page["title"],1)
        for block in page["blocks"]:
            if block["type"]=="paragraph":doc.add_paragraph(block["text"])
            else:
                if block["type"]=="figure":doc.add_picture(str(block["path"]),width=Inches(9.85 if land else 7.35))
                p=doc.add_paragraph(block["caption"]);p.paragraph_format.line_spacing=Pt(10)
                for r in p.runs:r.font.size=Pt(8.5)
                if block["type"]=="figure":continue
                frame=block["frame"];table=doc.add_table(rows=1,cols=len(frame.columns));table.style="Light Shading Accent 1";table.autofit=False
                weights=(.07,.055,.065,.075,.10,.10,.12,.12,.095,.10,.10)
                for col,weight in zip(table.columns,weights):col.width=Inches(10*weight)
                for cell,value in zip(table.rows[0].cells,frame.columns):cell.text=str(value)
                for values in frame.itertuples(index=False,name=None):
                    for cell,value in zip(table.add_row().cells,values):cell.text=str(value)
                for row in table.rows:
                    for cell,weight in zip(row.cells,weights):
                        cell.width=Inches(10*weight)
                        for p in cell.paragraphs:
                            p.paragraph_format.line_spacing=Pt(10);p.paragraph_format.space_after=Pt(0)
                            for r in p.runs:r.font.name="Arial";r.font.size=Pt(8.5)
    section(len(pages),False);doc.add_heading("Appendix — complete strategy definitions (final section)",1)
    for entry in entries:
        h=doc.add_paragraph();h.paragraph_format.keep_with_next=True;r=h.add_run(entry["title"]);r.bold=True;r.font.name="Arial";r.font.size=Pt(9.7)
        p=doc.add_paragraph("Origin: "+entry["origin"]+". "+entry["explanation"]);p.paragraph_format.keep_with_next=True;p.paragraph_format.line_spacing=Pt(10.6)
        for r in p.runs:r.font.size=Pt(9.2)
        p=doc.add_paragraph("Exact configuration: "+entry["exact_config"]);p.paragraph_format.keep_together=True;p.paragraph_format.line_spacing=Pt(8.1)
        for r in p.runs:r.font.name="Consolas";r.font.size=Pt(6.8)
    doc.save(path)


def _workbook(path,catalog,metrics,graph,bets,equity,entries,registry):
    ordered=metrics.set_index("strategy_id").reindex([c["strategy_id"] for c in catalog]).reset_index()
    desired=["strategy_id","name","strategy_name","origin","model_key","bet_count","hit_ratio","avg_winning_gain_units","avg_losing_loss_units","expected_gain_per_trade_units","expected_loss_per_trade_units","expected_net_profit_per_trade_units","realized_expectancy_per_trade_units","trading_frequency","profit","turnover","roi","expected_roi","return_on_initial_capital","max_drawdown","final_capital_sum_units","profit_factor","payoff_ratio","avg_stake_units","avg_selected_probability","avg_decimal_odds","roi_ci_low","roi_ci_high"]
    definitions=pd.DataFrame([("Capital","Each strategy independently starts 100,000 monetary units; no portfolio sum and no reset."),
        ("Expected components","Probability-weighted positive/negative P&L at recorded allocated stakes, conditional on the realised Kelly path."),
        ("Hit and frequency","Hit=wins/trades; frequency=trades/374 priced opportunities. Neither is full-match classification accuracy."),
        ("ROI and capital return","ROI=net profit/turnover; capital return=net profit/100000."),
        ("Undefined/infinity","Dash is unavailable, never invented zero; positive infinity retained where gains exist without losses."),
        ("Inference","All 84 recipes compare already examined 2025/26 data. No refit, fresh holdout or recommended winner."),
        ("Original policies","50 previously specified rules including 12 diagnostic thresholds; 34 newly exploratory rules. Old season/scope repetitions are consolidated."),
        ("Pending/cost studies","Pending future forecasts receive no invented settled account; earlier cost stress cases remain separate assumptions.")],columns=["Item","Definition"])
    sheets={"Read me":definitions,"Metrics overview":ordered[[c for c in desired if c in ordered]],"All metrics":ordered,
            "Complete catalog":pd.DataFrame(catalog),"Final strategy appendix":pd.DataFrame(entries),"Graph equity incl initial":graph,
            "Date equity":equity,"Every recorded bet":bets,"Input provenance":pd.DataFrame([{"path":p,**r} for p,r in registry.items()])}
    percent={"hit_ratio","trading_frequency","roi","expected_roi","return_on_initial_capital","max_drawdown","avg_selected_probability","roi_ci_low","roi_ci_high","trade_rate_all_predictions","active_date_fraction"}
    with pd.ExcelWriter(path,engine="xlsxwriter",engine_kwargs={"options":{"strings_to_urls":False}}) as writer:
        book=writer.book;head=book.add_format({"bold":True,"font_color":"white","bg_color":"#12364D","text_wrap":True});num=book.add_format({"num_format":"#,##0.0000"});pct=book.add_format({"num_format":"0.00%"})
        for name,frame in sheets.items():
            safe=frame.copy()
            for c in safe:
                if isinstance(safe[c].dtype,pd.DatetimeTZDtype):safe[c]=safe[c].dt.tz_convert(None)
                if safe[c].dtype=="object":safe[c]=safe[c].map(lambda v:_json(v) if isinstance(v,(dict,list,tuple)) else v)
                elif pd.api.types.is_numeric_dtype(safe[c]):safe[c]=safe[c].map(lambda v:"+∞" if v==float("inf") else "−∞" if v==float("-inf") else v)
            safe.to_excel(writer,sheet_name=name,index=False,na_rep="—");ws=writer.sheets[name];ws.write_row(0,0,list(safe.columns),head);ws.set_row(0,42);ws.freeze_panes(1,1);ws.autofilter(0,0,len(safe),len(safe.columns)-1)
            for j,c in enumerate(safe.columns):ws.set_column(j,j,55 if c in ("explanation","exact_config","Definition","path","name","strategy_name") else min(40,max(14,len(c)*.8)),pct if c in percent else num)
    return {key:len(value) for key,value in sheets.items()}


def build_report(output_dir:Path)->dict:
    output_dir=Path(output_dir).resolve();catalog=json.loads((output_dir/"catalog.json").read_text(encoding="utf-8"))
    metrics=pd.read_parquet(output_dir/"metrics.parquet");bets=pd.read_parquet(output_dir/"bets.parquet");equity=pd.read_parquet(output_dir/"equity.parquet");graph=pd.read_csv(output_dir/"graph_equity.csv")
    graph,equity,ids,anchor=_validate(catalog,metrics,graph,equity,bets);entries=_appendix(catalog)
    if [e["strategy_id"] for e in entries]!=ids:raise ValueError("Appendix IDs/order differ from catalog")
    # The final envelope is intentionally written after report generation.
    sources=[output_dir/name for name in ("catalog.json","metrics.parquet","bets.parquet","equity.parquet","graph_equity.csv","inputs_audit.json","source_provenance.json") if (output_dir/name).exists()]
    sources.extend([Path(__file__).resolve(),ROOT/"epl_strategy/report.py"])
    sources.extend(ROOT/name for name in ("epl_modeling/METHODOLOGY.md","epl_modeling/ABLATION_PROTOCOL.md","epl_strategy/METHODOLOGY.md","epl_strategy_lab/PROTOCOL.md"))
    sources=sorted(set(sources));registry={_relative(p):{"sha256":_sha(p),"bytes":p.stat().st_size} for p in sources}
    out=output_dir/"report";out.mkdir(exist_ok=True);figures=output_dir/"figures"
    graph_info=_graphs(catalog,graph,figures,entries);pages,display,origins=_body(catalog,metrics,ids,figures/"all_strategies_100k.png")
    display.to_csv(out/"all_strategy_display_metrics.csv",index=False)
    pd.DataFrame(entries).to_csv(out/"complete_strategy_appendix.csv",index=False)
    md=out/"strategy_lab_report.md";pdf=md.with_suffix(".pdf");docx=md.with_suffix(".docx")
    md.write_text(_markdown(pages,entries,out),encoding="utf-8");_pdf(pages,entries,pdf);_docx(pages,entries,docx)
    sheets=_workbook(out/"strategy_lab_overview.xlsx",catalog,metrics,graph,bets,equity,entries,registry)
    reader=PdfReader(pdf);text="\n".join(page.extract_text() or "" for page in reader.pages)
    appendix_text=text.split("Appendix — complete strategy definitions (final section)",1)
    if len(appendix_text)!=2 or any(" ".join(entry["title"].split()) not in " ".join(appendix_text[1].split()) for entry in entries):raise AssertionError("PDF omitted an appendix definition")
    if entries[-1]["strategy_id"] not in reader.pages[-1].extract_text():raise AssertionError("Appendix is not the final report content")
    changed=[_relative(p) for p in sources if _sha(p)!=registry[_relative(p)]["sha256"]]
    if changed:raise AssertionError("Protected inputs changed during reporting: "+str(changed))
    produced=[p for folder in (out,figures) for p in folder.rglob("*") if p.is_file() and p.name!="report_manifest.json" and "previews" not in p.parts]
    outputs={p.relative_to(output_dir).as_posix():{"sha256":_sha(p),"bytes":p.stat().st_size} for p in sorted(produced)}
    manifest={"status":"exploratory_draft","draft_date":DRAFT_DATE,"built_at_utc":datetime.now(timezone.utc).isoformat(),"strategy_count":len(ids),"catalog_ids":ids,
        "metrics_ids":metrics.strategy_id.tolist(),"appendix_ids":[e["strategy_id"] for e in entries],"graph":graph_info,"native_match_dates":113,"matches":374,
        "initial_bankroll_per_strategy":INITIAL,"initial_anchor_date":anchor.strftime("%Y-%m-%d"),"graph_point_count":len(graph),"origins":origins,
        "pdf_pages":len(reader.pages),"body_sections":len(pages),"appendix_is_last":True,"workbook_sheet_rows":sheets,"models_refitted":False,
        "consumed_source_files_unchanged":not changed,"inputs":registry,"outputs":outputs,
        "manifest_order":"Builder final manifest and independent_audit are intentionally not report inputs; no circular hashing.",
        "limits":"Already inspected common-season archived-price replay; separate independently funded accounts, not a portfolio or prospective performance."}
    (out/"report_manifest.json").write_text(json.dumps(manifest,ensure_ascii=False,indent=2,allow_nan=False),encoding="utf-8")
    return {"output_directory":str(out),"strategy_count":len(ids),"pdf_pages":len(reader.pages),"all_curves_initially_visible":len(ids),"appendix_entries":len(entries),"source_files_unchanged":not changed}


def main():
    parser=argparse.ArgumentParser(description=__doc__);parser.add_argument("--results",required=True,type=Path);args=parser.parse_args();print(json.dumps(build_report(args.results),indent=2))


if __name__=="__main__":main()
