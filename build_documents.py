"""Build the complete reading guide and accessible proposal exports."""
from pathlib import Path
import json
import re
from urllib.parse import quote
from xml.sax.saxutils import escape

import markdown
from bs4 import BeautifulSoup, NavigableString
from docx import Document
from docx.shared import Inches, Pt, RGBColor
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.opc.constants import RELATIONSHIP_TYPE as RT
from reportlab.lib import colors
from reportlab.lib.enums import TA_LEFT
from reportlab.lib.styles import getSampleStyleSheet, ParagraphStyle
from reportlab.lib.pagesizes import letter
from reportlab.pdfbase import pdfmetrics
from reportlab.pdfbase.ttfonts import TTFont
from reportlab.platypus import SimpleDocTemplate, Paragraph, Spacer, KeepTogether

ROOT = Path(__file__).resolve().parent
inventory = json.loads((ROOT / "paper_inventory.json").read_text(encoding="utf-8"))
groups = [
    ("Statistical arbitrage and trading decisions", "reading_stat_arb.md"),
    ("Filtering and machine learning", "reading_ml_filters.md"),
    ("Trends, risk and backtest validity", "reading_trends_risk.md"),
    ("Market structure and forecasting", "reading_market_structure.md"),
]
lines = [
    "# Reading guide: all 33 supplied PDFs",
    "",
    "**Completed September 25, 2026. Coverage: 33 documents, 874 PDF pages.**",
    "",
    "The complete texts were reviewed in parallel before the project proposal was selected. "
    "Methods, results, appendices and references were included; relevant figures and damaged "
    "equations were checked visually where noted. This is a literature review, not a replication "
    "of every empirical result or an independent verification of every proof. All seven scanned "
    "pages of Morrison and Pike were read visually. Frequency Arbitrage is a one-page teaching "
    "handout; the volatility-surface item is a slide deck. The folder also contains surveys, "
    "software/tutorial papers and an engineering application, so not every item supplies financial "
    "backtest evidence.",
    "",
    "Each annotation records the actual title, version, summary, page references and limitations. "
    "PDF page numbers include cover sheets. Sports applications are proposed methodological "
    "transfers, not findings established by the supplied papers.",
    "",
    "## Complete source inventory",
    "",
    "| # | Supplied file | PDF pages |",
    "|---|---|---:|",
]
for i, item in enumerate(inventory, 1):
    lines.append(f"| {i} | [{item['file']}](../papers/{quote(item['file'])}) | {item['pages']} |")
lines += [
    "",
    "## Synthesis for the Premier League project",
    "",
    "The strongest common foundation is a transparent relative-value signal, causal estimation, "
    "cost-sensitive decisions, finite-horizon exits and chronological evaluation. Financial-price "
    "mean reversion cannot be assumed for bounded match contracts. A bookmaker-derived reference "
    "probability is not an executable hedge, and a shrinking discrepancy can result from the "
    "reference moving rather than the held contract gaining value. The strongest project therefore "
    "tests actual prediction-market price adjustment and books cash P&L from entry asks and exit "
    "bids. FPL-related information belongs in a small, separately evaluated extension.",
    "",
    "The reviewed papers repeatedly distinguish forecasting from tradable profitability and "
    "warn about transaction costs, changing relationships, model search and optimistic backtests. "
    "A negative net result is informative. Complex copulas, wavelets, particle filters, deep "
    "learning and market-making models should not all be combined into one student project.",
    "",
]
for title, filename in groups:
    content = (ROOT / filename).read_text(encoding="utf-8").replace("../Papers/", "../papers/")
    content = content.split("\n", 1)[1].lstrip()
    content = re.sub(r"^## ", "### ", content, flags=re.M)
    lines += [f"## {title}", "", content, ""]
(ROOT / "reading_guide.md").write_text("\n".join(lines), encoding="utf-8")

source = (ROOT / "project_proposal.md").read_text(encoding="utf-8")
soup = BeautifulSoup(markdown.markdown(source), "html.parser")

# Word export with proper external hyperlinks and editable paragraphs.
doc = Document()
section = doc.sections[0]
section.top_margin = Inches(0.65)
section.bottom_margin = Inches(0.65)
section.left_margin = Inches(0.75)
section.right_margin = Inches(0.75)
normal = doc.styles["Normal"]
normal.font.name = "Arial"
normal.font.size = Pt(10.5)
normal.paragraph_format.space_after = Pt(6)
normal.paragraph_format.line_spacing = 1.08
for style_name in ["Title", "Heading 1", "Heading 2"]:
    doc.styles[style_name].font.name = "Arial"
    doc.styles[style_name].font.color.rgb = RGBColor.from_string("173D50")
doc.styles["Title"].font.size = Pt(21)
doc.styles["Heading 1"].font.size = Pt(12)

def add_inline(paragraph, node, bold=False, italic=False):
    if isinstance(node, NavigableString):
        run = paragraph.add_run(str(node))
        run.bold, run.italic = bold, italic
    elif node.name == "br":
        paragraph.add_run().add_break()
    elif node.name == "a":
        hyperlink = OxmlElement("w:hyperlink")
        hyperlink.set(qn("r:id"), paragraph.part.relate_to(node["href"], RT.HYPERLINK, is_external=True))
        run = OxmlElement("w:r")
        properties = OxmlElement("w:rPr")
        color = OxmlElement("w:color")
        color.set(qn("w:val"), "1D6482")
        properties.append(color)
        run.append(properties)
        text = OxmlElement("w:t")
        text.text = node.get_text()
        run.append(text)
        hyperlink.append(run)
        paragraph._p.append(hyperlink)
    else:
        for child in node.children:
            add_inline(paragraph, child, bold or node.name in ("b", "strong"), italic or node.name in ("i", "em"))

for element in soup.children:
    if isinstance(element, NavigableString):
        continue
    if element.name == "h1":
        doc.add_paragraph(element.get_text(), style="Title")
    elif element.name == "h2":
        doc.add_paragraph(element.get_text(), style="Heading 1")
    elif element.name == "p":
        p = doc.add_paragraph()
        add_inline(p, element)
    elif element.name in ("ol", "ul"):
        for li in element.find_all("li", recursive=False):
            p = doc.add_paragraph(style="List Number" if element.name == "ol" else "List Bullet")
            add_inline(p, li)
footer = section.footer.paragraphs[0]
footer.alignment = 2
footer.add_run("Statistical Arbitrage | Project Proposal | ")
field = OxmlElement("w:fldSimple")
field.set(qn("w:instr"), "PAGE")
footer._p.append(field)
doc.save(ROOT / "project_proposal.docx")

# PDF export, using embedded Unicode fonts for readable symbols and accents.
font_root = Path("C:/Windows/Fonts")
for name, file in [("ProposalArial", "arial.ttf"), ("ProposalArialBold", "arialbd.ttf"), ("ProposalArialItalic", "ariali.ttf"), ("ProposalArialBoldItalic", "arialbi.ttf")]:
    pdfmetrics.registerFont(TTFont(name, str(font_root / file)))
pdfmetrics.registerFontFamily("ProposalArial", normal="ProposalArial", bold="ProposalArialBold", italic="ProposalArialItalic", boldItalic="ProposalArialBoldItalic")
styles = getSampleStyleSheet()
styles.add(ParagraphStyle(name="ProposalBody", fontName="ProposalArial", fontSize=10, leading=14, spaceAfter=7, textColor=colors.HexColor("#202B32")))
styles.add(ParagraphStyle(name="ProposalTitle", fontName="ProposalArialBold", fontSize=20, leading=24, spaceAfter=11, textColor=colors.HexColor("#173D50")))
styles.add(ParagraphStyle(name="ProposalHeading", fontName="ProposalArialBold", fontSize=11.5, leading=15, spaceBefore=8, spaceAfter=5, keepWithNext=True, textColor=colors.HexColor("#173D50")))
styles.add(ParagraphStyle(name="ProposalRef", parent=styles["ProposalBody"], fontSize=9, leading=12, leftIndent=12, firstLineIndent=-12))

def pdf_inline(element):
    html = "".join(str(child) for child in element.children)
    html = html.replace("<strong>", "<b>").replace("</strong>", "</b>").replace("<em>", "<i>").replace("</em>", "</i>")
    html = re.sub(r'<a href="([^"]+)">', r'<a href="\1" color="#1D6482">', html)
    return html

story = []
for element in soup.children:
    if isinstance(element, NavigableString):
        continue
    if element.name in ("h1", "h2", "p"):
        sty = {"h1": "ProposalTitle", "h2": "ProposalHeading", "p": "ProposalBody"}[element.name]
        story.append(Paragraph(pdf_inline(element), styles[sty]))
    elif element.name in ("ol", "ul"):
        for i, li in enumerate(element.find_all("li", recursive=False), 1):
            story.append(Paragraph(f"{i}. " + pdf_inline(li), styles["ProposalRef"]))

def footer_pdf(canvas, document):
    canvas.saveState()
    canvas.setFont("ProposalArial", 8)
    canvas.setFillColor(colors.HexColor("#657781"))
    canvas.drawString(54, 26, "Statistical Arbitrage | Premier League Project Proposal")
    canvas.drawRightString(letter[0] - 54, 26, str(document.page))
    canvas.restoreState()

pdf = SimpleDocTemplate(str(ROOT / "project_proposal.pdf"), pagesize=letter, rightMargin=54, leftMargin=54, topMargin=42, bottomMargin=45, title="Cost-Aware Statistical Arbitrage in Premier League Betting Markets", author="")
pdf.build(story, onFirstPage=footer_pdf, onLaterPages=footer_pdf)
print("Built reading_guide.md, project_proposal.docx, project_proposal.pdf")
