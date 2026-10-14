"""Export readable views of released forecasts without changing their ledger."""
from __future__ import annotations

import argparse
from datetime import datetime, timezone
import json
from pathlib import Path

import pandas as pd
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

from epl_modeling.experiment import save_table, write_json
from .ledger import earliest_forecasts, read_forecasts, verify_ledger
from .train import ROOT, PROJECT, sha256


def export_overview(run_directory: Path) -> Path:
    run_directory = Path(run_directory).resolve()
    run = json.loads((run_directory / "run_summary.json").read_text(encoding="utf-8"))
    ledger_path = ROOT / "forecasts.jsonl"
    integrity = verify_ledger(ledger_path, "forecast")
    ledger = earliest_forecasts(read_forecasts(ledger_path))
    fixture_ids = set(pd.read_parquet(run_directory / "target_fixtures.parquet").match_id)
    # This view describes only forecasts available by the saved run's completion.
    cutoff = pd.Timestamp(run["created_at_utc"])
    records = [row for row in ledger if row["match_id"] in fixture_ids
               and pd.Timestamp(row["created_at_utc"]) <= cutoff]
    rows = []
    for record in records:
        p = record["probabilities"]
        pick = max(("H", "D", "A"), key=lambda label: p[label])
        quote = record.get("market_quote")
        rows.append({"match_id": record["match_id"], "kickoff_utc": record["kickoff_utc"],
                     "home_team": record["home_team"], "away_team": record["away_team"],
                     "feature_group": record.get("feature_group", "market_only"),
                     "model_id": record["model_id"],
                     "p_H": p["H"], "p_D": p["D"], "p_A": p["A"],
                     "most_likely_outcome": pick,
                     "forecast_created_at_utc": record["created_at_utc"],
                     "quote_observed_at_utc": quote.get("observed_at_utc") if quote else None,
                     "quote_publication_time_verified": False,
                     "settlement_status_at_run": "pending",
                     "model_sha256": record["model_sha256"],
                     "protocol_sha256": record["protocol_sha256"]})
    columns = ["match_id", "kickoff_utc", "home_team", "away_team", "feature_group", "model_id",
               "p_H", "p_D", "p_A", "most_likely_outcome", "forecast_created_at_utc",
               "quote_observed_at_utc", "quote_publication_time_verified", "settlement_status_at_run",
               "model_sha256", "protocol_sha256"]
    forecasts = pd.DataFrame(rows, columns=columns).sort_values(["kickoff_utc", "home_team", "model_id"])
    coverage = pd.read_parquet(run_directory / "forecast_coverage.parquet")
    released = {(row["match_id"], row["model_id"]) for row in records}
    coverage["released_by_run_completion"] = [(row.match_id, row.model_id) in released
                                              for row in coverage.itertuples()]
    identifiers = ["match_id", "kickoff_utc", "home_team", "away_team"]
    if len(forecasts):
        wide = forecasts.pivot(index=identifiers, columns="model_id", values=["p_H", "p_D", "p_A"])
        wide.columns = [f"{model_id.removeprefix('prospective_v1_')}_{probability}"
                        for probability, model_id in wide.columns]
        wide = wide.reset_index().sort_values(["kickoff_utc", "home_team"])
    else:
        wide = pd.DataFrame(columns=identifiers)
    snapshot = PROJECT / run["snapshot_path"].replace("\\", "/")
    model_dir = PROJECT / run["models_path"].replace("\\", "/")
    model_manifest = json.loads((model_dir / "manifest.json").read_text(encoding="utf-8"))
    model_choices = pd.DataFrame(model_manifest["models"])
    freshness = pd.DataFrame([{"source": name, "latest_completed_match_date": date}
                              for name, date in run["freshness"].items() if "last_completed_match_date" in name])
    notes = pd.DataFrame([
        ("View scope", "Earliest released model/fixture records available at saved run completion; a dated view, not a live feed."),
        ("Run date (UTC)", run["run_date"]),
        ("Snapshot id", run["snapshot_id"]),
        ("Released probability records", len(forecasts)),
        ("Distinct released fixtures", forecasts.match_id.nunique()),
        ("Frozen candidate models", len(model_manifest["models"])),
        ("Observed reference quote fixtures", run["observed_reference_quote_matches"]),
        ("Outcome status", "All released forecasts were pending at this run. No prospective performance or profit is available."),
        ("Probabilities", "H=home win; D=draw; A=away win. Each row sums to one. Most likely outcome is not a betting recommendation."),
        ("Market coverage", "Quote-dependent variants are held until a complete recent Bet365 triplet is observed. See Coverage."),
        ("Source clocks", "Capture times record our observations. Provider publication/update times remain unknown."),
        ("Model comparison", "All declared alternatives are retained; no new-outcome winner selection has occurred."),
        ("Fitting chronology", "2016/17-2023/24 training; 2024/25 tuning and refit; 2025/26 temperature calibration. No 2026/27 fitting outcomes."),
        ("Ledger hash at export", integrity["head_hash"]),
        ("Original run ledger hash", run["ledger_integrity"]["head_hash"]),
    ], columns=["item", "value"])
    output = run_directory / "overview"
    output.mkdir(exist_ok=False)
    save_table(forecasts, output, "released_forecasts")
    save_table(wide, output, "fixture_probability_comparison")
    workbook = output / "EPL_live_forecasts.xlsx"
    with pd.ExcelWriter(workbook, engine="openpyxl") as writer:
        for name, table in (("Read me", notes), ("Forecasts", forecasts), ("Fixture comparison", wide),
                            ("Coverage", coverage), ("Models", model_choices), ("Source freshness", freshness)):
            table.to_excel(writer, sheet_name=name, index=False)
            sheet = writer.sheets[name]
            sheet.freeze_panes = "A2"
            sheet.auto_filter.ref = sheet.dimensions
            for cell in sheet[1]:
                cell.font = Font(color="FFFFFF", bold=True)
                cell.fill = PatternFill("solid", fgColor="1E3A5F")
                cell.alignment = Alignment(wrap_text=True, vertical="top")
            for index, column in enumerate(table.columns, 1):
                width = 22 if str(column).startswith("p_") or "_p_" in str(column) else 28
                sheet.column_dimensions[get_column_letter(index)].width = width
                if str(column).startswith("p_") or "_p_" in str(column):
                    for cell in list(sheet.columns)[index - 1][1:]:
                        cell.number_format = "0.0%"
            if name == "Read me":
                sheet.column_dimensions["A"].width = 34
                sheet.column_dimensions["B"].width = 100
                for row in sheet.iter_rows(min_row=2):
                    row[1].alignment = Alignment(wrap_text=True, vertical="top")
                    sheet.row_dimensions[row[0].row].height = 33
    lines = ["# Released EPL forecast overview", "", f"Saved run: {run['run_date']} UTC.", "",
             f"{len(forecasts)} probability records for {forecasts.match_id.nunique()} future fixtures. "
             "All were pending at this run; no prospective accuracy, loss or profit is reported.", "",
             "[Excel workbook](EPL_live_forecasts.xlsx) · [All forecast probabilities](released_forecasts.csv) · "
             "[Fixture comparison](fixture_probability_comparison.csv)", "",
             "H, D and A denote home win, draw and away win. The four logistic alternatives remain separate; "
             "none has been selected using these future outcomes. This table records model estimates, not wagers.", "",
             "| Kickoff UTC | Fixture | Feature group | Home | Draw | Away |", "|---|---|---|---:|---:|---:|"]
    for row in forecasts.itertuples():
        lines.append(f"| {pd.Timestamp(row.kickoff_utc).strftime('%b %d %H:%M')} | {row.home_team} vs {row.away_team} | "
                     f"{row.feature_group} | {row.p_H:.1%} | {row.p_D:.1%} | {row.p_A:.1%} |")
    (output / "FORECASTS.md").write_text("\n".join(lines) + "\n", encoding="utf-8")
    inputs = [ledger_path, run_directory / "run_summary.json", run_directory / "target_fixtures.parquet",
              run_directory / "forecast_coverage.parquet", snapshot / "manifest.json", model_dir / "manifest.json"]
    manifest = {"exported_at_utc": datetime.now(timezone.utc).isoformat(),
                "view_cutoff_utc": run["created_at_utc"], "run_id": run["run_id"],
                "forecast_records": len(forecasts), "fixtures": int(forecasts.match_id.nunique()),
                "ledger_integrity": integrity, "input_sha256": {str(p.relative_to(PROJECT)): sha256(p) for p in inputs},
                "output_sha256": {p.name: sha256(p) for p in output.iterdir() if p.is_file()},
                "code_sha256": sha256(Path(__file__))}
    write_json(output / "overview_manifest.json", manifest)
    return output


def main():
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument("--run", type=Path, required=True)
    args = parser.parse_args()
    print(export_overview(args.run))


if __name__ == "__main__":
    main()
