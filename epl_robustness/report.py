"""Build a separate robustness companion from saved stress/diagnostic outputs.

Run: python -m epl_robustness.report --results PATH
No forecasts, selections, stakes, model fits or original reports are changed.
"""
from __future__ import annotations

import argparse
import hashlib
import html
import json
from datetime import datetime, timezone
from pathlib import Path

import matplotlib
matplotlib.use("Agg")
import matplotlib.pyplot as plt
import numpy as np
import pandas as pd
from docx import Document
from docx.enum.section import WD_ORIENT, WD_SECTION_START
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt
from PIL import Image as PILImage
from pypdf import PdfReader
from reportlab.lib import colors
from reportlab.lib.pagesizes import landscape, letter
from reportlab.lib.styles import ParagraphStyle
from reportlab.pdfbase import pdfmetrics
from reportlab.pdfbase.ttfonts import TTFont
from reportlab.platypus import BaseDocTemplate, Frame, Image, NextPageTemplate, PageBreak, PageTemplate, Paragraph, Spacer, Table, TableStyle

from epl_strategy.report import number, _label

ROOT=Path(__file__).resolve().parents[1]
TITLE="EPL strategy robustness — exploratory companion draft"
DATE="2026-10-14"
SCENARIOS={
    "baseline":"Base", "fee_0p25pct":"Fee\n0.25%", "fee_0p5pct":"Fee\n0.5%", "fee_1pct":"Fee\n1%", "fee_2pct":"Fee\n2%",
    "haircut_1pct":"Payout\n−1%", "haircut_2p5pct":"Payout\n−2.5%", "haircut_5pct":"Payout\n−5%",
    "combined_fee1pct_haircut2p5pct":"Fee 1% +\npayout −2.5%",
}
COMBINED="combined_fee1pct_haircut2p5pct"
TABLES=("cost_sensitivity","break_even","funding_details","diagnostic_availability","probability_bins","odds_bins","edge_bins",
        "selected_bet_scores","date_concentration","date_influence","season_fragility")


def _sha(path):
    digest=hashlib.sha256()
    with path.open("rb") as source:
        for block in iter(lambda:source.read(1024*1024),b""):digest.update(block)
    return digest.hexdigest()


def _path_label(path):
    try:return path.relative_to(ROOT).as_posix()
    except ValueError:return str(path)


def _read(results):
    tables={}
    for name in TABLES:
        aliases=(name,"availability") if name=="diagnostic_availability" else (name,)
        candidates=[results/(alias+".parquet") for alias in aliases]
        candidates += [results/"diagnostics"/(alias+".parquet") for alias in aliases]
        found=[p for p in candidates if p.exists()]
        if len(found)!=1:raise ValueError(f"Expected one source for {name}, found {found}")
        tables[name]=pd.read_parquet(found[0])
    return tables


def _focus(frame,which):
    if which=="first":part=frame.loc[frame.study.eq("first") & frame.scope.eq("test") & frame.split.eq("test")]
    else:part=frame.loc[frame.study.eq("ablation") & frame.scope.eq("pooled_seasons")]
    order=[c for c in ("policy","branch","model_type","feature_group","odds_mode","staking") if c in part]
    return part.sort_values(order,kind="stable").copy()


def _validate(tables):
    cost=tables["cost_sensitivity"]
    required={"scenario_id","fee_rate","payout_haircut","study","scope","strategy_id","episode_id","split","bet_count","hit_ratio",
              "expected_gain_total_units","expected_loss_total_units","expected_gain_per_trade_units","expected_loss_per_trade_units",
              "expected_net_profit_per_trade_units","avg_winning_gain_units","avg_losing_loss_units","trading_frequency",
              "realized_expectancy_per_trade_units","profit","turnover","roi","roi_on_cash_outlay","expected_roi",
              "max_drawdown","shadow_max_drawdown","funding_feasible"}
    if required-set(cost):raise ValueError("Missing stress columns: "+str(sorted(required-set(cost))))
    if set(cost.scenario_id)!=set(SCENARIOS):raise ValueError("Scenario set differs from the declared nine stress assumptions")
    baseline=cost.loc[cost.scenario_id.eq("baseline")]
    if len(baseline)!=158 or len(cost)!=158*9:raise ValueError("Expected158 baseline strategies and1422 scenario rows")
    keys=["study","scope","strategy_id","episode_id"]
    if cost.duplicated(keys+["scenario_id"]).any():raise ValueError("Duplicate scenario/strategy episode")
    if not cost.groupby(keys,dropna=False).size().eq(9).all():raise ValueError("An original scope is missing a scenario")
    if len(_focus(baseline,"first"))!=16 or len(_focus(baseline,"pooled"))!=22:raise ValueError("Incomplete first-test/pooled strategy coverage")
    if cost.loc[cost.funding_feasible.eq(False),"max_drawdown"].notna().any():raise ValueError("An unfundable shadow path has a conventional drawdown value")
    for _,part in cost.groupby(keys,dropna=False):
        if part.bet_count.nunique(dropna=False)!=1 or part.hit_ratio.nunique(dropna=False)!=1 or part.trading_frequency.nunique(dropna=False)!=1:
            raise ValueError("Selections or activity changed across fixed-exposure scenarios")
    return baseline


def _sensitivity(frame):
    rows=[]
    for _,part in frame.groupby("strategy_id",sort=False):
        if len(part)!=9:raise ValueError("Displayed strategy lacks nine scenarios")
        row={"Strategy":_label(part.iloc[0])}
        for scenario,label in SCENARIOS.items():
            item=part.loc[part.scenario_id.eq(scenario)].iloc[0]
            row[label]=number(item.roi,percent=True)+( "†" if item.funding_feasible==False else "")
        rows.append(row)
    return pd.DataFrame(rows)


def _stress_metrics(frame):
    rows=[]
    for _,row in frame.iterrows():
        funding="Yes / "+number(row.max_drawdown,percent=True) if row.funding_feasible else "No / —\nShadow "+number(row.shadow_max_drawdown,percent=True)
        rows.append({"Strategy":_label(row),"Trades":number(row.bet_count,0),"Hit %":number(row.hit_ratio,percent=True),
                     "Trade\nfreq. %":number(row.trading_frequency,percent=True),"Avg. +\ngain":number(row.avg_winning_gain_units),
                     "Avg. −\nloss size":number(row.avg_losing_loss_units),"Expected\ngain/trade":number(row.expected_gain_per_trade_units),
                     "Expected\nloss/trade":number(row.expected_loss_per_trade_units),"Realised\nnet/trade":number(row.realized_expectancy_per_trade_units),
                     "Fundable /\nmax. DD":funding})
    return pd.DataFrame(rows)


def _concentration(frame):
    rows=[]
    for _,row in frame.iterrows():
        interval=lambda low,high:"—" if pd.isna(low) or pd.isna(high) else f"[{number(low,percent=True)}, {number(high,percent=True)}]"
        rows.append({"Strategy":_label(row),"Base ROI":number(row.roi,percent=True),
                     "Top 1 date\npositive share":number(row.top1_positive_date_profit_share,percent=True),
                     "Top 5 dates\npositive share":number(row.top5_positive_date_profit_share,percent=True),
                     "Delete-one-date\nROI range":interval(row.leave_one_date_out_roi_min,row.leave_one_date_out_roi_max),
                     "Profitable\nseasons":number(row.profitable_season_count,0)+"/"+number(row.season_count,0),
                     "Delete-one-season\nROI range":interval(row.leave_one_season_out_roi_min,row.leave_one_season_out_roi_max)})
    return pd.DataFrame(rows)


def _tail(tables,out):
    odds=tables["odds_bins"]
    selected=odds.loc[odds.study.eq("first") & odds.scope.eq("test") & odds.branch.eq("calibration") &
                      odds.policy.eq("predeclared_5pct") & odds.odds_mode.eq("best_individual_book") & odds.staking.eq("fixed")].copy()
    selected=selected.sort_values("bin_index")
    if selected.empty or int(selected.bet_count.sum())!=333:raise ValueError("Missing the complete333-bet calibration tail illustration")
    tail=selected.loc[selected.bin_lower.ge(10)]
    if tail.empty or not tail.bin_lower.eq(10).any():raise ValueError("Saved odds bins do not support an exact odds≥10 illustration")
    n=int(tail.bet_count.sum());wins=int(tail.wins.sum())
    mean_p=float((tail.mean_probability*tail.bet_count).sum()/n)
    trade_table=pd.DataFrame([{"Odds bin":row.bin_label,"Trades":number(row.bet_count,0),"Wins":number(row.wins,0),
        "Mean p":number(row.mean_probability,percent=True),"Hit ratio":number(row.wins/row.bet_count if row.bet_count else None,percent=True),
        "Expected ROI":number(row.expected_roi,percent=True),"Realised ROI":number(row.roi,percent=True)} for _,row in selected.iterrows()])
    figures=out/"figures";figures.mkdir(exist_ok=True)
    fig,axes=plt.subplots(1,2,figsize=(12.2,3.8),layout="constrained")
    x=np.arange(len(selected));labels=selected.bin_label.astype(str).tolist();width=.36
    axes[0].bar(x-width/2,selected.mean_probability.to_numpy(float)*100,width,label="Mean forecast p",color="#24638C")
    hit=np.divide(selected.wins.to_numpy(float),selected.bet_count.to_numpy(float),out=np.full(len(selected),np.nan),where=selected.bet_count.to_numpy(float)>0)
    axes[0].bar(x+width/2,hit*100,width,label="Observed hit ratio",color="#C17B17")
    for j,count in enumerate(selected.bet_count):
        if count:axes[0].text(j,max(float(selected.mean_probability.iloc[j]),hit[j])*100+2,f"n={int(count)}"+("*" if count<30 else ""),ha="center",fontsize=8)
    axes[0].set_ylabel("Selected-outcome probability (%)");axes[0].set_title("Selected forecasts versus outcomes")
    axes[1].bar(x-width/2,selected.expected_roi.to_numpy(float)*100,width,label="Model-implied ROI",color="#24638C")
    axes[1].bar(x+width/2,selected.roi.to_numpy(float)*100,width,label="Realised ROI",color="#C17B17")
    axes[1].axhline(0,color="#555555",linewidth=.8);axes[1].set_ylabel("Net return / stake (%)");axes[1].set_title("Payoff weighting magnifies probability error")
    for ax in axes:
        ax.set_xticks(x,labels,rotation=25,ha="right");ax.set_xlabel("Archived decimal odds");ax.grid(axis="y",alpha=.2);ax.set_axisbelow(True)
        ax.spines[["top","right"]].set_visible(False);ax.legend(fontsize=8)
    fig.suptitle("First locked test: calibration / best book / fixed 5% policy; point estimates, not a new filter",fontsize=11)
    for suffix in (".png",".svg"):fig.savefig(figures/("selected_odds_tail"+suffix),dpi=220,bbox_inches="tight")
    plt.close(fig)
    return selected,trade_table,{"trades":n,"wins":wins,"mean_forecast_probability":mean_p},figures/"selected_odds_tail.png"


def _pages(tables,baseline,out,source_metadata):
    cost=tables["cost_sensitivity"];first=_focus(cost,"first");pooled=_focus(cost,"pooled")
    first_base=_focus(baseline,"first");pooled_base=_focus(baseline,"pooled")
    concentration=_focus(tables["date_concentration"],"pooled")
    if len(concentration)!=22:raise ValueError("Expected all22 pooled concentration rows")
    selected,tail_table,tail_facts,tail_figure=_tail(tables,out)
    displays={"first_roi_sensitivity":_sensitivity(first),"pooled_roi_sensitivity":_sensitivity(pooled),
              "first_combined_stress_metrics":_stress_metrics(_focus(cost.loc[cost.scenario_id.eq(COMBINED)],"first")),
              "pooled_combined_stress_metrics":_stress_metrics(_focus(cost.loc[cost.scenario_id.eq(COMBINED)],"pooled")),
              "selected_odds_tail":tail_table,"pooled_concentration":_concentration(concentration)}
    pages=[]
    def page(title,orientation="landscape"):
        item={"title":title,"orientation":orientation,"blocks":[]};pages.append(item);return item["blocks"]
    def p(blocks,text):blocks.append({"type":"paragraph","text":text})
    def table(blocks,key,caption,weights):blocks.append({"type":"table","frame":displays[key],"caption":caption,"weights":weights})
    sensitivity_weights=[.24]+[.084]*8+[.088]
    core_weights=[.25,.055,.065,.07,.075,.075,.10,.10,.10,.11]
    b=page(TITLE,"portrait")
    p(b,"October 14, 2026 | Separate extension of the frozen EPL study and the 172-row trading-metrics supplement. All analyses reuse already inspected history; no model, threshold, feature, outcome or selected trade is changed.")
    source_time=pd.Timestamp(source_metadata["started_at_utc"])
    local_time=source_time.tz_convert("America/New_York")
    p(b,f"Saved analysis run: {source_time.isoformat()} UTC; {local_time.strftime('%B %d, %Y at %I:%M %p %Z')} in New York. The source UTC timestamp and local draft date are recorded separately.")
    p(b,f"The stress table contains {len(cost):,} rows: {len(baseline)} original strategy/scope records under nine declared assumptions. The original 14 tuning-diagnostic rows are outside this stress extension. All 16 first locked-test policies and 22 pooled ablation variants remain visible, including cash policies. Every scenario retains hit ratio, expected gain/loss, frequency, realised returns and risk in the workbook.")
    p(b,"These are hypothetical stress assumptions, not observed bookmaker fees or attainable prices. With original stake s, odds O, fee rate c on every wager and winning-profit haircut h, winning net payoff is s·[(O−1)(1−h)−c], and losing net payoff is −s·(1+c). The haircut applies to winning net payout before fees, excluding returned stake; it is not a haircut of full decimal odds.")
    p(b,"Original selections, probabilities and allocated stakes stay fixed. There is no requalification, resizing or newly simulated Kelly sizing path. Fees are 0%, 0.25%, 0.5%, 1% or 2% of original stakes; payout haircuts are 1%, 2.5% or 5%; the combined stress uses a 1% fee plus a 2.5% haircut. Hit ratio and recorded trade frequency therefore remain unchanged by construction. Profitable-trade share remains separately defined from outcome hit ratio.")
    p(b,"Model-implied expected gain/loss are probability-weighted positive/negative outcome cashflow components at those recorded stakes. They are not unconditional forecasts of terminal wealth. ROI divides stressed profit by original wager turnover; the separate cash-outlay ROI includes hypothetical fees in its denominator. The same 5% qualification rule already embeds positive model-implied value, so large expected returns are not independent evidence of an edge.")
    p(b,"Funding is checked before date settlement: original date stakes plus fees must fit the stressed date-start cash. A failure leaves only an unfundable shadow cashflow calculation, marked †. Conventional max drawdown is then unavailable; shadow drawdown can exceed 100%. Each pooled season resets capital independently. No shadow path is presented as a feasible trading strategy or a fresh Kelly replay.")
    p(b,"Reported stress intervals, where available, condition on the recorded positions and use the same whole-date bootstrap draws across scenarios, stratified by reset episode. They omit fitting, selection, market execution and multiplicity uncertainty. Bin, deletion and concentration summaries are descriptive diagnostics; they do not create a new holdout or justify a newly selected tail filter.")
    b=page("First locked test: all 16 policies under every stress")
    p(b,"Cells show net profit / original wager turnover (%). Fee columns charge every original wager; payout columns reduce winning net payouts. A dash is undefined, including cash-policy ROI. † means the recorded exposure cannot be funded under that scenario.")
    table(b,"first_roi_sensitivity","Table R1. Complete 2025/26 locked-test policy grid; unchanged selections/stakes. These are assumptions applied to history, not actual provider fee schedules.",sensitivity_weights)
    b=page("Pooled ablation: all 22 variants under every stress")
    p(b,"Four exploratory outer seasons, 2022/23–2025/26; independently funded resets. Pooled profit and turnover add across seasons. † marks a failure in at least one episode; its value is a shadow cashflow return.")
    positive=pooled_base.loc[pooled_base.roi.gt(0)]
    spanning=int((positive.roi_ci_low.le(0) & positive.roi_ci_high.ge(0)).sum())
    p(b,f"For {spanning} of {len(positive)} positive baseline point returns, the conditional 95% ROI interval includes zero. Full scenario intervals remain in the workbook; these overlapping historical comparisons do not establish a newly selected winning strategy.")
    table(b,"pooled_roi_sensitivity","Table R2. Complete pooled variant/staking grid. No displayed scenario or strategy is selected for deployment from these results.",sensitivity_weights)
    b=page("First locked test: combined-stress trade metrics")
    p(b,"Combined assumption: 1% of each original stake plus a 2.5% haircut of winning profit. All money columns are units per trade; average +gain/−loss are realised conditional net-P&L magnitudes. Hit and frequency are percentages.")
    table(b,"first_combined_stress_metrics","Table R3. All 16 policies. Funding cell is Yes/conventional drawdown, or No/unavailable with separately labelled shadow drawdown. Totals, probability-weighted net expectancy and all other scenario metrics remain in the workbook.",core_weights)
    b=page("Pooled ablation: combined-stress trade metrics")
    p(b,"Same combined assumption and columns as R3. Pooled drawdown describes the worst separate episode; it is not a continuous bankroll path. Counterfactual recorded activity is not a claim of feasible executed activity.")
    table(b,"pooled_combined_stress_metrics","Table R4. Every pooled variant/staking rule. Missing per-trade means for no-trade policies remain unavailable, never invented zero payoffs.",core_weights)
    b=page("Selected-bet tails: expected and realised are different")
    p(b,f"Descriptive illustration: calibration/best-book/fixed 5% policy, with 333 recorded bets in the first test. In saved bins at odds ≥ 10, {tail_facts['trades']} bets win {tail_facts['wins']} times while their mean forecast probability is {number(tail_facts['mean_forecast_probability'],percent=True)}. This is a small retrospective tail, not evidence for a stable exclusion rule.")
    scores=tables["selected_bet_scores"]
    scores=scores.loc[scores.study.eq("first") & scores.scope.eq("test") & scores.branch.eq("calibration") & scores.policy.eq("predeclared_5pct") & scores.odds_mode.eq("best_individual_book") & scores.staking.eq("fixed")]
    if len(scores)!=1:raise ValueError("Missing unique selected-bet calibration scores")
    score=scores.iloc[0]
    p(b,f"Across these selected bets, mean forecast p is {number(score.mean_probability,percent=True)} versus hit ratio {number(score.hit_ratio,percent=True)}; binary Brier is {number(score.selected_binary_brier_score,4)}, log loss {number(score.selected_binary_log_loss,4)}, and ten-bin ECE {number(score.selected_binary_ece_10,4)}. Model-implied ROI is {number(score.expected_roi,percent=True)} versus realised {number(score.roi,percent=True)}. These conditional binary diagnostics differ from full-match three-outcome scores.")
    b.append({"type":"figure","path":tail_figure,"caption":"Figure R1. Saved selected-bet bins; point estimates. * denotes fewer than 30 trades, with no implied precision. High-odds payoff weights can enlarge modest probability errors; empty bins remain unobserved."})
    table(b,"selected_odds_tail","Table R5. Original baseline positions only. For a binary recorded wager, expected minus realised net P&L equals stake×odds×(p−outcome); summing gives an arithmetic decomposition, not a causal explanation or prospective result.",[.17,.08,.08,.16,.15,.18,.18])
    b=page("Date concentration and season fragility")
    p(b,"Top-date shares divide the largest positive date P&L sums by total positive date P&L, not by small overall net profit. Deletion ranges remove each entire episode/date or season while keeping all other recorded bets/stakes fixed; they are not refitted strategies or confidence intervals.")
    table(b,"pooled_concentration","Table R6. All 22 pooled baseline variants. Profitable seasons are observed signs. Overlapping strategies reuse fixtures; rows are not independent replications. Detailed date influence, season deletions and selected-bet scores remain in the workbook.",[.26,.09,.11,.11,.16,.08,.19])
    b=page("Coverage, empirical cost buffers and reproducibility","portrait")
    no_trades=int(baseline.bet_count.eq(0).sum());unfunded=int(cost.funding_feasible.eq(False).sum())
    funding_summary="All nine mild assumptions passed the funding check in this saved run; this does not verify price execution or future funding." if unfunded==0 else "Unfundable records remain shadow cashflow diagnostics, not feasible trading paths."
    p(b,f"The complete workbook contains all {len(cost):,} scenario rows and {len(tables['break_even'])} empirical break-even rows; {no_trades} baseline scopes hold no trades. It also retains {len(tables['funding_details']):,} per-scenario/episode funding summaries, diagnostic availability, probability/odds/edge bins, selected-bet scores, date influence and season fragility. Across all scopes/scenarios, {unfunded} rows have unfundable recorded exposure; these are not independent failure observations. {funding_summary}")
    p(b,"Historical break-even fee equals baseline net profit / turnover; the winning-profit haircut buffer equals baseline net profit / total winning net gains. These signed arithmetic buffers use observed payoffs and cannot be treated as forecasts of allowable future costs. A negative buffer means the original historical return was already below zero.")
    example=first_base.loc[first_base.branch.eq("calibration") & first_base.policy.eq("predeclared_5pct") & first_base.odds_mode.eq("best_individual_book") & first_base.staking.eq("fixed")]
    if len(example)!=1:raise ValueError("Missing unique baseline calibration example")
    e=example.iloc[0];fee=first.loc[first.strategy_id.eq(e.strategy_id) & first.scenario_id.eq("fee_0p5pct")]
    if len(fee)!=1:raise ValueError("Missing corresponding fee scenario")
    p(b,f"For the first-test calibration/best-book/fixed policy, historical ROI is {number(e.roi,percent=True)}. A hypothetical 0.5% charge on every original wager changes it to {number(fee.iloc[0].roi,percent=True)} without changing its selected outcomes or recorded frequency. This quantifies sensitivity under one stated assumption; it does not identify an actual fee or an executable advantage.")
    p(b,"Selected-outcome calibration is conditional on which bets were chosen. Bin hit ratios differ from full-match H/D/A accuracy and multiclass calibration. Model probability gaps and financial expectation gaps have different weighting; odds and stakes amplify some errors. Sparse tails, repeated clubs, overlapping strategies and already inspected outcomes restrict inference. The present-day club cohort retains its original survivor and source-coverage limits. No new model or betting filter is recommended from this extension.")
    p(b,"A dash denotes undefined/unavailable, and genuine infinity remains explicit. A no-trade policy has zero cash totals and activity when opportunities exist, while hit ratio, per-trade means and ROI remain undefined. Failed funding does not erase a recorded outcome, but it prevents interpreting its shadow path as a feasible wager sequence.")
    p(b,"This companion is separate from the seven-page draft and the 172-row supplement. Its manifest hashes every new analytic input and the unchanged original study/report sources. Rebuild with python -m epl_robustness.report --results PATH. Rendering does not collect data, fit models, change selection or overwrite prior results. Full-precision source tables remain beside the companion; display CSVs use the published rounding.")
    return pages,displays,tail_facts


def _markdown(pages,out):
    chunks=[]
    for i,page in enumerate(pages):
        chunks.extend([("# " if i==0 else "## ")+page["title"],""])
        for block in page["blocks"]:
            if block["type"]=="paragraph":chunks.extend([block["text"],""])
            elif block["type"]=="figure":chunks.extend([f"![{block['caption']}]({block['path'].relative_to(out).as_posix()})","",block["caption"],""])
            else:
                chunks.extend([block["caption"],""]);frame=block["frame"]
                chunks.append("| "+" | ".join(str(x).replace("\n","<br>") for x in frame.columns)+" |")
                chunks.append("| "+" | ".join("---" for _ in frame.columns)+" |")
                for row in frame.itertuples(index=False,name=None):chunks.append("| "+" | ".join(str(x).replace("\n","<br>") for x in row)+" |")
                chunks.append("")
    return "\n".join(chunks)


def _render_pdf(pages,path):
    fonts=Path("C:/Windows/Fonts");names={"body":"Times-Roman","bold":"Helvetica-Bold","sans":"Helvetica","italic":"Times-Italic"}
    for role,file in (("body","times.ttf"),("bold","arialbd.ttf"),("sans","arial.ttf"),("italic","timesi.ttf")):
        if (fonts/file).exists():name="Robustness"+role.title();pdfmetrics.registerFont(TTFont(name,str(fonts/file)));names[role]=name
    styles={"title":ParagraphStyle("RTitle",fontName=names["bold"],fontSize=14,leading=16,spaceAfter=7,textColor=colors.HexColor("#12364D")),
            "body":ParagraphStyle("RBody",fontName=names["body"],fontSize=10.2,leading=12.4,spaceAfter=7),
            "caption":ParagraphStyle("RCaption",fontName=names["italic"],fontSize=8.4,leading=9.8,spaceAfter=4),
            "cell":ParagraphStyle("RCell",fontName=names["sans"],fontSize=8,leading=9),
            "head":ParagraphStyle("RHead",fontName=names["bold"],fontSize=8,leading=9,textColor=colors.white)}
    def header(canvas,doc):
        w,h=canvas._pagesize;canvas.setFont(names["sans"],8);canvas.setFillColor(colors.HexColor("#52616B"))
        canvas.drawString(36,h-21,"EPL ROBUSTNESS — EXPLORATORY DRAFT | October 14, 2026");canvas.drawRightString(w-36,21,f"{doc.page} / {len(pages)}")
    doc=BaseDocTemplate(str(path),pagesize=letter,title=TITLE,author="Aman Dhillon and Panagiotis Housos")
    for orientation,size in (("portrait",letter),("landscape",landscape(letter))):
        w,h=size;doc.addPageTemplates(PageTemplate(id=orientation,pagesize=size,onPage=header,frames=[Frame(36,36,w-72,h-72,leftPadding=0,rightPadding=0,topPadding=0,bottomPadding=0)]))
    story=[]
    for i,page in enumerate(pages):
        if i:story.extend([NextPageTemplate(page["orientation"]),PageBreak()])
        width=(letter if page["orientation"]=="portrait" else landscape(letter))[0]-72
        story.append(Paragraph(html.escape(page["title"]),styles["title"]))
        for block in page["blocks"]:
            if block["type"]=="paragraph":story.append(Paragraph(html.escape(block["text"]),styles["body"]))
            elif block["type"]=="figure":
                with PILImage.open(block["path"]) as im:ratio=im.height/im.width
                story.extend([Image(str(block["path"]),width=width-10,height=(width-10)*ratio),Paragraph(html.escape(block["caption"]),styles["caption"])])
            else:
                story.append(Paragraph(html.escape(block["caption"]),styles["caption"]))
                frame=block["frame"];raw=[list(frame.columns)]+frame.values.tolist()
                cells=[[Paragraph(html.escape(str(v)).replace("\n","<br/>"),styles["head" if n==0 else "cell"]) for v in row] for n,row in enumerate(raw)]
                if abs(sum(block["weights"])-1)>1e-9:raise ValueError("Invalid report table widths")
                table=Table(cells,colWidths=[width*v for v in block["weights"]],repeatRows=1,hAlign="CENTER")
                table.setStyle(TableStyle([("BACKGROUND",(0,0),(-1,0),colors.HexColor("#12364D")),("ROWBACKGROUNDS",(0,1),(-1,-1),[colors.white,colors.HexColor("#F0F4F6")]),
                    ("VALIGN",(0,0),(-1,-1),"MIDDLE"),("LEFTPADDING",(0,0),(-1,-1),3),("RIGHTPADDING",(0,0),(-1,-1),3),("TOPPADDING",(0,0),(-1,-1),.6),("BOTTOMPADDING",(0,0),(-1,-1),.6)]))
                story.extend([table,Spacer(1,4)])
    doc.build(story)


def _render_docx(pages,path):
    doc=Document();normal=doc.styles["Normal"];normal.font.name="Times New Roman";normal.font.size=Pt(10.2)
    normal.paragraph_format.line_spacing=Pt(12.4);normal.paragraph_format.space_after=Pt(7)
    title=doc.styles["Heading 1"];title.font.name="Arial";title.font.size=Pt(14);title.paragraph_format.line_spacing=Pt(16);title.paragraph_format.space_before=Pt(0);title.paragraph_format.space_after=Pt(7)
    doc.core_properties.title=TITLE;doc.core_properties.author="Aman Dhillon and Panagiotis Housos"
    for i,page in enumerate(pages):
        s=doc.sections[0] if i==0 else doc.add_section(WD_SECTION_START.NEW_PAGE);land=page["orientation"]=="landscape"
        s.orientation=WD_ORIENT.LANDSCAPE if land else WD_ORIENT.PORTRAIT;s.page_width=Inches(11 if land else 8.5);s.page_height=Inches(8.5 if land else 11)
        s.top_margin=s.bottom_margin=s.left_margin=s.right_margin=Inches(.5);s.header_distance=s.footer_distance=Inches(.2)
        if i==0:
            r=s.header.paragraphs[0].add_run("EPL ROBUSTNESS — EXPLORATORY DRAFT | October 14, 2026");r.font.name="Arial";r.font.size=Pt(8)
            p=s.footer.paragraphs[0];p.alignment=2;field=OxmlElement("w:fldSimple");field.set(qn("w:instr"),"PAGE");p._p.append(field)
        doc.add_heading(page["title"],1)
        for block in page["blocks"]:
            if block["type"]=="paragraph":doc.add_paragraph(block["text"])
            else:
                if block["type"]=="figure":doc.add_picture(str(block["path"]),width=Inches(9.85 if land else 7.35))
                cap=doc.add_paragraph(block["caption"]);cap.paragraph_format.line_spacing=Pt(9.8);cap.paragraph_format.space_after=Pt(4)
                for r in cap.runs:r.font.size=Pt(8.4);r.italic=True
                if block["type"]=="figure":continue
                frame=block["frame"];table=doc.add_table(rows=1,cols=len(frame.columns));table.style="Light Shading Accent 1";table.autofit=False
                for column,weight in zip(table.columns,block["weights"]):column.width=Inches((10 if land else 7.5)*weight)
                for cell,value in zip(table.rows[0].cells,frame.columns):cell.text=str(value)
                for values in frame.itertuples(index=False,name=None):
                    for cell,value in zip(table.add_row().cells,values):cell.text=str(value)
                for row in table.rows:
                    for cell,weight in zip(row.cells,block["weights"]):
                        cell.width=Inches((10 if land else 7.5)*weight)
                        for p in cell.paragraphs:
                            p.paragraph_format.line_spacing=Pt(9);p.paragraph_format.space_after=Pt(0)
                            for r in p.runs:r.font.name="Arial";r.font.size=Pt(8)
    doc.save(path)


def _excel(path,tables,displays,registry):
    cost=tables["cost_sensitivity"]
    columns=["study","scope","episode_id","scenario_id","bet_count","hit_ratio","trading_frequency","avg_winning_gain_units","avg_losing_loss_units",
             "expected_gain_per_trade_units","expected_loss_per_trade_units","expected_net_profit_per_trade_units","realized_expectancy_per_trade_units",
             "expected_roi","roi","roi_on_cash_outlay","profit","turnover","funding_feasible","max_drawdown","shadow_max_drawdown",
             "profitable_trade_share","trade_rate_all_predictions","fee_units","cash_outlay","funding_breach_dates","roi_ci_low","roi_ci_high"]
    def view(frame):
        out=frame[[c for c in columns if c in frame]].copy();out.insert(2,"strategy_description",frame.apply(lambda row:_label(row).replace("\n"," | "),axis=1));return out
    readme=pd.DataFrame([("Status",TITLE),("Assumptions","Hypothetical fee on every stake; haircut of winning net payout, not actual bookmaker fees."),
        ("Fixed exposure","No requalification, stake resizing, model changes or new Kelly path."),("Unfundable records","Shadow cashflows only; conventional max drawdown unavailable."),
        ("Expectations","Probability-weighted positive/negative cash components at the actual original stake path, not terminal wealth."),
        ("Inference","Exploratory already-inspected history; no new selected winner or tail-filter recommendation."),
        ("ROI basis","Net profit / original wager turnover; cash-outlay ROI includes assumed fees."),
        ("remaining_signed_break_even_fee_rate_at_haircut","Total break-even fee c*(h)=(P0−hG)/T at the given haircut, not additional allowance above the current fee."),
        ("Undefined","Dash; empty-trade means/hit/ROI are unavailable. Cash totals/activity may be zero.")],columns=["Item","Definition"])
    sheets={"Read me":readme,"Scenario assumptions":cost[["scenario_id","fee_rate","payout_haircut","fee_basis","haircut_basis"]].drop_duplicates(),
            "Metrics overview":view(cost),"All cost metrics":cost,"First test metrics":view(_focus(cost,"first")),"Pooled metrics":view(_focus(cost,"pooled")),
            "First ROI grid":displays["first_roi_sensitivity"],"Pooled ROI grid":displays["pooled_roi_sensitivity"]}
    names={"break_even":"Empirical break even","funding_details":"Funding episodes","diagnostic_availability":"Diagnostic availability","probability_bins":"Probability bins",
           "odds_bins":"Odds bins","edge_bins":"Edge bins","selected_bet_scores":"Selected bet scores","date_concentration":"Date concentration","date_influence":"Date influence","season_fragility":"Season fragility"}
    sheets.update({names[k]:v for k,v in tables.items() if k in names})
    sheets["Input provenance"]=pd.DataFrame([{"path":p,**v} for p,v in registry.items()])
    percent={"fee_rate","payout_haircut","hit_ratio","trading_frequency","trade_rate_all_predictions","profitable_trade_share","expected_roi","roi","roi_on_cash_outlay",
             "max_drawdown","shadow_max_drawdown","roi_ci_low","roi_ci_high","hit_ratio_ci_low","hit_ratio_ci_high","mean_probability","turnover_share","season_roi","season_hit_ratio","season_expected_roi","roi_excluding_season"}
    with pd.ExcelWriter(path,engine="xlsxwriter",engine_kwargs={"options":{"strings_to_urls":False}}) as writer:
        book=writer.book;header=book.add_format({"bold":True,"font_color":"white","bg_color":"#12364D","text_wrap":True,"valign":"top"})
        numeric=book.add_format({"num_format":"#,##0.0000"});pct=book.add_format({"num_format":"0.00%"});integer=book.add_format({"num_format":"#,##0"})
        for name,frame in sheets.items():
            safe=frame.copy()
            for c in safe:
                if isinstance(safe[c].dtype,pd.DatetimeTZDtype):safe[c]=safe[c].dt.tz_convert(None)
                if safe[c].dtype=="object":safe[c]=safe[c].map(lambda v:json.dumps(v,default=str) if isinstance(v,(dict,list)) else v)
                elif pd.api.types.is_numeric_dtype(safe[c]):safe[c]=safe[c].map(lambda v:"+∞" if v==float("inf") else "−∞" if v==float("-inf") else v)
            safe.to_excel(writer,sheet_name=name,index=False,na_rep="—")
            ws=writer.sheets[name];ws.freeze_panes(1,2);ws.autofilter(0,0,len(safe),len(safe.columns)-1);ws.set_row(0,42)
            ws.write_row(0,0,list(safe.columns),header)
            for j,c in enumerate(safe.columns):
                fmt=pct if c in percent or c.endswith("_profit_share") else integer if c in ("bet_count","wins","winning_bet_count","funding_breach_dates","positive_date_count") else numeric
                width=48 if c in ("strategy_description","Definition","path") else min(38,max(14,len(str(c))*.75))
                ws.set_column(j,j,width,fmt)
    return {name:len(frame) for name,frame in sheets.items()}


def build_report(results:Path)->dict:
    results=Path(results).resolve()
    if not results.exists():raise FileNotFoundError(results)
    sources=[p for p in results.rglob("*") if p.is_file() and "report" not in p.relative_to(results).parts and "previews" not in p.parts]
    sources.extend(p for p in (ROOT/"epl_robustness").glob("*.py"))
    sources.extend(p for p in (ROOT/"epl_robustness").glob("*.md") if p.name!="README.md")
    sources.extend(p for p in (ROOT/"epl_modeling").glob("*.py"))
    sources.extend(p for p in (ROOT/"epl_modeling").glob("*.md"))
    for folder in (ROOT/"epl_report",ROOT/"epl_modeling/results",ROOT/"epl_modeling/ablation_results",ROOT/"epl_strategy"):
        if folder.exists():sources.extend(p for p in folder.rglob("*") if p.is_file() and "__pycache__" not in p.parts and "previews" not in p.parts and "tests" not in p.parts and p!=ROOT/"epl_strategy/README.md")
    sources=sorted(set(sources));registry={_path_label(p):{"sha256":_sha(p),"bytes":p.stat().st_size} for p in sources}
    tables=_read(results);baseline=_validate(tables);out=results/"report";out.mkdir(exist_ok=True)
    source_metadata=json.loads((results/"audit_manifest.json").read_text(encoding="utf-8"))
    pages,displays,tail_facts=_pages(tables,baseline,out,source_metadata)
    for name,frame in displays.items():frame.to_csv(out/(name+".csv"),index=False)
    md=out/"EPL_robustness_companion.md";pdf=md.with_suffix(".pdf");docx=md.with_suffix(".docx")
    md.write_text(_markdown(pages,out),encoding="utf-8");_render_pdf(pages,pdf);_render_docx(pages,docx)
    sheets=_excel(out/"EPL_robustness_metrics.xlsx",tables,displays,registry)
    reader=PdfReader(pdf)
    if len(reader.pages)!=8:raise AssertionError(f"Expected8 planned pages, rendered{len(reader.pages)}")
    word=Document(docx)
    if len(word.sections)!=8 or len(word.tables)!=6 or len(word.inline_shapes)!=1:raise AssertionError("Editable Word section/table/figure structure differs from report plan")
    changed=[_path_label(p) for p in sources if _sha(p)!=registry[_path_label(p)]["sha256"]]
    if changed:raise AssertionError("A protected source changed while rendering: "+str(changed))
    outputs={p.relative_to(out).as_posix():{"sha256":_sha(p),"bytes":p.stat().st_size} for p in sorted(out.rglob("*")) if p.is_file() and p.name!="report_manifest.json" and "previews" not in p.parts}
    manifest={"status":"exploratory_draft_companion","draft_date":DATE,"built_at_utc":datetime.now(timezone.utc).isoformat(),
        "source_run_started_at_utc":source_metadata["started_at_utc"],"source_run_started_at_new_york":pd.Timestamp(source_metadata["started_at_utc"]).tz_convert("America/New_York").isoformat(),
        "source_row_counts":{name:len(frame) for name,frame in tables.items()},"pdf_pages":len(reader.pages),"workbook_sheet_rows":sheets,
        "first_test_strategies":16,"pooled_ablation_strategies":22,"scenario_count":9,"tail_illustration":tail_facts,
        "models_refitted":False,"strategies_or_stakes_reselected":False,"original_studies_reports_and_sources_unchanged":not changed,
        "inputs":registry,"outputs":outputs,"assumption_limits":"Hypothetical fees and winning-profit haircuts; fixed recorded exposure; unfundable paths are shadow calculations.",
        "inference_limits":"Exploratory same-history conditional diagnostics; tails/concentration are not new trading filters or confirmatory evidence."}
    (out/"report_manifest.json").write_text(json.dumps(manifest,indent=2,ensure_ascii=False,allow_nan=False),encoding="utf-8")
    return {"output_directory":str(out),"pdf_pages":len(reader.pages),"scenario_rows":len(tables["cost_sensitivity"]),"inputs_hashed":len(registry),"source_files_unchanged":not changed}


def main():
    parser=argparse.ArgumentParser(description=__doc__);parser.add_argument("--results",type=Path,required=True)
    args=parser.parse_args();print(json.dumps(build_report(args.results),indent=2))


if __name__=="__main__":main()
