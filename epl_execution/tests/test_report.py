"""Coverage and presentation invariants for the combined execution report."""
import json

import pandas as pd
from docx import Document
from openpyxl import load_workbook
from pypdf import PdfReader

from epl_execution import report
from epl_strategy_lab.catalog import build_catalog


def test_appendix_keeps_every_exact_original_rule_and_explains_price_override():
    catalog = build_catalog()
    entries = report._appendix(catalog)
    assert [entry["strategy_id"] for entry in entries] == report.ALL_IDS
    assert [json.loads(entry["exact_config"]) for entry in entries] == catalog
    assert all("overrides only odds_mode to reference" in entry["explanation"] for entry in entries)
    assert all("unfunded diagnostic is a shadow cashflow" in entry["explanation"] for entry in entries)
    assert sum(entry["coverage"].startswith("Four-season") for entry in entries) == 56


def test_markdown_finishes_with_the_complete_appendix():
    entries = report._appendix(build_catalog())
    pages = [{"title": "Quote study", "orientation": "portrait", "blocks": [
        {"type": "paragraph", "text": "Reference replay and shadow attribution are separate."}]}]
    md = report._markdown(pages, entries)
    assert md.split("## " + report.APPENDIX_TITLE)[1].count("### S") == 84
    assert md.rstrip().endswith("```")
    assert md.index("Reference replay") < md.index(report.APPENDIX_TITLE)


def test_combined_pdf_and_word_keep_all84_definitions_last(tmp_path):
    entries = report._appendix(build_catalog())
    pages = [{"title": report.TITLE, "orientation": "portrait", "blocks": [
        {"type": "paragraph", "text": "Reference replay paths are funded. Unfunded shadows are attribution only."}]}]
    pdf = tmp_path / "report.pdf"; docx = tmp_path / "report.docx"
    report._pdf(pages, entries, pdf); report._docx(pages, entries, docx)
    reader = PdfReader(pdf)
    text = " ".join(" ".join((p.extract_text() or "").split()) for p in reader.pages)
    assert "QUOTE SENSITIVITY" in text
    assert all(" ".join(entry["title"].split()) in text for entry in entries)
    assert "S084" in reader.pages[-1].extract_text()
    paragraphs = [p.text for p in Document(docx).paragraphs]
    assert sum(p.startswith(tuple(report.ALL_IDS)) for p in paragraphs) == 84
    assert json.loads(paragraphs[-1].split(": ", 1)[1])["strategy_id"] == "S084"


def test_reference_graph_is_offline_all_visible_and_contains_no_shadow_curve(tmp_path):
    entries = report._appendix(build_catalog())
    graph = pd.DataFrame([{"strategy_id": sid, "match_date": date, "end_bankroll": cash, "is_initial": initial}
        for sid in report.COMPATIBLE_IDS for date, cash, initial in (
            (pd.Timestamp("2022-08-04"), 100000., True), (pd.Timestamp("2026-05-24"), 100001., False))])
    info = report._graphs({"catalog": build_catalog(), "graph": graph}, tmp_path, entries)
    assert info["trace_ids"] == report.COMPATIBLE_IDS
    assert info["initially_visible_traces"] == 56
    assert info["unfunded_shadow_paths_plotted"] is False
    text = (tmp_path / "all_reference_strategies_4seasons_100k.html").read_text(encoding="utf-8")
    assert "plotly.js" in text and "<script src=" not in text
    assert text.count("<article id='S") == 84
    assert text.count("<img alt='Preserved original") == 2
    assert text.index("<h2>" + report.APPENDIX_TITLE) > text.index("original all-84")


def test_workbook_has_all_metric_modes_diagnostics_and_external_ledger_hashes(tmp_path):
    metrics = pd.DataFrame({"execution_mode": report.MODES, "payoff_ratio": [2., 2., 2.],
        "roi": [.1, .1, .1], "profit": [1000., 1000., 1000.], "funding_feasible": [True, True, False]})
    minimal = pd.DataFrame({"value": [1]})
    data = {key: minimal for key in report.TABLE_NAMES}
    for key in ("metrics_single", "metrics_continuous", "metrics_seasonal"): data[key] = metrics
    data.update({"catalog": build_catalog(), "graph": minimal, "quotes": {"match_count": 1470},
        "legacy": {key: minimal for key in ("selection_comparison", "selection_global", "season_stability")}})
    data.update({"closing_quote_outcomes": pd.DataFrame({"verified_closing_line_value": [False]}),
        "archived_closing_reference": minimal, "closing_quotes": {"executable_quotes_verified": False}})
    ledger = tmp_path / "bets_reference_single.parquet"; minimal.to_parquet(ledger)
    out = tmp_path / "workbook.xlsx"
    rows = report._workbook(out, data, report._appendix(data["catalog"]), {"source": {"sha256": "a" * 64, "bytes": 1}}, tmp_path)
    assert rows["Single all modes"] == 3 and rows["Final84 appendix"] == 84
    assert rows["External ledger files"] == 1
    assert rows["Archived same-book closing"] == 1
    assert rows["Archived closing price panel"] == 1
    book = load_workbook(out, read_only=True)
    sheet = book["Single all modes"]
    assert [cell.value for cell in sheet[2]] == ["original", 2., .1, 1000., True]
    assert [cell.number_format for cell in sheet[2]][1:4] == ["#,##0.0000", "0.00%", "#,##0.0000"]
    assert book["External ledger files"][2][0].value == "../bets_reference_single.parquet"
    assert book["Archived same-book closing"][2][0].value is False
    book.close()


def test_paired_metrics_show_model_gain_loss_and_observed_payoffs_for_every_rule():
    row = {"strategy_id": "S017", "bet_count": 10, "hit_ratio": .4, "trading_frequency": .1,
        "expected_gain_per_trade_units": 25., "expected_loss_per_trade_units": 20.,
        "avg_winning_gain_units": 30., "avg_losing_loss_units": 10., "roi": .1, "profit": 1000., "max_drawdown": .05}
    frame = pd.DataFrame([{**row, "execution_mode": mode} for mode in report.MODES])
    displayed = report._paired_table(frame, ["S017"])
    assert len(displayed) == 1 and len(displayed.columns) == 13
    assert displayed.iloc[0]["Ref model\ngain/trade"] == "25.00"
    assert displayed.iloc[0]["Ref model\nloss/trade"] == "20.00"
    assert displayed.iloc[0]["Ref avg.\nwin gain"] == "30.00"
    assert displayed.iloc[0]["Ref avg.\nlosing loss"] == "10.00"


def test_scope_summary_keeps_unfunded_fixed_exposure_separate_from_reference_replays():
    data = {}
    for key in ("metrics_single", "metrics_continuous", "metrics_seasonal"):
        data[key] = pd.DataFrame([
            {"execution_mode": "original", "profit": 1., "funding_feasible": True},
            {"execution_mode": "reference_replay", "profit": -1., "funding_feasible": True},
            {"execution_mode": "fixed_exposure_reference", "profit": float("nan"), "funding_feasible": False}])
    summary = report._scope_summary(data)
    assert summary["Original\nprofitable"].eq("1").all()
    assert summary["Reference replay\nprofitable"].eq("0").all()
    assert summary["Reference replay\nunfunded"].eq("0").all()
    assert summary["Fixed-exposure\nunfunded shadows"].eq("1").all()
