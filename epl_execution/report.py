"""Combine the frozen strategy study with price sensitivity and exact attribution.

The final report reads completed tables. It does not rerun or select a model,
and every original strategy definition remains in its final appendix.
"""
from __future__ import annotations

import argparse
import base64
from datetime import datetime, timezone
import html
import json
from pathlib import Path
import shutil

import matplotlib
matplotlib.use("Agg")
import matplotlib.dates as mdates
import matplotlib.pyplot as plt
import numpy as np
import pandas as pd
import plotly.graph_objects as go
from docx import Document
from docx.enum.section import WD_ORIENT, WD_SECTION_START
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt
from PIL import Image as PILImage
from pypdf import PdfReader
from reportlab.lib import colors
from reportlab.lib.pagesizes import landscape, letter
from reportlab.lib.styles import ParagraphStyle
from reportlab.platypus import (BaseDocTemplate, Frame, Image, KeepTogether,
    NextPageTemplate, PageBreak, PageTemplate, Paragraph, Table, TableStyle)

from epl_strategy.report import number
from epl_strategy_lab import report as lab
from epl_strategy_validation import report as temporal

ROOT = Path(__file__).resolve().parents[1]
TEMPORAL = ROOT / "epl_strategy_validation/results/20261015_temporal_selection"
TITLE = "EPL strategy research: performance, selection and quote sensitivity"
APPENDIX_TITLE = temporal.APPENDIX_TITLE
INITIAL = 100000.0
ALL_IDS = temporal.ALL_IDS
COMPATIBLE_IDS = temporal.COMPATIBLE_IDS
MODES = ("original", "reference_replay", "fixed_exposure_reference")
BODY_WEIGHTS = temporal.BODY_WEIGHTS
TABLE_NAMES = ("metrics_single", "metrics_continuous", "metrics_seasonal",
    "execution_comparison", "payoff_attribution", "book_concentration", "quote_outcomes")
OPTIONAL_TABLES = ("closing_quote_outcomes", "archived_closing_reference")


def _load(output: Path) -> dict:
    data = {name: pd.read_parquet(output / (name + ".parquet")) for name in TABLE_NAMES}
    data["graph"] = pd.read_csv(output / "graph_equity_reference_continuous.csv", parse_dates=["match_date"])
    data["catalog"] = json.loads((output / "catalog.json").read_text(encoding="utf-8"))
    data["quotes"] = json.loads((output / "quote_summary.json").read_text(encoding="utf-8"))
    for name in OPTIONAL_TABLES:
        if (output / (name + ".parquet")).exists():
            data[name] = pd.read_parquet(output / (name + ".parquet"))
    if (output / "closing_quote_summary.json").exists():
        data["closing_quotes"] = json.loads((output / "closing_quote_summary.json").read_text(encoding="utf-8"))
    data["legacy"] = temporal._load(TEMPORAL)
    return data


def _validate(data: dict) -> None:
    if data["catalog"] != data["legacy"]["catalog"]:
        raise ValueError("The execution report changes a frozen strategy definition")
    temporal._validate(data["legacy"])
    for name, expected, count in (("metrics_single", ALL_IDS, 84),
            ("metrics_continuous", COMPATIBLE_IDS, 56),
            ("metrics_seasonal", COMPATIBLE_IDS, 224)):
        frame = data[name]
        if len(frame) != count * 3 or set(frame.execution_mode) != set(MODES):
            raise ValueError("Incomplete execution modes in " + name)
        if frame.duplicated(["account_id", "execution_mode"]).any():
            raise ValueError("Duplicate account/mode in " + name)
        for mode in MODES:
            part = frame.loc[frame.execution_mode.eq(mode)]
            if len(part) != count or set(part.strategy_id) != set(expected):
                raise ValueError("A strategy is omitted from " + name + ": " + mode)
        feasible = frame.loc[frame.execution_mode.ne("fixed_exposure_reference")]
        if not feasible.funding_feasible.eq(True).all():
            raise ValueError("A recorded original/reference replay is unfunded")
        if not np.allclose(feasible.initial_capital_sum_units, INITIAL, atol=1e-6, rtol=0):
            raise ValueError("Reference accounts must each start at exactly 100,000")
    if len(data["execution_comparison"]) != 364 or len(data["payoff_attribution"]) != 364:
        raise ValueError("Comparison and attribution must cover all 364 original accounts")
    graph = data["graph"]
    if not pd.api.types.is_bool_dtype(graph.is_initial):
        flags = graph.is_initial.astype(str).str.lower()
        if not flags.isin(["true", "false"]).all():
            raise ValueError("Invalid graph initial-point flag")
        graph["is_initial"] = flags.eq("true")
    if set(graph.strategy_id) != set(COMPATIBLE_IDS) or len(graph) != 56 * 454:
        raise ValueError("The reference graph must include all 56 complete curves")
    if graph.duplicated(["strategy_id", "match_date"]).any() or not np.isfinite(graph.end_bankroll).all() or graph.end_bankroll.lt(-1e-6).any():
        raise ValueError("Invalid reference graph balances")
    reference = data["metrics_continuous"].loc[lambda f: f.execution_mode.eq("reference_replay")].set_index("strategy_id")
    for sid, part in graph.groupby("strategy_id"):
        initial = part.loc[part.is_initial]
        if len(initial) != 1 or initial.match_date.iloc[0] != pd.Timestamp("2022-08-04") or not np.isclose(initial.end_bankroll.iloc[0], INITIAL, rtol=0, atol=1e-8):
            raise ValueError("Missing exact 100,000 reference anchor for " + sid)
        end = part.sort_values("match_date").end_bankroll.iloc[-1]
        if not np.isclose(end, INITIAL + reference.loc[sid, "profit"], rtol=0, atol=1e-5):
            raise ValueError("Reference graph endpoint differs from account profit")


def _appendix(catalog: list[dict]) -> list[dict]:
    entries = temporal._appendix(catalog)
    for entry in entries:
        entry["explanation"] += (" Reference replay overrides only odds_mode to reference (Bet365); "
            "probabilities and other original parameters stay unchanged, while eligibility, selection, consensus and stakes are recomputed. "
            "Fixed-exposure repricing keeps original bets/stakes; an unfunded diagnostic is a shadow cashflow.")
    return entries


def _graphs(data: dict, figures: Path, entries: list[dict]) -> dict:
    figures.mkdir(exist_ok=True)
    for name in ("all_strategies_4seasons_100k.png", "all_84_strategies_2025_100k.png"):
        shutil.copyfile(TEMPORAL / "figures" / name, figures / name)
    graph = data["graph"]
    catalog = [c for c in data["catalog"] if c["strategy_id"] in COMPATIBLE_IDS]
    styles = lab._styles(len(catalog))
    dash = {"solid": "-", "dash": "--", "dashdot": "-.", "dot": ":", "longdash": (0, (8, 3))}
    fig, ax = plt.subplots(figsize=(12.8, 7.2))
    fig.subplots_adjust(left=.085, right=.985, top=.91, bottom=.30)
    interactive = go.Figure()
    for index, config in enumerate(catalog):
        sid = config["strategy_id"]
        part = graph.loc[graph.strategy_id.eq(sid)].sort_values("match_date")
        color, line = styles[index]
        ax.step(part.match_date, part.end_bankroll, where="post", color=color,
            linestyle=dash[line], linewidth=.9, alpha=.9, label=sid)
        interactive.add_trace(go.Scatter(x=part.match_date.dt.strftime("%Y-%m-%d").tolist(),
            y=part.end_bankroll.astype(float).tolist(), mode="lines", name=sid, uid=sid, visible=True,
            line={"color": color, "dash": line, "width": 1.5, "shape": "hv"},
            customdata=[config["name"]] * len(part),
            hovertemplate=sid + " · %{customdata}<br>Date: %{x|%Y-%m-%d}<br>Balance: %{y:,.2f} monetary units<extra></extra>"))
    ax.set_title("All 56 Bet365 reference-price replays · initial 100,000 each · no account resets", fontsize=12)
    ax.set_xlabel("Calendar date (balance after date-level settlement)")
    ax.set_ylabel("Balance (monetary units)")
    ax.yaxis.set_major_formatter(matplotlib.ticker.StrMethodFormatter("{x:,.0f}"))
    ax.xaxis.set_major_locator(mdates.MonthLocator(interval=6))
    ax.xaxis.set_major_formatter(mdates.DateFormatter("%b %Y"))
    ax.grid(alpha=.2); ax.spines[["top", "right"]].set_visible(False)
    ax.set_xlim(graph.match_date.min(), graph.match_date.max())
    handles, labels = ax.get_legend_handles_labels()
    fig.legend(handles, labels, loc="lower center", ncol=10, fontsize=8, frameon=False,
        bbox_to_anchor=(.5, .014), handlelength=2, columnspacing=1.5, labelspacing=.65)
    with plt.rc_context({"svg.fonttype": "none"}):
        for suffix in (".png", ".svg"):
            fig.savefig(figures / ("all_reference_strategies_4seasons_100k" + suffix), dpi=240)
    plt.close(fig)
    interactive.update_layout(title={"text": "All 56 reference-price accounts · 100,000 each · no resets", "x": .01, "y": .99, "yanchor": "top"},
        template="plotly_white", height=1000, margin={"l": 85, "r": 35, "t": 115, "b": 240},
        xaxis={"title": "Calendar date (date-close settlement)", "type": "date"},
        yaxis={"title": "Balance (monetary units)", "type": "linear", "tickformat": ",.0f"}, hovermode="closest",
        legend={"title": {"text": "Every compatible strategy ID", "side": "top"}, "orientation": "h", "x": 0,
            "y": -.16, "entrywidth": 70, "entrywidthmode": "pixels", "font": {"size": 10}, "groupclick": "toggleitem"},
        updatemenus=[{"type": "buttons", "direction": "right", "x": 0, "y": 1.025, "yanchor": "bottom", "buttons": [
            {"label": "Show all", "method": "restyle", "args": [{"visible": True}]},
            {"label": "Hide all (keep legend)", "method": "restyle", "args": [{"visible": "legendonly"}]}]}])
    if len(interactive.data) != 56 or any(trace.visible is not True for trace in interactive.data):
        raise AssertionError("Every reference curve must be visible initially")
    parts = ["<!doctype html><html lang='en'><head><meta charset='utf-8'><meta name='viewport' content='width=device-width,initial-scale=1'><title>EPL quote sensitivity</title><style>body{font:15px system-ui;margin:24px;color:#173746}p{max-width:1100px;line-height:1.5}pre{white-space:pre-wrap;overflow-wrap:anywhere;font-size:12px;background:#f0f4f6;padding:12px}article{border-top:1px solid #cad5db;padding-top:12px}img{max-width:100%;height:auto}</style></head><body>",
        "<h1>Quote sensitivity: every compatible account starts at 100,000</h1><p>All 56 reference-price replay curves appear on one linear money-versus-time graph. Each account recomputes selections and stakes at archived Bet365 quotes, retains its frozen original probabilities and carries cash across four seasons without resets. Missing quote timestamps prevent a claim of executable trading. This page embeds its plotting library and works offline.</p>",
        interactive.to_html(full_html=False, include_plotlyjs=True, config={"responsive": True, "displaylogo": False})]
    for filename, title in (("all_strategies_4seasons_100k.png", "Preserved original 56-account four-season graph"),
            ("all_84_strategies_2025_100k.png", "Preserved original all-84 single-season graph")):
        parts += ["<h2>" + title + "</h2><img alt='" + title + "' src='data:image/png;base64," + base64.b64encode((figures / filename).read_bytes()).decode("ascii") + "'>"]
    parts += ["<p>Fixed original-bet/stake repricing is a separate attribution diagnostic. Unfunded shadow paths are excluded from these account graphs. The original 84-rule selection audit applies only to original recorded paths, not the new reference replays.</p><h2>" + APPENDIX_TITLE + "</h2>"]
    parts += ["<article id='" + e["strategy_id"] + "'><h3>" + html.escape(e["title"]) + "</h3><p>" + html.escape(e["explanation"]) + "</p><pre>" + html.escape(e["exact_config"]) + "</pre></article>" for e in entries]
    parts += ["</body></html>"]
    (figures / "all_reference_strategies_4seasons_100k.html").write_text("".join(parts), encoding="utf-8")
    return {"trace_ids": [trace.uid for trace in interactive.data], "initially_visible_traces": 56,
        "plotly_js_bundled": True, "preserved_original_graphs": 2, "linear_money_y_axis": True,
        "unfunded_shadow_paths_plotted": False}


def _paired_table(frame: pd.DataFrame, ids: list[str]) -> pd.DataFrame:
    original = frame.loc[frame.execution_mode.eq("original")].set_index("strategy_id").reindex(ids)
    reference = frame.loc[frame.execution_mode.eq("reference_replay")].set_index("strategy_id").reindex(ids)
    rows = []
    for sid in ids:
        old, ref = original.loc[sid], reference.loc[sid]
        rows.append({"ID": sid, "Trades\norig / ref": number(old.bet_count, 0) + " / " + number(ref.bet_count, 0),
            "Ref hit\n%": number(ref.hit_ratio, percent=True), "Ref trade\nfreq. %": number(ref.trading_frequency, percent=True),
            "Ref model\ngain/trade": number(ref.expected_gain_per_trade_units),
            "Ref model\nloss/trade": number(ref.expected_loss_per_trade_units),
            "Ref avg.\nwin gain": number(ref.avg_winning_gain_units), "Ref avg.\nlosing loss": number(ref.avg_losing_loss_units),
            "Orig\nROI %": number(old.roi, percent=True), "Ref\nROI %": number(ref.roi, percent=True),
            "Orig end\nbalance": number(INITIAL + old.profit, 0), "Ref end\nbalance": number(INITIAL + ref.profit, 0),
            "Ref max.\ndrawdown": number(ref.max_drawdown, percent=True)})
    return pd.DataFrame(rows)


def _scope_summary(data: dict) -> pd.DataFrame:
    rows = []
    for key, label, count in (("metrics_single", "2025/26", 84), ("metrics_continuous", "Four-season continuous", 56), ("metrics_seasonal", "Separate annual accounts", 224)):
        frame = data[key]
        orig = frame.loc[frame.execution_mode.eq("original")]
        ref = frame.loc[frame.execution_mode.eq("reference_replay")]
        fixed = frame.loc[frame.execution_mode.eq("fixed_exposure_reference")]
        rows.append({"Population": label, "Accounts": str(count), "Original\nprofitable": str(int(orig.profit.gt(0).sum())),
            "Reference replay\nprofitable": str(int(ref.profit.gt(0).sum())),
            "Reference replay\nunfunded": str(int((~ref.funding_feasible.astype(bool)).sum())),
            "Fixed-exposure\nunfunded shadows": str(int((~fixed.funding_feasible.astype(bool)).sum()))})
    return pd.DataFrame(rows)


def _body(data: dict, figures: Path) -> tuple[list[dict], dict[str, pd.DataFrame]]:
    old_pages, displays = temporal._body(data["legacy"], figures)
    # Relocate only figure pointers. The two original PNGs are bytewise copies.
    for page in old_pages:
        for block in page["blocks"]:
            if block["type"] == "figure": block["path"] = figures / block["path"].name
    new_pages = []
    def page(title, orientation="portrait"):
        item = {"title": title, "orientation": orientation, "blocks": []}; new_pages.append(item); return item["blocks"]
    def p(blocks, text): blocks.append({"type": "paragraph", "text": text})
    def table(blocks, frame, caption): blocks.append({"type": "table", "frame": frame, "caption": caption})
    summary = _scope_summary(data); displays["Execution scope summary"] = summary
    b = page(TITLE)
    p(b, "DRAFT FOR REVIEW | October 15, 2026 | EPL component of the statistical-arbitrage final project.")
    p(b, "The project now checks how much recorded strategy performance depends on the archived prices used to trade. The exact 84-strategy catalog, its all-84 graph, the continuous four-season graph for 56 compatible rules, annual diagnostics and the earlier conditional selection audit are retained. This report adds complete Bet365 reference-price replays, fixed-exposure quote attribution and bookmaker concentration. The final appendix still explains every strategy.")
    p(b, "There are 364 original account scopes: 84 common-season accounts, 56 uninterrupted four-season accounts and 224 separate annual accounts. Each is compared in three modes, producing 1,092 metric rows. An original account keeps its exact saved prices; a reference replay recomputes price-dependent choices and sizing using Bet365; fixed-exposure repricing retains the original selections and allocated stakes to isolate the recorded quote effect.")
    table(b, summary, "Table E1. Profitable means strictly positive final profit. Annual accounts are separate 100,000-unit diagnostic starts. Unfunded fixed-exposure shadows are not feasible investment accounts.")
    p(b, "Reference replay profits are realised historical outcomes under a different archived quote assumption. All learned probabilities and market inputs remain frozen. No model is refitted, no unseen outcome is introduced and no strategy is chosen from this comparison. The original 84-rule selection adjustment has not been rerun for reference prices and does not establish significance for the new replay paths.")
    b = page("Price replay and pure quote attribution answer different questions")
    p(b, "Every reference-price strategy starts with 100,000 over its available original history. Odds bands, maximum-edge outcome choice, strict thresholds, consensus votes and fractional-Kelly requests are recomputed at the reference quote. Fixed 1,000 requests and date-start capital rules are retained. All same-date requests are jointly allocated before any settlement, so same-date winnings cannot finance another trade. The continuous four-season account has no season reset, deposits, borrowing, interest or transaction fee assumption.")
    p(b, "A reference replay can trade different matches, select different outcomes and follow a different capital path. Its difference from the original is therefore not pure price uplift. The separate fixed-exposure diagnostic retains each original bet and allocated stake and replaces only its payoff odds. Quote uplift equals the sum over original winners of stake multiplied by original odds minus reference odds. Original profit equals fixed-exposure reference profit plus that realised quote uplift.")
    p(b, "Keeping old stake requests can exhaust the lower-payout shadow account. Funding is checked against cash at the start of every date, before that date's settlement. If original stakes cannot be funded, the diagnostic remains an arithmetic shadow: its raw cumulative profit and negative balances support attribution, while feasible terminal wealth, capital return and conventional maximum drawdown are unavailable. No negative shadow is plotted as a funded 100,000-unit strategy. Full reference replays resize requests to cash and remain feasible.")
    p(b, "All original probability inputs stay unchanged, including residual-model inputs and the raw de-vigged Bet365 probabilities used for market controls. The comparison isolates the quote panel used by the betting rule; it is not a new retail information set. Archived quote timestamps and simultaneous availability are unverified in both modes, and neither represents confirmed fills, account limits, liquidity or executable closing-line value.")
    b = page("Every reference-price account through four seasons", "landscape")
    p(b, "All 56 compatible rules appear on this one money-versus-time graph. Each starts at exactly 100,000 on August 4, 2022, settles the same 1,470 archived fixtures through May 24, 2026 and carries its own capital without resets. These are recomputed reference-price accounts, not fixed-exposure shadows.")
    b.append({"type": "figure", "path": figures / "all_reference_strategies_4seasons_100k.png",
        "caption": "Figure E1. Every S017–S072 reference replay on linear monetary and calendar-date axes; balances are separate accounts. The offline interactive version initially shows every trace and retains both original graphs and all strategy explanations."})
    for key, ids, label, population in (("metrics_single", ALL_IDS, "All-strategy 2025/26 price comparison", "374 common fixtures"),
            ("metrics_continuous", COMPATIBLE_IDS, "Four-season price comparison", "1,470 common fixtures, one uninterrupted account")):
        display = _paired_table(data[key], ids); displays[label] = display
        for start in range(0, len(ids), 28):
            b = page(label + ": " + ids[start] + "–" + ids[min(start + 27, len(ids) - 1)], "landscape")
            p(b, "Original and reference each begin at 100,000; " + population + ". Reference gain/loss expectations are probability-weighted model claims at allocated reference stakes; winning/losing means are observed net payoffs. Frequency uses the common priced universe. Undefined cash-strategy ratios appear as dashes.")
            table(b, display.iloc[start:start + 28], "Table E2. Complete original-versus-reference comparison in stable ID order. ROI is profit divided by stake turnover. Ending balance is initial 100,000 plus realised profit; drawdown includes the initial cash point. Full capital returns, intervals and all 1,092 metric rows are in execution_overview.xlsx.")
    b = page("Quote dependence, probability claims and bookmaker concentration", "landscape")
    quotes = data["quotes"]
    p(b, f"The quote inventory covers {number(quotes.get('match_count', 1470), 0)} fixtures and {number(quotes.get('outcome_quote_count', 4410), 0)} outcome-price comparisons. Best archived outcome prices are at least the Bet365 reference in every checked comparison: {str(quotes.get('all_best_at_least_reference', False))}. Provider quote timestamps do not verify an executable contemporaneous opportunity. The complete inventory, uplift distribution and missing-clock counts are retained in the workbook and quote_summary.json.")
    if "median_price_ratio_best_to_reference" in quotes:
        p(b, "Across all outcome-price cells, the median best/reference uplift is " + number(quotes["median_price_ratio_best_to_reference"] - 1, percent=True) +
            " and its mean is " + number(quotes["mean_price_ratio_best_to_reference"] - 1, percent=True) +
            ". This is a quote-cell distribution, not a strategy return or a selected-bet average. Winning-payoff uplift depends on which quoted outcomes were actually selected and staked.")
    p(b, "For each original selected outcome, let q be its frozen de-vigged reference-market probability. The market-implied component is stake × (q × odds − 1); the model probability premium is stake × odds × (p − q). Their sum equals the model expected profit. Replacing p with the realised win indicator gives actual profit as the same market component plus an outcome residual. These exact accounting identities distinguish quote effects from a model's probability claim and realised variation. They do not treat q as truth or demonstrate a predictive edge.")
    attribution = data["payoff_attribution"]
    examples = attribution.loc[(attribution.scope.astype(str).str.contains("single") & attribution.strategy_id.eq("S013")) |
        (attribution.scope.astype(str).str.contains("continuous") & attribution.strategy_id.isin(["S019", "S065"]))]
    if len(examples):
        cols = [c for c in ("scope", "strategy_id", "actual_profit", "model_expected_profit", "market_implied_expected_profit", "model_probability_premium", "realized_price_uplift_vs_reference") if c in examples]
        display = examples[cols].copy()
        display.columns = ["Scope", "ID", "Actual\nprofit", "Model expected\nprofit", "Market-implied\ncomponent", "Model probability\npremium", "Realised quote\nuplift"][:len(cols)]
        for c in display.columns[2:]: display[c] = display[c].map(number)
        displays["Attribution examples"] = display
        table(b, display, "Table E3. Illustrative fixed original-account decompositions, in monetary units; all 364 are available. S013 is the earlier calibration-selected flat-stake rule. S019 and S065 were previously discussed in seasonal stability checks. These examples do not select future winners.")
    concentration = data["book_concentration"]
    p(b, f"The bookmaker inventory retains {len(concentration):,} account/book rows. Bet and stake shares measure how strongly an original rule depends on individual providers; the reported Herfindahl indices sum squared shares. Signed profit shares can be negative or exceed 100% when providers' gains and losses offset, so they are not portfolio weights. Book concentration is distinct from quote availability, funding feasibility and provider execution capacity.")
    p(b, "Detailed reference and fixed-exposure bet/date ledgers are exported as CSV and Parquet beside the report. To keep the review workbook compact, it contains complete metrics, comparisons, funding indicators, payoff decomposition, book/quote diagnostics, previous selection tables, catalog and graph values; it links the external ledgers instead of duplicating every trade. Neither historical price sensitivity nor the retained conditional bootstrap constitutes an independent future test.")
    # Preserve the full temporal study while making the new findings the opening.
    old_pages[0]["title"] = "Preserved study: chronological coverage and account design"
    for page_ in old_pages:
        if page_["title"] == "Limitations, references and reproducibility":
            page_["blocks"].append({"type": "paragraph", "text": "This combined continuation is reproduced with python -m epl_execution.report --results PATH. Its execution_overview.xlsx contains all three price modes and all account scopes. The prior validation_overview.xlsx and original study remain frozen. epl_execution/PROTOCOL.md specifies quote replay, shadow funding and attribution. New reference replays do not inherit the original-path selection adjustment. The exact 84-strategy appendix remains the final section."})
    # Begin with the new price section, then retain every earlier body section.
    return new_pages + old_pages, displays


def _markdown(pages, entries) -> str:
    return temporal._markdown(pages, entries)


def _workbook(path: Path, data: dict, entries: list[dict], registry: dict, output: Path) -> dict:
    definitions = pd.DataFrame([
        ("Scope", "84 single-season + 56 continuous + 224 annual accounts, each in three modes = 1092 metric rows."),
        ("Original", "Exact preserved original account, quotes and allocated stakes."),
        ("Reference replay", "Bet365 price override; recompute price-dependent eligibility, selection, consensus, sizing and allocation. Frozen probabilities remain unchanged."),
        ("Fixed exposure", "Original bets and allocated stakes repriced at Bet365 for exact quote attribution. Unfunded paths are arithmetic shadows, not feasible terminal wealth."),
        ("Money", "Each account starts independently at 100000 monetary units. Continuous account has no resets or deposits. ROI divides by turnover, capital return by initial capital."),
        ("Expected payoffs", "Conditional model-probability claims at the saved stake path, distinct from observed winning/losing payoffs."),
        ("Selection inference", "Original 84-rule cashflow audit remains conditional and applies to original paths only, not new reference replays."),
        ("External ledgers", "Every reference/fixed-exposure bet and date ledger is CSV/Parquet beside this report. Detailed ledger paths are listed on the External ledger files sheet."),
        ("No execution proof", "Archived prices and missing quote clocks do not verify availability, fills or executable closing-line value.")], columns=["Item", "Definition"])
    ledger_files = sorted(p for p in output.glob("*.parquet") if p.name.startswith(("bets_", "equity_")))
    external = pd.DataFrame([{"path_relative_to_workbook": "../" + p.name,
        "sha256": lab._sha(p), "bytes": p.stat().st_size} for p in ledger_files])
    sheets = {"Read me": definitions, "Single all modes": data["metrics_single"],
        "Continuous all modes": data["metrics_continuous"], "Annual all modes": data["metrics_seasonal"],
        "Execution comparisons": data["execution_comparison"], "Payoff attribution": data["payoff_attribution"],
        "Book concentration": data["book_concentration"], "Quote outcome comparisons": data["quote_outcomes"],
        "Quote summary": pd.DataFrame([{"item": k, "value": lab._json(v) if isinstance(v, (dict, list)) else v} for k, v in data["quotes"].items()]),
        "Reference graph incl initial": data["graph"], "Original selection comparisons": data["legacy"]["selection_comparison"],
        "Original selection global": data["legacy"]["selection_global"], "Original season stability": data["legacy"]["season_stability"],
        "Complete84 catalog": pd.DataFrame(data["catalog"]), "Final84 appendix": pd.DataFrame(entries),
        "External ledger files": external,
        "Input provenance": pd.DataFrame([{"path": p, **record} for p, record in registry.items()])}
    for key, sheet in (("closing_quote_outcomes", "Archived same-book closing"),
            ("archived_closing_reference", "Archived closing price panel")):
        if key in data: sheets[sheet] = data[key]
    if "closing_quotes" in data:
        sheets["Archived closing summary"] = pd.DataFrame([{"item": k,
            "value": lab._json(v) if isinstance(v, (dict, list)) else v} for k, v in data["closing_quotes"].items()])
    percentages = {"hit_ratio", "trading_frequency", "trade_rate_all_predictions", "price_coverage", "active_date_fraction",
        "roi", "expected_roi", "return_on_initial_capital", "max_drawdown", "avg_selected_probability",
        "avg_model_expected_return_per_unit_staked", "avg_realized_return_per_unit_staked", "empirical_break_even_hit_ratio",
        "avg_winning_return_per_unit_staked", "avg_losing_return_per_unit_staked", "max_date_capital_exposure", "mean_date_capital_exposure",
        "date_return_mean", "date_return_sample_std", "bet_share", "stake_share", "signed_profit_share", "largest_book_stake_share",
        "mean_model_probability", "mean_reference_probability", "min_edge", "bankroll_fraction", "kelly_fraction", "max_bet_fraction", "date_cap_fraction"}
    with pd.ExcelWriter(path, engine="xlsxwriter", engine_kwargs={"options": {"strings_to_urls": False}}) as writer:
        book = writer.book
        head = book.add_format({"bold": True, "font_color": "white", "bg_color": "#12364D", "text_wrap": True})
        num = book.add_format({"num_format": "#,##0.0000"}); pct = book.add_format({"num_format": "0.00%"})
        for name, frame in sheets.items():
            safe = frame.copy()
            for column in safe:
                if isinstance(safe[column].dtype, pd.DatetimeTZDtype): safe[column] = safe[column].dt.tz_convert(None)
                if safe[column].dtype == "object": safe[column] = safe[column].map(lambda v: lab._json(v.tolist() if isinstance(v, np.ndarray) else v) if isinstance(v, (dict, list, tuple, np.ndarray)) else v)
                elif pd.api.types.is_numeric_dtype(safe[column]) and not pd.api.types.is_bool_dtype(safe[column]):
                    safe[column] = safe[column].map(lambda v: "+∞" if v == float("inf") else "−∞" if v == float("-inf") else v)
            safe.to_excel(writer, sheet_name=name, index=False, na_rep="—")
            worksheet = writer.sheets[name]; worksheet.write_row(0, 0, list(safe.columns), head); worksheet.set_row(0, 42)
            worksheet.freeze_panes(1, 1)
            if len(safe.columns): worksheet.autofilter(0, 0, len(safe), len(safe.columns) - 1)
            for index, column in enumerate(safe.columns):
                base = column.removesuffix("_ci_low").removesuffix("_ci_high")
                is_pct = base in percentages or column.endswith("_per_turnover")
                width = 65 if column in ("explanation", "exact_config", "Definition", "coverage", "path", "name", "strategy_name", "path_relative_to_workbook") else min(42, max(14, len(column) * .8))
                worksheet.set_column(index, index, width, pct if is_pct else num)
    if path.stat().st_size >= 50_000_000:
        raise AssertionError("The compact execution workbook exceeds its 50 MB review budget")
    return {name: len(frame) for name, frame in sheets.items()}



# PDF/Word layout adapted from the preserved temporal report renderer.
def _pdf(pages, entries, path: Path) -> None:
    fonts = lab._fonts()
    styles = {
        "title": ParagraphStyle("ValTitle", fontName=fonts["bold"], fontSize=14, leading=16, spaceAfter=7, textColor=colors.HexColor("#12364D")),
        "body": ParagraphStyle("ValBody", fontName=fonts["body"], fontSize=10.2, leading=12.4, spaceAfter=7),
        "caption": ParagraphStyle("ValCaption", fontName=fonts["body"], fontSize=8.5, leading=10, spaceAfter=4),
        "cell": ParagraphStyle("ValCell", fontName=fonts["sans"], fontSize=8.0, leading=9.3),
        "head": ParagraphStyle("ValHead", fontName=fonts["bold"], fontSize=8.0, leading=9.3, textColor=colors.white),
        "entrytitle": ParagraphStyle("ValEntryTitle", fontName=fonts["bold"], fontSize=9.7, leading=11.4, spaceBefore=5, spaceAfter=4),
        "entry": ParagraphStyle("ValEntry", fontName=fonts["body"], fontSize=9.2, leading=10.6, spaceAfter=4),
        "config": ParagraphStyle("ValConfig", fontName=fonts["mono"], fontSize=6.8, leading=8.1, spaceAfter=7, splitLongWords=True, wordWrap="CJK")}
    def header(canvas, doc):
        width, height = canvas._pagesize
        canvas.setFont(fonts["sans"], 8); canvas.setFillColor(colors.HexColor("#52616B"))
        canvas.drawString(36, height - 21, "EPL STRATEGY RESEARCH - QUOTE SENSITIVITY | EXPLORATORY | October 15, 2026")
        canvas.drawRightString(width - 36, 21, str(doc.page))
    doc = BaseDocTemplate(str(path), pagesize=letter, title=TITLE, author="Aman Dhillon and Panagiotis Housos")
    for orientation, size in (("portrait", letter), ("landscape", landscape(letter))):
        width, height = size
        doc.addPageTemplates(PageTemplate(id=orientation, pagesize=size, onPage=header,
            frames=[Frame(36, 36, width - 72, height - 72, leftPadding=0, rightPadding=0, topPadding=0, bottomPadding=0)]))
    story = []
    for index, page in enumerate(pages):
        if index: story.extend([NextPageTemplate(page["orientation"]), PageBreak()])
        width = (letter if page["orientation"] == "portrait" else landscape(letter))[0] - 72
        story.append(Paragraph(html.escape(page["title"]), styles["title"]))
        for block in page["blocks"]:
            if block["type"] == "paragraph": story.append(Paragraph(html.escape(block["text"]), styles["body"]))
            elif block["type"] == "figure":
                with PILImage.open(block["path"]) as picture: ratio = picture.height / picture.width
                story.extend([Image(str(block["path"]), width=width - 10, height=(width - 10) * ratio), Paragraph(html.escape(block["caption"]), styles["caption"])])
            else:
                story.append(Paragraph(html.escape(block["caption"]), styles["caption"]))
                frame = block["frame"]
                raw = [list(frame.columns)] + frame.values.tolist()
                cells = [[Paragraph(html.escape(str(value)).replace("\n", "<br/>"), styles["head" if row_index == 0 else "cell"]) for value in row] for row_index, row in enumerate(raw)]
                weights = BODY_WEIGHTS if page["title"].startswith(("Four-season trading", "Original 2025")) else [1 / len(frame.columns)] * len(frame.columns)
                table = Table(cells, colWidths=[width * weight for weight in weights], repeatRows=1)
                table.setStyle(TableStyle([("BACKGROUND", (0,0), (-1,0), colors.HexColor("#12364D")), ("ROWBACKGROUNDS", (0,1), (-1,-1), [colors.white, colors.HexColor("#F0F4F6")]), ("VALIGN", (0,0), (-1,-1), "MIDDLE"), ("LEFTPADDING", (0,0), (-1,-1), 3), ("RIGHTPADDING", (0,0), (-1,-1), 3), ("TOPPADDING", (0,0), (-1,-1), 1.1), ("BOTTOMPADDING", (0,0), (-1,-1), 1.1)]))
                story.append(table)
    story.extend([NextPageTemplate("portrait"), PageBreak(), Paragraph(APPENDIX_TITLE, styles["title"])])
    for entry in entries:
        story.append(KeepTogether([Paragraph(html.escape(entry["title"]), styles["entrytitle"]),
            Paragraph(html.escape("Origin: " + entry["origin"] + ". " + entry["explanation"]), styles["entry"]),
            Paragraph(html.escape("Exact original configuration: " + entry["exact_config"]), styles["config"])]))
    doc.build(story)

def _docx(pages, entries, path: Path) -> None:
    doc = Document()
    normal = doc.styles["Normal"]; normal.font.name = "Times New Roman"; normal.font.size = Pt(10.2)
    normal.paragraph_format.line_spacing = Pt(12.4); normal.paragraph_format.space_after = Pt(7)
    heading = doc.styles["Heading 1"]; heading.font.name = "Arial"; heading.font.size = Pt(14)
    heading.paragraph_format.space_before = Pt(0); heading.paragraph_format.space_after = Pt(7)
    doc.core_properties.title = TITLE; doc.core_properties.author = "Aman Dhillon and Panagiotis Housos"
    def section(index, wide):
        sec = doc.sections[0] if index == 0 else doc.add_section(WD_SECTION_START.NEW_PAGE)
        sec.orientation = WD_ORIENT.LANDSCAPE if wide else WD_ORIENT.PORTRAIT
        sec.page_width = Inches(11 if wide else 8.5); sec.page_height = Inches(8.5 if wide else 11)
        sec.top_margin = sec.bottom_margin = sec.left_margin = sec.right_margin = Inches(.5)
        sec.header_distance = sec.footer_distance = Inches(.2)
        if index == 0:
            run = sec.header.paragraphs[0].add_run("EPL STRATEGY RESEARCH - QUOTE SENSITIVITY | EXPLORATORY | October 15, 2026")
            run.font.name = "Arial"; run.font.size = Pt(8)
            footer = sec.footer.paragraphs[0]; footer.alignment = 2
            field = OxmlElement("w:fldSimple"); field.set(qn("w:instr"), "PAGE"); footer._p.append(field)
    for index, page in enumerate(pages):
        wide = page["orientation"] == "landscape"; section(index, wide); doc.add_heading(page["title"], 1)
        for block in page["blocks"]:
            if block["type"] == "paragraph": doc.add_paragraph(block["text"]); continue
            if block["type"] == "figure": doc.add_picture(str(block["path"]), width=Inches(9.85 if wide else 7.35))
            caption = doc.add_paragraph(block["caption"]); caption.paragraph_format.line_spacing = Pt(10)
            for run in caption.runs: run.font.size = Pt(8.5)
            if block["type"] == "figure": continue
            frame = block["frame"]; table = doc.add_table(rows=1, cols=len(frame.columns)); table.style = "Light Shading Accent 1"; table.autofit = False
            weights = BODY_WEIGHTS if page["title"].startswith(("Four-season trading", "Original 2025")) else [1 / len(frame.columns)] * len(frame.columns)
            width = 10 if wide else 7.5
            for col, weight in zip(table.columns, weights): col.width = Inches(width * weight)
            for cell, value in zip(table.rows[0].cells, frame.columns): cell.text = str(value)
            for values in frame.itertuples(index=False, name=None):
                for cell, value in zip(table.add_row().cells, values): cell.text = str(value)
            for row in table.rows:
                for cell, weight in zip(row.cells, weights):
                    cell.width = Inches(width * weight)
                    for paragraph in cell.paragraphs:
                        paragraph.paragraph_format.line_spacing = Pt(9.3); paragraph.paragraph_format.space_after = Pt(0)
                        for run in paragraph.runs: run.font.name = "Arial"; run.font.size = Pt(8)
    section(len(pages), False); doc.add_heading(APPENDIX_TITLE, 1)
    for entry in entries:
        title = doc.add_paragraph(); title.paragraph_format.keep_with_next = True
        run = title.add_run(entry["title"]); run.bold = True; run.font.name = "Arial"; run.font.size = Pt(9.7)
        paragraph = doc.add_paragraph("Origin: " + entry["origin"] + ". " + entry["explanation"])
        paragraph.paragraph_format.keep_with_next = True; paragraph.paragraph_format.line_spacing = Pt(10.6)
        for run in paragraph.runs: run.font.size = Pt(9.2)
        paragraph = doc.add_paragraph("Exact original configuration: " + entry["exact_config"])
        paragraph.paragraph_format.keep_together = True; paragraph.paragraph_format.line_spacing = Pt(8.1)
        for run in paragraph.runs: run.font.name = "Consolas"; run.font.size = Pt(6.8)
    doc.save(path)

def build_report(output_dir: Path) -> dict:
    output = Path(output_dir).resolve(); data = _load(output); _validate(data)
    entries = _appendix(data["catalog"])
    sources = [output / (name + ".parquet") for name in TABLE_NAMES]
    sources += [output / (name + ".parquet") for name in OPTIONAL_TABLES if (output / (name + ".parquet")).exists()]
    if (output / "closing_quote_summary.json").exists(): sources.append(output / "closing_quote_summary.json")
    sources += [output / name for name in ("catalog.json", "quote_summary.json", "graph_equity_reference_continuous.csv", "source_provenance.json") if (output / name).exists()]
    sources += [Path(__file__).resolve(), Path(temporal.__file__).resolve(), Path(lab.__file__).resolve(), ROOT / "epl_strategy/report.py"]
    sources += [TEMPORAL / name for name in ("catalog.json", "compatible_catalog.json", "metrics_continuous.parquet", "metrics_seasonal.parquet", "season_stability.parquet", "bets_continuous.parquet", "equity_continuous.parquet", "bets_seasonal.parquet", "equity_seasonal.parquet", "graph_equity_continuous.csv", "selection_comparison.parquet", "selection_global.parquet", "figures/all_strategies_4seasons_100k.png", "figures/all_84_strategies_2025_100k.png")]
    sources += [temporal.LEGACY / "metrics.parquet"]
    sources += [ROOT / name for name in ("epl_execution/PROTOCOL.md", "epl_strategy_validation/PROTOCOL.md", "epl_strategy_lab/PROTOCOL.md") if (ROOT / name).exists()]
    sources += [p for p in output.glob("*.parquet") if p.name.startswith(("bets_", "equity_"))]
    sources = sorted(set(sources))
    registry = {lab._relative(p): {"sha256": lab._sha(p), "bytes": p.stat().st_size} for p in sources}
    report = output / "report"; report.mkdir(exist_ok=True); figures = output / "figures"
    graph_info = _graphs(data, figures, entries)
    pages, displays = _body(data, figures)
    for name, frame in displays.items(): frame.to_csv(report / (name.lower().replace(" ", "_").replace("/", "_") + ".csv"), index=False)
    pd.DataFrame(entries).to_csv(report / "complete84_strategy_appendix.csv", index=False)
    md = report / "validation_report.md"; pdf = md.with_suffix(".pdf"); docx = md.with_suffix(".docx")
    md.write_text(_markdown(pages, entries), encoding="utf-8")
    _pdf(pages, entries, pdf); _docx(pages, entries, docx)
    sheets = _workbook(report / "execution_overview.xlsx", data, entries, registry, output)
    reader = PdfReader(pdf); extracted = "\n".join(p.extract_text() or "" for p in reader.pages)
    if APPENDIX_TITLE not in extracted: raise AssertionError("PDF omitted its final appendix")
    appendix_text = " ".join(extracted.split(APPENDIX_TITLE, 1)[1].split())
    if any(" ".join(e["title"].split()) not in appendix_text for e in entries): raise AssertionError("PDF omitted a strategy definition")
    if entries[-1]["strategy_id"] not in reader.pages[-1].extract_text(): raise AssertionError("Appendix is not the last report content")
    changed = [lab._relative(p) for p in sources if lab._sha(p) != registry[lab._relative(p)]["sha256"]]
    if changed: raise AssertionError("Consumed report inputs changed: " + str(changed))
    produced = [p for folder in (report, figures) for p in folder.rglob("*") if p.is_file() and p.name != "report_manifest.json" and "previews" not in p.parts]
    manifest = {"status": "retrospective_quote_sensitivity_combined_report", "draft_date_new_york": "2026-10-15",
        "built_at_utc": datetime.now(timezone.utc).isoformat(), "full_catalog_count": 84, "continuous_strategy_count": 56,
        "account_scopes": 364, "metric_rows": 1092, "execution_modes": list(MODES), "initial_capital_per_account": INITIAL,
        "graph": graph_info, "appendix_ids": [e["strategy_id"] for e in entries], "appendix_is_last": True,
        "pdf_pages": len(reader.pages), "body_sections": len(pages), "workbook_sheet_rows": sheets,
        "models_refitted": False, "consumed_source_files_unchanged": True, "inputs": registry,
        "outputs": {p.relative_to(output).as_posix(): {"sha256": lab._sha(p), "bytes": p.stat().st_size} for p in sorted(produced)},
        "manifest_order": "Final build envelope and independent audit are deliberately not consumed report inputs.",
        "limits": "Retrospective current-club cohort; missing quote clocks; fixed exposure may be unfunded; no new holdout or selection inference for reference replay."}
    (report / "report_manifest.json").write_text(json.dumps(manifest, ensure_ascii=False, indent=2, allow_nan=False), encoding="utf-8")
    return {"output_directory": str(report), "full_catalog_count": 84, "continuous_strategy_count": 56,
        "metric_rows": 1092, "pdf_pages": len(reader.pages), "appendix_entries": len(entries), "source_files_unchanged": True,
        "workbook_bytes": (report / "execution_overview.xlsx").stat().st_size}


def main() -> None:
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument("--results", required=True, type=Path)
    print(json.dumps(build_report(parser.parse_args().results), indent=2))


if __name__ == "__main__": main()
