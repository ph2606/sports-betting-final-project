"""Presentation invariants for a small prospective forecast evaluation."""
import json

import numpy as np
import pandas as pd
import pytest
from docx import Document
from openpyxl import load_workbook
from pypdf import PdfReader

from epl_prospective import report
from epl_strategy_lab.catalog import build_catalog


@pytest.fixture
def finished_data(monkeypatch):
    monkeypatch.setattr(report.execution, "_validate", lambda data: None)
    labels = ["team", "team_xg", "team_player", "combined"]
    audit = pd.DataFrame([{"match_id": "f" + str(i), "model_id": "m" + str(j),
        "feature_group": labels[j], "home_team": "Home " + str(i), "away_team": "Away " + str(i),
        "created_at_utc": "2026-10-14T09:22:52Z", "kickoff_utc": "2026-10-17T14:00:00Z",
        "settled_at_utc": "2026-10-20T22:30:00Z", "p_H": .7, "p_D": .2, "p_A": .1,
        "result": "H", "home_goals": 1, "away_goals": 0, "status": "resolved",
        "predicted_outcome": "H", "correct": True, "log_loss": -np.log(.7), "brier_score": .14}
        for i in range(10) for j in range(4)])
    scores = pd.DataFrame([{"model_id": "m" + str(i), "n_matches": 10, "accuracy": 1.,
        "accuracy_ci_low": .72, "accuracy_ci_high": 1., "log_loss": .3567,
        "log_loss_ci_low": .2, "log_loss_ci_high": .5, "brier_score": .14,
        "brier_score_ci_low": .1, "brier_score_ci_high": .2,
        "classwise_ece_5": .2, "classwise_ece_10": .2, "classwise_ece_15": .2} for i in range(4)])
    paired = pd.DataFrame([{"model_a": "m" + str(i), "model_b": "m" + str(j),
        "common_fixture_count": 10, "accuracy_diff": 0., "log_loss_diff": 0.,
        "log_loss_diff_ci_low": -.1, "log_loss_diff_ci_high": .1,
        "brier_score_diff": 0., "brier_score_diff_ci_low": -.1, "brier_score_diff_ci_high": .1}
        for i in range(4) for j in range(i + 1, 4)])
    return {"forecast_outcome_audit": audit, "model_scores": scores, "pairwise_scores": paired,
        "fixture_outcomes": audit.drop_duplicates("match_id"), "next_forecasts": pd.DataFrame(columns=["match_id", "model_id"]),
        "source_freshness": pd.DataFrame({"source": ["FPL"], "snapshot_role": ["results"],
            "observed_at_utc": ["2026-10-20T22:30:00Z"], "latest_completed_match_date": ["2026-10-19"],
            "source_capture_valid": [True], "core_source_status": ["valid"]}),
        "status": {"as_of_utc": "2026-10-20T22:30:00Z", "original_forecast_count": 40,
            "original_fixture_count": 10, "resolved_original_forecasts": 40, "pending_original_forecasts": 0,
            "resolved_original_fixtures": 10, "new_forecasts_released": 0, "new_target_fixtures": 0,
            "recorded_trade_count": 0, "prospective_betting_performance_available": False,
            "results_snapshot": "results capture", "player_snapshot": "player capture",
            "evaluation_path": "evaluation", "next_run_path": "next run"},
        "integrity": {"original_forecast_bytes_unchanged": True},
        "historical": {"catalog": build_catalog()}}


def test_original_batch_counts_ten_games_not_forty_independent_matches(finished_data):
    report._validate(finished_data)
    assert len(finished_data["forecast_outcome_audit"]) == 40
    assert finished_data["forecast_outcome_audit"].match_id.nunique() == 10


def test_reject_duplicate_original_model_fixture_identity(finished_data):
    frame = finished_data["forecast_outcome_audit"]
    frame.loc[1, "model_id"] = frame.loc[0, "model_id"]
    with pytest.raises(ValueError, match="identity is missing or duplicated"):
        report._validate(finished_data)


def test_reject_forecast_after_kickoff_or_invented_betting_performance(finished_data):
    frame = finished_data["forecast_outcome_audit"]
    frame.loc[0, "created_at_utc"] = "2026-10-18T00:00:00Z"
    with pytest.raises(ValueError, match="before kickoff"):
        report._validate(finished_data)
    frame.loc[0, "created_at_utc"] = "2026-10-14T09:22:52Z"
    finished_data["status"]["prospective_betting_performance_available"] = True
    with pytest.raises(ValueError, match="no recorded prospective betting"):
        report._validate(finished_data)


def test_reject_score_denominator_that_counts_forecasts_as_independent_games(finished_data):
    finished_data["model_scores"].loc[0, "n_matches"] = 40
    with pytest.raises(ValueError, match="score denominators"):
        report._validate(finished_data)


def test_score_tables_keep_paired_difference_direction_and_uncertainty(finished_data):
    scores = report._score_table(finished_data)
    paired = report._pairwise_table(finished_data)
    assert len(scores) == 4 and scores["Resolved\nfixtures"].eq("10").all()
    assert scores["95% Wilson\naccuracy interval"].eq("72.00% to 100.00%").all()
    assert len(paired) == 6 and paired["Common\nfixtures"].eq("10").all()
    assert "Log loss\nA minus B" in paired and "95% paired\nlog-loss interval" in paired


def test_workbook_preserves_every_forecast_probabilities_and_historical_link(tmp_path, finished_data):
    finished_data["original_forecast_coverage"] = pd.DataFrame({"status": ["withheld"], "reason": ["missing pre-match quote"]})
    finished_data["next_forecast_coverage"] = pd.DataFrame({"status": ["withheld"], "reason": ["missing pre-match quote"]})
    entries = report.execution._appendix(build_catalog())
    path = tmp_path / "prospective.xlsx"
    rows = report._workbook(path, finished_data, entries, {"input": {"sha256": "a" * 64, "bytes": 1}})
    assert rows["Every original forecast"] == 40 and rows["Original fixture outcomes"] == 10
    assert rows["Final84 appendix"] == 84
    assert rows["Original release exclusions"] == 1 and rows["Next release exclusions"] == 1
    book = load_workbook(path)
    sheet = book["Every original forecast"]
    names = [cell.value for cell in sheet[1]]
    assert sheet.cell(2, names.index("p_H") + 1).value == .7
    assert sheet.cell(2, names.index("p_H") + 1).number_format == "0.00%"
    readme = book["Read me"]
    hyperlink = [row[1].hyperlink for row in readme if row[0].value == "Historical metrics workbook"][0]
    assert hyperlink is not None and hyperlink.target.endswith("execution_overview.xlsx")
    assert book["Original release exclusions"][2][1].value == "missing pre-match quote"
    book.close()


def test_prospective_heading_and_exact84_appendix_remain_last_in_pdf_word_and_markdown(tmp_path):
    entries = report.execution._appendix(build_catalog())
    assert [json.loads(entry["exact_config"]) for entry in entries] == build_catalog()
    pages = [{"title": report.TITLE, "orientation": "portrait", "blocks": [
        {"type": "paragraph", "text": "Ten distinct fixtures. Prospective trading returns are unavailable."}]}]
    pdf = tmp_path / "report.pdf"; docx = tmp_path / "report.docx"
    report._pdf(pages, entries, pdf); report._docx(pages, entries, docx)
    reader = PdfReader(pdf)
    text = " ".join(" ".join((p.extract_text() or "").split()) for p in reader.pages)
    assert "FIRST PROSPECTIVE EVALUATION" in text
    assert "S084" in reader.pages[-1].extract_text()
    assert all(" ".join(entry["title"].split()) in text for entry in entries)
    paragraphs = [p.text for p in Document(docx).paragraphs]
    assert json.loads(paragraphs[-1].split(": ", 1)[1])["strategy_id"] == "S084"
    md = report._markdown(pages, entries)
    assert md.split("## " + report.APPENDIX_TITLE)[1].count("### S") == 84
    assert md.rstrip().endswith("```")
