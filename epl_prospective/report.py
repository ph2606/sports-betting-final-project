"""Present the first prospective forecast evaluation beside frozen research.

Read finished tables only. A forecast score is never converted into a bet or
bankroll return; the historical strategy definitions remain the final appendix.
"""
from __future__ import annotations

import argparse
from datetime import datetime, timezone
import html
import json
import os
from pathlib import Path
import shutil

import numpy as np
import pandas as pd
from docx import Document
from docx.enum.section import WD_ORIENT, WD_SECTION_START
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt
from PIL import Image as PILImage
from pypdf import PdfReader
from reportlab.lib import colors
from reportlab.lib.pagesizes import landscape, letter
from reportlab.lib.styles import ParagraphStyle
from reportlab.platypus import (BaseDocTemplate, Frame, Image, KeepTogether,
    NextPageTemplate, PageBreak, PageTemplate, Paragraph, Table, TableStyle)

from epl_execution import report as execution
from epl_strategy.report import number
from epl_strategy_lab import report as lab

ROOT = Path(__file__).resolve().parents[1]
HISTORICAL = ROOT / "epl_execution/results/20261015_quote_sensitivity"
TITLE = "EPL project: first prospective evaluation and strategy research"
APPENDIX_TITLE = execution.APPENDIX_TITLE
BODY_WEIGHTS = execution.BODY_WEIGHTS
TABLE_NAMES = ("source_freshness", "forecast_outcome_audit", "model_scores",
    "pairwise_scores", "fixture_outcomes", "next_forecasts")
OPTIONAL_TABLES = ("original_forecast_coverage", "next_forecast_coverage")
GRAPH_NAMES = ("all_reference_strategies_4seasons_100k.png",
    "all_strategies_4seasons_100k.png", "all_84_strategies_2025_100k.png")


def _load(output: Path) -> dict:
    data = {name: pd.read_parquet(output / (name + ".parquet")) for name in TABLE_NAMES}
    for name in OPTIONAL_TABLES:
        if (output / (name + ".parquet")).exists():
            data[name] = pd.read_parquet(output / (name + ".parquet"))
    # Keep report vocabulary consistent while preserving original export fields
    # in the source files and workbook.
    audit = data["forecast_outcome_audit"]
    if "result" not in audit and "actual_result" in audit: audit["result"] = audit.actual_result
    if "settled_at_utc" not in audit and "outcome_observed_at_utc" in audit: audit["settled_at_utc"] = audit.outcome_observed_at_utc
    data["status"] = json.loads((output / "prospective_status.json").read_text(encoding="utf-8"))
    data["integrity"] = json.loads((output / "ledger_integrity.json").read_text(encoding="utf-8"))
    data["historical"] = execution._load(HISTORICAL)
    return data


def _validate(data: dict) -> None:
    audit, status, scores = data["forecast_outcome_audit"], data["status"], data["model_scores"]
    required = ["match_id", "model_id", "feature_group", "p_H", "p_D", "p_A", "result", "created_at_utc", "kickoff_utc"]
    if set(required) - set(audit): raise ValueError("Missing prospective audit fields")
    if len(audit) != 40 or audit.match_id.nunique() != 10 or audit.model_id.nunique() != 4:
        raise ValueError("The original release must preserve 40 forecasts of ten fixtures from four models")
    if audit.duplicated(["match_id", "model_id"]).any() or not audit.groupby("match_id").size().eq(4).all():
        raise ValueError("A released model/fixture identity is missing or duplicated")
    probabilities = audit[["p_H", "p_D", "p_A"]].to_numpy(float)
    if not np.isfinite(probabilities).all() or (probabilities < 0).any() or (probabilities > 1).any() or not np.allclose(probabilities.sum(axis=1), 1, rtol=0, atol=1e-8):
        raise ValueError("Invalid released probabilities")
    created = pd.to_datetime(audit.created_at_utc, utc=True)
    kickoff = pd.to_datetime(audit.kickoff_utc, utc=True)
    if created.isna().any() or kickoff.isna().any() or not created.lt(kickoff).all():
        raise ValueError("A scored forecast was not recorded before kickoff")
    resolved = audit.result.isin(list("HDA"))
    if int(resolved.sum()) != int(status["resolved_original_forecasts"]) or int((~resolved).sum()) != int(status["pending_original_forecasts"]):
        raise ValueError("Prospective coverage counts do not reconcile")
    if audit.loc[resolved].match_id.nunique() != int(status["resolved_original_fixtures"]):
        raise ValueError("Resolved fixture denominator is inconsistent")
    if audit.loc[resolved].groupby("match_id").result.nunique().gt(1).any():
        raise ValueError("Models disagree on an observed fixture result")
    if resolved.any():
        if "settled_at_utc" not in audit: raise ValueError("A resolved forecast requires its observed settlement clock")
        observed = pd.to_datetime(audit.loc[resolved, "settled_at_utc"], utc=True)
        if observed.isna().any() or not observed.ge(kickoff.loc[resolved]).all():
            raise ValueError("A result was observed before kickoff")
        counts = audit.loc[resolved].groupby("model_id").size()
        if scores.model_id.duplicated().any() or set(scores.model_id) != set(counts.index):
            raise ValueError("Scores omit or duplicate a resolved model")
        observed_counts = scores.set_index("model_id").n_matches.reindex(counts.index)
        if not np.array_equal(observed_counts.to_numpy(), counts.to_numpy()):
            raise ValueError("Model score denominators do not match released resolved forecasts")
    elif len(scores):
        raise ValueError("Unresolved forecasts cannot have observed scores")
    if int(status["recorded_trade_count"]) != 0 or bool(status["prospective_betting_performance_available"]):
        raise ValueError("This release has no recorded prospective betting performance")
    if len(data["fixture_outcomes"]) != 10 or data["fixture_outcomes"].match_id.duplicated().any():
        raise ValueError("The fixture outcome table must preserve all ten targets")
    if set(data["fixture_outcomes"].match_id) != set(audit.match_id):
        raise ValueError("Outcome table targets differ from the original release")
    if len(data["next_forecasts"]) != int(status["new_forecasts_released"]):
        raise ValueError("Next-release count does not reconcile")
    execution._validate(data["historical"])


def _field(row, names, default=np.nan):
    for name in names:
        if name in row and pd.notna(row[name]): return row[name]
    return default


def _label(value) -> str:
    labels = {"team": "Team", "team_only": "Team", "team_xg": "Team + xG",
        "team_player": "Team + player", "team_players": "Team + player", "combined": "Combined"}
    return labels.get(str(value), str(value))


def _models(data: dict) -> dict:
    return {str(model): _label(group.feature_group.iloc[0]) for model, group in data["forecast_outcome_audit"].groupby("model_id")}


def _score_table(data: dict) -> pd.DataFrame:
    labels = _models(data); rows = []
    for _, row in data["model_scores"].iterrows():
        low = _field(row, ("accuracy_ci_low", "accuracy_wilson_ci_low", "accuracy_wilson_low", "wilson_accuracy_ci_low"))
        high = _field(row, ("accuracy_ci_high", "accuracy_wilson_ci_high", "accuracy_wilson_high", "wilson_accuracy_ci_high"))
        rows.append({"Model": labels.get(str(row.model_id), str(row.model_id)), "Resolved\nfixtures": number(row.n_matches, 0),
            "Accuracy": number(row.accuracy, percent=True), "95% Wilson\naccuracy interval": number(low, percent=True) + " to " + number(high, percent=True),
            "Log\nloss": number(row.log_loss, 4),
            "95% log-loss\ninterval": number(_field(row, ("log_loss_ci_low",)), 4) + " to " + number(_field(row, ("log_loss_ci_high",)), 4),
            "Brier\nscore": number(row.brier_score, 4),
            "95% Brier\ninterval": number(_field(row, ("brier_score_ci_low",)), 4) + " to " + number(_field(row, ("brier_score_ci_high",)), 4),
            "Classwise\nECE 5": number(_field(row, ("classwise_ece_5",)), 4),
            "Classwise\nECE 10": number(_field(row, ("classwise_ece_10",)), 4),
            "Classwise\nECE 15": number(_field(row, ("classwise_ece_15",)), 4)})
    return pd.DataFrame(rows)


def _pairwise_table(data: dict) -> pd.DataFrame:
    frame = data["pairwise_scores"]; labels = _models(data)
    if not len(frame): return pd.DataFrame(columns=["Comparison", "Status"])
    rows = []
    for _, row in frame.iterrows():
        rows.append({"Model A": labels.get(str(row.model_a), str(row.model_a)),
            "Model B": labels.get(str(row.model_b), str(row.model_b)),
            "Common\nfixtures": number(row.common_fixture_count, 0),
            "Log loss\nA minus B": number(row.log_loss_diff, 4),
            "95% paired\nlog-loss interval": number(row.log_loss_diff_ci_low, 4) + " to " + number(row.log_loss_diff_ci_high, 4),
            "Brier\nA minus B": number(row.brier_score_diff, 4),
            "95% paired\nBrier interval": number(row.brier_score_diff_ci_low, 4) + " to " + number(row.brier_score_diff_ci_high, 4),
            "Accuracy\nA minus B": number(row.accuracy_diff, percent=True)})
    return pd.DataFrame(rows)


def _freshness_table(data: dict) -> pd.DataFrame:
    frame = data["source_freshness"]
    preferred = ("source", "source_name", "snapshot_role", "observed_at_utc", "latest_completed_match_date", "source_capture_valid", "core_source_status", "capture_status", "usable", "status")
    columns = [c for c in preferred if c in frame]
    if not columns: columns = list(frame.columns[:6])
    display = frame[columns].copy()
    for c in display: display[c] = display[c].astype(str)
    display.columns = [c.replace("_", " ").replace("utc", "UTC") for c in columns]
    return display


def _fixture_table(data: dict) -> pd.DataFrame:
    frame = data["fixture_outcomes"].copy()
    if "result" not in frame and "actual_result" in frame: frame["result"] = frame.actual_result
    columns = [c for c in ("kickoff_utc", "home_team", "away_team", "home_goals", "away_goals", "result", "status") if c in frame]
    display = frame[columns].copy()
    for c in ("home_goals", "away_goals"):
        if c in display: display[c] = display[c].map(lambda v: number(v, 0))
    if "kickoff_utc" in display: display["kickoff_utc"] = pd.to_datetime(display.kickoff_utc, utc=True).dt.strftime("%b %d %H:%M")
    display.columns = [c.replace("_", " ").replace("utc", "UTC") for c in columns]
    return display


def _body(data: dict, figures: Path) -> tuple[list[dict], dict[str, pd.DataFrame]]:
    history, historical_displays = execution._body(data["historical"], figures)
    for page_ in history:
        for block in page_["blocks"]:
            if block["type"] == "figure": block["path"] = figures / block["path"].name
    history[0]["title"] = "Historical research: frozen strategy and quote-sensitivity study"
    history[0]["blocks"].insert(0, {"type": "paragraph", "text": "Historical view frozen on October 15, 2026. The following study retains its original dates, wording and observations; any prospective pending-status statement below describes that earlier date. Current settled coverage and subsequent forecasts are reported in the opening October 20 update."})
    pages = []; displays = {}; status = data["status"]
    def page(title, orientation="portrait"):
        item = {"title": title, "orientation": orientation, "blocks": []}; pages.append(item); return item["blocks"]
    def p(blocks, text): blocks.append({"type": "paragraph", "text": text})
    def table(blocks, frame, caption): blocks.append({"type": "table", "frame": frame, "caption": caption})
    b = page(TITLE)
    p(b, "DRAFT FOR REVIEW | October 20, 2026 | EPL component of the statistical-arbitrage final project.")
    p(b, "This update follows the forecasts recorded before the October 17–19 fixtures. It joins their frozen probabilities to newly observed finished results, reports the first prospective prediction scores, and documents source freshness and the next release. The complete historical strategy study and its 100,000-unit account graphs remain later in this report. The final appendix still explains every one of the 84 strategy configurations.")
    coverage = pd.DataFrame([
        ("Original released probability records", status.get("original_forecast_count", 40)),
        ("Distinct original fixtures", status.get("original_fixture_count", 10)),
        ("Resolved original probability records", status["resolved_original_forecasts"]),
        ("Resolved distinct original fixtures", status["resolved_original_fixtures"]),
        ("Pending original probability records", status["pending_original_forecasts"]),
        ("New probability records released", status["new_forecasts_released"]),
        ("New target fixtures", status["new_target_fixtures"]),
        ("Recorded prospective trades", status["recorded_trade_count"])], columns=["Coverage item", "Recorded count"])
    displays["Prospective coverage"] = coverage
    table(b, coverage, "Table P1. Evaluation as of " + str(status["as_of_utc"]) + ". Four model forecasts of the same game are correlated records, not four independent matches.")
    p(b, "The first release had no valid observed pre-match odds for its targets and did not record a betting policy or stakes. Prospective betting hit ratio, expected trading gain/loss, ROI, profit, drawdown and account wealth are unavailable. Zero recorded trades does not establish zero error or zero profit for an unobserved betting strategy. The historical money curves describe their separately labelled retrospective account replays.")
    b = page("First prospective prediction scores and their limits", "landscape")
    p(b, "Each model is scored only on its earliest valid pre-kickoff forecast and an identity-matched finished source result. Accuracy is the share whose maximum-probability H/D/A class was correct, with H/D/A tie order inherited from the frozen scorer. Lower log loss and Brier score are better. These are prediction scores, not a trading hit ratio or a betting return.")
    score = _score_table(data); displays["Prospective model scores"] = score
    if len(score): table(b, score, "Table P2. One row per released model. Wilson intervals summarize binomial accuracy uncertainty; their sampling assumptions are approximate for this small, dependent football sample. Ten fixtures cannot establish model superiority or stable calibration.")
    else: p(b, "No original forecast has an observed finished result. Accuracy, loss and calibration scores remain unavailable.")
    paired = _pairwise_table(data); displays["Paired score comparisons"] = paired
    if len(paired): table(b, paired, "Table P3. Every difference is model A minus model B on the same observed fixtures; negative loss differences favor A. Paired bootstrap intervals condition on this recorded batch and do not adjust for all comparisons or the wider research search. Full accuracy intervals and bootstrap settings are in pairwise_scores.")
    p(b, "The four alternatives share training history, fixture outcomes and much of their feature information. A larger point estimate is not independent confirmation. Classwise ECE depends strongly on bin choices when the sample is small; all declared 5-, 10- and 15-bin versions are retained. No candidate is selected, refitted or tuned using these results, and this first batch is not a validation of the retrospectively designed 84-rule trading catalog.")
    b = page("Every original target and observed result", "landscape")
    fixtures = _fixture_table(data); displays["Original target outcomes"] = fixtures
    table(b, fixtures, "Table P4. The ten original target identities remain in view, including pending records if any. Goals/result come from an observed finished fixture; elapsed kickoff alone never settles a prediction. H=home win, D=draw, A=away win. Kickoff is UTC.")
    p(b, "The prospective workbook preserves all 40 released H/D/A probability triplets, exact creation and kickoff clocks, predicted classes, per-forecast correctness, log loss and Brier score. Its fixture table provides the ten-match denominator. Settlement observation clocks describe when finished results were captured, rather than an inferred final-whistle time. Forecast and outcome ledgers remain separate append-only records.")
    p(b, "Ledger-chain checks detect internal edits, gaps and reordering relative to saved heads, while preserved repository versions provide additional release provenance. A locally recomputed chain is not independent proof of publication time. The integrity table and source manifest retain the exact paths, hashes and checks used by this update; no earlier forecast probability or settlement is replaced.")
    b = page("Source freshness and the next prospective release", "landscape")
    freshness = _freshness_table(data); displays["Source freshness"] = freshness
    table(b, freshness, "Table P5. Recent retrieval and fresh underlying results are separate facts. A capture's observation time is not the provider's publication clock; unresolved scheduled fixtures stay pending until a consistent finished result is observed.")
    p(b, "Results snapshot: " + str(status["results_snapshot"]) + ". Player-complete capture: " + str(status["player_snapshot"]) + ". Evaluation: " + str(status["evaluation_path"]) + ". Next-release run: " + str(status["next_run_path"]) + ". Exact local paths and capture metadata are in the workbook and source_provenance.json.")
    p(b, "The next release contains " + number(status["new_forecasts_released"], 0) + " probability records for " + number(status["new_target_fixtures"], 0) + " future targets. It uses the same frozen model bundles and the existing source/feature/quote gates. New forecasts are listed separately in next_forecasts and are not inserted into this completed batch's score denominator. Missing or stale quote inputs remain exclusions, without invented market forecasts or betting returns.")
    if "fpl_player_summaries_requested" in status:
        p(b, "The player-complete refresh captured " + number(status["fpl_player_summaries_downloaded"], 0) + " of " + number(status["fpl_player_summaries_requested"], 0) + " requested player summaries. Original and next-release coverage tables retain every released or withheld model/fixture variant and its stated exclusion reason in the workbook.")
    p(b, "The remainder preserves the full historical study: all 84 common-season configurations, all 56 compatible continuous four-season accounts, annual stability, selection diagnostics, quote sensitivity and complete trading metrics. These are retrospective results and retain their original dates and limitations. The 100,000 starts are hypothetical historical account starts; they do not describe prospective executed funds.")
    p(b, "Reproduce this combined report with python -m epl_prospective.report --results PATH. prospective_overview.xlsx contains the complete new record-level audit, fixture/model/pairwise scores, source freshness, next release and integrity records, plus a link to the frozen execution_overview.xlsx. The exact original 84-strategy appendix is the final section.")
    return pages + history, {**historical_displays, **displays}


def _markdown(pages, entries) -> str:
    return execution._markdown(pages, entries)


def _workbook(path: Path, data: dict, entries: list[dict], registry: dict) -> dict:
    readme = pd.DataFrame([
        ("Unit of observation", "40 original forecasts = four correlated model forecasts of ten fixtures. Model score denominators count resolved fixtures, not all 40 as independent games."),
        ("Trading performance", "No recorded policy/stakes or valid original pre-match quote ledger. Prospective betting returns, expected gain/loss and wealth are unavailable."),
        ("Interpretation", "First small prospective prediction batch. No model selection, tuning or superiority conclusion; ECE and intervals are uncertain."),
        ("Freshness", "Observation time measures retrieval, not publication. Latest completed match dates measure source content. Calendar passage never supplies outcomes."),
        ("Historical metrics workbook", lab._relative(HISTORICAL / "report/execution_overview.xlsx")),
        ("Historical graphs", "All 84 original common-season and both 56-account four-season 100000 graphs remain in the combined report. These are retrospective."),
        ("Original definitions", "All 84 strategy configurations remain exact and their explanations are the final report appendix.")], columns=["Item", "Definition"])
    sheets = {"Read me": readme,
        "Prospective status": pd.DataFrame([{"item": k, "value": lab._json(v) if isinstance(v, (dict, list)) else v} for k, v in data["status"].items()]),
        "Source freshness": data["source_freshness"], "Every original forecast": data["forecast_outcome_audit"],
        "Model scores": data["model_scores"], "Paired score comparisons": data["pairwise_scores"],
        "Original fixture outcomes": data["fixture_outcomes"], "Next release forecasts": data["next_forecasts"],
        "Ledger integrity": pd.DataFrame([{"item": k, "value": lab._json(v) if isinstance(v, (dict, list)) else v} for k, v in data["integrity"].items()]),
        "Complete84 catalog": pd.DataFrame(data["historical"]["catalog"]), "Final84 appendix": pd.DataFrame(entries),
        "Input provenance": pd.DataFrame([{"path": p, **item} for p, item in registry.items()])}
    for key, sheet_name in (("original_forecast_coverage", "Original release exclusions"),
            ("next_forecast_coverage", "Next release exclusions")):
        if key in data: sheets[sheet_name] = data[key]
    percentages = {"accuracy", "accuracy_ci_low", "accuracy_ci_high", "accuracy_wilson_ci_low", "accuracy_wilson_ci_high",
        "accuracy_wilson_low", "accuracy_wilson_high", "p_H", "p_D", "p_A"}
    with pd.ExcelWriter(path, engine="xlsxwriter", engine_kwargs={"options": {"strings_to_urls": False}}) as writer:
        book = writer.book; head = book.add_format({"bold": True, "font_color": "white", "bg_color": "#12364D", "text_wrap": True})
        num = book.add_format({"num_format": "#,##0.0000"}); pct = book.add_format({"num_format": "0.00%"})
        for name, frame in sheets.items():
            safe = frame.copy()
            for column in safe:
                if isinstance(safe[column].dtype, pd.DatetimeTZDtype): safe[column] = safe[column].dt.tz_convert(None)
                if safe[column].dtype == "object": safe[column] = safe[column].map(lambda v: lab._json(v.tolist() if isinstance(v, np.ndarray) else v) if isinstance(v, (dict, list, tuple, np.ndarray)) else v)
                elif pd.api.types.is_numeric_dtype(safe[column]) and not pd.api.types.is_bool_dtype(safe[column]):
                    safe[column] = safe[column].map(lambda v: "+∞" if pd.notna(v) and v == float("inf") else "−∞" if pd.notna(v) and v == float("-inf") else v)
            safe.to_excel(writer, sheet_name=name, index=False, na_rep="—")
            sheet = writer.sheets[name]; sheet.write_row(0, 0, list(safe.columns), head); sheet.set_row(0, 42); sheet.freeze_panes(1, 1)
            if name == "Read me":
                row_index = 1 + list(readme.Item).index("Historical metrics workbook")
                relative = os.path.relpath(HISTORICAL / "report/execution_overview.xlsx", path.parent).replace("\\", "/")
                sheet.write_url(row_index, 1, "external:" + relative, string="Open the frozen historical metrics workbook (repository files required)")
            if len(safe.columns): sheet.autofilter(0, 0, len(safe), len(safe.columns) - 1)
            for index, column in enumerate(safe.columns):
                width = 65 if column in ("Definition", "explanation", "exact_config", "path", "coverage", "value") else min(42, max(14, len(column) * .8))
                sheet.set_column(index, index, width, pct if column in percentages else num)
    return {name: len(frame) for name, frame in sheets.items()}



# Layout adapted from the preserved strategy report renderer.
def _pdf(pages, entries, path: Path) -> None:
    fonts = lab._fonts()
    styles = {
        "title": ParagraphStyle("ValTitle", fontName=fonts["bold"], fontSize=14, leading=16, spaceAfter=7, textColor=colors.HexColor("#12364D")),
        "body": ParagraphStyle("ValBody", fontName=fonts["body"], fontSize=10.2, leading=12.4, spaceAfter=7),
        "caption": ParagraphStyle("ValCaption", fontName=fonts["body"], fontSize=8.5, leading=10, spaceAfter=4),
        "cell": ParagraphStyle("ValCell", fontName=fonts["sans"], fontSize=8.0, leading=9.3),
        "head": ParagraphStyle("ValHead", fontName=fonts["bold"], fontSize=8.0, leading=9.3, textColor=colors.white),
        "entrytitle": ParagraphStyle("ValEntryTitle", fontName=fonts["bold"], fontSize=9.7, leading=11.4, spaceBefore=5, spaceAfter=4),
        "entry": ParagraphStyle("ValEntry", fontName=fonts["body"], fontSize=9.2, leading=10.6, spaceAfter=4),
        "config": ParagraphStyle("ValConfig", fontName=fonts["mono"], fontSize=6.8, leading=8.1, spaceAfter=7, splitLongWords=True, wordWrap="CJK")}
    def header(canvas, doc):
        width, height = canvas._pagesize
        canvas.setFont(fonts["sans"], 8); canvas.setFillColor(colors.HexColor("#52616B"))
        canvas.drawString(36, height - 21, "EPL PROJECT - FIRST PROSPECTIVE EVALUATION | October 20, 2026")
        canvas.drawRightString(width - 36, 21, str(doc.page))
    doc = BaseDocTemplate(str(path), pagesize=letter, title=TITLE, author="Aman Dhillon and Panagiotis Housos")
    for orientation, size in (("portrait", letter), ("landscape", landscape(letter))):
        width, height = size
        doc.addPageTemplates(PageTemplate(id=orientation, pagesize=size, onPage=header,
            frames=[Frame(36, 36, width - 72, height - 72, leftPadding=0, rightPadding=0, topPadding=0, bottomPadding=0)]))
    story = []
    for index, page in enumerate(pages):
        if index: story.extend([NextPageTemplate(page["orientation"]), PageBreak()])
        width = (letter if page["orientation"] == "portrait" else landscape(letter))[0] - 72
        story.append(Paragraph(html.escape(page["title"]), styles["title"]))
        for block in page["blocks"]:
            if block["type"] == "paragraph": story.append(Paragraph(html.escape(block["text"]), styles["body"]))
            elif block["type"] == "figure":
                with PILImage.open(block["path"]) as picture: ratio = picture.height / picture.width
                story.extend([Image(str(block["path"]), width=width - 10, height=(width - 10) * ratio), Paragraph(html.escape(block["caption"]), styles["caption"])])
            else:
                story.append(Paragraph(html.escape(block["caption"]), styles["caption"]))
                frame = block["frame"]
                raw = [list(frame.columns)] + frame.values.tolist()
                cells = [[Paragraph(html.escape(str(value)).replace("\n", "<br/>"), styles["head" if row_index == 0 else "cell"]) for value in row] for row_index, row in enumerate(raw)]
                weights = BODY_WEIGHTS if page["title"].startswith(("Four-season trading", "Original 2025")) else [1 / len(frame.columns)] * len(frame.columns)
                table = Table(cells, colWidths=[width * weight for weight in weights], repeatRows=1)
                table.setStyle(TableStyle([("BACKGROUND", (0,0), (-1,0), colors.HexColor("#12364D")), ("ROWBACKGROUNDS", (0,1), (-1,-1), [colors.white, colors.HexColor("#F0F4F6")]), ("VALIGN", (0,0), (-1,-1), "MIDDLE"), ("LEFTPADDING", (0,0), (-1,-1), 3), ("RIGHTPADDING", (0,0), (-1,-1), 3), ("TOPPADDING", (0,0), (-1,-1), 1.1), ("BOTTOMPADDING", (0,0), (-1,-1), 1.1)]))
                story.append(table)
    story.extend([NextPageTemplate("portrait"), PageBreak(), Paragraph(APPENDIX_TITLE, styles["title"])])
    for entry in entries:
        story.append(KeepTogether([Paragraph(html.escape(entry["title"]), styles["entrytitle"]),
            Paragraph(html.escape("Origin: " + entry["origin"] + ". " + entry["explanation"]), styles["entry"]),
            Paragraph(html.escape("Exact original configuration: " + entry["exact_config"]), styles["config"])]))
    doc.build(story)

def _docx(pages, entries, path: Path) -> None:
    doc = Document()
    normal = doc.styles["Normal"]; normal.font.name = "Times New Roman"; normal.font.size = Pt(10.2)
    normal.paragraph_format.line_spacing = Pt(12.4); normal.paragraph_format.space_after = Pt(7)
    heading = doc.styles["Heading 1"]; heading.font.name = "Arial"; heading.font.size = Pt(14)
    heading.paragraph_format.space_before = Pt(0); heading.paragraph_format.space_after = Pt(7)
    doc.core_properties.title = TITLE; doc.core_properties.author = "Aman Dhillon and Panagiotis Housos"
    def section(index, wide):
        sec = doc.sections[0] if index == 0 else doc.add_section(WD_SECTION_START.NEW_PAGE)
        sec.orientation = WD_ORIENT.LANDSCAPE if wide else WD_ORIENT.PORTRAIT
        sec.page_width = Inches(11 if wide else 8.5); sec.page_height = Inches(8.5 if wide else 11)
        sec.top_margin = sec.bottom_margin = sec.left_margin = sec.right_margin = Inches(.5)
        sec.header_distance = sec.footer_distance = Inches(.2)
        if index == 0:
            run = sec.header.paragraphs[0].add_run("EPL PROJECT - FIRST PROSPECTIVE EVALUATION | October 20, 2026")
            run.font.name = "Arial"; run.font.size = Pt(8)
            footer = sec.footer.paragraphs[0]; footer.alignment = 2
            field = OxmlElement("w:fldSimple"); field.set(qn("w:instr"), "PAGE"); footer._p.append(field)
    for index, page in enumerate(pages):
        wide = page["orientation"] == "landscape"; section(index, wide); doc.add_heading(page["title"], 1)
        for block in page["blocks"]:
            if block["type"] == "paragraph": doc.add_paragraph(block["text"]); continue
            if block["type"] == "figure": doc.add_picture(str(block["path"]), width=Inches(9.85 if wide else 7.35))
            caption = doc.add_paragraph(block["caption"]); caption.paragraph_format.line_spacing = Pt(10)
            for run in caption.runs: run.font.size = Pt(8.5)
            if block["type"] == "figure": continue
            frame = block["frame"]; table = doc.add_table(rows=1, cols=len(frame.columns)); table.style = "Light Shading Accent 1"; table.autofit = False
            weights = BODY_WEIGHTS if page["title"].startswith(("Four-season trading", "Original 2025")) else [1 / len(frame.columns)] * len(frame.columns)
            width = 10 if wide else 7.5
            for col, weight in zip(table.columns, weights): col.width = Inches(width * weight)
            for cell, value in zip(table.rows[0].cells, frame.columns): cell.text = str(value)
            for values in frame.itertuples(index=False, name=None):
                for cell, value in zip(table.add_row().cells, values): cell.text = str(value)
            for row in table.rows:
                for cell, weight in zip(row.cells, weights):
                    cell.width = Inches(width * weight)
                    for paragraph in cell.paragraphs:
                        paragraph.paragraph_format.line_spacing = Pt(9.3); paragraph.paragraph_format.space_after = Pt(0)
                        for run in paragraph.runs: run.font.name = "Arial"; run.font.size = Pt(8)
    section(len(pages), False); doc.add_heading(APPENDIX_TITLE, 1)
    for entry in entries:
        title = doc.add_paragraph(); title.paragraph_format.keep_with_next = True
        run = title.add_run(entry["title"]); run.bold = True; run.font.name = "Arial"; run.font.size = Pt(9.7)
        paragraph = doc.add_paragraph("Origin: " + entry["origin"] + ". " + entry["explanation"])
        paragraph.paragraph_format.keep_with_next = True; paragraph.paragraph_format.line_spacing = Pt(10.6)
        for run in paragraph.runs: run.font.size = Pt(9.2)
        paragraph = doc.add_paragraph("Exact original configuration: " + entry["exact_config"])
        paragraph.paragraph_format.keep_together = True; paragraph.paragraph_format.line_spacing = Pt(8.1)
        for run in paragraph.runs: run.font.name = "Consolas"; run.font.size = Pt(6.8)
    doc.save(path)

def build_report(output_dir: Path) -> dict:
    output = Path(output_dir).resolve(); data = _load(output); _validate(data)
    entries = execution._appendix(data["historical"]["catalog"])
    sources = [output / (name + ".parquet") for name in TABLE_NAMES]
    sources += [output / (name + ".parquet") for name in OPTIONAL_TABLES if (output / (name + ".parquet")).exists()]
    sources += [output / name for name in ("prospective_status.json", "ledger_integrity.json", "source_provenance.json")]
    # Preserve the exact transitive source inventory of the historical report.
    historical_manifest_path = HISTORICAL / "report/report_manifest.json"
    historical_manifest = json.loads(historical_manifest_path.read_text(encoding="utf-8"))
    sources += [ROOT / p for p in historical_manifest["inputs"]]
    sources += [historical_manifest_path, HISTORICAL / "report/execution_overview.xlsx"]
    sources += [HISTORICAL / "figures" / name for name in GRAPH_NAMES]
    sources += [Path(__file__).resolve(), Path(execution.__file__).resolve(), Path(lab.__file__).resolve()]
    for name in ("epl_prospective/PROTOCOL.md", "epl_live/OPERATIONAL_PROTOCOL.md"):
        if (ROOT / name).exists(): sources.append(ROOT / name)
    sources = sorted(set(sources))
    registry = {lab._relative(p): {"sha256": lab._sha(p), "bytes": p.stat().st_size} for p in sources}
    report = output / "report"; report.mkdir(exist_ok=True); figures = output / "figures"; figures.mkdir(exist_ok=True)
    for name in GRAPH_NAMES: shutil.copyfile(HISTORICAL / "figures" / name, figures / name)
    pages, displays = _body(data, figures)
    for name, frame in displays.items(): frame.to_csv(report / (name.lower().replace(" ", "_").replace("/", "_") + ".csv"), index=False)
    pd.DataFrame(entries).to_csv(report / "complete84_strategy_appendix.csv", index=False)
    md = report / "validation_report.md"; pdf = md.with_suffix(".pdf"); docx = md.with_suffix(".docx")
    md.write_text(_markdown(pages, entries), encoding="utf-8"); _pdf(pages, entries, pdf); _docx(pages, entries, docx)
    sheets = _workbook(report / "prospective_overview.xlsx", data, entries, registry)
    reader = PdfReader(pdf); text = "\n".join(p.extract_text() or "" for p in reader.pages)
    if APPENDIX_TITLE not in text: raise AssertionError("Combined report omitted its final appendix")
    final_text = " ".join(text.split(APPENDIX_TITLE, 1)[1].split())
    if any(" ".join(e["title"].split()) not in final_text for e in entries): raise AssertionError("A strategy explanation is missing")
    if "S084" not in reader.pages[-1].extract_text(): raise AssertionError("The full strategy appendix is not last")
    changed = [lab._relative(p) for p in sources if lab._sha(p) != registry[lab._relative(p)]["sha256"]]
    if changed: raise AssertionError("Report inputs changed during rendering: " + str(changed))
    preserved = {name: lab._sha(figures / name) == lab._sha(HISTORICAL / "figures" / name) for name in GRAPH_NAMES}
    if not all(preserved.values()): raise AssertionError("A historical graph was altered")
    produced = [p for folder in (report, figures) for p in folder.rglob("*") if p.is_file() and p.name != "report_manifest.json" and "previews" not in p.parts]
    manifest = {"status": "first_prospective_prediction_evaluation_with_frozen_strategy_research", "draft_date_new_york": "2026-10-20",
        "built_at_utc": datetime.now(timezone.utc).isoformat(), "pdf_pages": len(reader.pages), "body_sections": len(pages),
        "original_forecast_count": 40, "original_fixture_count": 10, "model_count": 4,
        "resolved_original_forecasts": int(data["status"]["resolved_original_forecasts"]),
        "recorded_trade_count": 0, "prospective_betting_performance_available": False,
        "historical_catalog_count": 84, "preserved_historical_graphs": preserved,
        "appendix_ids": [e["strategy_id"] for e in entries], "appendix_is_last": True,
        "workbook_sheet_rows": sheets, "consumed_source_files_unchanged": True, "inputs": registry,
        "outputs": {p.relative_to(output).as_posix(): {"sha256": lab._sha(p), "bytes": p.stat().st_size} for p in sorted(produced)},
        "manifest_order": "Final builder envelope and independent audit are deliberately not consumed report inputs.",
        "limits": "Ten distinct games with correlated model forecasts; uncertain calibration and pairwise intervals; no recorded prospective bets or trading returns; historical accounts remain retrospective."}
    (report / "report_manifest.json").write_text(json.dumps(manifest, ensure_ascii=False, indent=2, allow_nan=False), encoding="utf-8")
    return {"output_directory": str(report), "pdf_pages": len(reader.pages), "appendix_entries": 84,
        "historical_graphs_preserved": 3, "resolved_original_forecasts": manifest["resolved_original_forecasts"],
        "recorded_trade_count": 0, "source_files_unchanged": True, "workbook_bytes": (report / "prospective_overview.xlsx").stat().st_size}


def main() -> None:
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument("--results", required=True, type=Path)
    print(json.dumps(build_report(parser.parse_args().results), indent=2))


if __name__ == "__main__": main()
