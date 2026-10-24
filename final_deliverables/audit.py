"""Independent, read-only audit of the assembled NBA/EPL final artifacts."""
from __future__ import annotations

import argparse
from collections import Counter
import hashlib
import json
from pathlib import Path
import re

import fitz
import numpy as np
import pandas as pd
from pptx import Presentation

ROOT = Path(__file__).resolve().parents[1]
EPL = ROOT / "epl_player_models/results/20261020_player_models_review"
LAB = ROOT / "epl_strategy_lab/results/20261015T040052Z_100k84"
VALIDATION = ROOT / "epl_strategy_validation/results/20261015_temporal_selection"
NAMES = ("Aman Dhillon", "Panagiotis Housos", "Daniel Totouom-Tangho")


def require(condition, message):
    if not condition:
        raise AssertionError(message)


def read_json(path):
    return json.loads(Path(path).read_text(encoding="utf-8"))


def sha(path):
    return hashlib.sha256(Path(path).read_bytes()).hexdigest()


def norm(value):
    # Cambria/XeTeX can extract a printed hyphen as U+2011. This is the
    # same visible spelling, rather than a different supervisor name.
    value = str(value).translate({ord(c): "-" for c in "‐‑‒–—−"}).replace("\\%", "%")
    return re.sub(r"\s+", " ", value).strip()


def numeric(value):
    return re.sub(r"\s+", "", norm(value)).replace(",", "").replace("$", "").replace("+", "")


def near(actual, expected, label, atol=1e-10):
    require(np.isclose(float(actual), float(expected), atol=atol, rtol=1e-10),
        f"Incorrect {label}: {actual!r} versus source {expected!r}")


def formatted_present(row, value, kind="pct", label="table value"):
    if kind == "pct":
        expected = f"{float(value):.2%}"
    elif kind == "money2":
        expected = f"{float(value):,.2f}"
    elif kind == "money0":
        expected = f"{float(value):,.0f}"
    elif kind == "four":
        expected = f"{float(value):.4f}"
    else:
        expected = str(value)
    require(any(numeric(expected) in numeric(cell) for cell in row),
        f"Missing/incorrect {label}: expected {expected!r} in {row!r}")


def visible(spec):
    values = [spec["title"], spec.get("subtitle", ""), spec.get("kicker", "")]
    for element in spec.get("elements", []):
        kind = element.get("type")
        if kind == "text":
            values.append(element.get("text", ""))
        elif kind == "card":
            values.extend(element.get(key, "") for key in ("label", "value", "detail"))
        elif kind == "table":
            values.extend(element.get("headers", []))
            values.extend(value for row in element.get("rows", []) for value in row)
    return "\n".join(map(str, values))


def notebook_text(cell):
    return "\n".join("".join(o.get("text", [])) + "\n" + "".join(o.get("data", {}).get("text/plain", []))
        for o in cell.get("outputs", []))


def audit_source_evidence(nba, epl):
    notebook = read_json(ROOT / "ML Sports/Data.ipynb")
    original = notebook_text(notebook["cells"][19])
    pattern = r"(Calibration|Accuracy)\s+(LR|SVM)\s+(fixed|kelly)\s+final\s+\$\s*([\d,.]+)\s+ROI\s+([+-]?[\d.]+)%\s+bet\s+([\d.]+)%\s+won\s+([\d.]+)%"
    parsed = re.findall(pattern, original)
    require(len(parsed) == 4, "Cannot independently read four original NBA saved outcomes")
    saved = {(r["branch"], r["model"], r["staking"]): r for r in nba["betting_2018"]["records"]}
    for branch, model, staking, final, capital_return, frequency, hit in parsed:
        key = branch, model, staking
        require(key in saved, "A supplied NBA saved strategy was omitted")
        r = saved[key]
        near(r["final_capital_saved_rounded"], float(final.replace(",", "")), f"NBA {key} ending")
        near(r["capital_return_saved_rounded"], float(capital_return) / 100, f"NBA {key} capital return")
        near(r["trading_frequency_saved_rounded"], float(frequency) / 100, f"NBA {key} frequency")
        near(r["hit_ratio_saved_rounded"], float(hit) / 100, f"NBA {key} hit ratio")
        require(r["initial_capital"] == 10000, "NBA native initial capital changed")
    selection_text = notebook_text(notebook["cells"][16])
    for r in nba["training"]["model_selection_2017"]:
        for tag, key in (("B", "branch_B_ece"), ("C", "branch_C_accuracy")):
            match = re.search(rf"{tag}\s+{r['family']}\s+([\d.]+)", selection_text)
            require(match is not None, "Missing original NBA family-selection output")
            near(r[key], match.group(1), f"NBA {r['family']} {tag} separate selection metric")
    annual_text = notebook_text(notebook["cells"][23])
    annual_pattern = r"(20\d\d-\d\d)\s+n=\s*(\d+)\s+\|B\|=(\d+)\s+\|C\|=(\d+)\s+cal\s+([+-]?[\d.]+)%\s+acc\s+([+-]?[\d.]+)%"
    annual = re.findall(annual_pattern, annual_text)
    require(len(annual) == 9, "Cannot independently read nine saved NBA annual returns")
    annual_map = {r["season"]: r for r in nba["rolling_annual"]["records"]}
    for season, games, b, c, cal, acc in annual:
        row = annual_map[season]
        require(row["n_games"] == int(games), "NBA annual game count differs")
        near(row["calibration_return_on_initial_capital"], float(cal) / 100, "NBA annual calibration return")
        near(row["accuracy_return_on_initial_capital"], float(acc) / 100, "NBA annual accuracy return")
    require(nba["rolling_annual"]["calibration_wins_seasons"] == 3, "NBA annual winner count changed")
    require("logistic" in nba["rolling_annual"]["model_families"].lower(), "Annual extension family mislabeled")
    require(nba["ece_definition"]["penalty_value"] == 1.0 and nba["ece_definition"]["guard_minimum_nonempty_fraction"] == .8,
        "NBA occupancy penalty definition changed")
    require(nba["replay"]["all_four_saved_finals_match_to_cent"] is False,
        "Unreconciled NBA SVM was represented as fully reproduced")
    require(nba["replay"]["cent_matched_branches"] == ["Calibration LR"] and nba["replay"]["mismatched_branches"] == ["Accuracy SVM"],
        "NBA replay evidence classes changed")
    epl_scores = pd.read_csv(ROOT / "epl_modeling/results/classification_scores.csv")
    epl_scores = epl_scores.loc[epl_scores.split.eq("test") & epl_scores.scope.eq("matched_bookmaker_coverage")].set_index("model")
    for row in epl["first_experiment"]["test_probability_scores"]:
        source = epl_scores.loc[row["model"]]
        require(row["n_matches"] == 374, "Initial EPL comparison population changed")
        for key in ("accuracy", "classwise_ece_10", "log_loss", "brier_score"):
            near(row[key], source[key], "EPL initial " + key)
    primary = [row for row in epl["first_experiment"]["test_betting_native1000initial10flat"] if row["policy"] == "validation_selected"]
    require(len(primary) == 8 and all(row["bet_count"] == 0 and row["turnover"] == 0 and row["roi"] is None for row in primary),
        "EPL primary no-trade decision or undefined ROI changed")
    return {"nba_saved_capital_results": 4, "nba_separate_family_selection_rows": 4,
        "nba_annual_saved_rows": 9, "nba_lr_cent_matched": True, "nba_svm_reproduction_limit_retained": True,
        "epl_locked_score_rows": 4, "epl_primary_test_cash_variants": 8}


def table_rows(tables, key):
    require(key in tables and "headers" in tables[key] and "rows" in tables[key], "Missing numerical display table: " + key)
    return tables[key]["rows"]


def audit_tables(tables, nba, epl):
    capital = table_rows(tables, "nba_capital")
    require(len(capital) == 4, "NBA original capital table must contain four supplied strategies")
    for row, expected in zip(capital, nba["betting_2018"]["records"]):
        require(expected["branch"].lower() in " ".join(map(str, row)).lower(), "NBA capital row branch changed")
        formatted_present(row, expected["final_capital_saved_rounded"], "money2", "NBA original ending capital")
        formatted_present(row, expected["capital_return_saved_rounded"], "pct", "NBA capital return")
    selection = table_rows(tables, "nba_selection")
    require(len(selection) == 4, "NBA family-selection table omitted a model")
    for row, expected in zip(selection, nba["training"]["model_selection_2017"]):
        formatted_present(row, expected["branch_B_ece"], "pct", "NBA B calibration criterion")
        formatted_present(row, expected["branch_C_accuracy"], "pct", "NBA C accuracy criterion")
    verified = table_rows(tables, "nba_verified_scores")
    require(len(verified) == 2, "Only reconciled LR and direct market scores belong in verified NBA table")
    for row, expected in zip(verified, [r for r in nba["replay"]["classification_metrics"] if r["series"] != "Accuracy"]):
        formatted_present(row, expected["accuracy"], "pct", "Verified NBA accuracy")
        formatted_present(row, expected["log_loss"], "four", "Verified NBA log loss")
        formatted_present(row, expected["brier_binary"], "four", "Verified NBA binary Brier")
    verified_trades = table_rows(tables, "nba_verified_trades")
    require(len(verified_trades) == 2, "Only two cent-matched logistic ledgers support extra NBA trade statistics")
    for row, expected in zip(verified_trades, nba["replay"]["verified_calibration_lr_metrics"]):
        formatted_present(row, expected["n_bets"], "money0", "Verified NBA bet count")
        for key in ("hit_ratio", "trading_frequency", "roi_on_stakes", "max_drawdown"):
            formatted_present(row, expected[key], "pct", "Verified NBA " + key)
        for key in ("expected_gain_per_trade", "expected_loss_per_trade"):
            formatted_present(row, expected[key], "money2", "Verified NBA " + key)
    annual = table_rows(tables, "nba_annual")
    require(len(annual) == 9, "Annual NBA source table must contain nine independent seasons")
    for row, expected in zip(annual, nba["rolling_annual"]["records"]):
        require(expected["season"].replace("-", "/") in norm(" ".join(map(str, row))) or expected["season"] in norm(" ".join(map(str, row))),
            "NBA annual season order changed")
        formatted_present(row, expected["calibration_return_on_initial_capital"], "pct", "NBA annual calibration capital return")
        formatted_present(row, expected["accuracy_return_on_initial_capital"], "pct", "NBA annual accuracy capital return")
    epl_primary = table_rows(tables, "epl_primary_scores")
    require(len(epl_primary) == 4, "EPL first-stage score table omitted benchmark or learned model")
    for row, expected in zip(epl_primary, epl["first_experiment"]["test_probability_scores"]):
        formatted_present(row, expected["accuracy"], "pct", "EPL locked accuracy")
        formatted_present(row, expected["classwise_ece_10"], "pct", "EPL locked ECE")
        formatted_present(row, expected["log_loss"], "four", "EPL locked log loss")
    secondary = table_rows(tables, "epl_secondary")
    native_secondary = [r for r in epl["first_experiment"]["test_betting_native1000initial10flat"] if r["threshold"] == .05]
    require(len(secondary) == len(native_secondary) == 8, "EPL secondary comparison omitted a price/stake branch")
    for row, expected in zip(secondary, native_secondary):
        formatted_present(row, expected["bet_count"], "money0", "EPL secondary bet count")
        formatted_present(row, expected["profit"], "money2", "EPL secondary native-unit profit")
        for key in ("roi", "roi_ci_low", "roi_ci_high"):
            formatted_present(row, expected[key], "pct", "EPL secondary " + key)
    rolling = table_rows(tables, "epl_rolling")
    require(len(rolling) == 5, "EPL rolling comparison omitted a selected predictor")
    profiles = epl["rolling_and_player_extensions"]["main_model_profiles"]
    for row, key in zip(rolling, ("strongest_benchmark", "best_learned", "team", "player", "simplest")):
        expected = profiles[key]["metrics"]
        formatted_present(row, expected["accuracy"], "pct", "EPL rolling accuracy")
        formatted_present(row, expected["log_loss"], "four", "EPL rolling log loss")
        formatted_present(row, expected["brier_score"], "four", "EPL rolling Brier")
    representative = table_rows(tables, "epl_representative")
    require(len(representative) == 6, "EPL representative account table omitted a role")
    continuous = pd.concat([pd.read_csv(VALIDATION / "metrics_continuous.csv"), pd.read_csv(EPL / "metrics_continuous.csv")], ignore_index=True).set_index("strategy_id")
    for row in representative:
        source = continuous.loc[str(row[0])]
        for key in ("hit_ratio", "trading_frequency", "roi", "max_drawdown"):
            formatted_present(row, source[key], "pct", "EPL four-year " + key)
        for key in ("expected_gain_per_trade_units", "expected_loss_per_trade_units", "final_capital_sum_units"):
            formatted_present(row, source[key], "money2", "EPL four-year " + key)
    all_metrics = table_rows(tables, "epl_all_metrics")
    require(len(all_metrics) == 116, "Complete EPL metrics appendix must retain all 116 accounts")
    canonical = pd.concat([pd.read_csv(LAB / "metrics.csv"), pd.read_csv(EPL / "metrics_single.csv")], ignore_index=True).set_index("strategy_id")
    ids = [str(row[0]) for row in all_metrics]
    require(Counter(ids) == Counter(canonical.index), "EPL common-season metric IDs missing or duplicated")
    for row in all_metrics:
        source = canonical.loc[str(row[0])]
        formatted_present(row, source.final_capital_sum_units, "money2", "EPL final account")
        if pd.notna(source.hit_ratio):
            formatted_present(row, source.hit_ratio, "pct", "EPL trading hit ratio")
        formatted_present(row, source.trading_frequency, "pct", "EPL trade frequency")
        if pd.notna(source.roi):
            formatted_present(row, source.roi, "pct", "EPL stake ROI")
        formatted_present(row, source.max_drawdown, "pct", "EPL drawdown")
        for key in ("expected_gain_per_trade_units", "expected_loss_per_trade_units"):
            if pd.notna(source[key]):
                formatted_present(row, source[key], "money2", "EPL model payoff component")
    rules = table_rows(tables, "epl_all_rules")
    require(len(rules) == 116 and Counter(str(row[0]) for row in rules) == Counter(canonical.index),
        "Complete rule appendix omitted or duplicated a strategy")
    return {"nba_original_capital_rows": 4, "nba_branch_specific_selection_rows": 4,
        "nba_verified_predictive_rows": 2, "nba_verified_trading_rows": 2, "nba_independent_annual_rows": 9,
        "epl_first_probability_rows": 4, "epl_secondary_betting_rows": 8, "epl_rolling_model_rows": 5,
        "epl_four_year_representative_accounts": 6, "epl_complete_common_season_metric_rows": 116,
        "epl_complete_strategy_rule_rows": 116}


def expand_tex(path, stack=()):
    path = Path(path).resolve()
    require(path not in stack, "Recursive LaTeX input")
    source = path.read_text(encoding="utf-8")
    def include(match):
        # A macro definition can contain \input{tables/#3}; only literal
        # document inputs are filesystem dependencies at this stage.
        if "#" in match.group(1):
            return match.group(0)
        child = path.parent / match.group(1)
        if not child.suffix:
            child = child.with_suffix(".tex")
        return expand_tex(child, stack + (path,))
    return re.sub(r"\\input\{([^}]+)\}", include, source)


def page_spans_off_page(pdf):
    bad = []
    for number, page in enumerate(pdf, 1):
        for block in page.get_text("dict")["blocks"]:
            for line in block.get("lines", []):
                for span in line.get("spans", []):
                    b = fitz.Rect(span["bbox"])
                    if b.x0 < -2 or b.y0 < -2 or b.x1 > page.rect.width + 2 or b.y1 > page.rect.height + 2:
                        bad.append({"page": number, "text": span["text"], "bbox": list(b)})
    return bad


def audit_report(output):
    folder = output / "report"
    expanded = expand_tex(folder / "Final_Report.tex")
    sections = re.findall(r"\\section\{([^}]+)\}", expanded)
    require(sections[:3] == ["Common methodology", "NBA results", "Premier League results"],
        "Final report must order common methodology, NBA, then Premier League")
    logs = list(folder.glob("*.log"))
    require(logs, "Compiled LaTeX log is absent")
    log = next((p for p in logs if p.stem == "Final_Report"), logs[0]).read_text(encoding="utf-8", errors="replace")
    require(not re.search(r"(?:! LaTeX Error|Emergency stop|Fatal error|There were undefined references|Citation[^\n]+undefined|Reference[^\n]+undefined|Missing character:)", log),
        "LaTeX reports a fatal error, undefined reference or missing glyph")
    overfull = [float(x) for x in re.findall(r"Overfull \\[hv]box \(([\d.]+)pt too (?:wide|high)\)", log)]
    require(not any(value > 2.0 for value in overfull), "Material LaTeX overfull box: " + str(overfull))
    with fitz.open(folder / "Final_Report.pdf") as pdf:
        require(len(pdf) >= 10, "Combined report is unexpectedly incomplete")
        first = norm(pdf[0].get_text())
        for name in NAMES:
            require(name in first, "Name missing from report first page: " + name)
        require("Supervisor" in first, "Report cover omits supervisor label")
        text = "\n".join(page.get_text() for page in pdf)
        clean = norm(text)
        for section in sections[:3]:
            require(section in clean, "Compiled report omitted section: " + section)
        for final in (10097.74, 6500.40, 12502.98, 4488.59):
            require(numeric(f"{final:,.2f}") in numeric(text), "Report omitted a source-recorded NBA final")
        require("occup" in text.lower() and "penalt" in text.lower(), "NBA ECE occupancy sentinel not explained")
        require("does not" in text.lower() and "svm" in text.lower() and "reconcil" in text.lower(),
            "NBA SVM reproduction limit absent")
        require("hold cash" in text.lower(), "EPL primary cash policies not reported")
        ids = [f"S{i:03d}" for i in range(1, 85)] + [f"P{i:03d}" for i in range(1, 33)]
        for sid in ids:
            require(sid in text, "Compiled report omitted a strategy ID: " + sid)
        bad = page_spans_off_page(pdf)
        require(not bad, "Report text exceeds page boundary: " + str(bad[:5]))
        pages = len(pdf)
    return {"pages": pages, "cover_authors_and_supervisor": True,
        "main_section_order": sections[:3], "all116_strategy_ids_present": True,
        "latex_undefined_references": 0, "material_overfull_boxes": 0, "off_page_text_spans": 0}


def audit_presentation(output, nba, epl):
    folder = output / "presentation"
    specs = read_json(folder / "presentation_content.json")
    require(len(specs) == 23 and sum(not s.get("appendix", False) for s in specs) == 12,
        "Expected 12 main slides and 11 backup slides")
    main = specs[:12]
    require(all(not s.get("appendix", False) for s in main) and all(s.get("appendix", False) for s in specs[12:]),
        "Main slides and backup slides interleaved")
    require(all(s["title"].startswith("Common") for s in main[1:3]), "Common-methodology slide order changed")
    require(all(s["title"].startswith("NBA:") for s in main[3:6]), "NBA results must precede EPL")
    require(all(s["title"].startswith("EPL:") for s in main[6:]), "EPL results must be the last main block")
    require(all("EPL" not in s.get("kicker", "") for s in main[3:6]), "NBA pages retain an incorrect EPL section label")
    deck = Presentation(folder / "Final_Project_Presentation.pptx")
    require(len(deck.slides) == len(specs), "Editable slide count mismatch")
    first_text = norm("\n".join(shape.text for shape in deck.slides[0].shapes if shape.has_text_frame))
    for name in NAMES:
        require(name in first_text, "Name missing from first presentation slide: " + name)
    require("Supervisor" in first_text, "Presentation cover omits supervisor label")
    rendering = read_json(folder / "render_audit.json")
    native = rendering["native_layout"]
    require(native["text_overflow_count"] == 0, "Native slide text overflows: " + str(native.get("text_overflows", [])[:8]))
    require(not rendering["geometry"]["geometry_violations"], "Presentation has off-canvas shapes")
    for name, digest in rendering["sha256"].items():
        require(sha(folder / name) == digest, "Rendered presentation changed after layout audit")
    for i, (spec, ppt) in enumerate(zip(specs, deck.slides), 1):
        actual = norm("\n".join(s.text for s in ppt.shapes if s.has_text_frame))
        for text in visible(spec).splitlines():
            require(norm(text) in actual, f"Presentation dropped slide {i} text: {text!r}")
        require(spec.get("notes", "") in ppt.notes_slide.notes_text_frame.text, "Presentation lost full speaker notes")
    core_table = next(e for e in main[4]["elements"] if e.get("type") == "table")
    for row, source in zip(core_table["rows"], nba["betting_2018"]["records"]):
        formatted_present(row, source["final_capital_saved_rounded"], "money2", "Slide NBA saved final")
        formatted_present(row, source["capital_return_saved_rounded"], "pct", "Slide NBA capital return")
    require("annual" in visible(main[5]).lower() and "fresh" in visible(main[5]).lower(),
        "NBA annual restarts not made explicit")
    require("hold cash" in visible(main[6]).lower(), "EPL proposal slide omitted primary cash outcome")
    for spec, source in zip(main[8:], ["best_learned", "simplest", "team", "player"]):
        profile = epl["rolling_and_player_extensions"]["main_model_profiles"][source]
        body = visible(spec)
        require(pct(profile["metrics"]["accuracy"]) in body and f"{profile['metrics']['log_loss']:.4f}" in body,
            "EPL featured model's probability statistics differ")
        require("1,000" in body and "1%" in body and "initial" in body.lower(),
            "EPL featured stake and initial-capital basis omitted")
    configurations = {}
    for spec, ppt in zip(specs, deck.slides):
        if not spec["title"].startswith("EPL appendix: strategy rules"):
            continue
        for line in ppt.notes_slide.notes_text_frame.text.splitlines():
            if line.startswith("{"):
                config = json.loads(line)
                sid = config["strategy_id"]
                require(sid not in configurations, "Duplicate EPL exact appendix configuration")
                configurations[sid] = config
    canonical = {row["strategy_id"]: row for row in read_json(EPL / "all_strategy_catalog.json")}
    require(configurations == canonical, "Presentation appendix does not retain all 116 exact rules")
    with fitz.open(folder / "Final_Project_Presentation.pdf") as pdf:
        require(len(pdf) == 23, "PDF slide count differs from editable deck")
        for name in NAMES:
            require(name in norm(pdf[0].get_text()), "PDF presentation cover omits a required name")
        for i, (page, spec) in enumerate(zip(pdf, specs), 1):
            require(abs(page.rect.width / page.rect.height - 16 / 9) < .01, "Wrong slide aspect ratio")
            require(norm(spec["title"]) in norm(page.get_text()), f"PDF omitted slide {i} title")
        bad = page_spans_off_page(pdf)
        require(not bad, "Slide PDF text exceeds page boundary: " + str(bad[:5]))
    return {"main_slides": 12, "backup_slides": 11, "total_slides": 23,
        "cover_authors_and_supervisor": True, "main_order": ["Common methodology", "NBA results", "Premier League results"],
        "exact_epl_appendix_configurations": len(configurations), "native_text_overflows": 0,
        "off_page_text_spans": 0, "editable_text_and_speaker_notes_match": True}


def pct(value):
    return f"{float(value):.2%}"


def audit_graph_data():
    counts = {}
    metrics = {"2025": pd.concat([pd.read_csv(LAB / "metrics.csv"), pd.read_csv(EPL / "metrics_single.csv")], ignore_index=True),
        "continuous": pd.concat([pd.read_csv(VALIDATION / "metrics_continuous.csv"), pd.read_csv(EPL / "metrics_continuous.csv")], ignore_index=True)}
    for scope, expected in (("2025", 116), ("continuous", 88)):
        graph = pd.read_csv(EPL / f"graph_equity_{scope}_all.csv", parse_dates=["match_date"])
        require(graph.strategy_id.nunique() == expected, "EPL graph strategy count changed")
        require(not graph.duplicated(["strategy_id", "match_date"]).any(), "EPL graph duplicate date/strategy")
        anchors = graph.loc[graph.is_initial]
        require(len(anchors) == expected and anchors.end_bankroll.eq(100000).all(), "EPL graph initial capitals differ")
        source = metrics[scope].set_index("strategy_id")
        ends = graph.sort_values("match_date").groupby("strategy_id").tail(1).set_index("strategy_id").reindex(source.index)
        require(np.allclose(ends.end_bankroll, source.final_capital_sum_units, atol=1e-6, rtol=0), "EPL curve endpoints differ from saved metrics")
        counts[scope] = {"accounts": expected, "points": len(graph), "initial_capital": 100000, "endpoints_checked": expected}
    return counts


def audit(output: Path) -> dict:
    output = Path(output).resolve()
    nba = read_json(output / "data/nba_evidence.json")
    epl = read_json(output / "data/epl_evidence.json")
    method = read_json(output / "data/methodology_evidence.json")
    require(nba["authors"] == list(NAMES[:2]) and nba["supervisor"] == NAMES[2], "Author evidence differs from proposal/user")
    require(method["proposal"]["authors"] == list(NAMES[:2]), "Proposal authors mismatch")
    tables = read_json(output / "data/display_tables.json")
    return {"status": "passed", "read_only": True,
        "source_evidence": audit_source_evidence(nba, epl),
        "display_tables": audit_tables(tables, nba, epl),
        "epl_graphs": audit_graph_data(),
        "report": audit_report(output),
        "presentation": audit_presentation(output, nba, epl),
        "no_new_SVM_metrics_imputed": True, "annual_NBA_restarts_not_a_continuous_portfolio": True}


if __name__ == "__main__":
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument("--output", type=Path, required=True)
    args = parser.parse_args()
    print(json.dumps(audit(args.output), indent=2))
