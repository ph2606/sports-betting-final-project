"""Editable slide rendering and hidden, native Microsoft PowerPoint PDF export.

Coordinates are inches on a 16 by 9 canvas. Fonts are points. Text is never
automatically shrunk: the native PowerPoint layout check reports clipped text.
Only the presentation opened by this module is closed after PDF export.
"""
from __future__ import annotations

import hashlib
import json
import math
from pathlib import Path
from typing import Any

from PIL import Image
from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_CONNECTOR, MSO_SHAPE
from pptx.enum.text import MSO_ANCHOR, MSO_AUTO_SIZE, PP_ALIGN
from pptx.util import Inches, Pt

WIDTH, HEIGHT = 16.0, 9.0
NAVY, TEAL, MUTED, BACKGROUND = "11243A", "0D897D", "536578", "F7F8FA"
DEFAULT_KICKER = "NBA & EPL | FINAL PROJECT | 23 Oct 2026"
BODY_TOP, BODY_BOTTOM = 2.0, 8.25


def _color(value: str) -> RGBColor:
    return RGBColor.from_string(str(value).lstrip("#").upper())


def _number(value: Any, label: str) -> float:
    result = float(value)
    if not math.isfinite(result):
        raise ValueError(f"{label} must be finite")
    return result


def _box(element: dict, *, body: bool = True) -> tuple[float, float, float, float]:
    box = tuple(_number(element[key], key) for key in ("x", "y", "w", "h"))
    x, y, w, h = box
    is_line = element.get("type") == "line"
    if w < 0 or h < 0 or (not is_line and (w == 0 or h == 0)):
        raise ValueError(f"Invalid element dimensions: {box}")
    if x < -1e-6 or y < -1e-6 or x + w > WIDTH + 1e-6 or y + h > HEIGHT + 1e-6:
        raise ValueError(f"Element lies outside slide: {box}")
    if body and (y < BODY_TOP - 1e-6 or y + h > BODY_BOTTOM + 1e-6):
        raise ValueError(f"Body element must fit y={BODY_TOP}..{BODY_BOTTOM}: {box}")
    return box


def _text(slide, box, value, *, name, size=20, color=NAVY, bold=False,
          align="left", url=None, vertical="top"):
    x, y, w, h = box
    shape = slide.shapes.add_textbox(Inches(x), Inches(y), Inches(w), Inches(h))
    shape.name = name
    frame = shape.text_frame
    frame.clear()
    frame.auto_size = MSO_AUTO_SIZE.NONE
    frame.word_wrap = True
    frame.margin_left = frame.margin_right = 0
    frame.margin_top = frame.margin_bottom = 0
    frame.vertical_anchor = {"top": MSO_ANCHOR.TOP, "middle": MSO_ANCHOR.MIDDLE,
                             "bottom": MSO_ANCHOR.BOTTOM}[vertical]
    alignment = {"left": PP_ALIGN.LEFT, "center": PP_ALIGN.CENTER, "right": PP_ALIGN.RIGHT}
    if align not in alignment:
        raise ValueError(f"Unsupported alignment {align}")
    lines = str(value if value is not None else "").split("\n")
    for index, line in enumerate(lines):
        paragraph = frame.paragraphs[0] if index == 0 else frame.add_paragraph()
        paragraph.alignment = alignment[align]
        paragraph.space_before = paragraph.space_after = Pt(0)
        paragraph.line_spacing = 1.0
        run = paragraph.add_run()
        run.text = line
        run.font.name = "Arial"
        run.font.size = Pt(float(size))
        run.font.bold = bool(bold)
        run.font.color.rgb = _color(color)
        if url:
            run.hyperlink.address = str(url)
    return shape


def _rectangle(slide, box, *, name, fill="FFFFFF", line=None, line_width=0.8):
    x, y, w, h = box
    shape = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, Inches(x), Inches(y), Inches(w), Inches(h))
    shape.name = name
    shape.fill.solid()
    shape.fill.fore_color.rgb = _color(fill)
    if line:
        shape.line.color.rgb = _color(line)
        shape.line.width = Pt(line_width)
    else:
        shape.line.fill.background()
    return shape


def _table(slide, element: dict, prefix: str) -> None:
    """Use editable cell rectangles/text, with predictable per-cell layout checks."""
    x, y, w, h = _box(element)
    headers = list(element.get("headers", []))
    rows = [list(row) for row in element.get("rows", [])]
    if not headers or any(len(row) != len(headers) for row in rows):
        raise ValueError("Tables require nonempty headers and consistent row widths")
    widths = element.get("widths", [1 / len(headers)] * len(headers))
    if len(widths) != len(headers):
        raise ValueError("Table widths must match headers")
    widths = [_number(width, "table column fraction") for width in widths]
    if any(width <= 0 for width in widths) or not math.isclose(sum(widths), 1, abs_tol=1e-5):
        raise ValueError("Table widths must be positive fractions summing to one")
    row_height = _number(element.get("row_height", h / (len(rows) + 1)), "row_height")
    if row_height <= 0 or row_height * (len(rows) + 1) > h + 1e-6:
        raise ValueError("Table row heights exceed its declared box")
    pad_x, pad_y = 0.10, 0.08
    for row_index, row in enumerate([headers] + rows):
        top = y + row_index * row_height
        bg = NAVY if row_index == 0 else ("FFFFFF" if row_index % 2 else "EBF0F4")
        left = x
        for column_index, value in enumerate(row):
            cell_width = w * widths[column_index]
            if cell_width <= 2 * pad_x or row_height <= 2 * pad_y:
                raise ValueError("Table cell has no room for text after padding")
            name = f"{prefix}-r{row_index:02d}-c{column_index:02d}"
            _rectangle(slide, (left, top, cell_width, row_height), name=name + "-cell",
                       fill=bg, line="D8E0E7", line_width=0.5)
            aligns = element.get("aligns", [])
            align = aligns[column_index] if column_index < len(aligns) else "left"
            _text(slide, (left + pad_x, top + pad_y, cell_width - 2 * pad_x,
                          row_height - 2 * pad_y), value, name=name + "-text",
                  size=element.get("header_size", 14) if row_index == 0 else element.get("font_size", 16),
                  bold=row_index == 0, color="FFFFFF" if row_index == 0 else NAVY,
                  align=align, vertical="middle")
            left += cell_width


def _card(slide, element: dict, prefix: str) -> None:
    x, y, w, h = _box(element)
    if h < 1.20 or w < 1.5:
        raise ValueError("Cards require h>=1.20 and w>=1.5 inches")
    accent = element.get("accent", TEAL)
    _rectangle(slide, (x, y, w, h), name=prefix + "-background", fill="FFFFFF", line="D8E0E7")
    _rectangle(slide, (x, y, 0.055, h), name=prefix + "-accent", fill=accent)
    inner_x, inner_w = x + 0.17, w - 0.34
    _text(slide, (inner_x, y + 0.12, inner_w, 0.23), element.get("label", ""),
          name=prefix + "-label", size=12, color=MUTED, bold=True)
    _text(slide, (inner_x, y + 0.40, inner_w, 0.51), element.get("value", ""),
          name=prefix + "-value", size=30, color=accent, bold=True)
    detail_height = h - 1.02
    _text(slide, (inner_x, y + 0.98, inner_w, detail_height), element.get("detail", ""),
          name=prefix + "-detail", size=13, color=MUTED)


def _picture(slide, element: dict, prefix: str) -> None:
    x, y, w, h = _box(element)
    path = Path(element["path"]).resolve()
    if not path.is_file():
        raise FileNotFoundError(path)
    with Image.open(path) as picture:
        aspect = picture.width / picture.height
    picture_w = min(w, h * aspect)
    picture_h = picture_w / aspect
    shape = slide.shapes.add_picture(str(path), Inches(x + (w - picture_w) / 2),
                                    Inches(y + (h - picture_h) / 2),
                                    width=Inches(picture_w), height=Inches(picture_h))
    shape.name = prefix + "-image"


def _footer_sources(sources: list[str]) -> str:
    names = []
    for source in sources:
        name = str(source).replace("\\", "/").rstrip("/").rsplit("/", 1)[-1]
        if name not in names:
            names.append(name)
    value = "Sources: " + "; ".join(names) if names else "Sources and definitions in speaker notes."
    if len(value) > 190:
        value = value[:185].rsplit(";", 1)[0] + " ..."
    return value + " | Full paths in notes" if names else value


def _geometry_audit(prs) -> dict:
    elements, outside = [], []
    for slide_number, slide in enumerate(prs.slides, start=1):
        for shape in slide.shapes:
            row = {"slide": slide_number, "shape": shape.name,
                   "x": shape.left / 914400, "y": shape.top / 914400,
                   "w": shape.width / 914400, "h": shape.height / 914400,
                   "editable_text": bool(shape.has_text_frame),
                   "text": shape.text if shape.has_text_frame else None}
            elements.append(row)
            if row["x"] < -1e-5 or row["y"] < -1e-5 or row["x"] + row["w"] > WIDTH + 1e-5 or row["y"] + row["h"] > HEIGHT + 1e-5:
                outside.append(row)
    return {"canvas_inches": [WIDTH, HEIGHT], "slide_count": len(prs.slides),
            "shape_count": len(elements), "geometry_violations": outside, "shapes": elements}


def _native_pdf(pptx_path: Path, pdf_path: Path) -> dict:
    """Open our saved deck without a window, inspect native text layout, export PDF."""
    import pythoncom
    import win32com.client

    pythoncom.CoInitialize()
    app = presentation = None
    prior_alerts = None
    rows, overflows = [], []
    try:
        app = win32com.client.DispatchEx("PowerPoint.Application")
        prior_alerts = app.DisplayAlerts
        # ppAlertsNone. Do not set Visible=False: PowerPoint disallows hiding
        # its application object. WithWindow=False suppresses the deck window.
        app.DisplayAlerts = 1
        presentation = app.Presentations.Open(str(pptx_path.resolve()), -1, 0, 0)
        for slide_number in range(1, int(presentation.Slides.Count) + 1):
            native_slide = presentation.Slides(slide_number)
            for shape_number in range(1, int(native_slide.Shapes.Count) + 1):
                shape = native_slide.Shapes(shape_number)
                if not shape.HasTextFrame or not shape.TextFrame2.HasText:
                    continue
                frame, text_range = shape.TextFrame2, shape.TextFrame2.TextRange
                width = float(shape.Width) - float(frame.MarginLeft) - float(frame.MarginRight)
                height = float(shape.Height) - float(frame.MarginTop) - float(frame.MarginBottom)
                row = {"slide": slide_number, "shape": str(shape.Name),
                       "text": str(text_range.Text), "box_width_pt": width,
                       "box_height_pt": height, "text_width_pt": float(text_range.BoundWidth),
                       "text_height_pt": float(text_range.BoundHeight)}
                row["vertical_overflow"] = row["text_height_pt"] > height + 1.2
                row["horizontal_overflow"] = row["text_width_pt"] > width + 1.2
                rows.append(row)
                if row["vertical_overflow"] or row["horizontal_overflow"]:
                    overflows.append(row)
        # Native SaveAs ppSaveAsPDF=32. Produces a PDF of the actual editable
        # deck, so PDF and PowerPoint share one layout rather than two renderers.
        presentation.SaveAs(str(pdf_path.resolve()), 32)
        version = str(app.Version)
    finally:
        if presentation is not None:
            presentation.Close()
        if app is not None and prior_alerts is not None:
            app.DisplayAlerts = prior_alerts
        if app is not None and int(app.Presentations.Count) == 0:
            app.Quit()
        presentation = app = None
        pythoncom.CoUninitialize()
    if not pdf_path.is_file() or pdf_path.stat().st_size < 100:
        raise RuntimeError("Native PowerPoint PDF export did not create a PDF")
    return {"exporter": "Microsoft PowerPoint SaveAs ppSaveAsPDF=32", "version": version,
            "opened_with_window": False, "text_shapes_checked": len(rows),
            "text_overflow_count": len(overflows), "text_overflows": overflows,
            "text_measurements": rows}


def render(slides: list[dict], output: Path) -> dict:
    """Render editable PPTX and a PDF, writing a native layout audit alongside."""
    if not slides:
        raise ValueError("At least one slide is required")
    output = Path(output).resolve()
    output.mkdir(parents=True, exist_ok=True)
    pptx_path = output / "Final_Project_Presentation.pptx"
    pdf_path = output / "Final_Project_Presentation.pdf"
    prs = Presentation()
    prs.slide_width, prs.slide_height = Inches(WIDTH), Inches(HEIGHT)
    prs.core_properties.title = "Calibrated Machine Learning for Sports Betting: NBA and Premier League"
    prs.core_properties.subject = "Common methodology, NBA results, then Premier League results"
    prs.core_properties.author = "Aman Dhillon; Panagiotis Housos"
    prs.core_properties.keywords = "NBA; Premier League; statistical arbitrage; accuracy; calibration; final project"
    for slide_number, spec in enumerate(slides, start=1):
        slide = prs.slides.add_slide(prs.slide_layouts[6])
        slide.background.fill.solid()
        slide.background.fill.fore_color.rgb = _color(BACKGROUND)
        prefix = f"slide{slide_number:03d}"
        _text(slide, (0.65, 0.23, 14.7, 0.20), spec.get("kicker", DEFAULT_KICKER),
              name=prefix + "-kicker", size=10, color=TEAL, bold=True)
        _text(slide, (0.65, 0.55, 14.7, 0.55), spec["title"], name=prefix + "-title",
              size=30, color=NAVY, bold=True)
        _text(slide, (0.65, 1.22, 14.7, 0.60), spec.get("subtitle", ""),
              name=prefix + "-subtitle", size=15, color=MUTED)
        for index, element in enumerate(spec.get("elements", [])):
            kind = element.get("type", "text")
            element = {**element, "type": kind}
            box = _box(element)
            name = f"{prefix}-element{index:03d}-{kind}"
            if kind == "text":
                _text(slide, box, element.get("text", ""), name=name,
                      size=element.get("size", 20), color=element.get("color", NAVY),
                      bold=element.get("bold", False), align=element.get("align", "left"),
                      url=element.get("url"))
            elif kind == "rect":
                _rectangle(slide, box, name=name, fill=element.get("fill", "FFFFFF"),
                           line=element.get("line"))
            elif kind == "line":
                x, y, w, h = box
                shape = slide.shapes.add_connector(MSO_CONNECTOR.STRAIGHT,
                            Inches(x), Inches(y), Inches(x + w), Inches(y + h))
                shape.name = name
                shape.line.color.rgb = _color(element.get("color", "D8E0E7"))
                shape.line.width = Pt(element.get("width", 1))
            elif kind == "image":
                _picture(slide, element, name)
            elif kind == "table":
                _table(slide, element, name)
            elif kind == "card":
                _card(slide, element, name)
            else:
                raise ValueError(f"Unknown element type: {kind}")
        line = slide.shapes.add_connector(MSO_CONNECTOR.STRAIGHT, Inches(0.65), Inches(8.55), Inches(15.35), Inches(8.55))
        line.name = prefix + "-footer-rule"
        line.line.color.rgb = _color("D8E0E7")
        line.line.width = Pt(0.75)
        sources = [str(source) for source in spec.get("sources", [])]
        _text(slide, (0.65, 8.62, 13.9, 0.22), _footer_sources(sources),
              name=prefix + "-sources", size=8, color=MUTED)
        _text(slide, (14.75, 8.60, 0.60, 0.25), str(slide_number),
              name=prefix + "-page", size=10, color=MUTED, align="right")
        notes = str(spec.get("notes", ""))
        notes += "\n\nExact source paths:\n" + "\n".join(sources)
        if spec.get("appendix"):
            notes += "\n\nAppendix slide."
        slide.notes_slide.notes_text_frame.text = notes
    geometry = _geometry_audit(prs)
    if geometry["geometry_violations"]:
        raise ValueError(f"Slides contain off-canvas shapes: {geometry['geometry_violations']}")
    prs.save(str(pptx_path))
    native = _native_pdf(pptx_path, pdf_path)
    import fitz
    with fitz.open(pdf_path) as pdf:
        if len(pdf) != len(slides):
            raise RuntimeError("Native PDF page count differs from editable presentation")
        pdf_pages = len(pdf)
    result = {"pptx": str(pptx_path), "pdf": str(pdf_path), "pdf_pages": pdf_pages,
              "slide_count": len(slides), "geometry": geometry, "native_layout": native,
              "sha256": {path.name: hashlib.sha256(path.read_bytes()).hexdigest()
                         for path in (pptx_path, pdf_path)}}
    audit_path = output / "render_audit.json"
    audit_path.write_text(json.dumps(result, indent=2, ensure_ascii=False), encoding="utf-8")
    return result
