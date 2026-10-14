"""Render a separate strategy-metrics supplement from saved analytics tables.

Usage: python -m epl_strategy.report --results PATH
No model fitting, data acquisition, strategy selection or original-report edits.
"""
from __future__ import annotations

import argparse
import hashlib
import html
import json
import math
from datetime import datetime, timezone
from pathlib import Path

import pandas as pd
from docx import Document
from docx.enum.section import WD_ORIENT, WD_SECTION_START
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter
from pypdf import PdfReader
from reportlab.lib import colors
from reportlab.lib.pagesizes import landscape, letter
from reportlab.lib.styles import ParagraphStyle
from reportlab.pdfbase import pdfmetrics
from reportlab.pdfbase.ttfonts import TTFont
from reportlab.platypus import BaseDocTemplate, Frame, NextPageTemplate, PageBreak, PageTemplate, Paragraph, Spacer, Table, TableStyle

PROJECT_ROOT = Path(__file__).resolve().parents[1]
DRAFT_DATE = "2026-10-14"
TITLE = "EPL trading-strategy metrics supplement — draft for review"
REQUIRED = [
    "study", "scope", "strategy_id", "episode_id", "split", "outer_season", "policy", "staking",
    "bet_count", "winning_bet_count", "hit_ratio", "hit_ratio_ci_low", "hit_ratio_ci_high", "trading_frequency", "trade_rate_all_predictions",
    "avg_winning_gain_units", "avg_losing_loss_units", "expected_gain_total_units", "expected_loss_total_units",
    "expected_gain_per_trade_units", "expected_loss_per_trade_units", "expected_net_profit_total_units",
    "expected_net_profit_per_trade_units", "realized_expectancy_per_trade_units", "expected_roi", "profit", "turnover",
    "roi", "roi_ci_low", "roi_ci_high", "profit_factor", "payoff_ratio", "max_drawdown", "return_on_initial_capital",
]


def _missing(value):
    return value is None or (not isinstance(value, (list, dict)) and bool(pd.isna(value)))


def number(value, places=2, percent=False):
    """Human rendering: unavailable stays unavailable; genuine infinity stays infinite."""
    if _missing(value):
        return "—"
    value = float(value)
    if math.isinf(value):
        return "+∞" if value > 0 else "−∞"
    if percent:
        value *= 100
    return f"{value:,.{places}f}{'%' if percent else ''}".replace("-", "−")


def _interval(low, high, percentage=True, suffix=True):
    if _missing(low) or _missing(high):
        return "—"
    if percentage and not suffix:
        return f"[{number(float(low)*100)}, {number(float(high)*100)}]"
    return f"[{number(low,percent=percentage)}, {number(high,percent=percentage)}]"


def _label(row):
    if str(row.get("study", "")).lower().startswith("first") or str(row.get("branch", "")) in ("accuracy", "calibration"):
        branch = {"accuracy": "Accuracy", "calibration": "Calibration"}.get(str(row.get("branch")), str(row.get("branch", "")))
        policy = "Primary" if row.get("policy")=="validation_selected" else "Fixed 5%" if row.get("policy")=="predeclared_5pct" else str(row.get("policy", ""))
        if bool(row.get("is_tuning_diagnostic",False)):
            policy="Tuning threshold "+number(row.get("threshold"),percent=True)
        odds = "Best book" if row.get("odds_mode")=="best_individual_book" else "Bet365" if row.get("odds_mode")=="bet365_reference" else str(row.get("odds_mode", ""))
        stake = "Fixed" if row.get("staking")=="fixed" else "Quarter Kelly" if row.get("staking")=="kelly" else str(row.get("staking", ""))
        return f"{branch} / {policy}\n{odds} / {stake}"
    group = {"team":"Team", "team_xg":"Team+xG", "team_player":"Team+players", "combined":"Combined"}.get(str(row.get("feature_group")), str(row.get("feature_group", "")))
    model = str(row.get("model_type", ""))
    if row.get("feature_group")=="class_frequency":
        label = "Class frequency"
    elif model=="market_temperature":
        label = "Temperature market"
    elif row.get("feature_group")=="market_only":
        label = "Raw market"
    else:
        label = {"logistic":"Logistic", "market_residual":"Market residual"}.get(model, model)+": "+group
    return label+" / "+("Fixed" if row.get("staking")=="fixed" else "Quarter Kelly")


def _selected_groups(metrics):
    """Strict group selection; unexpected schemas fail instead of silently dropping rows."""
    diagnostic = metrics.get("is_tuning_diagnostic", pd.Series(False,index=metrics.index)).fillna(False).astype(bool)
    first = metrics.loc[metrics.study.eq("first") & metrics.scope.eq("test") & metrics.split.eq("test") & ~diagnostic].copy()
    pooled = metrics.loc[metrics.study.eq("ablation") & metrics.scope.eq("pooled_seasons") & ~diagnostic].copy()
    if len(first)!=16 or len(pooled)!=22:
        raise ValueError(f"Expected 16 first locked-test rows and 22 pooled ablation rows, found {len(first)}/{len(pooled)}; study/scope values: {metrics[['study','scope']].drop_duplicates().to_dict('records')}")
    if first.strategy_id.duplicated().any() or pooled.strategy_id.duplicated().any():
        raise ValueError("Duplicate displayed strategies would double-count a report row")
    first = first.sort_values(["policy","branch","odds_mode","staking"],kind="stable")
    pooled = pooled.sort_values(["model_type","feature_group","staking"],kind="stable")
    return first, pooled


def _performance(frame):
    rows=[]
    for _,row in frame.iterrows():
        hit=number(row.hit_ratio,percent=True).rstrip("%")
        bounds=_interval(row.hit_ratio_ci_low,row.hit_ratio_ci_high,suffix=False)
        if bounds!="—":
            hit+="\n"+bounds
        rows.append({"Strategy":_label(row),"Trades":number(row.bet_count,0),"Hit %\n[95% CI]":hit,
                     "Trade\nfreq. %":number(row.trading_frequency,percent=True).rstrip("%"),
                     "Avg. win\ngain":number(row.avg_winning_gain_units),"Avg. loss\nsize":number(row.avg_losing_loss_units),
                     "Expected\ngain/trade":number(row.expected_gain_per_trade_units),
                     "Expected\nloss/trade":number(row.expected_loss_per_trade_units),
                     "Expected\nnet/trade":number(row.expected_net_profit_per_trade_units),
                     "Realised\nnet/trade":number(row.realized_expectancy_per_trade_units)})
    return pd.DataFrame(rows)


def _risk(frame):
    return pd.DataFrame([{"Strategy":_label(row),"Net profit":number(row.profit),"Turnover":number(row.turnover),
                          "ROI":number(row.roi,percent=True),"95% ROI CI":_interval(row.roi_ci_low,row.roi_ci_high),
                          "Expected\nROI":number(row.expected_roi,percent=True),"Profit\nfactor":number(row.profit_factor),
                          "Payoff\nratio":number(row.payoff_ratio),"Max.\ndrawdown":number(row.max_drawdown,percent=True),
                          "Capital\nreturn":number(row.return_on_initial_capital,percent=True)} for _,row in frame.iterrows()])


def _definitions():
    rows=[
        ("hit_ratio","Winning settled bets / settled bets; conditional hit ratio, not H/D/A classification accuracy.","Fraction"),
        ("avg_winning_gain_units","Mean positive realised net P&L, conditional on a winning trade.","Units / winning trade"),
        ("avg_losing_loss_units","Mean magnitude of realised negative net P&L, conditional on a losing trade; reported positively.","Units / losing trade"),
        ("expected_gain_total_units","Σ s p (O−1): model-implied expected positive net-P&L component.","Units"),
        ("expected_loss_total_units","Σ s (1−p): model-implied expected loss component, positive magnitude.","Units"),
        ("expected_net_profit_total_units","Expected gain minus expected loss = Σ s (pO−1). Conditional on actual allocated stakes, including the realised bankroll-dependent Kelly path; not unconditional terminal wealth.","Units"),
        ("expected_gain_per_trade_units / expected_loss_per_trade_units","Respective expected component totals / number of recorded trades.","Units / trade"),
        ("expected_net_profit_per_trade_units","Model-implied expected net P&L total / trades; not observed average winning payoff.","Units / trade"),
        ("realized_expectancy_per_trade_units","Realised net profit / trades.","Units / trade"),
        ("expected_roi","Model-implied expected net P&L / actual turnover.","Fraction"),
        ("trading_frequency","Trades / matched-valid-price fixtures; report price coverage alongside this conditional rate.","Fraction"),
        ("trade_rate_all_predictions","Trades / all predicted fixtures, including missing-price opportunities.","Fraction"),
        ("active_date_fraction","Dates with trades / eligible match dates; no intraday timing is inferred.","Fraction"),
        ("bets_per_calendar_week","Trades / (sum of inclusive first-to-last priced-match calendar spans within each reset episode / 7). Off-season gaps between pooled episodes are excluded.","Trades / week"),
        ("profit / turnover / roi","Realised net P&L / total stakes / net P&L divided by total stakes, respectively.","Units / units / fraction"),
        ("profit_factor","Total positive realised P&L / magnitude of total negative realised P&L. +∞ requires gains and no losses.","Ratio"),
        ("payoff_ratio","Mean winning gain / mean losing loss magnitude; undefined if the required conditional mean is unavailable.","Ratio"),
        ("max_drawdown","Largest decline from the date-close bankroll running peak, including initial capital. Pooled value is worst episode, not a stitched path.","Fraction"),
        ("initial_capital_sum_units / final_capital_sum_units","Sum of separate episode opening/closing capital; pooled resets remain separate.","Units"),
        ("return_on_initial_capital","Net P&L / sum of initial episode capital; differs from ROI on turnover.","Fraction"),
        ("avg_stake_units","Mean actual stake after commitment/capital constraints.","Units / trade"),
        ("avg_selected_probability / avg_decimal_odds","Unweighted means over selected trades; probabilities are model forecasts and odds are archived decimal prices.","Fraction / decimal odds"),
        ("price_coverage / predictions_without_valid_prices","Matched valid price fixtures / all predictions; missing-price opportunities are not economic abstentions.","Fraction / count"),
        ("date_return_mean / date_return_sample_std","Mean and sample standard deviation of returns on matched dates, including dates with no trades.","Fraction"),
        ("date_return_mean_to_std_unannualized","Unannualized mean / sample standard deviation on irregular match dates; no risk-free adjustment. This is not an annualized Sharpe ratio.","Ratio"),
        ("max_date_capital_exposure / mean_date_capital_exposure","Date turnover / date-start bankroll; summarised across matched match dates.","Fraction"),
        ("empirical_break_even_hit_ratio","Mean losing-loss magnitude / (mean winning gain + mean losing-loss magnitude). Descriptive for this observed payoff mixture, not a universal changing-odds threshold.","Fraction"),
        ("roi_ci_low/high; hit_ratio_ci_low/high","Conditional 95% whole-date bootstrap bounds: 1,000 draws; pooled ablation resamples dates within each season. No model refitting.","Fraction"),
        ("is_tuning_diagnostic","Strategy-validation threshold comparisons are tuning diagnostics, not additional held-out evidence.","Boolean"),
        ("Missing and infinity policy","— means undefined/unavailable, never zero. Empty-trade cash totals and rates can be zero when their denominators exist.","Display rule"),
    ]
    return pd.DataFrame(rows,columns=["Metric","Definition","Unit"])


def _pages(first,pooled,metrics,prospective,arbitrage):
    tables={"first_trade_metrics":_performance(first),"first_risk_metrics":_risk(first),
            "pooled_ablation_trade_metrics":_performance(pooled),"pooled_ablation_risk_metrics":_risk(pooled)}
    pages=[]
    def page(title,orientation="landscape"):
        item={"title":title,"orientation":orientation,"blocks":[]};pages.append(item);return item["blocks"]
    def paragraph(blocks,text):blocks.append({"type":"paragraph","text":text})
    def table(blocks,key,caption):blocks.append({"type":"table","frame":tables[key],"caption":caption})
    b=page(TITLE,"portrait")
    paragraph(b,"October 14, 2026 | Companion to the unchanged seven-page EPL project draft. Historical archived-price analytics; no model refitting or new strategy selection.")
    paragraph(b,f"This supplement exposes requested trade outcomes, forecast-implied expectations, activity and bankroll risk. The complete workbook/CSV retains {len(metrics)} strategy-period rows, including validation, test, historical monitoring, individual ablation seasons, pooled ablation summaries and tuning diagnostics. The displayed tables include all {len(first)} locked-test strategies and all {len(pooled)} pooled ablation variants/staking rules; cash policies are retained.")
    paragraph(b,"For stake s, forecast win probability p and decimal odds O, expected gain is s·p·(O−1), expected loss is s·(1−p), and expected net P&L is s·(pO−1). Gain/loss expectations are probability-weighted positive/negative P&L components, distinct from realised conditional winning/losing payoffs. Per-trade components divide totals by trades. They condition on actual allocated stakes, including the realised bankroll-dependent Kelly path; they are not unconditional predicted terminal wealth or verified true expectations.")
    paragraph(b,"Hit ratio is winning bets divided by settled bets. Trading frequency divides bets by matched fixtures with valid prices; the workbook also gives bets divided by all predicted fixtures, active-date fraction and bets per calendar week. Missing quotes must be distinguished from abstention. Classification accuracy is a separate H/D/A metric and cannot substitute for trade hit ratio.")
    paragraph(b,"ROI is net profit divided by turnover. Capital return divides profit by the sum of episode initial capital. Profit factor compares total gains with total losses; payoff ratio compares average winning gain with average losing loss. Drawdown is measured at date close, including initial capital. Pooled ablation capital resets remain separate: cash totals may be added, but the displayed pooled drawdown is the worst episode, not a continuous investment track record.")
    paragraph(b,"A dash means undefined or unavailable. A cash policy can have zero profit, turnover and trade frequency while its hit ratio, ROI and per-trade averages remain undefined. Genuine positive infinity is displayed explicitly for profit factor when gains exist and losses do not. Pending forecasts supply no settled hit rate or trading profit.")
    paragraph(b,"95% hit-ratio/ROI intervals use 1,000 whole-match-date bootstrap draws; pooled draws stratify by season. They condition on recorded wagers and fitted predictions, omit fitting/selection uncertainty and lack multiplicity adjustment. New interpolated-percentile intervals can differ from original noninterpolating bounds; original outputs remain unchanged. The strictly positive qualification rule pO−1>0.05 builds positive model-implied expected ROI into the fixed 5% policy; this is not independent evidence of an edge.")
    example=first.loc[first.branch.eq("calibration") & first.policy.eq("predeclared_5pct") & first.odds_mode.eq("best_individual_book") & first.staking.eq("fixed")]
    if len(example)!=1:raise ValueError("Missing unique calibration/best-book/fixed example")
    e=example.iloc[0]
    paragraph(b,f"For example, the calibration branch's best-book fixed policy wins {number(e.winning_bet_count,0)} of {number(e.bet_count,0)} trades ({number(e.hit_ratio,percent=True)}), with mean winning gain {number(e.avg_winning_gain_units)} and mean losing loss {number(e.avg_losing_loss_units)} units. Model-implied net expectancy is {number(e.expected_net_profit_per_trade_units,4)} units/trade ({number(e.expected_roi,percent=True)} expected ROI), versus realised {number(e.realized_expectancy_per_trade_units,4)} units/trade ({number(e.roi,percent=True)} ROI). These are different quantities.")
    b=page("First study: complete locked-test trade metrics")
    paragraph(b,"2025/26: the primary validation-selected policy and separately declared 5% policy remain distinct. Gain/loss/expectancy columns use bankroll units; hit and frequency use percentages.")
    table(b,"first_trade_metrics","Table S1. Every locked-test strategy. Hit bounds are conditional 95% intervals; frequency uses matched valid prices. Expected components use recorded actual stakes.")
    b=page("First study: complete locked-test return and risk metrics")
    paragraph(b,"Cash policies remain in the table. ROI measures return per unit staked; capital return measures profit relative to initial bankroll, so the two percentages answer different questions.")
    table(b,"first_risk_metrics","Table S2. Same strategies and period as S1. Conditional bounds resample observed dates; undefined cash-policy ROI remains a dash.")
    b=page("Ablation: complete pooled trade metrics")
    paragraph(b,"Four exploratory outer seasons, 2022/23–2025/26; all source groups, benchmarks, fixed stakes and quarter Kelly are retained. The same inspected history is not a fresh holdout.")
    table(b,"pooled_ablation_trade_metrics","Table S3. Every pooled variant/staking rule. Components pool recorded cash totals; hit ratio pools wins/bets, not seasonal averages. Probability quality and availability claims remain separate.")
    b=page("Ablation: complete pooled return and risk metrics")
    positive=pooled.loc[pooled.roi.gt(0)]
    spanning=int((positive.roi_ci_low.le(0) & positive.roi_ci_high.ge(0)).sum())
    paragraph(b,f"Each outer season starts a separate bankroll. Pooled ROI uses total profit / turnover; capital return uses total opening capital. {spanning} of {len(positive)} positive point-ROI intervals include zero; conditional comparisons do not identify a reliable winning strategy.")
    table(b,"pooled_ablation_risk_metrics","Table S4. Same pooled variants as S3. Max drawdown is the worst seasonal episode; no stitched pooled equity path or annualized Sharpe is claimed.")
    b=page("Coverage, pending opportunities and audit trail","portrait")
    prospective_models=prospective.groupby("model_id",dropna=False)[["released_forecast_count","quote_backed_forecast_count","resolved_forecast_count"]].max()
    released=int(prospective_models.released_forecast_count.sum())
    quoted=int(prospective_models.quote_backed_forecast_count.sum())
    resolved=int(prospective_models.resolved_forecast_count.sum())
    paragraph(b,f"The prospective table contains {len(prospective)} declared strategy/staking status slots, each summarising a model's future opportunities; these are not {len(prospective)} fixtures or executed trades. Across distinct models it records {released} released forecasts, {quoted} quote-backed forecasts and {resolved} resolved forecasts. Trading performance remains unavailable without recorded quote-backed positions. Unknown trade counts, profits and expectations stay unavailable rather than becoming zero. Full per-slot reasons remain in the workbook and CSV.")
    if arbitrage is not None:
        paragraph(b,f"The separate arbitrage screen contains {len(arbitrage)} source rows. The arithmetic sum of inverse best-outcome prices can describe a candidate; it does not prove synchronized observation, accepted stakes, compatible settlement, fees or execution. Screen statuses are retained separately from strategy wagers and are not credited as trading profit.")
    paragraph(b,f"Validation and threshold-grid diagnostics are labelled explicitly and excluded from the two held-out/pool displays. Historical monitoring is retrospective. All {len(metrics)} underlying rows remain available without selecting the best strategy after observing returns. The supplement neither changes decisions nor promises future profit.")
    paragraph(b,"Open EPL_strategy_metrics.xlsx for all strategies, source scopes, prospective status, arbitrage status and metric definitions. strategy_metrics_full.csv retains every source metric at original precision. The four display CSVs contain the rounded published tables. No fee or fill assumptions are silently introduced.")
    paragraph(b,"The companion report_manifest.json lists input/output SHA-256 hashes and verifies that original study outputs and the seven-page draft stayed unchanged. Regenerate this supplement with python -m epl_strategy.report --results PATH. The report reads the saved analytics; it does not retrieve data, train estimators, tune thresholds or rewrite historical results.")
    paragraph(b,"Definitions and interval conventions follow epl_strategy/METHODOLOGY.md. Winning-trade share, conditional mean win/loss and profit-factor terminology is consistent with the official MetaTrader report reference linked there; the probability-weighted components remain separately defined in this project.")
    return pages,tables


def _markdown(pages):
    chunks=[]
    for i,page in enumerate(pages):
        chunks.extend([("# " if i==0 else "## ")+page["title"],""])
        for block in page["blocks"]:
            if block["type"]=="paragraph":chunks.extend([block["text"],""])
            else:
                chunks.extend([block["caption"],""])
                frame=block["frame"]
                chunks.append("| "+" | ".join(c.replace("\n","<br>") for c in frame.columns)+" |")
                chunks.append("| "+" | ".join("---" for _ in frame.columns)+" |")
                for row in frame.itertuples(index=False,name=None):chunks.append("| "+" | ".join(str(v).replace("\n","<br>") for v in row)+" |")
                chunks.append("")
    return "\n".join(chunks)


def _pdf(pages,path):
    names={"body":"Times-Roman","bold":"Helvetica-Bold","sans":"Helvetica","italic":"Times-Italic"}
    fonts=Path("C:/Windows/Fonts")
    for role,file in (("body","times.ttf"),("bold","arialbd.ttf"),("sans","arial.ttf"),("italic","timesi.ttf")):
        if (fonts/file).exists():
            font="Metrics"+role.title();pdfmetrics.registerFont(TTFont(font,str(fonts/file)));names[role]=font
    styles={
        "title":ParagraphStyle("MetricsTitle",fontName=names["bold"],fontSize=14,leading=16,spaceAfter=7,textColor=colors.HexColor("#12364D")),
        "body":ParagraphStyle("MetricsBody",fontName=names["body"],fontSize=10.2,leading=12.4,spaceAfter=7),
        "caption":ParagraphStyle("MetricsCaption",fontName=names["italic"],fontSize=8.6,leading=10,spaceAfter=4),
        "cell":ParagraphStyle("MetricsCell",fontName=names["sans"],fontSize=8.2,leading=9.2),
        "head":ParagraphStyle("MetricsHead",fontName=names["bold"],fontSize=8.2,leading=9.2,textColor=colors.white),
    }
    def header(canvas,doc):
        w,h=canvas._pagesize;canvas.setFont(names["sans"],8);canvas.setFillColor(colors.HexColor("#52616B"))
        canvas.drawString(36,h-21,"EPL STRATEGY METRICS — DRAFT FOR REVIEW | October 14, 2026")
        canvas.drawRightString(w-36,21,f"{doc.page} / {len(pages)}")
    doc=BaseDocTemplate(str(path),pagesize=letter,title=TITLE,author="Aman Dhillon and Panagiotis Housos")
    for orientation,size in (("portrait",letter),("landscape",landscape(letter))):
        w,h=size;doc.addPageTemplates(PageTemplate(id=orientation,pagesize=size,onPage=header,
              frames=[Frame(36,36,w-72,h-72,leftPadding=0,rightPadding=0,topPadding=0,bottomPadding=0)]))
    story=[]
    for i,page in enumerate(pages):
        if i:story.extend([NextPageTemplate(page["orientation"]),PageBreak()])
        width=(letter if page["orientation"]=="portrait" else landscape(letter))[0]-72
        story.append(Paragraph(html.escape(page["title"]),styles["title"]))
        for block in page["blocks"]:
            if block["type"]=="paragraph":story.append(Paragraph(html.escape(block["text"]),styles["body"]))
            else:
                story.append(Paragraph(html.escape(block["caption"]),styles["caption"]))
                frame=block["frame"];raw=[list(frame.columns)]+frame.values.tolist()
                cells=[[Paragraph(html.escape(str(v)).replace("\n","<br/>"),styles["head" if i==0 else "cell"]) for v in row] for i,row in enumerate(raw)]
                weights=[.23,.055,.13,.075,.085,.085,.085,.085,.085,.085] if "Trades" in frame else [.25,.09,.10,.06,.13,.075,.075,.075,.07,.075]
                table=Table(cells,colWidths=[width*v for v in weights],repeatRows=1,hAlign="CENTER")
                table.setStyle(TableStyle([("BACKGROUND",(0,0),(-1,0),colors.HexColor("#12364D")),
                    ("ROWBACKGROUNDS",(0,1),(-1,-1),[colors.white,colors.HexColor("#F0F4F6")]),
                    ("VALIGN",(0,0),(-1,-1),"MIDDLE"),("LEFTPADDING",(0,0),(-1,-1),3),("RIGHTPADDING",(0,0),(-1,-1),3),
                    ("TOPPADDING",(0,0),(-1,-1),.5),("BOTTOMPADDING",(0,0),(-1,-1),.5)]))
                story.extend([table,Spacer(1,4)])
    doc.build(story)


def _docx(pages,path):
    doc=Document();normal=doc.styles["Normal"];normal.font.name="Times New Roman";normal.font.size=Pt(10.2)
    normal.paragraph_format.line_spacing=Pt(12.4);normal.paragraph_format.space_after=Pt(7)
    heading=doc.styles["Heading 1"];heading.font.name="Arial";heading.font.size=Pt(14)
    heading.paragraph_format.space_before=Pt(0);heading.paragraph_format.space_after=Pt(7);heading.paragraph_format.line_spacing=Pt(16)
    doc.core_properties.title=TITLE;doc.core_properties.author="Aman Dhillon and Panagiotis Housos"
    for i,page in enumerate(pages):
        s=doc.sections[0] if i==0 else doc.add_section(WD_SECTION_START.NEW_PAGE)
        is_land=page["orientation"]=="landscape";s.orientation=WD_ORIENT.LANDSCAPE if is_land else WD_ORIENT.PORTRAIT
        s.page_width=Inches(11 if is_land else 8.5);s.page_height=Inches(8.5 if is_land else 11)
        s.top_margin=s.bottom_margin=s.left_margin=s.right_margin=Inches(.5);s.header_distance=s.footer_distance=Inches(.2)
        if i==0:
            r=s.header.paragraphs[0].add_run("EPL STRATEGY METRICS — DRAFT FOR REVIEW | October 14, 2026");r.font.name="Arial";r.font.size=Pt(8)
            p=s.footer.paragraphs[0];p.alignment=2;f=OxmlElement("w:fldSimple");f.set(qn("w:instr"),"PAGE");p._p.append(f)
        doc.add_heading(page["title"],level=1)
        for block in page["blocks"]:
            if block["type"]=="paragraph":doc.add_paragraph(block["text"])
            else:
                p=doc.add_paragraph(block["caption"]);p.paragraph_format.line_spacing=Pt(10);p.paragraph_format.space_after=Pt(4)
                for r in p.runs:r.font.size=Pt(8.6);r.italic=True
                frame=block["frame"];table=doc.add_table(rows=1,cols=len(frame.columns));table.style="Light Shading Accent 1";table.autofit=False
                weights=[.23,.055,.13,.075,.085,.085,.085,.085,.085,.085] if "Trades" in frame else [.25,.09,.10,.06,.13,.075,.075,.075,.07,.075]
                for column,weight in zip(table.columns,weights):column.width=Inches((10 if is_land else 7.5)*weight)
                for cell,value in zip(table.rows[0].cells,frame.columns):cell.text=str(value)
                for values in frame.itertuples(index=False,name=None):
                    for cell,value in zip(table.add_row().cells,values):cell.text=str(value)
                for row in table.rows:
                    for cell,weight in zip(row.cells,weights):
                        cell.width=Inches((10 if is_land else 7.5)*weight)
                        for p in cell.paragraphs:
                            p.paragraph_format.space_after=Pt(0);p.paragraph_format.line_spacing=Pt(9.2)
                            for r in p.runs:r.font.name="Arial";r.font.size=Pt(8.2)
                        margins=OxmlElement("w:tcMar")
                        for side in ("top","bottom"):
                            v=OxmlElement("w:"+side);v.set(qn("w:w"),"10");v.set(qn("w:type"),"dxa");margins.append(v)
                        cell._tc.get_or_add_tcPr().append(margins)
    doc.save(path)


def _excel(path,metrics,prospective,arbitrage,definitions,first,pooled,registry):
    readme=pd.DataFrame([
        ("Status",TITLE),("Date",DRAFT_DATE),("Rows",len(metrics)),
        ("Expected quantities","Model-implied components at recorded stakes/archived prices, not true expectations."),
        ("Undefined values","Dash; genuine infinity is explicit. Cash totals/rates can be zero."),
        ("Pooled bankrolls","Season capital resets remain separate. Pooled drawdown is worst episode."),
        ("Intervals","Conditional 95% whole-date bootstrap, 1,000 draws; pooled draws stratify by season."),
        ("Interpretation","Retrospective replay; no executable-profit, player-availability or prospective-performance claim."),
    ],columns=["Item","Description"])
    core=["study","scope","episode_id","bet_count","winning_bet_count","hit_ratio","trading_frequency","avg_winning_gain_units","avg_losing_loss_units",
          "expected_gain_per_trade_units","expected_loss_per_trade_units","expected_net_profit_per_trade_units","realized_expectancy_per_trade_units",
          "expected_roi","roi","profit","turnover","profit_factor","payoff_ratio","max_drawdown","return_on_initial_capital",
          "trade_rate_all_predictions","active_date_fraction","bets_per_calendar_week","avg_stake_units","avg_selected_probability","avg_decimal_odds",
          "roi_ci_low","roi_ci_high","hit_ratio_ci_low","hit_ratio_ci_high","risk_basis","threshold","is_tuning_diagnostic","strategy_id"]
    def view(frame):
        result=frame[[c for c in core if c in frame]].copy()
        result.insert(2,"strategy_description",frame.apply(lambda row:_label(row).replace("\n"," | "),axis=1))
        return result
    sheets={"Read me":readme,"Metrics overview":view(metrics),"All strategies":metrics,"First locked test":view(first),"Ablation pooled":view(pooled)}
    # Scope sheets retain source scope labels; sanitise only Excel sheet names.
    for scope,part in metrics.groupby("scope",sort=True,dropna=False):
        name="Scope "+str(scope)
        for bad in "[]:*?/\\":name=name.replace(bad,"_")
        name=name[:31];base=name;n=2
        while name in sheets:name=base[:27]+f"_{n}";n+=1
        sheets[name]=view(part)
    sheets.update({"Prospective status":prospective,"Metric definitions":definitions,
                   "Input provenance":pd.DataFrame([{"path":p,**record} for p,record in registry.items()])})
    if arbitrage is not None:sheets["Arbitrage screen"]=arbitrage
    with pd.ExcelWriter(path,engine="openpyxl") as writer:
        for sheet,frame in sheets.items():
            safe=frame.copy()
            for c in safe:
                if safe[c].dtype=="object":safe[c]=safe[c].map(lambda x: json.dumps(x,default=str) if isinstance(x,(list,dict)) else x)
                if pd.api.types.is_numeric_dtype(safe[c]):safe[c]=safe[c].map(lambda x:"+∞" if x==float("inf") else "−∞" if x==float("-inf") else x)
            safe.to_excel(writer,sheet_name=sheet,index=False,na_rep="—",inf_rep="+∞")
            ws=writer.sheets[sheet];ws.freeze_panes="A2";ws.auto_filter.ref=ws.dimensions
            ws.sheet_properties.pageSetUpPr.fitToPage=True;ws.page_setup.orientation="landscape";ws.page_setup.paperSize=ws.PAPERSIZE_A4
            ws.page_setup.fitToWidth=1;ws.page_setup.fitToHeight=0
            for cell in ws[1]:cell.font=Font(bold=True,color="FFFFFF");cell.fill=PatternFill("solid",fgColor="12364D");cell.alignment=Alignment(wrap_text=True,vertical="center")
            ws.row_dimensions[1].height=36
            for j,c in enumerate(frame.columns,start=1):
                width=min(44,max(13,len(str(c))*.85))
                if c in ("Description","Definition","reason","path"):width=65
                if c=="strategy_description":width=48
                ws.column_dimensions[get_column_letter(j)].width=width
                percent_col=c in {"hit_ratio","hit_ratio_ci_low","hit_ratio_ci_high","trading_frequency","trade_rate_all_predictions","active_date_fraction","expected_roi","roi","roi_ci_low","roi_ci_high","max_drawdown","return_on_initial_capital","avg_selected_probability"}
                for cell in list(ws.columns)[j-1][1:]:
                    cell.alignment=Alignment(vertical="top",wrap_text=isinstance(cell.value,str) and len(cell.value)>65)
                    if isinstance(cell.value,(int,float)) and not isinstance(cell.value,bool):cell.number_format="0.00%" if percent_col else "#,##0.0000"


def _sha(path):
    return hashlib.sha256(path.read_bytes()).hexdigest()


def build_report(results: Path) -> dict:
    results=Path(results).resolve()
    metric_path=results/"strategy_metrics.parquet";status_path=results/"prospective_status.parquet"
    if not metric_path.exists() or not status_path.exists():raise FileNotFoundError("Expected strategy_metrics.parquet and prospective_status.parquet in --results")
    # Capture provenance before rendering and verify every recorded input afterwards.
    paths=[p for p in results.iterdir() if p.is_file()]
    paths.append(Path(__file__).resolve())
    if (PROJECT_ROOT/"epl_strategy/METHODOLOGY.md").exists():paths.append(PROJECT_ROOT/"epl_strategy/METHODOLOGY.md")
    for folder in (PROJECT_ROOT/"epl_report",PROJECT_ROOT/"epl_modeling/results",PROJECT_ROOT/"epl_modeling/ablation_results"):
        if folder.exists():paths.extend(p for p in folder.rglob("*") if p.is_file() and "__pycache__" not in p.parts and "previews" not in p.parts)
    paths=sorted(set(paths))
    def label(path):
        try:return path.relative_to(PROJECT_ROOT).as_posix()
        except ValueError:return str(path)
    registry={label(p):{"sha256":_sha(p),"bytes":p.stat().st_size} for p in paths}
    metrics=pd.read_parquet(metric_path);prospective=pd.read_parquet(status_path)
    missing=set(REQUIRED)-set(metrics.columns)
    if missing:raise ValueError("Missing strategy metrics fields: "+", ".join(sorted(missing)))
    arbitrage_path=results/"arbitrage_screen.parquet"
    arbitrage=pd.read_parquet(arbitrage_path) if arbitrage_path.exists() else None
    first,pooled=_selected_groups(metrics)
    definitions=_definitions();pages,tables=_pages(first,pooled,metrics,prospective,arbitrage)
    out=results/"report";out.mkdir(exist_ok=True)
    metrics.to_csv(out/"strategy_metrics_full.csv",index=False)
    prospective.to_csv(out/"prospective_status_full.csv",index=False)
    if arbitrage is not None:arbitrage.to_csv(out/"arbitrage_screen_full.csv",index=False)
    definitions.to_csv(out/"metric_definitions.csv",index=False)
    for name,frame in tables.items():frame.to_csv(out/(name+".csv"),index=False)
    md=out/"EPL_strategy_metrics_supplement.md";pdf=md.with_suffix(".pdf");docx=md.with_suffix(".docx")
    md.write_text(_markdown(pages),encoding="utf-8")
    _pdf(pages,pdf);_docx(pages,docx)
    workbook=out/"EPL_strategy_metrics.xlsx";_excel(workbook,metrics,prospective,arbitrage,definitions,first,pooled,registry)
    reader=PdfReader(pdf)
    if len(reader.pages)!=len(pages):raise AssertionError(f"Supplement overflow: {len(reader.pages)} pages versus {len(pages)} planned")
    doc=Document(docx)
    if len(doc.tables)!=4 or len(doc.sections)!=len(pages):raise AssertionError("Editable document table/section count mismatch")
    changed=[label(p) for p in paths if _sha(p)!=registry[label(p)]["sha256"]]
    if changed:raise AssertionError("Report inputs changed while rendering: "+str(changed))
    produced=[p for p in out.iterdir() if p.is_file() and p.name!="report_manifest.json"]
    outputs={p.name:{"sha256":_sha(p),"bytes":p.stat().st_size} for p in sorted(produced)}
    manifest={"status":"draft_companion_supplement_not_final_submission","draft_date":DRAFT_DATE,
              "built_at_utc":datetime.now(timezone.utc).isoformat(),"strategy_period_rows":len(metrics),
              "scope_counts":{str(k):int(v) for k,v in metrics.groupby("scope",dropna=False).size().items()},
              "displayed_first_locked_test_rows":len(first),"displayed_pooled_ablation_rows":len(pooled),
              "prospective_status_slots":len(prospective),"pdf_pages":len(reader.pages),"models_retrained":False,
              "strategies_reselected":False,"original_reports_and_study_sources_unchanged":not changed,
              "inputs":registry,"outputs":outputs,
              "expected_metric_basis":"Model-implied probability-weighted P&L components at recorded actual stakes/archived odds.",
              "pooled_risk_basis":"Separate episode capital resets; pooled drawdown is worst episode, not stitched equity.",
              "limits":"Retrospective conditional diagnostics; pending statuses do not establish prospective or executable trading returns."}
    (out/"report_manifest.json").write_text(json.dumps(manifest,indent=2,ensure_ascii=False,allow_nan=False),encoding="utf-8")
    return {"output_directory":str(out),"strategy_rows":len(metrics),"pdf_pages":len(reader.pages),"inputs_hashed":len(registry),"source_files_unchanged":not changed}


def main():
    parser=argparse.ArgumentParser(description=__doc__)
    parser.add_argument("--results",type=Path,required=True,help="Saved analytics directory; new files go in its report/ subdirectory.")
    args=parser.parse_args()
    print(json.dumps(build_report(args.results),indent=2))


if __name__=="__main__":main()
