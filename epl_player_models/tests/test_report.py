"""The expanded report must retain every account and exact strategy definition."""
import json

import numpy as np
import pandas as pd
import pytest
from docx import Document
from openpyxl import load_workbook
from pypdf import PdfReader

from epl_player_models import report
from epl_strategy_lab.catalog import build_catalog


def catalog():
    rows = []
    for group in report.GROUPS:
        for family in report.FAMILIES:
            for price in ("best", "reference"):
                for staking in ("fixed", "kelly"):
                    rows.append({"strategy_id": f"P{len(rows) + 1:03d}", "model_key": "pm_" + family + "_" + group,
                        "model_label": group + " / " + family, "feature_group": group, "model_type": family,
                        "name": group + " / " + family + " / " + price + " / " + staking,
                        "odds_mode": price, "staking": staking, "fixed_stake": 1000., "min_edge": .05,
                        "kelly_fraction": .25, "max_bet_fraction": .02})
    return rows


def graph(ids, date_count, anchor, last):
    dates = pd.date_range(pd.Timestamp(anchor) + pd.Timedelta(days=1), last, periods=date_count).normalize()
    return pd.DataFrame([{"strategy_id": sid, "match_date": date, "end_bankroll": 100000., "is_initial": initial}
        for sid in ids for date, initial in [(pd.Timestamp(anchor), True)] + [(d, False) for d in dates]])


@pytest.fixture
def finished_data(monkeypatch):
    monkeypatch.setattr(report.prior, "_validate", lambda data: None)
    metric = {"initial_capital_sum_units": 100000., "profit": 0., "bet_count": 0, "hit_ratio": np.nan,
        "trading_frequency": 0., "avg_winning_gain_units": np.nan, "avg_losing_loss_units": np.nan,
        "expected_gain_per_trade_units": np.nan, "expected_loss_per_trade_units": np.nan,
        "roi": np.nan, "return_on_initial_capital": 0., "max_drawdown": 0.}
    single = pd.DataFrame([{**metric, "strategy_id": sid, "prediction_count": 374, "matched_valid_odds_count": 374} for sid in report.NEW_IDS])
    continuous = single.assign(prediction_count=1470, matched_valid_odds_count=1470)
    annual = pd.concat([single.assign(season=season) for season in report.SEASONS], ignore_index=True)
    models = [{"model_key": "pm_" + family + "_" + group, "feature_group": group, "model_type": family,
        "label": group + " / " + family, "n_features": 1} for group in report.GROUPS for family in report.FAMILIES]
    pooled = pd.DataFrame([{**m, "outer_season": "pooled", "n_matches": 1470, "accuracy": .5,
        "log_loss": 1., "brier_score": .6, "classwise_ece_5": .1, "classwise_ece_10": .1, "classwise_ece_15": .1} for m in models])
    return {"new_strategy_catalog": catalog(), "model_catalog": {"models": models,
        "groups": {group: ["feature"] for group in report.GROUPS}, "chronological_roles": {}},
        "metrics_single": single, "metrics_continuous": continuous, "metrics_seasonal": annual,
        "model_metrics_pooled": pooled, "model_metrics_seasonal": pd.concat([pooled.assign(outer_season=s, n_matches=n)
            for s, n in zip(report.SEASONS, (368, 360, 368, 374))], ignore_index=True),
        "predictions": pd.DataFrame({"p_H": [.5], "p_D": [.3], "p_A": [.2]}),
        "feature_coverage": pd.DataFrame({"season": ["2022-23"], "observed_matches": [368]}),
        "paired_feature_comparisons": pd.DataFrame([{"feature_group": "player_form", "model_type": "logistic",
            "reference_group": "team_control", "reference_model_type": "logistic", "outer_season": "pooled",
            "metric": metric, "difference": -.01, "ci_lower": -.03, "ci_upper": .01, "n_matches": 1470}
            for metric in ("log_loss", "brier_score", "accuracy")]),
        "future_predictions": pd.DataFrame({"match_id": ["future"], "p_H": [.5], "p_D": [.3], "p_A": [.2]}),
        "study_status": {"report_date_new_york": "2026-10-20"},
        "features": pd.DataFrame({"feature": [1.5]}), "future_features": pd.DataFrame({"feature": [1.6]}),
        "graph_2025": graph(report.ALL_IDS, 113, "2025-08-14", "2026-05-24"),
        "graph_continuous": graph(report.CONTINUOUS_IDS, 453, "2022-08-04", "2026-05-24"),
        "prior": {"historical": {"catalog": build_catalog()}}}


def test_all_original_and_new_strategies_share_graphs_and_exact_initial_cash(finished_data):
    report._validate(finished_data)
    assert finished_data["graph_2025"].strategy_id.nunique() == 116
    assert finished_data["graph_continuous"].strategy_id.nunique() == 88


def test_reject_missing_curve_or_initial_cash_reset(finished_data):
    frame = finished_data["graph_2025"]
    finished_data["graph_2025"] = frame.loc[frame.strategy_id.ne("P032")].copy()
    with pytest.raises(ValueError, match="omits a strategy"):
        report._validate(finished_data)
    finished_data["graph_2025"] = frame
    frame.loc[frame.is_initial & frame.strategy_id.eq("P001"), "end_bankroll"] = 400000.
    with pytest.raises(ValueError, match="100,000 initial anchor"):
        report._validate(finished_data)


def test_reject_an_unmatched_strategy_population_or_endpoint(finished_data):
    finished_data["metrics_single"].loc[0, "matched_valid_odds_count"] = 373
    with pytest.raises(ValueError, match="common fixture population"):
        report._validate(finished_data)
    finished_data["metrics_single"].loc[0, "matched_valid_odds_count"] = 374
    finished_data["metrics_single"].loc[0, "profit"] = 1.
    with pytest.raises(ValueError, match="endpoint differs"):
        report._validate(finished_data)


def test_expanded_appendix_preserves_exact_old_and_new_definitions(finished_data):
    entries = report._appendix(finished_data)
    assert [e["strategy_id"] for e in entries] == report.ALL_IDS
    assert [json.loads(e["exact_config"]) for e in entries[:84]] == build_catalog()
    assert [json.loads(e["exact_config"]) for e in entries[84:]] == catalog()
    assert entries[-1]["strategy_id"] == "P032"
    assert "strictly above 5%" in entries[-1]["explanation"]
    assert "0.25" in entries[-1]["explanation"] and "0.02" in entries[-1]["explanation"]


def test_new_strategy_metrics_keep_every_requested_payoff_and_frequency(finished_data):
    display = report._strategy_table(finished_data["metrics_continuous"])
    assert display.ID.tolist() == report.NEW_IDS
    assert "Hit\n%" in display and "Trade\nfreq. %" in display
    assert "Model expected\ngain/trade" in display and "Model expected\nloss/trade" in display
    assert "Avg. winning\ngain" in display and "Avg. losing\nloss" in display
    paired = report._paired_table(finished_data)
    assert len(paired) == 1 and paired.iloc[0].Metric == "log_loss"
    assert paired.iloc[0].Difference == "−0.0100"


def test_coverage_summary_uses_only_evaluated_fixture_histories(finished_data):
    finished_data["predictions"] = pd.DataFrame({"match_id": ["m1", "m2"]})
    finished_data["feature_coverage"] = pd.DataFrame({"match_id": ["m1", "m2", "future"],
        "season": ["2022-23"] * 3, "side": ["home"] * 3, "observed_matches5": [5, 3, 0],
        "observed_matches10": [10, 8, 0], "xi_selected_players": [11, 9, 0], "prior_positive_pool_count": [20, 18, 0]})
    summary = report._coverage_table(finished_data)
    assert len(summary) == 1 and summary.iloc[0]["Target\nsides"] == "2"
    assert summary.iloc[0]["Mean observed\nprior slots / 5"] == "4.00"
    assert summary.iloc[0]["Full prior\nprojected XI"] == "50.00%"


def test_workbook_contains_features_all_metrics_future_and_external_links(tmp_path, finished_data):
    entries = report._appendix(finished_data)
    ledger = tmp_path / "bets_continuous.parquet"
    pd.DataFrame({"stake": [1000.]}).to_parquet(ledger)
    path = tmp_path / "player.xlsx"
    rows = report._workbook(path, finished_data, entries, {"input": {"sha256": "a" * 64, "bytes": 1}}, tmp_path)
    assert rows["New continuous metrics"] == 32 and rows["New annual metrics"] == 128
    assert rows["Final116 appendix"] == 116 and rows["Engineered match features"] == 1
    assert rows["Future match features"] == 1 and rows["External files"] == 1
    book = load_workbook(path)
    assert sum(row[1].hyperlink is not None for row in book["Read me"]) == 2
    sheet = book["Every model prediction"]
    assert sheet[2][0].value == .5 and sheet[2][0].number_format == "0.00%"
    sheet = book["New continuous metrics"]
    cols = [c.value for c in sheet[1]]
    assert sheet.cell(2, cols.index("profit") + 1).number_format == "#,##0.0000"
    book.close()


def test_both_offline_graphs_show_every_trace_and_end_with_all116_explanations(tmp_path, finished_data):
    data = dict(finished_data)
    for key in ("graph_2025", "graph_continuous"):
        frame = data[key]
        data[key] = frame.loc[frame.is_initial | frame.match_date.eq(frame.match_date.max())].copy()
    info = report._graphs(data, tmp_path, report._appendix(data))
    assert info["2025"]["trace_ids"] == report.ALL_IDS and info["2025"]["initially_visible_traces"] == 116
    assert info["continuous"]["trace_ids"] == report.CONTINUOUS_IDS and info["continuous"]["initially_visible_traces"] == 88
    for key in ("2025", "continuous"):
        text = (tmp_path / ("all_strategies_" + key + "_100k.html")).read_text(encoding="utf-8")
        assert "plotly.js" in text and "<script src=" not in text
        assert text.count("<article id='") == 116
        assert "<article id='P032'>" in text


def test_pdf_word_markdown_end_with_p032_after_all116_entries(tmp_path, finished_data):
    entries = report._appendix(finished_data)
    pages = [{"title": report.TITLE, "orientation": "portrait", "blocks": [
        {"type": "paragraph", "text": "Player-focused models. All new redesign results remain retrospective."}]}]
    pdf = tmp_path / "report.pdf"; docx = tmp_path / "report.docx"
    report._pdf(pages, entries, pdf); report._docx(pages, entries, docx)
    reader = PdfReader(pdf)
    text = " ".join(" ".join((p.extract_text() or "").split()) for p in reader.pages)
    assert "PLAYER MODELS AND COMPLETE STRATEGY COMPARISON" in text
    assert all(" ".join(e["title"].split()) in text for e in entries)
    assert "P032" in reader.pages[-1].extract_text()
    paragraphs = [p.text for p in Document(docx).paragraphs]
    assert json.loads(paragraphs[-1].split(": ", 1)[1])["strategy_id"] == "P032"
    md = report._markdown(pages, entries)
    appendix = md.split("## " + report.APPENDIX_TITLE)[1]
    assert appendix.count("### S") == 84 and appendix.count("### P") == 32
