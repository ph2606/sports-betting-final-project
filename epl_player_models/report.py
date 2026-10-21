"""Player-focused forecasts, complete account comparisons and final rule appendix.

Read completed tables only. Historical reports and their sources stay frozen;
new model redesigns are retrospective, while new future forecasts stay pending.
"""
from __future__ import annotations

import argparse
import base64
from datetime import datetime, timezone
import html
import json
import os
from pathlib import Path
import shutil

import matplotlib
matplotlib.use("Agg")
import matplotlib.dates as mdates
import matplotlib.pyplot as plt
from matplotlib.colors import to_hex
import numpy as np
import pandas as pd
import plotly.graph_objects as go
from docx import Document
from docx.enum.section import WD_ORIENT, WD_SECTION_START
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt
from PIL import Image as PILImage
from pypdf import PdfReader
from reportlab.lib import colors
from reportlab.lib.pagesizes import landscape, letter
from reportlab.lib.styles import ParagraphStyle
from reportlab.platypus import (BaseDocTemplate, Frame, Image, KeepTogether,
    NextPageTemplate, PageBreak, PageTemplate, Paragraph, Table, TableStyle)

from epl_prospective import report as prior
from epl_strategy.report import number
from epl_strategy_lab import report as lab

ROOT = Path(__file__).resolve().parents[1]
PRIOR = ROOT / "epl_prospective/results/20261020_first_outcomes"
TITLE = "EPL project: player-focused models and complete strategy comparison"
APPENDIX_TITLE = prior.APPENDIX_TITLE
BODY_WEIGHTS = prior.BODY_WEIGHTS
OLD_IDS = [f"S{i:03d}" for i in range(1, 85)]
NEW_IDS = [f"P{i:03d}" for i in range(1, 33)]
ALL_IDS = OLD_IDS + NEW_IDS
CONTINUOUS_IDS = [f"S{i:03d}" for i in range(17, 73)] + NEW_IDS
GROUPS = ("team_control", "player_form", "player_squad", "team_player_squad")
FAMILIES = ("logistic", "hist_boost")
SEASONS = ("2022-23", "2023-24", "2024-25", "2025-26")
TABLE_NAMES = ("predictions", "model_metrics_pooled", "model_metrics_seasonal",
    "paired_feature_comparisons", "feature_coverage", "metrics_continuous", "metrics_single",
    "metrics_seasonal", "future_predictions")
OPTIONAL_TABLES = ("features", "future_features", "parameters", "fold_parameters", "tuning_trials", "model_fit_audit", "season_stability",
    "fitted_choices", "chronological_roles", "logistic_coefficients")


def _load(output: Path) -> dict:
    data = {name: pd.read_parquet(output / (name + ".parquet")) for name in TABLE_NAMES}
    for name in OPTIONAL_TABLES:
        if (output / (name + ".parquet")).exists(): data[name] = pd.read_parquet(output / (name + ".parquet"))
    for name in ("model_catalog", "new_strategy_catalog", "study_status"):
        data[name] = json.loads((output / (name + ".json")).read_text(encoding="utf-8"))
    for name in ("2025", "continuous"):
        data["graph_" + name] = pd.read_csv(output / ("graph_equity_" + name + "_all.csv"), parse_dates=["match_date"])
    data["prior"] = prior._load(PRIOR)
    return data


def _validate_graph(graph: pd.DataFrame, ids: list[str], *, date_count: int, anchor: str) -> None:
    if not pd.api.types.is_bool_dtype(graph.is_initial):
        flags = graph.is_initial.astype(str).str.lower()
        if not flags.isin(["true", "false"]).all(): raise ValueError("Invalid graph initial-point flag")
        graph["is_initial"] = flags.eq("true")
    if set(graph.strategy_id) != set(ids) or len(graph) != len(ids) * (date_count + 1):
        raise ValueError("A combined graph omits a strategy or settlement date")
    if graph.duplicated(["strategy_id", "match_date"]).any() or not np.isfinite(graph.end_bankroll).all() or graph.end_bankroll.lt(-1e-6).any():
        raise ValueError("Invalid combined graph cash path")
    expected_dates = set(graph.loc[~graph.is_initial, "match_date"])
    if len(expected_dates) != date_count: raise ValueError("Incorrect common graph date universe")
    for sid, part in graph.groupby("strategy_id"):
        first = part.loc[part.is_initial]
        if len(first) != 1 or first.match_date.iloc[0] != pd.Timestamp(anchor) or not np.isclose(first.end_bankroll.iloc[0], 100000., atol=1e-8, rtol=0):
            raise ValueError("Every strategy requires one exact 100,000 initial anchor")
        if set(part.loc[~part.is_initial, "match_date"]) != expected_dates:
            raise ValueError("A combined curve uses a different settlement calendar")


def _validate(data: dict) -> None:
    catalog = data["new_strategy_catalog"]
    if [c["strategy_id"] for c in catalog] != NEW_IDS: raise ValueError("New strategies must preserve P001–P032 in order")
    expected = {(g, f) for g in GROUPS for f in FAMILIES}
    learned = data["model_metrics_pooled"]
    actual = set(learned.loc[learned.feature_group.isin(GROUPS), ["feature_group", "model_type"]].itertuples(index=False, name=None))
    if not expected.issubset(actual): raise ValueError("A declared player-model alternative is missing")
    for key, rows, matches in (("metrics_single", 32, 374), ("metrics_continuous", 32, 1470), ("metrics_seasonal", 128, None)):
        frame = data[key]
        if len(frame) != rows or set(frame.strategy_id) != set(NEW_IDS): raise ValueError("Incomplete new strategy metrics: " + key)
        if key != "metrics_seasonal" and frame.strategy_id.duplicated().any(): raise ValueError("Repeated new strategy metric")
        if not np.allclose(frame.initial_capital_sum_units, 100000., rtol=0, atol=1e-6): raise ValueError("Each new account must start with 100,000")
        if matches is not None and (not frame.prediction_count.eq(matches).all() or not frame.matched_valid_odds_count.eq(matches).all()):
            raise ValueError("A new strategy silently changes its common fixture population")
    _validate_graph(data["graph_2025"], ALL_IDS, date_count=113, anchor="2025-08-14")
    _validate_graph(data["graph_continuous"], CONTINUOUS_IDS, date_count=453, anchor="2022-08-04")
    for graph_key, metric_key in (("graph_2025", "metrics_single"), ("graph_continuous", "metrics_continuous")):
        graph = data[graph_key].sort_values("match_date").groupby("strategy_id").tail(1).set_index("strategy_id")
        metrics = data[metric_key].set_index("strategy_id").reindex(NEW_IDS)
        if not np.allclose(graph.reindex(NEW_IDS).end_bankroll, 100000. + metrics.profit, atol=1e-5, rtol=0):
            raise ValueError("A new graph endpoint differs from its account profit")
    prior._validate(data["prior"])


def _group_label(group: str) -> str:
    return {"team_control": "Team control", "player_form": "Player form", "player_squad": "Player squad",
        "team_player_squad": "Team + player squad"}.get(str(group), str(group))


def _family_label(family: str) -> str:
    return {"logistic": "Logistic", "hist_boost": "Hist. boosted trees"}.get(str(family), str(family))


def _label_map(data: dict) -> dict:
    return {str(row["model_key"]): str(row.get("label", row["model_key"])) for row in data["model_catalog"].get("models", [])}


def _row_label(row, labels: dict) -> str:
    if "label" in row and pd.notna(row["label"]): return str(row["label"])
    if "model_key" in row and str(row["model_key"]) in labels: return labels[str(row["model_key"])]
    return _group_label(row.get("feature_group", "")) + " / " + _family_label(row.get("model_type", ""))


def _appendix(data: dict) -> list[dict]:
    entries = prior.execution._appendix(data["prior"]["historical"]["catalog"])
    labels = _label_map(data)
    for config in data["new_strategy_catalog"]:
        sid = config["strategy_id"]; model = str(config.get("model_label", labels.get(config.get("model_key"), config.get("model_key", sid))))
        quote = "best archived eligible individual-book prices" if config["odds_mode"] == "best" else "archived Bet365 reference prices"
        sizing = ("Request a flat stake of 1,000 monetary units." if config["staking"] == "fixed" else
            "Request B × min(0.02, 0.25 × max((p × odds − 1)/(odds − 1), 0)), where B is the date-start cash balance.")
        explanation = (model + ". Use " + quote + ". Choose the H/D/A outcome with the largest p × odds − 1, with H/D/A tie order; "
            "trade only when that edge is strictly above 5%. " + sizing +
            " Allocate every same-date request before settlement and scale proportionally to available cash. No same-date winning proceeds fund another request. "
            "Each continuous account starts once at 100,000 across 2022/23–2025/26; the common-2025 and annual diagnostic accounts start independently at 100,000. "
            "Probabilities use the relevant chronologically fitted and calibrated player-model fold. This retrospective rule is not an executed prospective bet.")
        entries.append({"strategy_id": sid, "title": sid + " — " + str(config.get("name", model + " | " + config["odds_mode"] + " | " + config["staking"])),
            "origin": "new_player_model_strategy", "coverage": "Four chronological evaluation seasons and the separate common 2025/26 account.",
            "explanation": explanation, "exact_config": lab._json(config)})
    if [e["strategy_id"] for e in entries] != ALL_IDS: raise ValueError("The final appendix must contain every original and new strategy")
    return entries


def _model_table(data: dict) -> pd.DataFrame:
    labels = _label_map(data); rows = []
    for _, row in data["model_metrics_pooled"].iterrows():
        rows.append({"Model": _row_label(row, labels), "Matches": number(row.n_matches, 0),
            "Accuracy": number(row.accuracy, percent=True), "Log loss": number(row.log_loss, 4),
            "Brier": number(row.brier_score, 4), "ECE 5": number(row.get("classwise_ece_5", np.nan), 4),
            "ECE 10": number(row.get("classwise_ece_10", np.nan), 4), "ECE 15": number(row.get("classwise_ece_15", np.nan), 4)})
    return pd.DataFrame(rows)


def _season_model_table(data: dict) -> pd.DataFrame:
    labels = _label_map(data); frame = data["model_metrics_seasonal"].copy()
    frame["_label"] = frame.apply(lambda r: _row_label(r, labels), axis=1)
    rows = []
    for label, part in frame.groupby("_label", sort=False):
        row = {"Model": label}
        for metric, title in (("log_loss", "LL"), ("brier_score", "Brier")):
            for season in SEASONS:
                values = part.loc[part.outer_season.eq(season), metric]
                row[season[:4] + "\n" + title] = number(values.iloc[0], 4) if len(values) == 1 else "—"
        rows.append(row)
    return pd.DataFrame(rows)


def _coverage_table(data: dict) -> pd.DataFrame:
    frame = data["feature_coverage"].copy()
    if "match_id" in frame and "match_id" in data["predictions"]:
        frame = frame.loc[frame.match_id.isin(set(data["predictions"].match_id))]
    if {"season", "side", "observed_matches5", "observed_matches10", "xi_selected_players"}.issubset(frame):
        frame = frame.loc[frame.season.astype(str).isin(SEASONS)]
        rows = []
        for (season, side), part in frame.groupby(["season", "side"], sort=True):
            rows.append({"Season": season, "Side": side, "Target\nsides": str(len(part)),
                "Mean observed\nprior slots / 5": number(part.observed_matches5.mean()),
                "Mean observed\nprior slots / 10": number(part.observed_matches10.mean()),
                "Full prior\nprojected XI": number(part.xi_selected_players.eq(11).mean(), percent=True),
                "Min selected\nplayers": number(part.xi_selected_players.min(), 0),
                "Mean prior\npositive pool": number(part.prior_positive_pool_count.mean()) if "prior_positive_pool_count" in part else "—"})
        return pd.DataFrame(rows)
    preferred = [c for c in ("season", "side", "target_sides", "sides_with_history", "full_prior_xi", "observed_matches") if c in frame]
    return frame[preferred].head(16) if preferred else frame.iloc[:12, :7]


def _paired_table(data: dict) -> pd.DataFrame:
    frame = data["paired_feature_comparisons"]
    if "outer_season" in frame: frame = frame.loc[frame.outer_season.astype(str).eq("pooled")]
    if "metric" in frame: frame = frame.loc[frame.metric.eq("log_loss")]
    rows = []
    for _, row in frame.iterrows():
        group = row.get("candidate_feature_group", row.get("feature_group", ""))
        family = row.get("candidate_model_type", row.get("model_type", ""))
        ref_group = row.get("reference_feature_group", row.get("reference_group", ""))
        ref_family = row.get("reference_model_type", "")
        rows.append({"Candidate": _group_label(group) + " / " + _family_label(family),
            "Reference": _group_label(ref_group) + (" / " + _family_label(ref_family) if ref_family else ""),
            "Metric": str(row.metric), "Difference": number(row.difference, 4),
            "95% CI\nlower": number(row.ci_lower, 4), "95% CI\nupper": number(row.ci_upper, 4),
            "Matches": number(row.n_matches, 0)})
    return pd.DataFrame(rows)


def _strategy_table(frame: pd.DataFrame) -> pd.DataFrame:
    ordered = frame.set_index("strategy_id").reindex(NEW_IDS)
    return pd.DataFrame([{"ID": sid, "Trades": number(r.bet_count, 0), "Hit\n%": number(r.hit_ratio, percent=True),
        "Trade\nfreq. %": number(r.trading_frequency, percent=True), "Avg. winning\ngain": number(r.avg_winning_gain_units),
        "Avg. losing\nloss": number(r.avg_losing_loss_units), "Model expected\ngain/trade": number(r.expected_gain_per_trade_units),
        "Model expected\nloss/trade": number(r.expected_loss_per_trade_units), "ROI on\nstakes": number(r.roi, percent=True),
        "Return on\n100,000": number(r.return_on_initial_capital, percent=True), "Max.\ndrawdown": number(r.max_drawdown, percent=True)}
        for sid, r in ordered.iterrows()])


def _stability_table(data: dict) -> pd.DataFrame:
    frame = data["metrics_seasonal"].copy()
    season_column = "season" if "season" in frame else "outer_season"
    rows = []
    for sid, part in frame.groupby("strategy_id", sort=False):
        row = {"ID": sid, "Positive\nseasons": str(int(part.profit.gt(0).sum())) + " / 4"}
        for season in SEASONS:
            values = part.loc[part[season_column].eq(season), "return_on_initial_capital"]
            row[season + "\nreturn"] = number(values.iloc[0], percent=True) if len(values) == 1 else "—"
        row["Worst annual\nreturn"] = number(part.return_on_initial_capital.min(), percent=True)
        row["Best annual\nreturn"] = number(part.return_on_initial_capital.max(), percent=True)
        rows.append(row)
    return pd.DataFrame(rows).set_index("ID").reindex(NEW_IDS).reset_index()


def _graphs(data: dict, figures: Path, entries: list[dict]) -> dict:
    figures.mkdir(exist_ok=True)
    for name in prior.GRAPH_NAMES: shutil.copyfile(PRIOR / "figures" / name, figures / name)
    catalog = data["prior"]["historical"]["catalog"] + data["new_strategy_catalog"]
    configs = {c["strategy_id"]: c for c in catalog}
    palette = plt.get_cmap("tab20").colors
    lines = ["solid", "dash", "dashdot", "dot", "longdash", "longdashdot"]
    styles = {sid: (to_hex(palette[i % 20]), lines[(i // 20) % len(lines)]) for i, sid in enumerate(ALL_IDS)}
    dash = {"solid": "-", "dash": "--", "dashdot": "-.", "dot": ":", "longdash": (0, (8, 3)), "longdashdot": (0, (8, 3, 1, 3))}
    appendix_html = "".join("<article id='" + e["strategy_id"] + "'><h3>" + html.escape(e["title"]) + "</h3><p>" + html.escape(e["explanation"]) + "</p><pre>" + html.escape(e["exact_config"]) + "</pre></article>" for e in entries)
    info = {}
    for key, ids, title in (("2025", ALL_IDS, "All 116 strategies · 2025/26 · initial 100,000 each"),
            ("continuous", CONTINUOUS_IDS, "All 88 compatible strategies · four seasons · initial 100,000 each")):
        graph = data["graph_" + key]
        fig, ax = plt.subplots(figsize=(13.6, 7.8)); fig.subplots_adjust(left=.085, right=.985, top=.91, bottom=.35)
        interactive = go.Figure()
        for sid in ids:
            config = configs[sid]; part = graph.loc[graph.strategy_id.eq(sid)].sort_values("match_date")
            color, line = styles[sid]
            ax.step(part.match_date, part.end_bankroll, where="post", color=color, linestyle=dash[line], linewidth=.85, alpha=.88, label=sid)
            interactive.add_trace(go.Scatter(x=part.match_date.dt.strftime("%Y-%m-%d").tolist(), y=part.end_bankroll.astype(float).tolist(),
                mode="lines", name=sid, uid=sid, visible=True, line={"color": color, "dash": line, "width": 1.4, "shape": "hv"},
                customdata=[str(config.get("name", config.get("model_label", sid)))] * len(part),
                hovertemplate=sid + " · %{customdata}<br>Date: %{x|%Y-%m-%d}<br>Balance: %{y:,.2f} monetary units<extra></extra>"))
        ax.set_title(title, fontsize=12); ax.set_xlabel("Calendar date (balance after date-level settlement)"); ax.set_ylabel("Balance (monetary units)")
        ax.yaxis.set_major_formatter(matplotlib.ticker.StrMethodFormatter("{x:,.0f}"))
        ax.xaxis.set_major_locator(mdates.MonthLocator(interval=2 if key == "2025" else 6)); ax.xaxis.set_major_formatter(mdates.DateFormatter("%b %Y"))
        ax.grid(alpha=.2); ax.spines[["top", "right"]].set_visible(False); ax.set_xlim(graph.match_date.min(), graph.match_date.max())
        handles, names = ax.get_legend_handles_labels()
        fig.legend(handles, names, loc="lower center", ncol=12 if key == "2025" else 11, fontsize=7.3, frameon=False,
            bbox_to_anchor=(.5, .012), handlelength=2, columnspacing=1.4, labelspacing=.65)
        stem = "all_strategies_" + key + "_100k"
        with plt.rc_context({"svg.fonttype": "none"}):
            for suffix in (".png", ".svg"): fig.savefig(figures / (stem + suffix), dpi=240)
        plt.close(fig)
        interactive.update_layout(title={"text": title, "x": .01, "y": .99, "yanchor": "top"}, template="plotly_white", height=1120,
            margin={"l": 85, "r": 35, "t": 115, "b": 340}, xaxis={"title": "Calendar date (date-close settlement)", "type": "date"},
            yaxis={"title": "Balance (monetary units)", "type": "linear", "tickformat": ",.0f"}, hovermode="closest",
            legend={"title": {"text": "Every strategy ID", "side": "top"}, "orientation": "h", "x": 0, "y": -.16,
                "entrywidth": 70, "entrywidthmode": "pixels", "font": {"size": 10}, "groupclick": "toggleitem"},
            updatemenus=[{"type": "buttons", "direction": "right", "x": 0, "y": 1.025, "yanchor": "bottom", "buttons": [
                {"label": "Show all", "method": "restyle", "args": [{"visible": True}]},
                {"label": "Hide all (keep legend)", "method": "restyle", "args": [{"visible": "legendonly"}]}]}])
        if len(interactive.data) != len(ids) or any(t.visible is not True for t in interactive.data): raise AssertionError("Every strategy must initially be visible")
        historical_images = "".join("<h2>Preserved historical graph</h2><img alt='Preserved historical account graph' src='data:image/png;base64," + base64.b64encode((figures / name).read_bytes()).decode("ascii") + "'>" for name in prior.GRAPH_NAMES)
        page = ("<!doctype html><html lang='en'><head><meta charset='utf-8'><meta name='viewport' content='width=device-width,initial-scale=1'><title>EPL player-model strategy comparison</title><style>body{font:15px system-ui;margin:24px;color:#173746}p{max-width:1100px;line-height:1.5}pre{white-space:pre-wrap;overflow-wrap:anywhere;font-size:12px;background:#f0f4f6;padding:12px}article{border-top:1px solid #cad5db;padding-top:12px}img{max-width:100%;height:auto}</style></head><body><h1>" + html.escape(title) +
            "</h1><p>Every curve is initially visible on one linear money-versus-time graph. These are separate hypothetical retrospective accounts, each starting at 100,000. They are not added into a portfolio. No trade or executable quote is claimed. The four-season view omits only the 28 earlier first-study recipes with no compatible saved forecasts before 2025. Model/strategy redesign remains exploratory; 2025 is not a fresh holdout.</p>" +
            interactive.to_html(full_html=False, include_plotlyjs=True, config={"responsive": True, "displaylogo": False}) + historical_images +
            "<h2>" + APPENDIX_TITLE + "</h2>" + appendix_html + "</body></html>")
        (figures / (stem + ".html")).write_text(page, encoding="utf-8")
        info[key] = {"trace_ids": ids, "initially_visible_traces": len(ids), "plotly_js_bundled": True, "linear_money_y_axis": True}
    return info


def _body(data: dict, figures: Path) -> tuple[list[dict], dict[str, pd.DataFrame]]:
    history, old_displays = prior._body(data["prior"], figures)
    for page_ in history:
        for block in page_["blocks"]:
            if block["type"] == "figure": block["path"] = figures / block["path"].name
    history[0]["title"] = "Preserved research snapshot: first prospective evaluation"
    history[0]["blocks"].insert(0, {"type": "paragraph", "text": "The following research snapshot is preserved from the October 20 first-outcome report. Its 84-strategy references describe the original catalog. The new player-model study above adds P001–P032; the expanded 116-rule appendix remains the final section of this combined report."})
    pages = []; displays = {}; status = data["study_status"]
    def page(title, orientation="portrait"):
        item = {"title": title, "orientation": orientation, "blocks": []}; pages.append(item); return item["blocks"]
    def p(blocks, text): blocks.append({"type": "paragraph", "text": text})
    def table(blocks, frame, caption): blocks.append({"type": "table", "frame": frame, "caption": caption})
    b = page(TITLE)
    p(b, "DRAFT FOR REVIEW | " + str(status.get("report_date_new_york", "2026-10-20")) + " | EPL component of the statistical-arbitrage final project.")
    p(b, "This study tests player-focused predictors directly rather than treating player data only as a small supplement to team statistics. Four declared feature blocks—team control, player form, player squad, and team plus player squad—are each fitted with multinomial logistic regression and histogram gradient boosting. The eight learned alternatives are compared on the same four chronological evaluation seasons, with market and earlier-study benchmarks kept visible.")
    p(b, "The new trading comparison adds 32 fully specified rules P001–P032: eight models × two price panels × two sizing rules. All start at 100,000. The full common-2025 graph contains the original 84 plus all 32 new rules on one set of axes; the continuous four-season graph contains all 56 compatible old rules plus the new 32. The final appendix explains every one of the 116 strategies, retaining every original configuration unchanged.")
    p(b, "These are new retrospective redesigns of already examined history. Chronological fitting protects each individual fold's information boundary but does not make 2025/26 a fresh holdout or establish a winning strategy. Prediction accuracy is separate from trade hit ratio, and model expected payoffs are conditional claims rather than demonstrated future returns. Newly generated player-model future probabilities are a separate forecast cohort and do not replace the earlier frozen prospective ledger.")
    groups = data["model_catalog"].get("groups", {})
    feature_rows = [{"Feature block": _group_label(group), "Declared inputs": len(columns), "Models": "Logistic + hist. boosted trees"} for group, columns in groups.items()]
    if feature_rows: table(b, pd.DataFrame(feature_rows), "Table M1. Exact ordered feature lists and model specifications are in the workbook. Counts precede fitted preprocessing and its missingness indicators.")
    b = page("What the player features observe, and what they cannot know")
    p(b, "Player form is computed from past player performance. The richer squad block uses a rule-based projected XI, positional summaries, prior minutes and squad depth derived from strictly earlier observations. The projected XI is a historical selection rule, not the target match's announced lineup or future minutes. Current roster labels, later transfers, target appearances and season-final aggregates cannot supply earlier predictors. Missing or incomplete player history remains visible through the declared coverage fields.")
    p(b, "Every fold keeps training, hyperparameter tuning, probability calibration and evaluation chronologically separate. Imputation and scaling are fitted only on the permitted fitting block. Histogram boosting's random internal early stopping is disabled; tuning uses the declared earlier season. The same feature/source identity and calendar rules apply to both model families and every block, so within-family comparisons isolate feature additions more clearly than comparing different algorithms.")
    p(b, "The primary test population is the same 1,470 quoted fixtures involving the present-day 2026/27 club cohort: 368 in 2022/23, 360 in 2023/24, 368 in 2024/25 and 374 in 2025/26. This cohort is not the full historical league and retains survivor/selection limitations. Earlier completed matches may update later lagged features; their outcomes cannot update the already frozen model parameters within that evaluation season. No 2026/27 label is used to fit or select these models.")
    p(b, "Coverage and role completeness are diagnostics, not performance-based filters. The complete feature frame, per-season coverage, exact fitted parameters and chronology are retained alongside probabilities. Source-specific fantasy definitions and position labels can change across seasons; historical player contributions and projected depth do not verify injuries, availability, transfers or a starting XI known before the target.")
    p(b, "When native historical role fields are unavailable, a season-fixed archived element_type supplies a disclosed position proxy. Its original publication clock is unverified. Position-based squad features therefore do not establish a fully timestamp-verified historical lineup feed. The player-form block does not require these role labels.")
    coverage = _coverage_table(data); displays["Player history coverage"] = coverage
    b = page("Player-history coverage and projected squad completeness", "landscape")
    table(b, coverage, "Table M1b. Coverage of the evaluated target-side histories where detailed rows are available. Prior observed fixture slots are not imputed target appearances. Every detailed source record and cutoff remains in the workbook; a complete projected XI is not confirmed player availability.")
    p(b, "The projection uses a fixed 1 goalkeeper, 4 defenders, 4 midfielders and 2 forwards, chosen by strictly prior minutes. It excludes candidates whose known historical transfer evidence contradicts the target club. It does not use the target match's minutes, target lineup or later player membership. Sparse positions remain visibly incomplete rather than acquiring invented players; differences between observed historical roles and future match participation remain a limitation.")
    pooled = _model_table(data); displays["Player model pooled scores"] = pooled
    b = page("Every model and benchmark on the same quoted fixtures", "landscape")
    table(b, pooled, "Table M2. Pooled chronological evaluation. Lower log loss/Brier are better. Accuracy is a classification score, not a betting hit ratio. ECE at 5/10/15 bins is descriptive. Shared fixtures and overlapping predictors make the rows dependent; no model is selected from this table.")
    seasonal = _season_model_table(data); displays["Player model seasonal losses"] = seasonal
    b = page("Predictive losses through the four evaluation seasons", "landscape")
    table(b, seasonal, "Table M3. All saved models and benchmarks, in source order. LL=multiclass log loss; Brier is the unscaled three-class squared probability error. Every annual accuracy/calibration score and sample count is in the workbook. A strong pooled score need not persist across years.")
    paired = _paired_table(data); displays["Player feature paired comparisons"] = paired
    for start in range(0, len(paired), 24):
        b = page("Paired log-loss feature comparisons" + (" (continued)" if start else ""), "landscape")
        table(b, paired.iloc[start:start + 24], "Table M4. Candidate minus reference on matched fixtures; negative log-loss differences favor the candidate. Whole-date bootstrap intervals are descriptive, condition on the recorded fits and do not adjust for this enlarged search. Every annual and pooled log-loss, Brier and accuracy comparison remains in the workbook.")
    for key, title, name in (("2025", "Every strategy on one common-season money graph", "all_strategies_2025_100k.png"),
            ("continuous", "Every compatible strategy on one continuous money graph", "all_strategies_continuous_100k.png")):
        b = page(title, "landscape")
        p(b, "All " + str(len(ALL_IDS) if key == "2025" else len(CONTINUOUS_IDS)) + " strategies are visible on the same axes, each as its own 100,000-unit account. Calendar time is horizontal and money is vertical. The one-season comparison shares 374 fixtures; the four-season comparison shares 1,470 fixtures without resets. IDs map to the final complete appendix.")
        b.append({"type": "figure", "path": figures / name, "caption": "Figure M" + ("1" if key == "2025" else "2") + ". No account balances are added together. Same-date requests are allocated before settlement. The offline HTML embeds all traces and its library, supports hover/legend controls, and initially shows every strategy."})
    for metric, label in (("metrics_continuous", "New strategies: continuous four-season accounts"), ("metrics_single", "New strategies: common 2025/26 accounts")):
        display = _strategy_table(data[metric]); displays[label] = display
        b = page(label + " | P001–P032", "landscape")
        table(b, display, "Table M5. Every new strategy in stable ID order. Frequency=trades/common priced fixtures. Expected gain/loss are model-probability-weighted components at allocated stakes; observed winning/losing means describe realised net payoffs. ROI divides profit by stake turnover; capital return divides by 100,000. All additional payoff, exposure and conditional-interval statistics are in the workbook.")
    stability = _stability_table(data); displays["Player strategy annual stability"] = stability
    b = page("Annual starts expose the stability of each new rule", "landscape")
    table(b, stability, "Table M6. Each annual diagnostic independently starts at 100,000; these are not deposits or resets in the continuous graph. A zero-profit season is not counted positive. Seasonal consistency does not remove retrospective design or prove future profitability.")
    b = page("New player-model future probabilities are a separate cohort")
    future = data["future_predictions"]
    n_fixtures = future.match_id.nunique() if "match_id" in future else 0
    p(b, "The new player-model generation contains " + number(len(future), 0) + " future probability records for " + number(n_fixtures, 0) + " targets. They belong to this model/feature protocol, not the original four-model prospective comparison. Their H/D/A probabilities, target identities, model/feature/source hashes and generation clocks are preserved in future_predictions and its provenance. Existing released forecasts, observed outcomes and their first ten-game scores remain unchanged.")
    p(b, "These future player-model probabilities are unscored at this report's cut-off. Later results must be observed under their own immutable cohort and identity checks. Missing pre-kickoff prices or a trade ledger leave prospective trading gain/loss, hit ratio, ROI, drawdown and wealth unavailable. The retrospective P001–P032 account paths are not prospective executed returns, and no real wager is placed.")
    p(b, "The remainder preserves the prior research snapshot and all original graphs. player_models_overview.xlsx contains predictions, engineered features, feature lists, model/strategy metrics, coverage, future records and exact configurations. Detailed bet/date ledgers are separate CSV/Parquet files with hashes and workbook links. Reproduce this report with python -m epl_player_models.report --results PATH. The final expanded appendix contains exact S001–S084 and P001–P032 definitions.")
    return pages + history, {**old_displays, **displays}


def _markdown(pages, entries) -> str:
    return prior._markdown(pages, entries)


def _workbook(path: Path, data: dict, entries: list[dict], registry: dict, output: Path) -> dict:
    readme = pd.DataFrame([
        ("Study", "Eight learned models: four declared feature blocks × logistic/histogram boosting, four chronological evaluation seasons. New retrospective redesign, not a fresh 2025 holdout."),
        ("Accounts", "32 new rules, 192 metric rows: 32 continuous + 32 common-2025 + 128 separate annual starts. Each starts at 100000; date-level cash prevents borrowing and same-date financing."),
        ("Payoffs", "Hit ratio and realised winning/losing means describe selected bets. Expected gain/loss use model probabilities at allocated stakes and are conditional claims."),
        ("Projection", "Projected XI/position/depth features use strictly prior observations, not target announced lineups, future minutes or current-roster backfilling."),
        ("Future generation", "New player-model probabilities are a separate pending protocol/cohort. The earlier prospective ledger and ten-game scores are unchanged."),
        ("Historical metrics workbook", lab._relative(prior.HISTORICAL / "report/execution_overview.xlsx")),
        ("Prior prospective workbook", lab._relative(PRIOR / "report/prospective_overview.xlsx")),
        ("External ledgers", "Detailed new bet/date ledgers are CSV/Parquet beside the report. External files sheet contains paths and SHA256 hashes.")], columns=["Item", "Definition"])
    ledger_files = sorted(p for p in output.glob("*.parquet") if p.name.startswith(("bets_", "equity_")))
    external = pd.DataFrame([{"path_relative_to_workbook": "../" + p.name, "sha256": lab._sha(p), "bytes": p.stat().st_size} for p in ledger_files])
    feature_list = pd.DataFrame([{"feature_group": group, "feature_order": i, "feature": col} for group, columns in data["model_catalog"].get("groups", {}).items() for i, col in enumerate(columns)])
    sheets = {"Read me": readme, "Study status": pd.DataFrame([{"item": k, "value": lab._json(v) if isinstance(v, (dict, list)) else v} for k, v in data["study_status"].items()]),
        "Model catalog": pd.DataFrame(data["model_catalog"].get("models", [])), "Ordered feature lists": feature_list,
        "Model chronology": pd.DataFrame([{"item": k, "value": lab._json(v)} for k, v in data["model_catalog"].get("chronological_roles", {}).items()]) if isinstance(data["model_catalog"].get("chronological_roles", {}), dict) else pd.DataFrame(data["model_catalog"]["chronological_roles"]),
        "Every model prediction": data["predictions"], "Model pooled scores": data["model_metrics_pooled"], "Model annual scores": data["model_metrics_seasonal"],
        "Paired feature comparisons": data["paired_feature_comparisons"], "Feature coverage": data["feature_coverage"],
        "New continuous metrics": data["metrics_continuous"], "New 2025 metrics": data["metrics_single"], "New annual metrics": data["metrics_seasonal"],
        "New strategy catalog": pd.DataFrame(data["new_strategy_catalog"]), "Future player forecasts": data["future_predictions"],
        "Combined 2025 graph": data["graph_2025"], "Combined continuous graph": data["graph_continuous"],
        "Final116 appendix": pd.DataFrame(entries), "External files": external,
        "Input provenance": pd.DataFrame([{"path": p, **item} for p, item in registry.items()])}
    optional_names = {"features": "Engineered match features", "future_features": "Future match features", "parameters": "Fitted parameters",
        "fold_parameters": "Fold parameters", "tuning_trials": "Every tuning trial", "model_fit_audit": "Model fitting audit", "season_stability": "Strategy season stability",
        "fitted_choices": "Selected fold settings", "chronological_roles": "Chronological fitting roles", "logistic_coefficients": "Logistic coefficients"}
    for key, name in optional_names.items():
        if key in data: sheets[name] = data[key]
    percentages = {"accuracy", "hit_ratio", "trading_frequency", "trade_rate_all_predictions", "price_coverage", "active_date_fraction",
        "roi", "expected_roi", "return_on_initial_capital", "max_drawdown", "avg_selected_probability", "p_H", "p_D", "p_A",
        "max_date_capital_exposure", "mean_date_capital_exposure", "min_edge", "kelly_fraction", "max_bet_fraction", "bankroll_fraction"}
    with pd.ExcelWriter(path, engine="xlsxwriter", engine_kwargs={"options": {"strings_to_urls": False}}) as writer:
        book = writer.book; head = book.add_format({"bold": True, "font_color": "white", "bg_color": "#12364D", "text_wrap": True})
        num = book.add_format({"num_format": "#,##0.0000"}); pct = book.add_format({"num_format": "0.00%"})
        for name, frame in sheets.items():
            safe = frame.copy()
            for column in safe:
                if isinstance(safe[column].dtype, pd.DatetimeTZDtype): safe[column] = safe[column].dt.tz_convert(None)
                if safe[column].dtype == "object": safe[column] = safe[column].map(lambda v: lab._json(v.tolist() if isinstance(v, np.ndarray) else v) if isinstance(v, (dict, list, tuple, np.ndarray)) else v)
                elif pd.api.types.is_numeric_dtype(safe[column]) and not pd.api.types.is_bool_dtype(safe[column]):
                    safe[column] = safe[column].map(lambda v: "+∞" if pd.notna(v) and v == float("inf") else "−∞" if pd.notna(v) and v == float("-inf") else v)
            safe.to_excel(writer, sheet_name=name, index=False, na_rep="—")
            sheet = writer.sheets[name]; sheet.write_row(0, 0, list(safe.columns), head); sheet.set_row(0, 42); sheet.freeze_panes(1, 1)
            if len(safe.columns): sheet.autofilter(0, 0, len(safe), len(safe.columns) - 1)
            for index, column in enumerate(safe.columns):
                base = column.removesuffix("_ci_low").removesuffix("_ci_high")
                sheet.set_column(index, index, 65 if column in ("Definition", "explanation", "exact_config", "path", "feature", "coverage", "value") else min(42, max(14, len(column) * .8)), pct if base in percentages else num)
            if name == "Read me":
                for label, target in (("Historical metrics workbook", prior.HISTORICAL / "report/execution_overview.xlsx"), ("Prior prospective workbook", PRIOR / "report/prospective_overview.xlsx")):
                    row_index = 1 + list(readme.Item).index(label)
                    relative = os.path.relpath(target, path.parent).replace("\\", "/")
                    sheet.write_url(row_index, 1, "external:" + relative, string="Open " + label.lower() + " (repository files required)")
    if path.stat().st_size >= 90_000_000: raise AssertionError("The review workbook exceeds its 90 MB budget")
    return {name: len(frame) for name, frame in sheets.items()}



# Layout adapted from the preserved research-report renderer.
def _pdf(pages, entries, path: Path) -> None:
    fonts = lab._fonts()
    styles = {
        "title": ParagraphStyle("ValTitle", fontName=fonts["bold"], fontSize=14, leading=16, spaceAfter=7, textColor=colors.HexColor("#12364D")),
        "body": ParagraphStyle("ValBody", fontName=fonts["body"], fontSize=10.2, leading=12.4, spaceAfter=7),
        "caption": ParagraphStyle("ValCaption", fontName=fonts["body"], fontSize=8.5, leading=10, spaceAfter=4),
        "cell": ParagraphStyle("ValCell", fontName=fonts["sans"], fontSize=8.0, leading=9.3),
        "head": ParagraphStyle("ValHead", fontName=fonts["bold"], fontSize=8.0, leading=9.3, textColor=colors.white),
        "entrytitle": ParagraphStyle("ValEntryTitle", fontName=fonts["bold"], fontSize=9.7, leading=11.4, spaceBefore=5, spaceAfter=4),
        "entry": ParagraphStyle("ValEntry", fontName=fonts["body"], fontSize=9.2, leading=10.6, spaceAfter=4),
        "config": ParagraphStyle("ValConfig", fontName=fonts["mono"], fontSize=6.8, leading=8.1, spaceAfter=7, splitLongWords=True, wordWrap="CJK")}
    def header(canvas, doc):
        width, height = canvas._pagesize
        canvas.setFont(fonts["sans"], 8); canvas.setFillColor(colors.HexColor("#52616B"))
        canvas.drawString(36, height - 21, "EPL PROJECT - PLAYER MODELS AND COMPLETE STRATEGY COMPARISON")
        canvas.drawRightString(width - 36, 21, str(doc.page))
    doc = BaseDocTemplate(str(path), pagesize=letter, title=TITLE, author="Aman Dhillon and Panagiotis Housos")
    for orientation, size in (("portrait", letter), ("landscape", landscape(letter))):
        width, height = size
        doc.addPageTemplates(PageTemplate(id=orientation, pagesize=size, onPage=header,
            frames=[Frame(36, 36, width - 72, height - 72, leftPadding=0, rightPadding=0, topPadding=0, bottomPadding=0)]))
    story = []
    for index, page in enumerate(pages):
        if index: story.extend([NextPageTemplate(page["orientation"]), PageBreak()])
        width = (letter if page["orientation"] == "portrait" else landscape(letter))[0] - 72
        story.append(Paragraph(html.escape(page["title"]), styles["title"]))
        for block in page["blocks"]:
            if block["type"] == "paragraph": story.append(Paragraph(html.escape(block["text"]), styles["body"]))
            elif block["type"] == "figure":
                with PILImage.open(block["path"]) as picture: ratio = picture.height / picture.width
                story.extend([Image(str(block["path"]), width=width - 10, height=(width - 10) * ratio), Paragraph(html.escape(block["caption"]), styles["caption"])])
            else:
                story.append(Paragraph(html.escape(block["caption"]), styles["caption"]))
                frame = block["frame"]
                raw = [list(frame.columns)] + frame.values.tolist()
                cells = [[Paragraph(html.escape(str(value)).replace("\n", "<br/>"), styles["head" if row_index == 0 else "cell"]) for value in row] for row_index, row in enumerate(raw)]
                weights = BODY_WEIGHTS if page["title"].startswith(("Four-season trading", "Original 2025")) else [1 / len(frame.columns)] * len(frame.columns)
                table = Table(cells, colWidths=[width * weight for weight in weights], repeatRows=1)
                table.setStyle(TableStyle([("BACKGROUND", (0,0), (-1,0), colors.HexColor("#12364D")), ("ROWBACKGROUNDS", (0,1), (-1,-1), [colors.white, colors.HexColor("#F0F4F6")]), ("VALIGN", (0,0), (-1,-1), "MIDDLE"), ("LEFTPADDING", (0,0), (-1,-1), 3), ("RIGHTPADDING", (0,0), (-1,-1), 3), ("TOPPADDING", (0,0), (-1,-1), 1.1), ("BOTTOMPADDING", (0,0), (-1,-1), 1.1)]))
                story.append(table)
    story.extend([NextPageTemplate("portrait"), PageBreak(), Paragraph(APPENDIX_TITLE, styles["title"])])
    for entry in entries:
        story.append(KeepTogether([Paragraph(html.escape(entry["title"]), styles["entrytitle"]),
            Paragraph(html.escape("Origin: " + entry["origin"] + ". " + entry["explanation"]), styles["entry"]),
            Paragraph(html.escape("Exact original configuration: " + entry["exact_config"]), styles["config"])]))
    doc.build(story)

def _docx(pages, entries, path: Path) -> None:
    doc = Document()
    normal = doc.styles["Normal"]; normal.font.name = "Times New Roman"; normal.font.size = Pt(10.2)
    normal.paragraph_format.line_spacing = Pt(12.4); normal.paragraph_format.space_after = Pt(7)
    heading = doc.styles["Heading 1"]; heading.font.name = "Arial"; heading.font.size = Pt(14)
    heading.paragraph_format.space_before = Pt(0); heading.paragraph_format.space_after = Pt(7)
    doc.core_properties.title = TITLE; doc.core_properties.author = "Aman Dhillon and Panagiotis Housos"
    def section(index, wide):
        sec = doc.sections[0] if index == 0 else doc.add_section(WD_SECTION_START.NEW_PAGE)
        sec.orientation = WD_ORIENT.LANDSCAPE if wide else WD_ORIENT.PORTRAIT
        sec.page_width = Inches(11 if wide else 8.5); sec.page_height = Inches(8.5 if wide else 11)
        sec.top_margin = sec.bottom_margin = sec.left_margin = sec.right_margin = Inches(.5)
        sec.header_distance = sec.footer_distance = Inches(.2)
        if index == 0:
            run = sec.header.paragraphs[0].add_run("EPL PROJECT - PLAYER MODELS AND COMPLETE STRATEGY COMPARISON")
            run.font.name = "Arial"; run.font.size = Pt(8)
            footer = sec.footer.paragraphs[0]; footer.alignment = 2
            field = OxmlElement("w:fldSimple"); field.set(qn("w:instr"), "PAGE"); footer._p.append(field)
    for index, page in enumerate(pages):
        wide = page["orientation"] == "landscape"; section(index, wide); doc.add_heading(page["title"], 1)
        for block in page["blocks"]:
            if block["type"] == "paragraph": doc.add_paragraph(block["text"]); continue
            if block["type"] == "figure": doc.add_picture(str(block["path"]), width=Inches(9.85 if wide else 7.35))
            caption = doc.add_paragraph(block["caption"]); caption.paragraph_format.line_spacing = Pt(10)
            for run in caption.runs: run.font.size = Pt(8.5)
            if block["type"] == "figure": continue
            frame = block["frame"]; table = doc.add_table(rows=1, cols=len(frame.columns)); table.style = "Light Shading Accent 1"; table.autofit = False
            weights = BODY_WEIGHTS if page["title"].startswith(("Four-season trading", "Original 2025")) else [1 / len(frame.columns)] * len(frame.columns)
            width = 10 if wide else 7.5
            for col, weight in zip(table.columns, weights): col.width = Inches(width * weight)
            for cell, value in zip(table.rows[0].cells, frame.columns): cell.text = str(value)
            for values in frame.itertuples(index=False, name=None):
                for cell, value in zip(table.add_row().cells, values): cell.text = str(value)
            for row in table.rows:
                for cell, weight in zip(row.cells, weights):
                    cell.width = Inches(width * weight)
                    for paragraph in cell.paragraphs:
                        paragraph.paragraph_format.line_spacing = Pt(9.3); paragraph.paragraph_format.space_after = Pt(0)
                        for run in paragraph.runs: run.font.name = "Arial"; run.font.size = Pt(8)
    section(len(pages), False); doc.add_heading(APPENDIX_TITLE, 1)
    for entry in entries:
        title = doc.add_paragraph(); title.paragraph_format.keep_with_next = True
        run = title.add_run(entry["title"]); run.bold = True; run.font.name = "Arial"; run.font.size = Pt(9.7)
        paragraph = doc.add_paragraph("Origin: " + entry["origin"] + ". " + entry["explanation"])
        paragraph.paragraph_format.keep_with_next = True; paragraph.paragraph_format.line_spacing = Pt(10.6)
        for run in paragraph.runs: run.font.size = Pt(9.2)
        paragraph = doc.add_paragraph("Exact original configuration: " + entry["exact_config"])
        paragraph.paragraph_format.keep_together = True; paragraph.paragraph_format.line_spacing = Pt(8.1)
        for run in paragraph.runs: run.font.name = "Consolas"; run.font.size = Pt(6.8)
    doc.save(path)

def build_report(output_dir: Path) -> dict:
    output = Path(output_dir).resolve(); data = _load(output); _validate(data); entries = _appendix(data)
    sources = [output / (name + ".parquet") for name in TABLE_NAMES]
    sources += [output / (name + ".parquet") for name in OPTIONAL_TABLES if (output / (name + ".parquet")).exists()]
    sources += [output / (name + ".json") for name in ("model_catalog", "new_strategy_catalog", "study_status", "source_provenance")]
    sources += [output / ("graph_equity_" + name + "_all.csv") for name in ("2025", "continuous")]
    sources += [p for p in output.glob("*.parquet") if p.name.startswith(("bets_", "equity_"))]
    old_manifest_path = PRIOR / "report/report_manifest.json"
    old_manifest = json.loads(old_manifest_path.read_text(encoding="utf-8"))
    sources += [ROOT / p for p in old_manifest["inputs"]]
    sources += [old_manifest_path, PRIOR / "report/prospective_overview.xlsx", prior.HISTORICAL / "report/execution_overview.xlsx"]
    sources += [PRIOR / "figures" / name for name in prior.GRAPH_NAMES]
    sources += [Path(__file__).resolve(), Path(prior.__file__).resolve(), Path(lab.__file__).resolve()]
    for name in ("epl_player_models/PROTOCOL.md", "epl_prospective/PROTOCOL.md"):
        if (ROOT / name).exists(): sources.append(ROOT / name)
    sources = sorted(set(sources))
    registry = {lab._relative(p): {"sha256": lab._sha(p), "bytes": p.stat().st_size} for p in sources}
    report = output / "report"; report.mkdir(exist_ok=True); figures = output / "figures"
    graph_info = _graphs(data, figures, entries)
    pages, displays = _body(data, figures)
    for name, frame in displays.items(): frame.to_csv(report / (name.lower().replace(" ", "_").replace("/", "_").replace(":", "") + ".csv"), index=False)
    pd.DataFrame(entries).to_csv(report / "complete116_strategy_appendix.csv", index=False)
    md = report / "validation_report.md"; pdf = md.with_suffix(".pdf"); docx = md.with_suffix(".docx")
    md.write_text(_markdown(pages, entries), encoding="utf-8"); _pdf(pages, entries, pdf); _docx(pages, entries, docx)
    sheets = _workbook(report / "player_models_overview.xlsx", data, entries, registry, output)
    reader = PdfReader(pdf); text = "\n".join(page.extract_text() or "" for page in reader.pages)
    if APPENDIX_TITLE not in text: raise AssertionError("The expanded final appendix is absent")
    final_text = " ".join(text.split(APPENDIX_TITLE, 1)[1].split())
    if any(" ".join(e["title"].split()) not in final_text for e in entries): raise AssertionError("An original or new strategy explanation is missing")
    if NEW_IDS[-1] not in reader.pages[-1].extract_text(): raise AssertionError("P032 must end the final strategy appendix")
    changed = [lab._relative(p) for p in sources if lab._sha(p) != registry[lab._relative(p)]["sha256"]]
    if changed: raise AssertionError("Report inputs changed during rendering: " + str(changed))
    preserved = {name: lab._sha(figures / name) == lab._sha(PRIOR / "figures" / name) for name in prior.GRAPH_NAMES}
    if not all(preserved.values()): raise AssertionError("A historical graph changed")
    produced = [p for folder in (report, figures) for p in folder.rglob("*") if p.is_file() and p.name != "report_manifest.json" and "previews" not in p.parts]
    manifest = {"status": "retrospective_player_model_study_with_separate_pending_future_cohort", "built_at_utc": datetime.now(timezone.utc).isoformat(),
        "pdf_pages": len(reader.pages), "body_sections": len(pages), "learned_models": 8, "new_strategies": 32,
        "full_catalog_count": len(entries), "continuous_strategy_count": len(CONTINUOUS_IDS), "new_account_metric_rows": 192,
        "graphs": graph_info, "preserved_historical_graphs": preserved,
        "appendix_ids": [e["strategy_id"] for e in entries], "appendix_is_last": True, "workbook_sheet_rows": sheets,
        "consumed_source_files_unchanged": True, "inputs": registry,
        "outputs": {p.relative_to(output).as_posix(): {"sha256": lab._sha(p), "bytes": p.stat().st_size} for p in sorted(produced)},
        "manifest_order": "Final builder envelope and independent audit are deliberately not consumed report inputs.",
        "limits": "New retrospective redesign on already examined history; projected XI is past-only, not target lineup; archived execution unverified; new future cohort separate and unscored."}
    (report / "report_manifest.json").write_text(json.dumps(manifest, ensure_ascii=False, indent=2, allow_nan=False), encoding="utf-8")
    return {"output_directory": str(report), "pdf_pages": len(reader.pages), "appendix_entries": len(entries),
        "all_visible_traces_2025": len(ALL_IDS), "all_visible_traces_continuous": len(CONTINUOUS_IDS),
        "source_files_unchanged": True, "workbook_bytes": (report / "player_models_overview.xlsx").stat().st_size}


def main() -> None:
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument("--results", required=True, type=Path)
    print(json.dumps(build_report(parser.parse_args().results), indent=2))


if __name__ == "__main__": main()
