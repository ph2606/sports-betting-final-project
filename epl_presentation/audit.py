"""Independent checks of saved slide content against the frozen project exports.

This module reads existing evidence only. It does not fit, release, or settle
anything, and it writes no file when called directly.
"""
from __future__ import annotations

from collections import Counter
import hashlib
import json
from pathlib import Path
import re
import xml.etree.ElementTree as ET

import fitz
import numpy as np
import pandas as pd
from pptx import Presentation

ROOT = Path(__file__).resolve().parents[1]
STUDY = ROOT / "epl_player_models/results/20261020_player_models_review"
FEATURED = ("S029", "S033", "S017", "P009", "S060", "P028")


def _require(condition, message):
    if not condition:
        raise AssertionError(message)


def _json(path):
    return json.loads(Path(path).read_text(encoding="utf-8"))


def _sha(path):
    return hashlib.sha256(Path(path).read_bytes()).hexdigest()


def _normal(value):
    return re.sub(r"\s+", " ", str(value)).strip()


def _visible(slide):
    values = [slide["title"], slide.get("subtitle", "")]
    for element in slide.get("elements", []):
        if element.get("type") == "text":
            values.append(element.get("text", ""))
        elif element.get("type") == "card":
            values.extend(element.get(key, "") for key in ("label", "value", "detail"))
        elif element.get("type") == "table":
            values.extend(element.get("headers", []))
            values.extend(value for row in element.get("rows", []) for value in row)
    return "\n".join(map(str, values))


def _slide(slides, title):
    matches = [slide for slide in slides if slide["title"] == title]
    _require(len(matches) == 1, f"Expected exactly one slide titled {title!r}")
    return matches[0]


def _table(slide):
    tables = [element for element in slide["elements"] if element.get("type") == "table"]
    _require(len(tables) == 1, f"Expected one table: {slide['title']}")
    return tables[0]


def _same_frame(actual, expected, name):
    try:
        pd.testing.assert_frame_equal(actual.reset_index(drop=True), expected.reset_index(drop=True),
                                      check_dtype=False, check_categorical=False,
                                      atol=1e-8, rtol=1e-12)
    except AssertionError as error:
        raise AssertionError(f"Saved {name} differs from frozen evidence: {error}") from error


def _money(value):
    return f"{float(value):,.0f}"


def _pct(value):
    return f"{float(value):.2%}"


def _dec(value, digits=4):
    return f"{float(value):.{digits}f}"


def _audit_sources(data):
    ranking = data["ranking"].sort_values("log_loss")
    _require(len(ranking) == 19 and ranking.model_key.nunique() == 19,
             "Expected 19 unique comparable probability series")
    _require(ranking.n_matches.eq(1470).all(), "Probability ranks use different populations")
    _require(int(ranking.learned_predictor.sum()) == 16, "Expected 16 learned probability series")
    _require(ranking.iloc[0].model_key == "baseline_market_only", "Raw-market rank changed")
    _require(ranking.loc[ranking.learned_predictor].iloc[0].model_key == "residual_team_players",
             "Best learned predictor changed")
    _require(data["single"].final_capital_sum_units.idxmax() == "P028", "Single-season capital leader changed")
    _require(data["continuous"].final_capital_sum_units.idxmax() == "S060", "Continuous capital leader changed")
    _require(len(data["single"]) == 116 and data["single"].index.is_unique,
             "Expected 116 unique single-season accounts")
    _require(len(data["continuous"]) == 88 and data["continuous"].index.is_unique,
             "Expected 88 unique continuous accounts")
    return {"probability_series": 19, "learned_predictors": 16, "matches_per_series": 1470,
            "best_learned_by_log_loss": "residual_team_players",
            "strongest_probability_benchmark": "baseline_market_only",
            "single_season_capital_leader": "P028", "continuous_capital_leader": "S060"}


def _audit_graphs(data, output):
    result = {}
    for scope, count, points, stem in (("single", 116, 114, "all_accounts_2025"),
                                      ("continuous", 88, 454, "all_accounts_continuous")):
        graph, metrics = data["graphs"][scope], data[scope]
        _require(graph.strategy_id.nunique() == count and len(graph) == count * points,
                 f"Incorrect {scope} graph curve/point counts")
        _require(set(graph.strategy_id) == set(metrics.index), f"Missing {scope} curve IDs")
        _require(not graph.duplicated(["strategy_id", "match_date"]).any(), f"Duplicate {scope} curve dates")
        _require(graph.groupby("strategy_id").size().eq(points).all(), f"Unequal {scope} date grids")
        anchors = graph.loc[graph.is_initial]
        _require(len(anchors) == count and anchors.end_bankroll.eq(100000).all(),
                 f"Incorrect {scope} initial-capital anchors")
        ends = graph.sort_values("match_date").groupby("strategy_id").tail(1).set_index("strategy_id")
        ends = ends.reindex(metrics.index)
        _require(np.allclose(ends.end_bankroll, metrics.final_capital_sum_units, atol=1e-6, rtol=0),
                 f"{scope} graph endpoints disagree with account metrics")
        svg = ET.parse(output / "figures" / (stem + ".svg"))
        groups = svg.findall(".//{http://www.w3.org/2000/svg}g")
        identifiers = [element.attrib.get("id", "") for element in groups]
        strategy_ids = [value.removeprefix("strategy-") for value in identifiers if value.startswith("strategy-")]
        if strategy_ids:
            _require(Counter(strategy_ids) == Counter(metrics.index), f"SVG omits/duplicates {scope} strategies")
        else:
            line_groups = [value for value in identifiers if value.startswith("line2d_")]
            _require(len(line_groups) >= count, f"SVG has fewer paths than {scope} strategies")
        _require((output / "figures" / (stem + ".png")).stat().st_size > 10000,
                 f"Missing/empty {scope} graph image")
        result[scope] = {"strategies": count, "points": len(graph), "points_per_curve": points,
                         "all_initial_balances": 100000, "endpoints_reconciled": count,
                         "svg_strategy_groups_verified": len(strategy_ids)}
    for name in ("all_strategies_2025_100k.html", "all_strategies_continuous_100k.html"):
        _require(_sha(output / name) == _sha(STUDY / "figures" / name),
                 f"Companion HTML differs from original: {name}")
    return result


def _audit_profiles(slides, data):
    roles = (("best_learned", "S029", "Best learned predictor: market residual + team/FPL"),
             ("simplest", "S033", "Simplest model: historical class frequencies"),
             ("team", "S017", "Team-data model: regularized logistic regression"),
             ("player", "P009", "Player-only model: recent FPL form"))
    for key, sid, title in roles:
        slide = _slide(slides, title)
        metrics = data["models"]["profiles"][key]["metrics"]
        cards = [element for element in slide["elements"] if element.get("type") == "card"]
        expected = [_pct(metrics["accuracy"]), _dec(metrics["log_loss"]),
                    _dec(metrics["brier_score"]), _pct(metrics["classwise_ece_10"])]
        _require([card["value"] for card in cards] == expected, f"Incorrect predictive statistics: {key}")
        account = data["continuous"].loc[sid]
        expected_rows = [["Ending account", _money(account.final_capital_sum_units)],
                         ["Return on starting 100K", _pct(account.return_on_initial_capital)],
                         ["Hit ratio / trade frequency", _pct(account.hit_ratio) + " / " + _pct(account.trading_frequency)],
                         ["ROI on total stakes", _pct(account.roi)], ["Maximum drawdown", _pct(account.max_drawdown)]]
        _require(_table(slide)["rows"] == expected_rows, f"Incorrect {sid} profile account statistics")
    scorecard = _table(_slide(slides, "The same-scope trading scorecard"))["rows"]
    _require(len(scorecard) == len(FEATURED), "Featured scorecard omitted a policy")
    for values, sid in zip(scorecard, FEATURED):
        account = data["continuous"].loc[sid]
        expected = [_money(account.final_capital_sum_units), _pct(account.hit_ratio), _pct(account.trading_frequency),
                    _dec(account.expected_gain_per_trade_units, 2), _dec(account.expected_loss_per_trade_units, 2),
                    _pct(account.roi), _pct(account.max_drawdown)]
        _require(str(values[0]).startswith(sid + " | ") and values[1:] == expected,
                 f"Featured scorecard differs from {sid} continuous evidence")
    return {"profiles_checked": 4, "featured_continuous_account_rows_checked": 6,
            "expected_components_kept_distinct_from_realized_payoffs": True}


def _audit_catalog(slides, pptx, data, output):
    catalog = data["catalog"]
    _require(_json(output / "complete_strategy_catalog.json") == catalog, "Full strategy catalog changed")
    _require(catalog == _json(STUDY / "all_strategy_catalog.json"), "Catalog is not the saved canonical catalog")
    expected = {row["strategy_id"]: row for row in catalog}
    visible_ids, notes_configs = [], {}
    catalog_slides = []
    for number, (spec, native) in enumerate(zip(slides, pptx.slides), 1):
        if not spec["title"].startswith("Strategy appendix | "):
            continue
        catalog_slides.append(number)
        _require(spec.get("appendix") is True, "Catalog page is not marked appendix")
        rows = _table(spec)["rows"]
        visible_ids.extend(row[0] for row in rows)
        notes = native.notes_slide.notes_text_frame.text
        for line in notes.splitlines():
            if line.startswith("{"):
                config = json.loads(line)
                sid = config["strategy_id"]
                _require(sid not in notes_configs, f"Duplicate exact note configuration: {sid}")
                notes_configs[sid] = config
        native_text = "\n".join(shape.text for shape in native.shapes if shape.has_text_frame)
        for row in rows:
            _require(row[0] in native_text, f"Catalog ID absent from visible slide: {row[0]}")
    _require(Counter(visible_ids) == Counter(expected.keys()), "Appendix must show every strategy ID exactly once")
    _require(notes_configs == expected, "Speaker notes do not retain all exact saved configurations")
    _require(catalog_slides and catalog_slides == list(range(catalog_slides[0], len(slides) + 1)),
             "Complete strategy appendix must be the final slide block")
    return {"catalog_ids_visible_once": len(visible_ids), "exact_note_configurations": len(notes_configs),
            "catalog_appendix_slides": len(catalog_slides), "final_appendix_preserved": True}


def _audit_forecasts():
    old_path, outcome_path = ROOT / "epl_live/forecasts.jsonl", ROOT / "epl_live/outcomes.jsonl"
    new_path = STUDY / "forecasts.jsonl"
    old = [_json_line for _json_line in map(json.loads, old_path.read_text(encoding="utf-8").splitlines())]
    outcomes = list(map(json.loads, outcome_path.read_text(encoding="utf-8").splitlines()))
    new = list(map(json.loads, new_path.read_text(encoding="utf-8").splitlines()))
    _require((len(old), len(outcomes), len(new)) == (80, 10, 80), "Existing forecast/outcome counts changed")
    original = ROOT / "epl_player_models/results/20261020_player_models/forecasts.jsonl"
    _require(_sha(original) == _sha(new_path), "Reviewed forecast ledger is not an exact copy")
    old_records, new_records = [row["record"] for row in old], [row["record"] for row in new]
    settled_ids = {row["record"]["match_id"] for row in outcomes}
    resolved = [row for row in old_records if row["match_id"] in settled_ids]
    pending_old = [row for row in old_records if row["match_id"] not in settled_ids]
    _require(len(resolved) == 40 and len(pending_old) == 40, "Unexpected old-model pending/resolved counts")
    _require(len({row["model_id"] for row in resolved}) == 4, "Unexpected resolved-model count")
    pending_ids = {row["match_id"] for row in pending_old}
    _require(len(pending_ids) == 10 and {row["match_id"] for row in new_records} == pending_ids,
             "Old/new pending records do not refer to the same ten fixtures")
    _require(not settled_ids.intersection(pending_ids), "Future player records already have saved outcomes")
    keys = [(row["model_id"], row["match_id"], row["update_version"]) for row in old_records + new_records]
    _require(len(set(keys)) == 160, "Duplicate unique forecast keys counted")
    _require(all(row.get("market_quote") is None for row in old_records + new_records),
             "Forecast quote availability changed")
    return {"older_forecasts": 80, "saved_outcomes": 10, "player_study_forecasts": 80,
            "resolved_forecasts": 40, "unique_pending_forecasts": 120,
            "independent_resolved_fixtures": 10, "independent_pending_fixtures": 10,
            "review_copy_counted_once": True, "no_forecast_or_outcome_written": True,
            "source_ledger_sha256": {str(p.relative_to(ROOT)): _sha(p) for p in (old_path, outcome_path, new_path)}}


def _audit_render(output, slides, pptx):
    rendering = _json(output / "render_audit.json")
    _require(rendering["slide_count"] == len(slides) == len(pptx.slides), "PPTX slide counts differ")
    _require(not rendering["geometry"]["geometry_violations"], "Off-canvas PowerPoint shapes")
    native = rendering["native_layout"]
    _require(native["text_overflow_count"] == 0,
             "Native PowerPoint text overflows: " + json.dumps(native.get("text_overflows", [])[:8]))
    _require(native["text_shapes_checked"] > len(slides) * 5, "Native layout checked too little content")
    for name, digest in rendering["sha256"].items():
        _require(_sha(output / name) == digest, f"Rendered file changed after layout measurement: {name}")
    for number, (spec, native_slide) in enumerate(zip(slides, pptx.slides), 1):
        visible = _normal("\n".join(shape.text for shape in native_slide.shapes if shape.has_text_frame))
        for text in _visible(spec).splitlines():
            _require(_normal(text) in visible, f"PPTX lost slide {number} text: {text!r}")
        _require(spec.get("notes", "") in native_slide.notes_slide.notes_text_frame.text,
                 f"Speaker notes differ on slide {number}")
    off_page = []
    with fitz.open(output / "EPL_Project_Presentation.pdf") as pdf:
        _require(len(pdf) == len(slides), "PDF page count differs from PPTX")
        for number, (page, spec) in enumerate(zip(pdf, slides), 1):
            _require(abs(page.rect.width / page.rect.height - 16 / 9) < .01,
                     f"Wrong PDF slide aspect ratio: {number}")
            _require(_normal(spec["title"]) in _normal(page.get_text()), f"PDF lost slide {number} title")
            for block in page.get_text("dict")["blocks"]:
                for line in block.get("lines", []):
                    for span in line.get("spans", []):
                        box = fitz.Rect(span["bbox"])
                        if box.x0 < -2 or box.y0 < -2 or box.x1 > page.rect.width + 2 or box.y1 > page.rect.height + 2:
                            off_page.append({"slide": number, "text": span["text"], "bbox": list(box)})
        _require(not off_page, f"PDF text crosses the page: {off_page[:5]}")
        pages = len(pdf)
    return {"editable_slides": len(slides), "native_pdf_pages": pages,
            "native_text_shapes_checked": native["text_shapes_checked"],
            "native_text_overflows": 0, "off_page_pdf_text_spans": 0,
            "pptx_visible_content_and_notes_reconciled": True}


def audit_presentation(output, *, data=None):
    """Validate one saved presentation; raise on mismatches and return evidence."""
    if data is None:
        from .build import load
        data = load()
    output = Path(output).resolve()
    slides = _json(output / "presentation_content.json")
    _require(isinstance(slides, list) and len(slides) >= 18, "Missing presentation slide content")
    intro = _visible(slides[0]).lower()
    _require(slides[0]["title"] == "EPL: what we have implemented", "First page must introduce implemented work")
    _require("what is already built" in intro and "116" in intro and "19" in intro,
             "Implemented introduction lacks the completed system")
    _require(not re.search(r"\b(?:proposal|proposed|initial idea|will implement)\b", intro),
             "First page describes a proposal instead of implemented work")
    for spec in slides:
        for source in spec.get("sources", []):
            _require((ROOT / source).is_file(), f"Slide source path does not exist: {source}")
    _same_frame(pd.read_csv(output / "strategy_metrics_single.csv"), data["single"].reset_index(), "single metrics")
    _same_frame(pd.read_csv(output / "strategy_metrics_continuous.csv"), data["continuous"].reset_index(), "continuous metrics")
    _same_frame(pd.read_csv(output / "model_probability_scores.csv"), data["ranking"], "19 probability scores")
    sources = _audit_sources(data)
    graphs = _audit_graphs(data, output)
    profiles = _audit_profiles(slides, data)
    pptx = Presentation(output / "EPL_Project_Presentation.pptx")
    catalog = _audit_catalog(slides, pptx, data, output)
    forecasts = _audit_forecasts()
    rendering = _audit_render(output, slides, pptx)
    workbook = pd.ExcelFile(output / "EPL_Presentation_Data.xlsx")
    _require({"All19 probability scores", "All116 2025 metrics", "All88 continuous metrics",
              "Complete116 configurations", "Graph 2025", "Graph continuous"}.issubset(workbook.sheet_names),
             "Companion workbook omitted required source tables")
    prospective = _slide(slides, "A prospective pipeline exists; trading P&L does not")
    _require("120 pending" in _visible(prospective) and "40 earlier-model + 80 new-model forecasts" in _visible(prospective),
             "Prospective slide double-counts or omits forecasts")
    _require("31.27%" in _visible(prospective) and "83.18%" in _visible(prospective),
             "First-cohort uncertainty omitted")
    return {"status": "passed", "read_only": True, "implemented_only_introduction": True,
            "sources_and_rankings": sources, "graphs": graphs, "featured_profiles": profiles,
            "strategy_catalog": catalog, "forecasts": forecasts, "rendering": rendering,
            "saved_metric_rows_reconciled": {"single": 116, "continuous": 88, "probability_scores": 19},
            "companion_workbook_required_sheets": True}
