"""Read-only audit of the six requested slides and full strategy appendix."""
from __future__ import annotations

from collections import Counter
import json
from pathlib import Path
import re
import xml.etree.ElementTree as ET

import numpy as np
import pandas as pd
from pptx import Presentation

from .audit import (_audit_render, _audit_sources, _dec, _json, _money, _normal,
                    _pct, _require, _same_frame, _sha, _table, _visible)
from .build import ROOT, STUDY
from .simple_rules import describe_strategy, model_name

ORIGINAL = ROOT / "epl_presentation/results/20261021_implemented"
ROLES = (("best_learned", "S029"), ("simplest", "S033"),
         ("team", "S017"), ("player", "P009"))
SHARED_FIELDS = ("model_key", "selection_rule", "min_edge", "allowed_outcomes",
                 "min_odds", "max_odds", "consensus_min", "consensus_model_keys",
                 "cash_only", "odds_mode", "date_cap_fraction")


def _expand_ids(value):
    """Expand displayed slash-separated IDs or an inclusive cash-rule range."""
    value = str(value).strip()
    match = re.fullmatch(r"([SP])(\d{3})\s*[-\u2013]\s*([SP])(\d{3})", value)
    if match:
        prefix, start, other_prefix, stop = match.groups()
        _require(prefix == other_prefix and int(start) <= int(stop), "Invalid displayed ID range")
        return [f"{prefix}{number:03d}" for number in range(int(start), int(stop) + 1)]
    ids = re.split(r"\s*/\s*", value)
    _require(all(re.fullmatch(r"[SP]\d{3}", sid) for sid in ids), f"Unrecognized strategy IDs: {value}")
    return ids


def _audit_profiles(slides, data):
    catalog = {row["strategy_id"]: row for row in data["catalog"]}
    for spec, (key, sid) in zip(slides[2:6], ROLES):
        model = data["models"]["profiles"][key]["metrics"]
        account, config = data["continuous"].loc[sid], catalog[sid]
        _require(sid in spec.get("kicker", ""), f"Incorrect featured strategy: {sid}")
        _require("2022/23" in spec["subtitle"] and "2025/26" in spec["subtitle"]
                 and "1,470" in spec["subtitle"], f"Unclear four-year profile population: {sid}")
        visible = _visible(spec)
        expected_forecast = "Forecast accuracy: " + _pct(model["accuracy"]) + "  |  Prediction loss: " + _dec(model["log_loss"])
        _require(expected_forecast in visible, f"Forecast statistics do not match source: {sid}")
        cards = [element for element in spec["elements"] if element.get("type") == "card"]
        _require([card["value"] for card in cards] ==
                 [_money(account.final_capital_sum_units), _pct(account.return_on_initial_capital)],
                 f"Wrong ending capital or capital return: {sid}")
        expected_rows = [
            ["Bets / trading frequency", _money(account.bet_count) + " / " + _pct(account.trading_frequency)],
            ["Hit ratio: bets won", _pct(account.hit_ratio)],
            ["Expected gain per trade", _dec(account.expected_gain_per_trade_units, 2)],
            ["Expected loss per trade", _dec(account.expected_loss_per_trade_units, 2)],
            ["Profit / ROI on total stakes", _money(account.profit) + " / " + _pct(account.roi)],
            ["Largest drop from a prior peak", _pct(account.max_drawdown)],
        ]
        _require(_table(spec)["rows"] == expected_rows, f"Featured trading statistics differ: {sid}")
        _require(config["staking"] == "fixed" and config["fixed_stake"] == 1000
                 and config["min_edge"] == .05 and config["selection_rule"] == "max_edge"
                 and config["allowed_outcomes"] == list("HDA") and config["date_cap_fraction"] is None,
                 f"Featured simple rule does not describe its actual configuration: {sid}")
        _require("above 5%" in visible and "1,000 per bet: 1% of the initial 100,000" in visible,
                 f"Entry threshold or initial-capital staking distinction missing: {sid}")
        _require("opening cash" in visible.lower() and "not realized average wins/losses" in visible,
                 f"Cash or expected-payoff qualification missing: {sid}")
        notes = spec.get("notes", "")
        _require("reduced proportionally" in notes and "same-date" in notes,
                 f"Joint same-date cash allocation omitted: {sid}")
        configs = [json.loads(line) for line in notes.splitlines() if line.startswith("{")]
        _require(configs == [config], f"Featured full configuration changed: {sid}")
    best = _visible(slides[2]).lower()
    _require("among the learned models" in best and "lowest prediction error" in best,
             "Best predictor requires an explicit probability-loss criterion")
    _require("not the highest ending account" in slides[2]["notes"] and
             "raw market benchmark has lower prediction loss" in slides[2]["notes"],
             "Best learned role is conflated with profit or the raw-market benchmark")
    return {"profiles": [sid for _, sid in ROLES], "forecast_statistic_pairs": 4,
            "capital_statistic_pairs": 4, "trading_statistic_rows": 24,
            "all_fixed_stakes": 1000, "stake_basis": "one_percent_of_initial_100000",
            "common_profile_population": 1470, "joint_cash_reductions_explained": True}


def _audit_graph(slides, output, data):
    spec = slides[1]
    _require("116" in spec["title"] and "2025/26" in spec["subtitle"]
             and "100,000" in spec["subtitle"], "Required common-season money graph omitted")
    for suffix in ("png", "svg"):
        name = "all_accounts_2025." + suffix
        _require(_sha(output / "figures" / name) == _sha(ORIGINAL / "figures" / name),
                 f"Previously audited graph changed: {name}")
    images = [element for element in spec["elements"] if element.get("type") == "image"]
    _require(len(images) == 1 and _sha(images[0]["path"]) == _sha(output / "figures/all_accounts_2025.png"),
             "Slide uses a different money graph")
    graph = data["graphs"]["single"]
    source = pd.read_csv(ROOT / STUDY / "graph_equity_2025_all.csv", parse_dates=["match_date"])
    _same_frame(graph, source, "common-season graph data")
    _require(len(graph) == 116 * 114 and graph.strategy_id.nunique() == 116,
             "Graph does not cover all 116 accounts on 114 shared dates")
    _require(not graph.duplicated(["strategy_id", "match_date"]).any(), "Duplicate graph strategy dates")
    starts = graph.loc[graph.is_initial]
    _require(len(starts) == 116 and starts.end_bankroll.eq(100000).all(), "Incorrect graph starting balances")
    ends = graph.sort_values("match_date").groupby("strategy_id").tail(1).set_index("strategy_id")
    ends = ends.reindex(data["single"].index)
    _require(np.allclose(ends.end_bankroll, data["single"].final_capital_sum_units, atol=1e-6, rtol=0),
             "Graph ending balances differ from account metrics")
    tree = ET.parse(output / "figures/all_accounts_2025.svg")
    ids = [node.attrib["id"].removeprefix("strategy-")
           for node in tree.findall(".//{http://www.w3.org/2000/svg}g")
           if node.attrib.get("id", "").startswith("strategy-")]
    _require(Counter(ids) == Counter(data["single"].index), "SVG omits or duplicates strategy curves")
    return {"accounts": 116, "starting_capital": 100000, "graph_points": len(graph),
            "svg_strategy_paths_verified": 116, "account_endpoints_reconciled": 116,
            "same_audited_graph_as_previous_presentation": True}


def _check_compact_row(values, configs):
    """Check each displayed grouping against the actual settings, including caps."""
    _require(len(values) == 5, "Compact appendix rows need five columns")
    first = configs[0]
    _require(values[1] == model_name(first["model_key"]), "Wrong compact model label")
    if first["cash_only"]:
        _require(all(row["cash_only"] and row["model_key"] == first["model_key"] for row in configs),
                 "Cash range includes a trading strategy")
        _require(values[2:] == ["No bets; keep cash", "0", "No betting"], "Cash range shows exposure")
        return
    _require(len(configs) <= 2 and all(all(row[key] == first[key] for key in SHARED_FIELDS) for row in configs),
             "Paired strategy IDs differ in more than their staking settings")
    entry = values[2]
    if first["selection_rule"] == "max_edge":
        _require(entry.splitlines()[0] == f"Estimated return > {100 * first['min_edge']:g}%",
                 "Wrong displayed return threshold")
        if first["allowed_outcomes"] != list("HDA"):
            _require("Only " + "/".join(first["allowed_outcomes"]) in entry, "Outcome restriction omitted")
    elif first["selection_rule"] == "fixed_outcome":
        name = {"H": "home win", "D": "draw", "A": "away win"}[first["allowed_outcomes"][0]]
        _require(entry.startswith("Always " + name), "Wrong unconditional outcome")
    elif first["selection_rule"] == "market_favorite":
        _require(entry.startswith("Bookmaker's most likely result"), "Favorite rule changed")
    elif first["selection_rule"] == "market_underdog":
        _require(entry.startswith("Bookmaker's least likely result"), "Underdog rule changed")
    elif first["selection_rule"] == "model_favorite":
        _require(entry.startswith("Model's most likely result"), "Model-favorite rule changed")
    else:
        raise AssertionError("Unrecognized catalog entry rule")
    if first["min_odds"] is not None or first["max_odds"] is not None:
        _require("Odds " + str(first["min_odds"] or "any") + "\u2013" + str(first["max_odds"] or "any") in entry,
                 "Compact odds eligibility filter omitted")
    if first["consensus_min"] is not None:
        _require(f"Positive return: {first['consensus_min']}/{len(first['consensus_model_keys'])} models" in entry,
                 "Compact consensus restriction omitted")
    stake_lines = values[3].splitlines()
    _require(len(stake_lines) == len(configs), "A paired strategy stake is missing")
    for line, config in zip(stake_lines, configs):
        _require(line.startswith(config["strategy_id"] + ": "), "Stake is assigned to wrong ID")
        if config["staking"] == "fixed":
            expected = _money(config["fixed_stake"])
        elif config["staking"] == "dynamic_fraction":
            expected = f"{100 * config['bankroll_fraction']:g}% of current cash"
        elif config["staking"] == "kelly":
            multiplier = {.125: "1/8", .25: "1/4", .5: "1/2", 1.: "Full"}.get(config["kelly_fraction"], str(config["kelly_fraction"]))
            expected = multiplier + f" Kelly; \u2264{100 * config['max_bet_fraction']:g}% cash"
        else:
            raise AssertionError("Unknown compact staking type")
        _require(line.split(": ", 1)[1] == expected, "Compact stake multiplier or cap changed")
    price = "Best odds" if first["odds_mode"] == "best" else "Bet365"
    budget = "Cash available" if first["date_cap_fraction"] is None else f"Date cap {100 * first['date_cap_fraction']:g}% cash"
    _require(values[4] == price + "\n" + budget, "Price source or combined date budget changed")


def _audit_appendix(slides, pptx, output, data):
    catalog = data["catalog"]
    _require(_json(output / "complete_strategy_catalog.json") == catalog, "Saved full catalog changed")
    _require(catalog == _json(ROOT / STUDY / "all_strategy_catalog.json"), "Catalog differs from frozen study")
    expected = {row["strategy_id"]: row for row in catalog}
    seen, exact, groups = [], {}, 0
    _require(slides[6]["title"] == "Appendix: how much capital each rule uses", "Capital guide missing")
    guide = _visible(slides[6])
    _require("INITIAL 100,000" in guide and "Cash 80,000" in guide and "800" in guide
             and "2,500" in guide and "2,000" in guide and "never borrow" in guide,
             "Guide omits fixed/dynamic/Kelly amount or joint funding explanation")
    for spec, native in zip(slides[7:], list(pptx.slides)[7:]):
        _require(spec["title"].startswith("Appendix: strategy rules ("), "Unexpected extra appendix page")
        for row in _table(spec)["rows"]:
            ids = _expand_ids(row[0])
            _require(all(sid in expected for sid in ids), "Unknown visible strategy ID")
            _check_compact_row(row, [expected[sid] for sid in ids])
            seen.extend(ids)
            groups += 1
        for line in native.notes_slide.notes_text_frame.text.splitlines():
            if line.startswith("{"):
                config = json.loads(line)
                sid = config["strategy_id"]
                _require(sid not in exact, "Duplicate complete appendix configuration")
                exact[sid] = config
    _require(Counter(seen) == Counter(expected.keys()), "Appendix must cover all 116 IDs exactly once")
    _require(exact == expected, "Appendix notes do not preserve all 116 exact configurations")
    rules = pd.read_csv(output / "EPL_Presentation_Short_Rules.csv")
    canonical_rules = pd.DataFrame([{"strategy_id": row["strategy_id"], **describe_strategy(row)} for row in catalog])
    _same_frame(rules, canonical_rules, "all 116 plain-language rules")
    return {"individual_strategies_covered_once": 116, "full_note_configurations": len(exact),
            "compact_groups_checked": groups, "capital_guide_pages": 1,
            "catalog_pages": len(slides) - 7, "plain_language_rule_rows_reconciled": 116,
            "thresholds_outcomes_odds_consensus_stakes_and_date_caps_reconciled": True}


def audit(output, *, data):
    """Audit one short presentation without mutating any project artifact."""
    output = Path(output).resolve()
    slides = _json(output / "presentation_content.json")
    _require(len(slides) > 7, "Missing six-slide presentation and strategy appendix")
    _require(sum(not slide.get("appendix", False) for slide in slides) == 6
             and not any(slide.get("appendix", False) for slide in slides[:6])
             and all(slide.get("appendix", False) for slide in slides[6:]),
             "Presentation must have exactly six main pages, followed by its appendix")
    _require(slides[0]["title"] == "What we have implemented", "Opening page is not implemented-only")
    intro = _visible(slides[0]).lower()
    _require("collected and checked" in intro and "built models" in intro and "116" in intro,
             "Opening page omits completed system")
    _require(not re.search(r"\b(?:proposal|proposed|will implement|initial idea)\b", intro), "Intro describes an unimplemented idea")
    for spec in slides:
        for source in spec.get("sources", []):
            _require((ROOT / source).is_file(), f"Unknown slide source: {source}")
    source_checks = _audit_sources(data)
    profiles = _audit_profiles(slides, data)
    graph = _audit_graph(slides, output, data)
    pptx = Presentation(output / "EPL_Project_Presentation.pptx")
    appendix = _audit_appendix(slides, pptx, output, data)
    rendering = _audit_render(output, slides, pptx)
    return {"status": "passed", "read_only": True, "main_slides": 6,
            "appendix_slides": len(slides) - 6, "implemented_only_intro": True,
            "sources_and_rankings": source_checks, "featured_models_and_rules": profiles,
            "all_strategy_graph": graph, "strategy_appendix": appendix, "rendering": rendering,
            "no_fitting_forecasting_settlement_or_trading_performed": True}
