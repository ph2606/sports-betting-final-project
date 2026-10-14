"""Render the EPL project draft from immutable saved experiment artifacts.

Run from final_project: python epl_report/build_report.py
No models are fitted, no source/result files are changed, and no NBA results
are generated. All numerical report facts are extracted from saved tables.
"""
from __future__ import annotations

import hashlib
import html
import io
import json
import re
from datetime import datetime, timezone
from pathlib import Path

import pandas as pd
from PIL import Image as PILImage
from pypdf import PdfReader
from docx import Document
from docx.enum.section import WD_ORIENT, WD_SECTION_START
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor
from reportlab.lib import colors
from reportlab.lib.pagesizes import letter, landscape
from reportlab.lib.styles import ParagraphStyle
from reportlab.pdfbase import pdfmetrics
from reportlab.pdfbase.ttfonts import TTFont
from reportlab.platypus import BaseDocTemplate, Frame, Image, NextPageTemplate, PageBreak, PageTemplate, Paragraph, Spacer, Table, TableStyle

ROOT = Path(__file__).resolve().parents[1]
OUT = Path(__file__).resolve().parent
DRAFT_DATE = "2026-10-14"
TITLE = "English Premier League: probability calibration, player features and archived betting prices"
LIVE_RELEASE = "20261014T092230.366245Z"
LIVE_EVALUATION = "20261014T092819.822459Z"


class Inputs:
    def __init__(self):
        self.records = {}

    def bytes(self, path):
        content = path.read_bytes()
        self.records[path.relative_to(ROOT).as_posix()] = {"sha256": hashlib.sha256(content).hexdigest(), "bytes": len(content)}
        return content

    def text(self, path):
        return self.bytes(path).decode("utf-8")

    def json(self, path):
        return json.loads(self.text(path))

    def all_tables(self, folder):
        tables = {}
        for csv in sorted(folder.glob("*.csv")):
            try:
                original = pd.read_csv(io.BytesIO(self.bytes(csv)), low_memory=False)
            except pd.errors.EmptyDataError:
                original = pd.DataFrame()
            self.records[csv.relative_to(ROOT).as_posix()].update(rows=len(original), columns=len(original.columns), read_as="csv")
            parquet = csv.with_suffix(".parquet")
            if parquet.exists():
                frame = pd.read_parquet(io.BytesIO(self.bytes(parquet)))
                self.records[parquet.relative_to(ROOT).as_posix()].update(rows=len(frame), columns=len(frame.columns), read_as="parquet")
                assert len(frame)==len(original) and list(frame.columns)==list(original.columns), f"Inconsistent source table copies: {csv}"
            else:
                frame = original
            tables[csv.stem] = frame
        return tables


def one(frame, **conditions):
    mask = pd.Series(True, index=frame.index)
    for column, value in conditions.items():
        mask &= frame[column].eq(value)
    selected = frame.loc[mask]
    assert len(selected)==1, f"Expected unique result {conditions}, found {len(selected)}"
    return selected.iloc[0]


def decimal(value, places=4):
    return f"{float(value):.{places}f}"


def percent(value, signed=False):
    return f"{float(value)*100:{'+' if signed else ''}.2f}%".replace("-", "−")


def interval(low, high, percentage=False):
    formatter = percent if percentage else lambda number: decimal(number)
    return f"[{formatter(low)}, {formatter(high)}]"


def documents(inputs):
    first = ROOT / "epl_modeling/results"
    second = ROOT / "epl_modeling/ablation_results"
    a, b = inputs.all_tables(first), inputs.all_tables(second)
    first_json = {path.stem: inputs.json(path) for path in sorted(first.glob("*.json"))}
    second_json = {path.stem: inputs.json(path) for path in sorted(second.glob("*.json"))}
    # These are report/protocol inputs, not model imports or rerun instructions.
    for relative in ("epl_modeling/METHODOLOGY.md", "epl_modeling/ABLATION_PROTOCOL.md",
                     "epl_modeling/PROSPECTIVE_PROTOCOL.md", "epl_modeling/FEATURE_COVERAGE.md",
                     "epl_data/SOURCE_GUIDE.md", "epl_data/metadata/fpl_notes.md", "epl_data/metadata/advanced_notes.md"):
        inputs.text(ROOT / relative)
    proposal = PdfReader(io.BytesIO(inputs.bytes(ROOT / "ML_Sports_Betting_Project_Proposal.pdf")))
    assert "September 30, 2026" in " ".join(page.extract_text() or "" for page in proposal.pages)
    membership = pd.read_csv(io.BytesIO(inputs.bytes(ROOT / "epl_data/metadata/current_teams.csv")))
    assert len(membership)==20
    run = first_json["run_metadata"]
    ablation_run = second_json["run_metadata"]
    repair = second_json["source_reconciliation_summary"]
    scores = a["classification_scores"]
    svm = one(scores, model="accuracy", split="test", scope="matched_bookmaker_coverage")
    mlp = one(scores, model="calibration", split="test", scope="matched_bookmaker_coverage")
    market = one(scores, model="bet365_devig", split="test", scope="matched_bookmaker_coverage")
    frequency = one(scores, model="training_class_frequency", split="test", scope="matched_bookmaker_coverage")
    assert int(svm.n_matches)==int(mlp.n_matches)==int(market.n_matches)==int(frequency.n_matches)
    test_n = int(svm.n_matches)
    ci = a["paired_test_metric_intervals"]
    ece_ci = one(ci, metric="classwise_ece_10")
    loss_ci = one(ci, metric="log_loss")
    accuracy_ci = one(ci, metric="accuracy")
    bets = a["betting_summary"]
    secondary = bets[(bets.split=="test") & (bets.policy=="predeclared_5pct")].copy()
    primary = bets[(bets.split=="test") & (bets.policy=="validation_selected")]
    assert primary.bet_count.eq(0).all() and primary.turnover.eq(0).all()
    best_mlp = one(secondary, branch="calibration", odds_mode="best_individual_book", staking="fixed")
    best_svm = one(secondary, branch="accuracy", odds_mode="best_individual_book", staking="fixed")
    table1 = pd.DataFrame([
        {"Predictor": label, "Accuracy (%)": decimal(row.accuracy*100,2), "CW-ECE": decimal(row.classwise_ece_10),
         "Log loss": decimal(row.log_loss), "Brier": decimal(row.brier_score)}
        for label, row in (("Accuracy-selected SVM",svm), ("ECE-selected MLP",mlp),
                          ("Normalized Bet365",market), ("Training frequency",frequency))])
    rows = []
    for branch in ("accuracy", "calibration"):
        for odds in ("best_individual_book", "bet365_reference"):
            for staking in ("fixed", "kelly"):
                row = one(secondary, branch=branch, odds_mode=odds, staking=staking)
                rows.append({"Branch": "Accuracy" if branch=="accuracy" else "Calibration",
                             "Prices": "Best book" if odds=="best_individual_book" else "Bet365",
                             "Stake": "Fixed" if staking=="fixed" else "¼ Kelly",
                             "Bets": str(int(row.bet_count)), "Profit (units)": decimal(row.profit,1),
                             "ROI": percent(row.roi), "Conditional 95% ROI CI": interval(row.roi_ci_low,row.roi_ci_high,True)})
    table2 = pd.DataFrame(rows)
    pooled = b["classification_scores"][(b["classification_scores"].outer_season=="pooled") & (b["classification_scores"].scope=="common_odds_matches")]
    raw = one(pooled, model_type="baseline", feature_group="market_only")
    temperature = one(pooled, model_type="market_temperature", feature_group="market_only")
    group_labels = {"team":"Team", "team_xg":"Team + xG", "team_player":"Team + players", "combined":"Combined"}
    rows3 = []
    for model, groups in (("baseline",["market_only"]), ("market_temperature",["market_only"]),
                          ("logistic",list(group_labels)), ("market_residual",list(group_labels)),
                          ("baseline",["class_frequency"])):
        for group in groups:
            row = one(pooled, model_type=model, feature_group=group)
            label = "Raw market" if model=="baseline" and group=="market_only" else "Temperature market" if model=="market_temperature" else "Frequency" if group=="class_frequency" else ("LR: " if model=="logistic" else "Residual: ")+group_labels[group]
            rows3.append({"Variant": label, "Accuracy (%)": decimal(row.accuracy*100,2), "CW-ECE": decimal(row.classwise_ece_10),
                          "Log loss": decimal(row.log_loss), "Brier": decimal(row.brier_score)})
    table3 = pd.DataFrame(rows3)
    paired = b["paired_comparisons"]
    player_cis = {}
    for model in ("logistic", "market_residual"):
        player_cis[model] = one(paired, outer_season="pooled", scope="common_odds_matches", metric="log_loss",
                               model_type=model, feature_group="team_player", reference_group="team", reference_model_type=model)
    logistic_xg = one(paired, outer_season="pooled", scope="common_odds_matches", metric="log_loss", model_type="logistic",
                      feature_group="team_xg", reference_group="team", reference_model_type="logistic")
    best_residual = pooled[pooled.model_type.eq("market_residual")].sort_values("log_loss").iloc[0]
    residual_temp = one(paired, outer_season="pooled", scope="common_odds_matches", metric="log_loss", model_type="market_residual",
                        feature_group=best_residual.feature_group, reference_group="market_only", reference_model_type="market_temperature")
    evaluation = b["chronological_roles"][b["chronological_roles"].role.eq("evaluation")]
    pooled_n = int(evaluation.n_matches.sum())
    complete_n = int(evaluation.complete_history_matches.sum())
    assert pooled_n==int(raw.n_matches)==int(ablation_run["common_sample_matches"])
    fixed = b["betting_summary"][b["betting_summary"].staking.eq("fixed")]
    pooled_bets = fixed.groupby(["model_type","feature_group"])[["profit","turnover"]].sum()
    pooled_bets["roi"] = pooled_bets.profit/pooled_bets.turnover
    roi_min,roi_max = float(pooled_bets.roi.min()),float(pooled_bets.roi.max())
    assert b["betting_summary"].roi_ci_low.isna().all() and b["betting_summary"].roi_ci_high.isna().all()
    first_features = first_json["feature_metadata"]
    source_matches_before = int(first_features["coverage"]["fpl"]["exact_linked_fixtures"])
    source_matches_after = int(repair["eligible_exact_fixtures"])
    cohorts = "; ".join(membership.team.tolist())
    repair_by_season = ", ".join(f"{value['repaired_missing_opponent']:,} in {season}" for season,value in repair["by_season"].items() if value["repaired_missing_opponent"])
    first_families = first_json["locked_choices"]["families"]
    selected = first_json["selected_features"]
    assert first_families=={"accuracy":"svm","calibration":"mlp"}
    assert len(selected["accuracy"])==2 and len(selected["calibration"])==1
    figures = {
        "coverage": ROOT / "epl_modeling/ablation_results/ablation_history_coverage.png",
        "reliability": ROOT / "epl_modeling/results/test_reliability.png",
        "paired": ROOT / "epl_modeling/ablation_results/ablation_paired_logloss.png",
        "roi": ROOT / "epl_modeling/ablation_results/ablation_betting_roi.png",
    }
    for path in figures.values():
        inputs.bytes(path)
        inputs.bytes(path.with_suffix(".svg"))
    live_folder = ROOT / "epl_live/runs" / LIVE_RELEASE
    live = inputs.json(live_folder / "run_summary.json")
    live_evaluation = inputs.json(ROOT / "epl_live/evaluations" / LIVE_EVALUATION / "evaluation_summary.json")
    live_targets = pd.read_csv(io.BytesIO(inputs.bytes(live_folder / "target_fixtures.csv")))
    live_coverage = pd.read_csv(io.BytesIO(inputs.bytes(live_folder / "forecast_coverage.csv")))
    ledger = [json.loads(line)["record"] for line in inputs.text(ROOT / "epl_live/forecasts.jsonl").splitlines() if line.strip()]
    release_records = [record for record in ledger if record["snapshot_id"]==live["snapshot_id"]]
    released_n = len(release_records)
    released_models = sorted({record["model_id"] for record in release_records})
    assert released_n==live["ledger_append"]["appended"]==live_evaluation["forecast_records"]
    assert all("logistic" in model for model in released_models)
    assert live_evaluation["resolved_prediction_count"]==0 and not live_evaluation["prospective_performance_available"]
    assert len(live_targets)==live["target_fixture_count"]
    kickoff = pd.to_datetime(live_targets.kickoff_utc,utc=True)
    release_span=f"October {kickoff.min().day}–{kickoff.max().day}, 2026"
    withheld_n=int(live_coverage.status.eq("withheld").sum())
    facts = {"first_test_matches":test_n, "first_family_choices":first_families,
             "first_primary_bets":int(primary.bet_count.sum()), "first_secondary_mlp_fixed_roi":float(best_mlp.roi),
             "first_secondary_svm_fixed_roi":float(best_svm.roi), "ablation_matches":pooled_n,
             "ablation_models":int(ablation_run["n_models"]), "ablation_complete_history_matches":complete_n,
             "fpl_repaired_rows":int(repair["missing_opponent_repaired"]), "fpl_fixture_coverage_before":source_matches_before,
             "fpl_fixture_coverage_after":source_matches_after, "fpl_added_fixtures":source_matches_after-source_matches_before,
             "ablation_raw_market_log_loss":float(raw.log_loss), "ablation_raw_market_brier":float(raw.brier_score),
             "ablation_fixed_roi_min":roi_min,"ablation_fixed_roi_max":roi_max,
             "pending_forecast_release":{"run_id":LIVE_RELEASE,"forecasts":released_n,"model_alternatives":len(released_models),
                                         "target_fixtures":len(live_targets),"withheld_quote_dependent_forecasts":withheld_n,
                                         "settled_fixtures":int(live_evaluation["settled_fixtures"]),"performance_available":False},
             "first_experiment_preserved":bool(second_json["first_experiment_integrity"]["preserved"])}
    tables = {"table1_first_probability_scores":table1,"table2_first_secondary_betting":table2,"table3_pooled_ablation_scores":table3}
    pages = []
    def page(title, orientation="portrait"):
        item={"title":title,"orientation":orientation,"blocks":[]};pages.append(item);return item["blocks"]
    def p(blocks,text):blocks.append({"type":"paragraph","text":text})
    def h(blocks,text):blocks.append({"type":"heading","text":text})
    def figure(blocks,key,caption):blocks.append({"type":"figure","path":figures[key],"caption":caption})
    def table(blocks,key,caption):blocks.append({"type":"table","frame":tables[key],"caption":caption})

    blocks=page(TITLE)
    p(blocks,"DRAFT FOR REVIEW — EPL section only | October 14, 2026 | Statistical Arbitrage, Fall 2026")
    h(blocks,"Abstract")
    p(blocks,f"This section studies whether probability-calibration selection improves football forecasts and archived-price betting outcomes, and whether historical player contributions add information beyond team statistics and bookmaker probabilities. The locked first experiment evaluates {test_n} fixtures in 2025/26; a later exploratory ablation compares fixed source groups over {pooled_n:,} matches. Calibration selection lowers estimated calibration error relative to accuracy selection, but the primary betting policy holds cash. The market benchmark has the strongest pooled proper probability scores. The evidence supports a reproducible forecasting study, not an established tradable betting edge.")
    h(blocks,"1. Research question and reference")
    p(blocks,"The September 30 proposal motivates two opportunities: value betting when estimated outcome probabilities exceed price-implied probabilities, and mathematical arbitrage when the best prices across outcomes appear inconsistent. For the EPL contribution, the outcome is home win, draw or away win. We ask whether selecting models by classwise expected calibration error (ECE), rather than accuracy, changes probability quality and hypothetical betting returns, then separately examine xG and player features. These are empirical questions; good probability calibration does not itself determine an advantageous price.")
    p(blocks,"Walsh and Joshi [1] motivate opponent-relative season statistics, feature screening and calibration-oriented model selection. Their 2025 corrigendum [2] reports feature-engineering errors and replaces the original profitability results: corrected average ROI is −9.77% for calibration selection versus −26.78% for accuracy selection. It also states that Kelly staking can fail with a calibrated model. We use the corrected reference and make no profitability guarantee.")
    h(blocks,"2. Population and scope")
    p(blocks,f"The population is historical EPL fixtures involving at least one of the {len(membership)} clubs eligible in 2026/27, as verified in the saved membership list [8]. The cohort comprises {cohorts}. West Ham, Burnley and Wolverhampton Wanderers remain legitimate historical opponents but are excluded as focal clubs. Selecting today's clubs retrospectively creates survivor/selection bias: the population is not the complete historical league. All results in this section are EPL results; no NBA experiment is supplied.")
    p(blocks,f"The {test_n}-fixture test omits six games between excluded clubs, rather than losing otherwise eligible observations.")

    blocks=page("Data, identity and temporal controls")
    h(blocks,"3. Sources and source histories")
    p(blocks,"The saved acquisition snapshot is October 6, 2026; both retrospective experiments were run on October 13.")
    p(blocks,f"Football-Data E0 files [4] provide fixture outcomes, goals, shots, corners, fouls, cards and bookmaker triplets. The modeling frame contains {first_features['rows']:,} existing cohort matches; {first_features['history_scope']['full_raw_e0_matches']:,} saved whole-league matches reconstruct both clubs' basic histories, including excluded-opponent games. Understat [6] supplies displayed team xG through exact integrated fixture links. Official FPL [5] and a version-pinned public archive [7] supply actual-club player-fixture contributions. ESPN profiles, injury/valuation snapshots and other acquired descriptive sources support data exploration but do not enter historical predictors.")
    p(blocks,f"The original early FPL table lacked opponent names on {repair['missing_opponent_repaired']:,} player rows ({repair_by_season}). A separate read-only reconciliation required one exact season/date/actual-club/venue E0 candidate, with known names, season-local IDs and same-fixture peers agreeing. It repaired all those names without conflicting accepted fixtures and increased usable FPL coverage from {source_matches_before:,} to {source_matches_after:,} fixtures, adding {source_matches_after-source_matches_before:,}. Official current-season history supersedes overlapping archive rows; {repair['assistant_manager_rows_excluded']} assistant-manager records are excluded. Original acquisition and first-study artifacts are preserved.")
    h(blocks,"4. Features and time discipline")
    p(blocks,"Team inputs are season-to-date own-minus-opponent box-score averages, prior-season EPL win percentage, and five-/ten-match form. Promoted clubs have explicit missing-history flags. Source-specific xG and FPL aggregates are lagged from earlier completed fixtures; same-calendar-date outcomes cannot update each other. Missing source observations remain missing with observed-count controls. Neither the target score/xG nor its realised lineup/minutes enters its own prediction. Current roster, injury, ownership and price snapshots are excluded. Lagged FPL contributions measure past performance, not future player availability.")
    p(blocks,"Chronological separation also governs fitted preprocessing, feature selection, regularization and calibration. Match-date ordering still does not prove that retrospective statistics were published before an archived quote: provider corrections and original publication clocks are unverified. Figure 1 describes source support rather than available starting lineups.")
    p(blocks,"Median imputation with missingness indicators and standardization use fitted training values. Screening thresholds are 40% missingness, BH-adjusted KS q<0.01 with distance>0.15, and absolute Spearman correlation≥0.9. Internal search scores are optimization diagnostics, not unbiased nested estimates. Base estimators/preprocessing are refitted through 2021/22 and frozen before calibration. One positive temperature minimizes calibration-season log loss without changing the underlying features.")
    figure(blocks,"coverage",f"Figure 1. Reused ablation source-support figure. The shared complete-history sensitivity has {complete_n} matches: both clubs have at least eight observed xG and FPL slots among their last ten earlier E0 games, with fixed performance inputs observed. This is coverage, not roster availability.")

    blocks=page("Locked first-stage probability comparison","landscape")
    h(blocks,"5. Chronology, model selection and metrics")
    p(blocks,f"Training spans 2014/15–2020/21; 2021/22 tunes hyperparameters, 2022/23 fits held-out temperatures [3], 2023/24 selects model families, and 2024/25 selects betting thresholds. The 2025/26 test follows all locked decisions. Each branch considers logistic regression, random forest, SVM and MLP with {run['trials_per_family_per_branch']} seeded TPE trials per family. Training-only missingness, KS drift and Spearman redundancy screening precede an 18-candidate forward search. The resulting accuracy branch uses an SVM with two corners/shot predictors; the ECE branch uses an MLP with one corners predictor. No player/xG predictor survives this limited search, so it cannot establish incremental player benefit.")
    p(blocks,"Both branches receive scalar temperature calibration on the same separate season; this compares selection procedures, not calibrated versus deliberately uncalibrated models. Probabilities are ordered H/D/A and sum to one. Accuracy is the argmax hit rate. Ten-bin classwise ECE averages weighted absolute probability/frequency gaps across all three one-vs-rest classes. Log loss is the mean negative log probability of the realised class; multiclass Brier is the mean sum of three squared probability errors, ranging from zero to two. Lower losses/ECE are better. The reference market probabilities normalize inverse Bet365 prices; training class frequencies provide a resolution-poor baseline.")
    table(blocks,"table1_first_probability_scores",f"Table 1. Locked 2025/26 test scores on the same {test_n} complete-reference-price fixtures. ECE is on [0,1]; Brier is the unscaled multiclass sum.")
    p(blocks,f"The calibration-minus-accuracy ECE difference is {decimal(ece_ci.difference_calibration_minus_accuracy)}, with conditional 95% interval {interval(ece_ci.ci_lower,ece_ci.ci_upper)}. Its log-loss difference is {decimal(loss_ci.difference_calibration_minus_accuracy)}, interval {interval(loss_ci.ci_lower,loss_ci.ci_upper)}. Both include zero. The accuracy difference is {decimal(accuracy_ci.difference_calibration_minus_accuracy*100,2)} percentage points, interval [{decimal(accuracy_ci.ci_lower*100,2)}, {decimal(accuracy_ci.ci_upper*100,2)}] points. The 1,000 date draws condition on fitted choices and omit selection uncertainty.")
    figure(blocks,"reliability","Figure 2. Reused locked-test reliability curves; the diagonal denotes equality between forecast probability and empirical outcome frequency. Only plotted bins with at least five matches are shown; scores use all matches. Sparse-bin deviations are uncertain.")

    blocks=page("Archived-price replay and the first betting result")
    h(blocks,"6. Prices, staking and arbitrage mathematics")
    p(blocks,"For decimal price O and forecast p, estimated return per unit stake is pO−1. The secondary threshold is 5% estimated return, not a five-percentage-point probability gap. Each fixture contributes at most one chosen outcome. Reference-book prices are compared with the best complete individual-book triplets. Aggregate maxima/averages, exchanges and unreliable Pinnacle records from July 23, 2025 are excluded under source-quality rules [4]. Preclosing means an earlier archived snapshot, not a verified opening or contemporaneous quote.")
    p(blocks,"Fixed staking uses 10 units from 1,000 initial units. Quarter Kelly uses one-quarter of max(0,(pO−1)/(O−1)), capped at 2% of date-start capital per bet. Same-date commitments are scaled before settlement; that day's winnings cannot finance another commitment. ROI is profit divided by total stakes, not profit divided by initial bankroll. Price shopping can alter both the chosen outcome and qualifying bets, so its contrast measures a policy difference rather than an isolated causal price effect.")
    h(blocks,"7. Primary decision and predeclared secondary replay")
    p(blocks,"The primary policy chooses its threshold using only 2024/25: at least thirty validation bets and positive fixed-stake ROI are required. Neither branch qualifies. Both therefore hold cash on the locked test, with zero turnover and undefined ROI. The separately predeclared 5% policy is reported transparently as a secondary comparison; it does not replace the failed primary policy.")
    table(blocks,"table2_first_secondary_betting",f"Table 2. Secondary 5% archived-price replay on {test_n} fixtures. Profit is in bankroll units; ROI intervals use 1,000 whole-date bootstrap draws conditional on recorded bets and stakes.")
    p(blocks,f"The calibration model's best-book fixed-stake point ROI is {percent(best_mlp.roi)}, versus {percent(best_svm.roi)} for accuracy selection; their intervals cross zero. The small positive point estimate is not reliable evidence of an edge. Kelly does not create information absent from the forecasts and increases sensitivity to probability errors.")
    p(blocks,"A mathematical arbitrage screen checks 1/O_H+1/O_D+1/O_A<1 for best outcome prices, with equal-payout stakes proportional to inverse odds. This arithmetic requires compatible contracts and simultaneous accepted prices. The locked test has no such archived candidate. All test quotes lack verified compatible observation clocks; simultaneous availability, limits, fees and accepted execution cannot be established. No arbitrage profit is credited. A feature before kickoff may still postdate an earlier Friday/Tuesday quote, so this is a hypothetical archived-price replay.")

    blocks=page("Exploratory fixed-group source ablation","landscape")
    h(blocks,"8. Repaired-source study and conditional contrasts")
    p(blocks,f"The second study evaluates four source groups—Team ({ablation_run['feature_group_counts']['team']} inputs), Team+xG ({ablation_run['feature_group_counts']['team_xg']}), Team+players ({ablation_run['feature_group_counts']['team_player']}) and Combined ({ablation_run['feature_group_counts']['combined']})—with logistic and market-residual models. Across four outer seasons, 2022/23–2025/26, this yields {ablation_run['n_models']} fitted group models and {pooled_n:,} common quoted-price matches. Each fold trains/tunes before a separate calibration season, then evaluates the next season; regularization uses the fixed three-value grid. All groups share source-coverage controls. Market-residual scores add learned corrections to log normalized bookmaker probabilities; raw and temperature-only markets separate feature information from recalibration.")
    p(blocks,f"This is exploratory retrospective work: 2025/26 was already inspected in stage one, and correct rolling fitting order cannot make it a fresh confirmatory holdout. Within logistic models, adding xG changes pooled log loss by {decimal(logistic_xg.difference)}, interval {interval(logistic_xg.ci_lower,logistic_xg.ci_upper)}. Player-minus-team changes are {decimal(player_cis['logistic'].difference)} for logistic, interval {interval(player_cis['logistic'].ci_lower,player_cis['logistic'].ci_upper)}, and {decimal(player_cis['market_residual'].difference)} for residual models, interval {interval(player_cis['market_residual'].ci_lower,player_cis['market_residual'].ci_upper)}. Both player intervals include zero. These conditional, unadjusted comparisons do not establish incremental player benefit.")
    figure(blocks,"paired","Figure 3. Reused seasonal paired log-loss contrasts against the same model type's Team group. Negative values favor the added inputs. Conditional 95% intervals use 1,000 whole-date draws within season; pooled draws stratify by season. They omit tuning/fitting uncertainty and have no multiple-comparison adjustment.")
    market_discussion=f"The raw market has the lowest observed pooled proper scores: log loss {decimal(raw.log_loss)} and Brier {decimal(raw.brier_score)}. The lowest-loss residual variant has point log loss {decimal(residual_temp.difference)} higher than temperature-only market. The {complete_n}-match complete-history sensitivity changes descriptive rankings, illustrating population dependence rather than validating a chosen production model."

    blocks=page("Four-season probability and return summaries","landscape")
    table(blocks,"table3_pooled_ablation_scores",f"Table 3. Pooled scores on {pooled_n:,} identical fixtures. LR denotes logistic regression. Pooled ECE is recomputed from predictions, not averaged across seasonal ECE values.")
    figure(blocks,"roi",f"Figure 4. Reused fixed-5% ablation replay. Profit/turnover combines separate season bankroll resets; it is not a continuous track record. Fixed-stake point ROI ranges {percent(roi_min)} to {percent(roi_max)}. This run estimated no betting ROI intervals; blank bounds do not imply certainty or a reliable edge.")

    blocks=page("Discussion, limits and prospective evaluation plan")
    h(blocks,"9. Evidence and limitations")
    p(blocks,market_discussion)
    p(blocks,"The first study observes an accuracy/calibration trade-off between selected procedures, while its admissibility rule rejects active primary betting. The later study finds modest, model-dependent source-feature changes and stronger overall proper scores for bookmaker probabilities. Neither establishes a dependable incremental player advantage or profitable deployable system. ECE alone is insufficient: class frequencies can attain low error while lacking useful discrimination. Differences also combine feature and model-family choices, so they cannot isolate a causal effect of calibration itself.")
    p(blocks,"The fixed present-day cohort restricts generalization. Historical opponents have incomplete xG/FPL histories outside saved cohort-linked fixtures; changes in source definitions and fantasy scoring can produce drift. Displayed team xG is distinct from additive shot/player xG. Retrospective publication/revision clocks, synchronized prices, fill acceptance, costs and limits are unavailable. Date bootstrap intervals do not cover all dependence or model-selection uncertainty. Exploring many variants after inspecting outcomes adds multiplicity and prevents a new confirmatory interpretation.")
    h(blocks,"10. Prospective protocol and pending release")
    p(blocks,"The saved prospective protocol defines an append-only forecast ledger after its committed cutoff. It requires pre-kickoff H/D/A probabilities, code/configuration hashes, latest contributing source dates and actual source-availability/quote clocks; missing forecasts and rejected fills must also be retained. All eight declared alternatives and market benchmarks remain fixed, with 2025/26 reassigned to calibration rather than reused as a new test. Primary evaluation would use log loss on common valid forecasts, with Brier, ECE, accuracy and separately identified paper-trading summaries.")
    p(blocks,f"Operational update, October 14: the committed ledger contains {released_n} forecasts from {len(released_models)} logistic alternatives for {len(live_targets)} fixtures on {release_span}. Market-dependent alternatives are withheld: {withheld_n} forecast slots lack suitable EPL quotes. Outcomes remain pending, so no prospective performance is reported. The next step is fresh quote capture and settlement of these future matches, retaining missingness and evaluating without outcome-driven protocol changes.")
    h(blocks,"References and reproducibility")
    references = [
        ("[1] Walsh, C. & Joshi, A. (2024). Machine learning for sports betting: Should model selection be based on accuracy or calibration? Machine Learning with Applications, 16, 100539.","https://doi.org/10.1016/j.mlwa.2024.100539"),
        ("[2] Walsh, C. & Joshi, A. (2025). Corrigendum to the 2024 sports-betting article. Machine Learning with Applications, 19, 100627.","https://doi.org/10.1016/j.mlwa.2025.100627"),
        ("[3] Guo, C., Pleiss, G., Sun, Y. & Weinberger, K. Q. (2017). On Calibration of Modern Neural Networks. PMLR, 70, 1321–1330.","https://proceedings.mlr.press/v70/guo17a.html"),
        ("[4] Football-Data. Historical results/odds definitions and timing/reliability notes.","https://www.football-data.co.uk/data.php"),
        ("[5] Premier League. Official Fantasy Premier League data service.","https://fantasy.premierleague.com/api/bootstrap-static/"),
        ("[6] Understat. EPL expected-goals data.","https://understat.com/league/EPL/2026"),
        ("[7] Anand, V. Fantasy Premier League public historical archive; snapshot revision recorded in source manifests.","https://github.com/vaastav/Fantasy-Premier-League"),
        ("[8] Premier League. Official 2026/27 membership reference recorded in current_teams.csv.",membership.membership_source.iloc[0]),
    ]
    for title,url in references:
        blocks.append({"type":"reference","text":title,"url":url})
    p(blocks,"Reproduce the draft with epl_report/build_report.py. Source tables, protocols, figures and hashes are listed in report_manifest.json; no report step retrains models. The proposal and immutable first/ablation outputs remain separate from this editable draft.")
    return pages, tables, facts, inputs


def markdown(pages):
    chunks=[]
    for number,page in enumerate(pages,start=1):
        chunks.extend([("# " if number==1 else "## ")+page["title"], ""])
        for block in page["blocks"]:
            if block["type"]=="paragraph":chunks.extend([block["text"],""])
            elif block["type"]=="heading":chunks.extend(["### "+block["text"],""])
            elif block["type"]=="figure":
                path="../"+block["path"].relative_to(ROOT).as_posix()
                chunks.extend([f"![{block['caption']}]({path})","",block["caption"],""])
            elif block["type"]=="table":
                chunks.extend([block["caption"],""])
                frame=block["frame"]
                chunks.append("| "+" | ".join(frame.columns)+" |")
                chunks.append("| "+" | ".join("---" for _ in frame.columns)+" |")
                for row in frame.itertuples(index=False,name=None):chunks.append("| "+" | ".join(str(value) for value in row)+" |")
                chunks.append("")
            else:chunks.extend([f"{block['text']} [{block['url']}]({block['url']})",""])
    return "\n".join(chunks)


def render_pdf(pages,path):
    fonts=Path("C:/Windows/Fonts")
    for name,file in (("Body","times.ttf"),("BodyBold","timesbd.ttf"),("BodyItalic","timesi.ttf"),("Sans","arial.ttf"),("SansBold","arialbd.ttf")):
        pdfmetrics.registerFont(TTFont(name,str(fonts/file)))
    styles={
        "body":ParagraphStyle("body",fontName="Body",fontSize=10.4,leading=12.5,spaceAfter=5),
        "title":ParagraphStyle("title",fontName="SansBold",fontSize=16,leading=18,spaceAfter=8,textColor=colors.HexColor("#12364D")),
        "heading":ParagraphStyle("heading",fontName="SansBold",fontSize=11.3,leading=13,spaceBefore=4,spaceAfter=4),
        "caption":ParagraphStyle("caption",fontName="BodyItalic",fontSize=8.8,leading=10.4,spaceAfter=4),
        "ref":ParagraphStyle("ref",fontName="Body",fontSize=8.5,leading=10.1,spaceAfter=4),
        "cell":ParagraphStyle("cell",fontName="Sans",fontSize=8.6,leading=9.6),
        "headcell":ParagraphStyle("headcell",fontName="SansBold",fontSize=8.6,leading=9.6,textColor=colors.white),
    }
    def header(canvas,doc):
        w,h=canvas._pagesize
        canvas.setFont("Sans",8)
        canvas.setFillColor(colors.HexColor("#52616B"))
        canvas.drawString(36,h-21,"EPL PROJECT SECTION — DRAFT FOR REVIEW | October 14, 2026")
        canvas.drawRightString(w-36,21,f"{doc.page} / {len(pages)}")
    doc=BaseDocTemplate(str(path),pagesize=letter,leftMargin=36,rightMargin=36,topMargin=36,bottomMargin=36,
                        title=TITLE+" — draft",author="Aman Dhillon and Panagiotis Housos")
    templates=[]
    for orientation,size in (("portrait",letter),("landscape",landscape(letter))):
        w,h=size
        templates.append(PageTemplate(id=orientation,frames=[Frame(36,36,w-72,h-72,leftPadding=0,rightPadding=0,topPadding=0,bottomPadding=0)],pagesize=size,onPage=header))
    doc.addPageTemplates(templates)
    story=[]
    for number,page in enumerate(pages):
        if number:
            story.extend([NextPageTemplate(page["orientation"]),PageBreak()])
        width=letter[0]-72 if page["orientation"]=="portrait" else landscape(letter)[0]-72
        story.append(Paragraph(html.escape(page["title"]),styles["title"]))
        for block in page["blocks"]:
            kind=block["type"]
            if kind in ("paragraph","heading"):
                story.append(Paragraph(html.escape(block["text"]),styles["body" if kind=="paragraph" else "heading"]))
            elif kind=="reference":
                text=html.escape(block["text"])+f' <link href="{html.escape(block["url"])}" color="#174D72">{html.escape(block["url"])}</link>'
                story.append(Paragraph(text,styles["ref"]))
            elif kind=="figure":
                source=PILImage.open(block["path"])
                figure_width=width-10
                story.append(Image(str(block["path"]),width=figure_width,height=figure_width*source.height/source.width))
                story.append(Paragraph(html.escape(block["caption"]),styles["caption"]))
            else:
                story.append(Paragraph(html.escape(block["caption"]),styles["caption"]))
                frame=block["frame"]
                raw=[list(frame.columns)]+frame.astype(str).values.tolist()
                data=[[Paragraph(html.escape(str(value)),styles["headcell" if row==0 else "cell"]) for value in line] for row,line in enumerate(raw)]
                if len(frame.columns)==7:weights=[.14,.13,.10,.06,.12,.09,.36]
                else:weights=[.36,.16,.16,.16,.16]
                table=Table(data,colWidths=[width*weight for weight in weights],repeatRows=1,hAlign="CENTER")
                table.setStyle(TableStyle([("BACKGROUND",(0,0),(-1,0),colors.HexColor("#12364D")),
                                           ("ROWBACKGROUNDS",(0,1),(-1,-1),[colors.white,colors.HexColor("#F0F4F6")]),
                                           ("VALIGN",(0,0),(-1,-1),"MIDDLE"),("LEFTPADDING",(0,0),(-1,-1),4),
                                           ("RIGHTPADDING",(0,0),(-1,-1),4),("TOPPADDING",(0,0),(-1,-1),.7),
                                           ("BOTTOMPADDING",(0,0),(-1,-1),.7),("LINEBELOW",(0,-1),(-1,-1),.5,colors.HexColor("#C7D0D5"))]))
                story.extend([table,Spacer(1,4)])
    doc.build(story)


def render_docx(pages,path):
    document=Document()
    normal=document.styles["Normal"]
    normal.font.name="Times New Roman";normal.font.size=Pt(10.4)
    normal.paragraph_format.line_spacing=Pt(12.5);normal.paragraph_format.space_after=Pt(5)
    for style in ("Heading 1","Heading 2"):
        document.styles[style].font.name="Arial";document.styles[style].font.color.rgb=RGBColor.from_string("12364D")
    document.styles["Heading 1"].font.size=Pt(16)
    document.styles["Heading 2"].font.size=Pt(11.3)
    document.styles["Heading 1"].paragraph_format.line_spacing=Pt(18)
    document.styles["Heading 1"].paragraph_format.space_before=Pt(0)
    document.styles["Heading 1"].paragraph_format.space_after=Pt(8)
    document.styles["Heading 2"].paragraph_format.line_spacing=Pt(13)
    document.styles["Heading 2"].paragraph_format.space_before=Pt(4)
    document.styles["Heading 2"].paragraph_format.space_after=Pt(4)
    document.core_properties.title=TITLE+" — draft"
    document.core_properties.author="Aman Dhillon and Panagiotis Housos"
    for number,page in enumerate(pages):
        section=document.sections[0] if number==0 else document.add_section(WD_SECTION_START.NEW_PAGE)
        section.orientation=WD_ORIENT.LANDSCAPE if page["orientation"]=="landscape" else WD_ORIENT.PORTRAIT
        section.page_width=Inches(11 if page["orientation"]=="landscape" else 8.5)
        section.page_height=Inches(8.5 if page["orientation"]=="landscape" else 11)
        section.top_margin=section.bottom_margin=section.left_margin=section.right_margin=Inches(.5)
        section.header_distance=section.footer_distance=Inches(.2)
        if number==0:
            header=section.header.paragraphs[0]
            run=header.add_run("EPL PROJECT SECTION — DRAFT FOR REVIEW | October 14, 2026");run.font.name="Arial";run.font.size=Pt(8)
            footer=section.footer.paragraphs[0];footer.alignment=2
            field=OxmlElement("w:fldSimple");field.set(qn("w:instr"),"PAGE");footer._p.append(field)
        document.add_heading(page["title"],level=1)
        for block in page["blocks"]:
            kind=block["type"]
            if kind=="paragraph":document.add_paragraph(block["text"])
            elif kind=="heading":document.add_heading(block["text"],level=2)
            elif kind=="reference":
                paragraph=document.add_paragraph(block["text"]+" "+block["url"])
                for run in paragraph.runs:run.font.size=Pt(8.5)
                paragraph.paragraph_format.line_spacing=Pt(10.1)
            elif kind=="figure":
                document.add_picture(str(block["path"]),width=Inches(9.85 if page["orientation"]=="landscape" else 7.35))
                caption=document.add_paragraph(block["caption"])
                for run in caption.runs:run.italic=True;run.font.size=Pt(8.8)
                caption.paragraph_format.line_spacing=Pt(10.4)
            else:
                caption=document.add_paragraph(block["caption"])
                for run in caption.runs:run.italic=True;run.font.size=Pt(8.8)
                caption.paragraph_format.line_spacing=Pt(10.4)
                frame=block["frame"]
                table=document.add_table(rows=1,cols=len(frame.columns));table.style="Light Shading Accent 1"
                for cell,value in zip(table.rows[0].cells,frame.columns):cell.text=str(value)
                for values in frame.itertuples(index=False,name=None):
                    for cell,value in zip(table.add_row().cells,values):cell.text=str(value)
                for row in table.rows:
                    for cell in row.cells:
                        for paragraph in cell.paragraphs:
                            paragraph.paragraph_format.space_after=Pt(0);paragraph.paragraph_format.line_spacing=Pt(9.6)
                            for run in paragraph.runs:run.font.name="Arial";run.font.size=Pt(8.6)
                        margins=OxmlElement("w:tcMar")
                        for side in ("top","bottom"):
                            margin=OxmlElement("w:"+side);margin.set(qn("w:w"),"14");margin.set(qn("w:type"),"dxa");margins.append(margin)
                        cell._tc.get_or_add_tcPr().append(margins)
    document.save(path)


def main():
    OUT.mkdir(parents=True,exist_ok=True)
    inputs=Inputs()
    inputs.bytes(Path(__file__).resolve())
    pages,tables,facts,inputs=documents(inputs)
    text=markdown(pages)
    narrative=" ".join(block.get("text","") for page in pages for block in page["blocks"] if block["type"] in ("paragraph","heading"))
    words=len(re.findall(r"\b[\w]+(?:[’'-][\w]+)*\b",narrative))
    assert 1800<=words<=2300,f"Narrative outside requested word target: {words}"
    for name,frame in tables.items():
        (OUT/"tables").mkdir(exist_ok=True)
        frame.to_csv(OUT/"tables"/(name+".csv"),index=False)
    md,pdf,docx=[OUT/("EPL_project_draft"+suffix) for suffix in (".md",".pdf",".docx")]
    md.write_text(text,encoding="utf-8")
    render_pdf(pages,pdf);render_docx(pages,docx)
    reader=PdfReader(pdf)
    assert len(reader.pages)<=7,f"PDF exceeded seven planned pages: {len(reader.pages)}"
    assert len(reader.pages)==len(pages),"A planned report section was omitted or overflowed"
    extracted="\n".join(page.extract_text() or "" for page in reader.pages)
    assert "DRAFT FOR REVIEW" in extracted and "17,465" in extracted and "1,470" in extracted
    source_unchanged=all(hashlib.sha256((ROOT/path).read_bytes()).hexdigest()==record["sha256"] for path,record in inputs.records.items())
    assert source_unchanged,"An input file changed during report build"
    outputs={path.relative_to(OUT).as_posix():{"sha256":hashlib.sha256(path.read_bytes()).hexdigest(),"bytes":path.stat().st_size}
             for path in [md,pdf,docx,OUT/"README.md",OUT/"requirements.txt"]+sorted((OUT/"tables").glob("*.csv"))}
    manifest={"artifact_status":"draft_for_review_not_final_submission","draft_date":DRAFT_DATE,
              "built_at_utc":datetime.now(timezone.utc).isoformat(),"data_snapshot_as_of":"2026-10-06",
              "study_run_date":"2026-10-13","narrative_word_count":words,"narrative_whitespace_tokens":len(narrative.split()),
              "word_count_scope":"Body paragraphs/headings only; punctuation-aware tokens; excludes captions, table cells and references.",
              "pdf_extracted_tokens_including_tables_references_headers":len(extracted.split()),"pdf_pages":len(reader.pages),
              "planned_pages":len(pages),"research_figures":4,"key_tables":3,"models_retrained":False,
              "input_files_unchanged":source_unchanged,"derived_facts":facts,"inputs":inputs.records,"outputs":outputs,
              "corrigendum_summary_words":len(next(block["text"] for page in pages for block in page["blocks"] if "Walsh and Joshi [1]" in block.get("text","")).split()),"corrigendum_quotes":0,
              "primary_references_verified_on_client_date":"2026-10-14",
              "figure_policy":"Original saved figures reused unchanged; landscape pages preserve readability.",
              "limits":"Stage two is exploratory; the dated operational update contains pending forecasts, with no prospective performance."}
    (OUT/"report_manifest.json").write_text(json.dumps(manifest,indent=2,ensure_ascii=False,allow_nan=False),encoding="utf-8")
    print(json.dumps({"narrative_words":words,"pdf_pages":len(reader.pages),"inputs_hashed":len(inputs.records),"outputs":list(outputs)},indent=2))


if __name__=="__main__":
    main()
