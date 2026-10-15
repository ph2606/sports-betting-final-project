"""Readable temporal/selection companion, complete curves and final rule appendix.

Read finished tables only.  The builder writes its envelope after reporting so
the envelope and independent audit are deliberately not consumed inputs.
"""
from __future__ import annotations

import argparse
import base64
from datetime import datetime, timezone
import html
import json
from pathlib import Path

import matplotlib
matplotlib.use("Agg")
import matplotlib.dates as mdates
import matplotlib.pyplot as plt
import numpy as np
import pandas as pd
import plotly.graph_objects as go
from docx import Document
from docx.enum.section import WD_ORIENT, WD_SECTION_START
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt
from PIL import Image as PILImage
from pypdf import PdfReader
from reportlab.lib import colors
from reportlab.lib.pagesizes import landscape, letter
from reportlab.lib.styles import ParagraphStyle
from reportlab.platypus import (BaseDocTemplate, Frame, Image, KeepTogether,
    NextPageTemplate, PageBreak, PageTemplate, Paragraph, Table, TableStyle)

from epl_strategy.report import number
from epl_strategy_lab import report as lab

ROOT = Path(__file__).resolve().parents[1]
LEGACY = ROOT / "epl_strategy_lab/results/20261015T040052Z_100k84"
TITLE = "EPL strategy validation: four seasons and selection uncertainty"
INITIAL = 100000.0
COMPATIBLE_IDS = [f"S{i:03d}" for i in range(17, 73)]
ALL_IDS = [f"S{i:03d}" for i in range(1, 85)]
APPENDIX_TITLE = "Appendix — every strategy explained (final section)"
SEASONS = ["2022-23", "2023-24", "2024-25", "2025-26"]
BODY_WEIGHTS = (.06, .055, .06, .07, .105, .105, .12, .12, .10, .105, .10)


def _load(output: Path) -> dict:
    data = {name: pd.read_parquet(output / (name + ".parquet")) for name in (
        "metrics_continuous", "metrics_seasonal", "season_stability",
        "bets_continuous", "equity_continuous", "bets_seasonal", "equity_seasonal",
        "selection_comparison", "selection_global")}
    data["catalog"] = json.loads((output / "catalog.json").read_text(encoding="utf-8"))
    data["compatible_catalog"] = json.loads((output / "compatible_catalog.json").read_text(encoding="utf-8"))
    data["graph"] = pd.read_csv(output / "graph_equity_continuous.csv", parse_dates=["match_date"])
    data["legacy_metrics"] = pd.read_parquet(LEGACY / "metrics.parquet")
    return data


def _validate(data: dict) -> pd.Timestamp:
    if [c["strategy_id"] for c in data["catalog"]] != ALL_IDS:
        raise ValueError("Full catalog must preserve S001–S084 in order")
    if [c["strategy_id"] for c in data["compatible_catalog"]] != COMPATIBLE_IDS:
        raise ValueError("Four-season catalog must preserve S017–S072 only")
    if data["compatible_catalog"] != [c for c in data["catalog"] if c["strategy_id"] in COMPATIBLE_IDS]:
        raise ValueError("The temporal catalog changes a frozen rule")
    expected = set(COMPATIBLE_IDS)
    for name in ("metrics_continuous", "season_stability"):
        frame = data[name]
        if len(frame) != 56 or frame.strategy_id.duplicated().any() or set(frame.strategy_id) != expected:
            raise ValueError(name + " is not one-to-one with all 56 compatible strategies")
    metric = data["metrics_continuous"].set_index("strategy_id").reindex(COMPATIBLE_IDS)
    if not metric.prediction_count.eq(1470).all() or not metric.matched_valid_odds_count.eq(1470).all():
        raise ValueError("Continuous metrics do not share all 1,470 opportunities")
    if not np.allclose(metric.initial_capital_sum_units, INITIAL, rtol=0, atol=1e-6):
        raise ValueError("A continuous strategy resets its 100,000 account")
    seasonal = data["metrics_seasonal"]
    if len(seasonal) != 224 or set(seasonal.strategy_id) != expected:
        raise ValueError("Expected four annual reset rows for every compatible strategy")
    if not seasonal.groupby("strategy_id").size().eq(4).all():
        raise ValueError("Missing annual strategy metrics")
    if not np.allclose(seasonal.initial_capital_sum_units, INITIAL, rtol=0, atol=1e-6):
        raise ValueError("Annual comparisons must each start at 100,000")
    equity = data["equity_continuous"].copy()
    equity["match_date"] = pd.to_datetime(equity.match_date)
    graph = data["graph"]
    if not pd.api.types.is_bool_dtype(graph.is_initial):
        flags = graph.is_initial.astype(str).str.lower()
        if not flags.isin(["true", "false"]).all():
            raise ValueError("Invalid initial-point flags")
        graph["is_initial"] = flags.eq("true")
    if set(graph.strategy_id) != expected or set(equity.strategy_id) != expected:
        raise ValueError("Graph or equity omits a compatible strategy")
    if graph.duplicated(["strategy_id", "match_date"]).any():
        raise ValueError("Repeated plotted date")
    dates = sorted(equity.match_date.unique())
    if len(dates) != 453 or pd.Timestamp(dates[0]) != pd.Timestamp("2022-08-05") or pd.Timestamp(dates[-1]) != pd.Timestamp("2026-05-24"):
        raise ValueError("Unexpected four-season settlement window")
    anchor = pd.Timestamp(dates[0]) - pd.Timedelta(days=1)
    if len(graph) != 56 * 454 or not np.isfinite(graph.end_bankroll).all() or graph.end_bankroll.lt(-1e-6).any():
        raise ValueError("Invalid graph balances or incomplete graph")
    for sid, part in graph.groupby("strategy_id", sort=False):
        initial = part.loc[part.is_initial]
        if len(initial) != 1 or initial.match_date.iloc[0] != anchor or not np.isclose(initial.end_bankroll.iloc[0], INITIAL, rtol=0, atol=1e-8):
            raise ValueError("Missing exact 100,000 pre-match anchor for " + sid)
        if set(part.loc[~part.is_initial, "match_date"]) != set(pd.to_datetime(dates)):
            raise ValueError("A graph curve omits a settlement date: " + sid)
    joined = graph.loc[~graph.is_initial].merge(equity[["strategy_id", "match_date", "end_bankroll"]], on=["strategy_id", "match_date"], validate="one_to_one", suffixes=("_graph", "_ledger"))
    if len(joined) != len(equity) or not np.allclose(joined.end_bankroll_graph, joined.end_bankroll_ledger, rtol=0, atol=1e-6):
        raise ValueError("Plotted balances differ from the saved equity ledger")
    final = graph.sort_values("match_date").groupby("strategy_id").tail(1).set_index("strategy_id").end_bankroll.reindex(COMPATIBLE_IDS)
    if not np.allclose(final, INITIAL + metric.profit, rtol=0, atol=1e-5):
        raise ValueError("Graph endpoints do not reconcile with continuous profit")
    bets = data["bets_continuous"]
    counts = bets.groupby("strategy_id").size().reindex(COMPATIBLE_IDS, fill_value=0)
    if not np.array_equal(counts.to_numpy(), metric.bet_count.to_numpy()):
        raise ValueError("Continuous trade counts differ from the bet ledger")
    comparison, global_ = data["selection_comparison"], data["selection_global"]
    if len(comparison) != 672 or len(global_) != 8:
        raise ValueError("Selection tables must contain 84 × 2 × 4 comparisons and eight global rows")
    if comparison.duplicated(["strategy_id", "benchmark_id", "block_length"]).any() or set(comparison.strategy_id) != set(ALL_IDS):
        raise ValueError("Selection diagnostics omit or duplicate a catalog rule")
    if not comparison.groupby(["benchmark_id", "block_length"]).size().eq(84).all():
        raise ValueError("Each conditional comparison must retain all 84 rules")
    return anchor


def _appendix(catalog: list[dict]) -> list[dict]:
    entries = lab._appendix(catalog)
    for entry in entries:
        compatible = entry["strategy_id"] in COMPATIBLE_IDS
        entry["coverage"] = ("Four-season rolling replay: 2022/23–2025/26; also included in the original 84-strategy 2025/26 comparison."
            if compatible else "Original 2025/26 comparison only. No matching earlier frozen predictions exist for this first-study recipe, so no earlier curve is fabricated.")
        entry["explanation"] = entry["coverage"] + " " + entry["explanation"]
    return entries


def _graphs(catalog: list[dict], graph: pd.DataFrame, figures: Path, entries: list[dict]) -> dict:
    figures.mkdir(exist_ok=True)
    styles = lab._styles(len(catalog))
    fig, ax = plt.subplots(figsize=(12.8, 7.2))
    fig.subplots_adjust(left=.085, right=.985, top=.91, bottom=.30)
    dash = {"solid": "-", "dash": "--", "dashdot": "-.", "dot": ":", "longdash": (0, (8, 3))}
    interactive = go.Figure()
    for index, config in enumerate(catalog):
        sid = config["strategy_id"]
        part = graph.loc[graph.strategy_id.eq(sid)].sort_values("match_date")
        color, line = styles[index]
        ax.step(part.match_date, part.end_bankroll, where="post", color=color, linestyle=dash[line], linewidth=.9, alpha=.9, label=sid)
        interactive.add_trace(go.Scatter(x=part.match_date.dt.strftime("%Y-%m-%d").tolist(), y=part.end_bankroll.astype(float).tolist(),
            mode="lines", name=sid, uid=sid, visible=True, legendgroup=sid,
            line={"color": color, "dash": line, "width": 1.5, "shape": "hv"},
            customdata=[[config["name"], bool(flag)] for flag in part.is_initial],
            hovertemplate=sid + " · %{customdata[0]}<br>Date: %{x|%Y-%m-%d}<br>Balance: %{y:,.2f} monetary units<extra></extra>"))
    ax.set_title("All 56 four-season strategies · initial 100,000 each · no account resets", fontsize=12)
    ax.set_xlabel("Calendar date (balance after date-level settlement)")
    ax.set_ylabel("Balance (monetary units)")
    ax.yaxis.set_major_formatter(matplotlib.ticker.StrMethodFormatter("{x:,.0f}"))
    ax.xaxis.set_major_locator(mdates.MonthLocator(interval=6))
    ax.xaxis.set_major_formatter(mdates.DateFormatter("%b %Y"))
    ax.grid(alpha=.2)
    ax.spines[["top", "right"]].set_visible(False)
    ax.set_xlim(graph.match_date.min(), graph.match_date.max())
    handles, labels = ax.get_legend_handles_labels()
    fig.legend(handles, labels, loc="lower center", ncol=10, fontsize=8, frameon=False,
        bbox_to_anchor=(.5, .014), handlelength=2, columnspacing=1.5, labelspacing=.65)
    with plt.rc_context({"svg.fonttype": "none"}):
        for suffix in (".png", ".svg"):
            fig.savefig(figures / ("all_strategies_4seasons_100k" + suffix), dpi=240)
    plt.close(fig)
    interactive.update_layout(title={"text": "All 56 four-season accounts · 100,000 each · no resets", "x": .01, "y": .99, "yanchor": "top"},
        template="plotly_white", height=1000, margin={"l": 85, "r": 35, "t": 115, "b": 240},
        xaxis={"title": "Calendar date (date-close settlement)", "type": "date"},
        yaxis={"title": "Balance (monetary units)", "type": "linear", "tickformat": ",.0f"}, hovermode="closest",
        legend={"title": {"text": "Every strategy ID", "side": "top"}, "orientation": "h", "x": 0, "xanchor": "left",
            "y": -.16, "yanchor": "top", "entrywidth": 70, "entrywidthmode": "pixels", "font": {"size": 10}, "groupclick": "toggleitem"},
        updatemenus=[{"type": "buttons", "direction": "right", "x": 0, "xanchor": "left", "y": 1.025, "yanchor": "bottom", "buttons": [
            {"label": "Show all", "method": "restyle", "args": [{"visible": True}]},
            {"label": "Hide all (keep legend)", "method": "restyle", "args": [{"visible": "legendonly"}]}]}])
    if len(interactive.data) != 56 or any(trace.visible is not True for trace in interactive.data):
        raise AssertionError("Every temporal curve must be visible initially")
    legacy_png = LEGACY / "figures/all_strategies_100k.png"
    legacy_copy = figures / "all_84_strategies_2025_100k.png"
    legacy_copy.write_bytes(legacy_png.read_bytes())
    appendix_html = "".join("<article id='" + e["strategy_id"] + "'><h3>" + html.escape(e["title"]) + "</h3><p>" + html.escape(e["explanation"]) + "</p><pre>" + html.escape(e["exact_config"]) + "</pre></article>" for e in entries)
    chart = interactive.to_html(full_html=False, include_plotlyjs=True, config={"responsive": True, "displaylogo": False})
    embedded_old = base64.b64encode(legacy_png.read_bytes()).decode("ascii")
    page = "<!doctype html><html lang='en'><head><meta charset='utf-8'><meta name='viewport' content='width=device-width,initial-scale=1'><title>EPL strategy temporal validation</title><style>body{font:15px system-ui;margin:24px;color:#173746}p{max-width:1100px;line-height:1.5}pre{white-space:pre-wrap;overflow-wrap:anywhere;font-size:12px;background:#f0f4f6;padding:12px}article{border-top:1px solid #cad5db;padding-top:12px}img{max-width:100%;height:auto}h1{font-size:27px}</style></head><body><h1>Four seasons: every compatible strategy starts at 100,000</h1><p>All 56 strategies S017–S072 are initially visible on one linear-money chart. Accounts run continuously through all 1,470 common archived fixtures and 453 settlement dates. Each is a separate account; there is no portfolio sum or season reset. Click legend entries to explore. This HTML embeds its plotting library and original chart and works offline.</p>" + chart + "<h2>Full catalog coverage: original 84-strategy graph</h2><p>The first-study recipes S001–S016 and S073–S084 have frozen predictions only for 2025/26. Their earlier returns are unavailable. The original complete 84-account comparison is retained below, with its common 100,000 starting point; it is a separate one-season replay.</p><img alt='All 84 strategy accounts in the original 2025/26 replay' src='data:image/png;base64," + embedded_old + "'><h2>Interpretation</h2><p>These strategies and outcomes had already been examined. Four chronological evaluation seasons do not erase retrospective strategy design. Archived quote timing, simultaneous availability and fills are unverified. Model expected payoffs are conditional claims, not demonstrated future returns. Selection-adjusted diagnostics in the accompanying report condition on recorded cash paths and concern the 84-rule 2025/26 comparison.</p><h2>" + APPENDIX_TITLE + "</h2>" + appendix_html + "</body></html>"
    (figures / "all_strategies_4seasons_100k.html").write_text(page, encoding="utf-8")
    return {"trace_ids": [trace.uid for trace in interactive.data], "initially_visible_traces": len(interactive.data),
        "plotly_js_bundled": True, "legacy_84_chart_embedded": True, "linear_money_y_axis": True}


def _metrics_table(frame: pd.DataFrame, ids: list[str]) -> pd.DataFrame:
    ordered = frame.set_index("strategy_id").reindex(ids)
    return pd.DataFrame([{"ID": sid, "Trades": number(row.bet_count, 0), "Hit\n%": number(row.hit_ratio, percent=True),
        "Trade\nfreq. %": number(row.trading_frequency, percent=True), "Avg. winning\ngain": number(row.avg_winning_gain_units),
        "Avg. losing\nloss": number(row.avg_losing_loss_units), "Model expected\ngain/trade": number(row.expected_gain_per_trade_units),
        "Model expected\nloss/trade": number(row.expected_loss_per_trade_units), "ROI on\nstakes": number(row.roi, percent=True),
        "Return on\n100,000": number(row.return_on_initial_capital, percent=True), "Max.\ndrawdown": number(row.max_drawdown, percent=True)}
        for sid, row in ordered.iterrows()])


def _stability_table(frame: pd.DataFrame) -> pd.DataFrame:
    rows = []
    for sid, row in frame.set_index("strategy_id").reindex(COMPATIBLE_IDS).iterrows():
        rows.append({"ID": sid, "+ve\nyears": number(row.positive_seasons, 0),
            "2022/23\nreturn": number(row.season_return_2022_23, percent=True), "2023/24\nreturn": number(row.season_return_2023_24, percent=True),
            "2024/25\nreturn": number(row.season_return_2024_25, percent=True), "2025/26\nreturn": number(row.season_return_2025_26, percent=True),
            "Worst year\nreturn": number(row.worst_season_return, percent=True), "Best year\nreturn": number(row.best_season_return, percent=True),
            "Pooled reset\nROI": number(row.pooled_reset_roi, percent=True), "Drop-one-year\nmin ROI": number(row.loo_min_reset_roi, percent=True),
            "Drop-one-year\nmax ROI": number(row.loo_max_reset_roi, percent=True)})
    return pd.DataFrame(rows)


def _selection_table(frame: pd.DataFrame) -> pd.DataFrame:
    rows = []
    for row in frame.itertuples(index=False):
        rows.append({"Benchmark": str(row.benchmark_id), "Block\ndates": number(row.block_length, 0), "Rules": number(row.family_size, 0),
            "Date\ncount": number(row.matched_match_dates, 0), "Highest mean\nID": str(row.best_observed_strategy_id),
            "Highest excess\nmean/date": number(row.max_observed_mean_date_excess_profit_units), "Global\none-sided p": number(row.global_one_sided_p, 4),
            "Simultaneous\ncritical units": number(row.simultaneous_95_critical_mean_date_units), "Replicates": number(row.bootstrap_samples, 0),
            "Distinct\npaths": number(row.distinct_recorded_paths, 0), "Reject at\n5%?": "Yes" if row.reject_global_at_05 else "No"})
    return pd.DataFrame(rows)


def _body(data: dict, figures: Path) -> tuple[list[dict], dict[str, pd.DataFrame]]:
    pages = []
    displays = {"Continuous metrics": _metrics_table(data["metrics_continuous"], COMPATIBLE_IDS),
        "Annual stability": _stability_table(data["season_stability"]),
        "Original 84 metrics": _metrics_table(data["legacy_metrics"], ALL_IDS),
        "Selection globals": _selection_table(data["selection_global"])}
    def page(title, orientation="portrait"):
        item = {"title": title, "orientation": orientation, "blocks": []}; pages.append(item); return item["blocks"]
    def p(blocks, text): blocks.append({"type": "paragraph", "text": text})
    def table(blocks, frame, caption): blocks.append({"type": "table", "frame": frame, "caption": caption})
    b = page(TITLE)
    p(b, "DRAFT FOR REVIEW | October 15, 2026 | EPL component of the statistical-arbitrage final project.")
    p(b, "This continuation checks whether the strategy results persist through time and how much apparent success survives comparison with the whole strategy catalog. It adds a continuous four-season account replay for all 56 rules with compatible saved predictions, annual reset diagnostics, and conditional multiple-strategy uncertainty for the original full 84-rule common season. No model is retrained and no earlier strategy result is overwritten.")
    p(b, "The four-season universe is the same 1,470 common quoted cohort fixtures for every compatible rule: 368 in 2022/23, 360 in 2023/24, 368 in 2024/25 and 374 in 2025/26. There are 453 settlement dates from August 5, 2022 through May 24, 2026. Each continuous account begins at exactly 100,000 on August 4, 2022. Its cash carries across summers and seasons without a reset, deposit, borrowing or interest.")
    p(b, "The full 84-rule catalog remains intact. S017–S072 use rolling-study probability series and can be evaluated across all four seasons. S001–S016 and S073–S084 depend on the first study, whose matching saved test predictions cover 2025/26 only. Those 28 strategies retain their original one-season graph and complete metrics; no earlier probability or return is invented. The final appendix explains all 84 strategies and states each one's available coverage.")
    p(b, "These are retrospective exploratory results. Although each rolling prediction comes from a chronologically earlier fitted model, strategy recipes and this enlarged catalog were created after outcomes had been examined. Extending the calendar and adjusting a conditional comparison does not produce a fresh holdout, a deployable betting system or a validated winning strategy. Historical match selection uses the current 2026/27 club cohort and therefore has survivor/selection limitations.")
    continuous = data["metrics_continuous"]
    best = continuous.loc[continuous.return_on_initial_capital.idxmax()]
    worst = continuous.loc[continuous.return_on_initial_capital.idxmin()]
    stability = data["season_stability"]
    globals_ = data["selection_global"]
    adjusted = data["selection_comparison"].family_adjusted_one_sided_p
    b = page("Findings from the completed replay")
    p(b, f"Of the 56 continuous four-season accounts, {int(continuous.profit.gt(0).sum())} finish above their 100,000 starting balance. Endpoints range from {number(INITIAL + worst.profit)} to {number(INITIAL + best.profit)} monetary units. The largest historical endpoint belongs to {best.strategy_id}, with a capital return of {number(best.return_on_initial_capital, percent=True)}, wager-turnover ROI of {number(best.roi, percent=True)} and date-close maximum drawdown of {number(best.max_drawdown, percent=True)}. These identify recorded outcomes; they do not select a recommended future strategy.")
    p(b, f"With annual accounts restarted separately at 100,000, {int(stability.positive_seasons.eq(4).sum())} of 56 rules show positive profit in all four seasons, while {int(stability.positive_seasons.le(1).sum())} are profitable in at most one. The annual-return and drop-one-year tables expose which recorded results are consistent across these periods and which rely on a narrower interval. Neither consistency nor a larger four-year endpoint corrects the retrospective design of the catalog.")
    rejection_count = int(globals_.reject_global_at_05.sum())
    p(b, f"In the separate complete 84-rule 2025/26 comparison, {rejection_count} of the eight benchmark/block maximum-statistic families reject their conditional null at 5%. The smallest global one-sided p-value is {number(globals_.global_one_sided_p.min(), 4)}, and the smallest rule-level family-adjusted p-value is {number(adjusted.min(), 4)}. All eight block/benchmark sensitivities and all 672 individual rows are retained. These diagnostics assess fixed recorded cash paths and have the limitations described below.")
    p(b, "The tables distinguish model expected gain and loss from realised average winning and losing payoffs. A favorable model expectation alone does not demonstrate an edge: probability errors are amplified by high decimal odds, and a realised bankroll-dependent sizing path is already conditioned on past outcomes. Frequency, drawdown, exposure and payoff uncertainty matter alongside an ending balance.")
    b = page("One continuous 100,000-unit account per compatible strategy", "landscape")
    p(b, "All 56 compatible rules appear on this single chart. Every account starts at 100,000; account balances are not added together. Curves continue through four seasons and off-season gaps without resetting. IDs map to the complete final appendix.")
    b.append({"type": "figure", "path": figures / "all_strategies_4seasons_100k.png", "caption": "Figure 1. All S017–S072 balances on linear monetary and calendar-date axes. Step lines record end-of-date cash after allocating every same-date request before settlement. The offline HTML version retains every visible trace and provides hover/legend controls."})
    for start in (0, 28):
        b = page(f"Four-season trading metrics: {COMPATIBLE_IDS[start]}–{COMPATIBLE_IDS[start+27]}", "landscape")
        p(b, "One uninterrupted account; 1,470 priced opportunities per strategy. Expected gain/loss are model-implied components conditional on recorded stakes. Winning and losing means are realised net payoffs per trade; losing loss is a positive magnitude.")
        table(b, displays["Continuous metrics"].iloc[start:start+28], "Table 1. Every compatible rule, in catalog order. Frequency=trades/1,470 fixtures; ROI=profit/turnover; capital return=profit/100,000. Undefined values appear as dashes. Full metrics and confidence intervals are in the workbook.")
    b = page("Separate annual starts test stability across seasons")
    p(b, "The continuous chart is the account history requested by the project. To separate season conditions from inherited capital, a second calculation gives each rule a fresh 100,000 at the beginning of each season. These are four independent diagnostic accounts, not four deposits into the continuous account. Annual performance is therefore not spliced into the main graph.")
    p(b, "Annual capital returns divide each season's realised profit by its own 100,000 start. The positive-season count uses strictly positive profit, with a zero-profit year not counted as positive. Worst and best annual returns show the range. Pooled reset ROI sums the four annual profits and divides by their combined turnover; it is neither a compounded return nor the continuous account's ROI.")
    p(b, "Drop-one-year minimum and maximum ROI recompute that pooled reset ROI after removing each season in turn. This checks how strongly the recorded pooled result depends on one year. It does not select a rule, exclude an inconvenient season or create independent validation. Returns from fixed and Kelly rules can differ between the continuous account and annual accounts because each Kelly path depends on its own realised starting cash.")
    for start in (0, 28):
        b = page(f"Annual stability for every rule: {COMPATIBLE_IDS[start]}–{COMPATIBLE_IDS[start+27]}", "landscape")
        table(b, displays["Annual stability"].iloc[start:start+28], "Table 2. Each annual column uses a separate 100,000-unit start. Pooled and drop-one-year columns use wager-turnover ROI from reset accounts. All 224 annual strategy rows, full payoffs and activity metrics are in the workbook.")
    b = page("Selection uncertainty across the entire original catalog")
    p(b, "The conditional selection diagnostic compares all 84 strategies over their shared 2025/26 history: 374 fixtures and 113 settlement dates. This is deliberately a separate population from the 56-rule four-season replay. It asks whether observed mean date-level excess profit remains unusual when the maximum across all 84 recorded rules is considered, rather than reading one attractive unadjusted curve in isolation.")
    p(b, "There are two benchmarks: cash, with zero date-level P&L, and S042, the original market-favorite flat-1,000 rule. Each comparison uses paired same-date strategy-minus-benchmark cashflows. The S042 account is itself stochastic and is not a no-risk cash proxy. Conditional on the saved paths, circular blocks of 1, 3, 5 and 10 consecutive settlement dates preserve different amounts of local date dependence. The block unit is a recorded settlement date, not a calendar day.")
    p(b, "For each benchmark and block length separately, 19,999 joint resamples center each rule's observed excess cashflow under its null, apply the same date indices to every rule, and retain the largest one-sided mean statistic across all 84 rules. The comparison table gives raw and family-adjusted one-sided p-values, pointwise intervals and a common simultaneous interval critical value. The global table below tests the maximum of that one declared family. Duplicate and cash paths remain in the catalog.")
    p(b, "These are conditional resampling diagnostics inspired by maximum-statistic data-snooping adjustments, not a full White Reality Check or guaranteed family-wise error control. Strategies, fitted probabilities and the realised bankroll-dependent stake path are held fixed; a counterfactual bootstrap Kelly path is not regenerated. Changing bankroll scale and one-season nonstationarity weaken stationarity assumptions. The earlier wider research search is uncounted. The eight benchmark/block families are separate sensitivity analyses, with no joint eight-family guarantee or chosen block length.")
    b = page("Global conditional selection comparisons", "landscape")
    table(b, displays["Selection globals"], "Table 3. Eight separate one-sided 84-rule maximum-mean diagnostics on the original 2025/26 date cashflows. Units are excess monetary P&L per settlement date. Statistical rejection, if present, remains conditional on recorded paths and the assumptions stated above.")
    p(b, "The complete 672 rule/benchmark/block comparison rows are in the workbook and selection_comparison.csv. Raw, adjusted and simultaneous quantities are retained for every rule; the report does not suppress unsuccessful comparisons or select the most favorable block sensitivity.")
    b = page("Full catalog coverage: all 84 original account curves", "landscape")
    p(b, "This preserved comparison uses all 374 common 2025/26 fixtures, August 15, 2025 through May 24, 2026, with an August 14 initial 100,000 point for every rule. It includes the 28 first-study recipes unavailable for the earlier rolling seasons and all eight original cash-policy traces.")
    b.append({"type": "figure", "path": figures / "all_84_strategies_2025_100k.png", "caption": "Figure 2. The complete original 84-strategy graph, copied byte for byte from the frozen strategy-laboratory artifact. This separate one-season replay is the population used by the conditional selection diagnostic."})
    for start in (0, 28, 56):
        b = page(f"Original 2025/26 metrics: {ALL_IDS[start]}–{ALL_IDS[start+27]}", "landscape")
        table(b, displays["Original 84 metrics"].iloc[start:start+28], "Table 4. Complete original common-season metrics for all 84 catalog IDs. These remain frozen, with frequency based on 374 fixtures and each account beginning at 100,000. These rows do not describe the four-season account.")
    b = page("Payoff definitions and cash accounting")
    p(b, "For selected win probability p, stake s and decimal odds O, expected gain is p·s·(O−1), expected loss magnitude is (1−p)·s, and their difference is model expected net profit. These components use allocated realised stakes. They are model-implied claims conditional on the observed account path, not unbiased future investment expectations or unconditional forecasts of terminal capital.")
    p(b, "A win returns stake×odds including principal; a loss returns zero. Realised net profit is gross return minus stake. Hit ratio is wins/trades, distinct from full-match classification accuracy. Average winning gain excludes returned principal; average losing loss is positive. Frequency is trades divided by valid priced opportunities. ROI divides profit by turnover; capital return divides by the stated starting capital. Drawdown includes the initial 100,000 and subsequent date-close peaks.")
    p(b, "Fixed stakes request 1,000; dynamic and fractional-Kelly rules use the balance at the start of that settlement date. All same-date requests are allocated together and uniformly scaled to available cash and any declared date budget. Same-date winnings cannot fund another request. Outcome and inclusive odds filters apply before choice; edge thresholds are strict, ties follow H, D, A, and at most one outcome is selected per match. The final appendix records every model, threshold, eligibility condition, sizing fraction and cap.")
    p(b, "Whole-date bootstrap intervals in the ordinary metric tables condition on the recorded model, choices and realised stake path and do not adjust for the enlarged search. They differ from the separately labelled 84-rule selection diagnostics. Neither interval mechanism resizes a counterfactual Kelly account. No-trade totals/activity are zero; undefined per-trade means, hit ratios and ROI remain unavailable.")
    b = page("Limitations, references and reproducibility")
    p(b, "Archived best-individual-book prices and reference prices do not verify provider publication clocks, simultaneous executable availability, accepted fills, liquidity, account limits or actual transaction fees. The new account comparison adds no interest or fee assumption. The separate frozen robustness study contains explicit hypothetical payout/fee shocks. These results do not establish current bookmaker arbitrage or live profitability.")
    p(b, "Chronological fitting comes from the frozen rolling-origin ablation protocol. Each outer-season prediction uses a prior training/tuning/calibration history, while the expanded rule catalog remains retrospective. No future prospective fixture is treated as settled. The append-only forecast ledger and its later outcome evaluations remain separate from this historical study.")
    p(b, "Method/source references: epl_modeling/ABLATION_PROTOCOL.md; epl_strategy/METHODOLOGY.md; epl_strategy_lab/PROTOCOL.md; epl_strategy_validation/PROTOCOL.md. The circular-block construction follows the fixed-block resampling approach described by Politis and Romano, Stanford technical report, https://statistics.stanford.edu/technical-reports/circular-block-resampling-procedure-stationary-data. The maximum-statistic data-snooping concept is associated with White (2000), https://doi.org/10.1111/1468-0262.00152. The implementation here has the narrower conditional scope stated above.")
    p(b, "Reproduce this report from finished tables with python -m epl_strategy_validation.report --results PATH. validation_overview.xlsx contains continuous/annual metrics, every cashflow, season stability, all selection comparisons, exact catalog configurations and input hashes. The report manifest verifies consumed sources were unchanged during reporting. The final builder envelope is written after reporting; the independent audit is intentionally not a report input. This avoids a circular envelope dependency.")
    p(b, "The appendix follows and is the final report section. Each of the 84 stable strategy IDs has its own explanation, exact original configuration and explicit coverage note. A missing earlier first-study curve denotes unavailable saved predictions, never an assumed zero return.")
    return pages, displays


def _markdown(pages, entries) -> str:
    chunks = []
    for index, page in enumerate(pages):
        chunks.extend([("# " if index == 0 else "## ") + page["title"], ""])
        for block in page["blocks"]:
            if block["type"] == "paragraph":
                chunks.extend([block["text"], ""])
            elif block["type"] == "figure":
                chunks.extend(["![Complete strategy account graph](../figures/" + block["path"].name + ")", "", block["caption"], ""])
            else:
                frame = block["frame"]
                chunks.extend([block["caption"], "", "| " + " | ".join(str(c).replace("\n", "<br>") for c in frame.columns) + " |", "| " + " | ".join("---" for _ in frame.columns) + " |"])
                for row in frame.itertuples(index=False, name=None):
                    chunks.append("| " + " | ".join(str(v) for v in row) + " |")
                chunks.append("")
    chunks.extend(["## " + APPENDIX_TITLE, ""])
    for entry in entries:
        chunks.extend(["### " + entry["title"], "", "Origin: " + entry["origin"] + ".", "", entry["explanation"], "", "Exact original configuration:", "", "```json", entry["exact_config"], "```", ""])
    return "\n".join(chunks)


def _pdf(pages, entries, path: Path) -> None:
    fonts = lab._fonts()
    styles = {
        "title": ParagraphStyle("ValTitle", fontName=fonts["bold"], fontSize=14, leading=16, spaceAfter=7, textColor=colors.HexColor("#12364D")),
        "body": ParagraphStyle("ValBody", fontName=fonts["body"], fontSize=10.2, leading=12.4, spaceAfter=7),
        "caption": ParagraphStyle("ValCaption", fontName=fonts["body"], fontSize=8.5, leading=10, spaceAfter=4),
        "cell": ParagraphStyle("ValCell", fontName=fonts["sans"], fontSize=8.0, leading=9.3),
        "head": ParagraphStyle("ValHead", fontName=fonts["bold"], fontSize=8.0, leading=9.3, textColor=colors.white),
        "entrytitle": ParagraphStyle("ValEntryTitle", fontName=fonts["bold"], fontSize=9.7, leading=11.4, spaceBefore=5, spaceAfter=4),
        "entry": ParagraphStyle("ValEntry", fontName=fonts["body"], fontSize=9.2, leading=10.6, spaceAfter=4),
        "config": ParagraphStyle("ValConfig", fontName=fonts["mono"], fontSize=6.8, leading=8.1, spaceAfter=7, splitLongWords=True, wordWrap="CJK")}
    def header(canvas, doc):
        width, height = canvas._pagesize
        canvas.setFont(fonts["sans"], 8); canvas.setFillColor(colors.HexColor("#52616B"))
        canvas.drawString(36, height - 21, "EPL TEMPORAL AND SELECTION VALIDATION — EXPLORATORY | October 15, 2026")
        canvas.drawRightString(width - 36, 21, str(doc.page))
    doc = BaseDocTemplate(str(path), pagesize=letter, title=TITLE, author="Aman Dhillon and Panagiotis Housos")
    for orientation, size in (("portrait", letter), ("landscape", landscape(letter))):
        width, height = size
        doc.addPageTemplates(PageTemplate(id=orientation, pagesize=size, onPage=header,
            frames=[Frame(36, 36, width - 72, height - 72, leftPadding=0, rightPadding=0, topPadding=0, bottomPadding=0)]))
    story = []
    for index, page in enumerate(pages):
        if index: story.extend([NextPageTemplate(page["orientation"]), PageBreak()])
        width = (letter if page["orientation"] == "portrait" else landscape(letter))[0] - 72
        story.append(Paragraph(html.escape(page["title"]), styles["title"]))
        for block in page["blocks"]:
            if block["type"] == "paragraph": story.append(Paragraph(html.escape(block["text"]), styles["body"]))
            elif block["type"] == "figure":
                with PILImage.open(block["path"]) as picture: ratio = picture.height / picture.width
                story.extend([Image(str(block["path"]), width=width - 10, height=(width - 10) * ratio), Paragraph(html.escape(block["caption"]), styles["caption"])])
            else:
                story.append(Paragraph(html.escape(block["caption"]), styles["caption"]))
                frame = block["frame"]
                raw = [list(frame.columns)] + frame.values.tolist()
                cells = [[Paragraph(html.escape(str(value)).replace("\n", "<br/>"), styles["head" if row_index == 0 else "cell"]) for value in row] for row_index, row in enumerate(raw)]
                weights = BODY_WEIGHTS if page["title"].startswith(("Four-season trading", "Original 2025")) else [1 / len(frame.columns)] * len(frame.columns)
                table = Table(cells, colWidths=[width * weight for weight in weights], repeatRows=1)
                table.setStyle(TableStyle([("BACKGROUND", (0,0), (-1,0), colors.HexColor("#12364D")), ("ROWBACKGROUNDS", (0,1), (-1,-1), [colors.white, colors.HexColor("#F0F4F6")]), ("VALIGN", (0,0), (-1,-1), "MIDDLE"), ("LEFTPADDING", (0,0), (-1,-1), 3), ("RIGHTPADDING", (0,0), (-1,-1), 3), ("TOPPADDING", (0,0), (-1,-1), 1.1), ("BOTTOMPADDING", (0,0), (-1,-1), 1.1)]))
                story.append(table)
    story.extend([NextPageTemplate("portrait"), PageBreak(), Paragraph(APPENDIX_TITLE, styles["title"])])
    for entry in entries:
        story.append(KeepTogether([Paragraph(html.escape(entry["title"]), styles["entrytitle"]),
            Paragraph(html.escape("Origin: " + entry["origin"] + ". " + entry["explanation"]), styles["entry"]),
            Paragraph(html.escape("Exact original configuration: " + entry["exact_config"]), styles["config"])]))
    doc.build(story)


def _docx(pages, entries, path: Path) -> None:
    doc = Document()
    normal = doc.styles["Normal"]; normal.font.name = "Times New Roman"; normal.font.size = Pt(10.2)
    normal.paragraph_format.line_spacing = Pt(12.4); normal.paragraph_format.space_after = Pt(7)
    heading = doc.styles["Heading 1"]; heading.font.name = "Arial"; heading.font.size = Pt(14)
    heading.paragraph_format.space_before = Pt(0); heading.paragraph_format.space_after = Pt(7)
    doc.core_properties.title = TITLE; doc.core_properties.author = "Aman Dhillon and Panagiotis Housos"
    def section(index, wide):
        sec = doc.sections[0] if index == 0 else doc.add_section(WD_SECTION_START.NEW_PAGE)
        sec.orientation = WD_ORIENT.LANDSCAPE if wide else WD_ORIENT.PORTRAIT
        sec.page_width = Inches(11 if wide else 8.5); sec.page_height = Inches(8.5 if wide else 11)
        sec.top_margin = sec.bottom_margin = sec.left_margin = sec.right_margin = Inches(.5)
        sec.header_distance = sec.footer_distance = Inches(.2)
        if index == 0:
            run = sec.header.paragraphs[0].add_run("EPL TEMPORAL AND SELECTION VALIDATION — EXPLORATORY | October 15, 2026")
            run.font.name = "Arial"; run.font.size = Pt(8)
            footer = sec.footer.paragraphs[0]; footer.alignment = 2
            field = OxmlElement("w:fldSimple"); field.set(qn("w:instr"), "PAGE"); footer._p.append(field)
    for index, page in enumerate(pages):
        wide = page["orientation"] == "landscape"; section(index, wide); doc.add_heading(page["title"], 1)
        for block in page["blocks"]:
            if block["type"] == "paragraph": doc.add_paragraph(block["text"]); continue
            if block["type"] == "figure": doc.add_picture(str(block["path"]), width=Inches(9.85 if wide else 7.35))
            caption = doc.add_paragraph(block["caption"]); caption.paragraph_format.line_spacing = Pt(10)
            for run in caption.runs: run.font.size = Pt(8.5)
            if block["type"] == "figure": continue
            frame = block["frame"]; table = doc.add_table(rows=1, cols=len(frame.columns)); table.style = "Light Shading Accent 1"; table.autofit = False
            weights = BODY_WEIGHTS if page["title"].startswith(("Four-season trading", "Original 2025")) else [1 / len(frame.columns)] * len(frame.columns)
            width = 10 if wide else 7.5
            for col, weight in zip(table.columns, weights): col.width = Inches(width * weight)
            for cell, value in zip(table.rows[0].cells, frame.columns): cell.text = str(value)
            for values in frame.itertuples(index=False, name=None):
                for cell, value in zip(table.add_row().cells, values): cell.text = str(value)
            for row in table.rows:
                for cell, weight in zip(row.cells, weights):
                    cell.width = Inches(width * weight)
                    for paragraph in cell.paragraphs:
                        paragraph.paragraph_format.line_spacing = Pt(9.3); paragraph.paragraph_format.space_after = Pt(0)
                        for run in paragraph.runs: run.font.name = "Arial"; run.font.size = Pt(8)
    section(len(pages), False); doc.add_heading(APPENDIX_TITLE, 1)
    for entry in entries:
        title = doc.add_paragraph(); title.paragraph_format.keep_with_next = True
        run = title.add_run(entry["title"]); run.bold = True; run.font.name = "Arial"; run.font.size = Pt(9.7)
        paragraph = doc.add_paragraph("Origin: " + entry["origin"] + ". " + entry["explanation"])
        paragraph.paragraph_format.keep_with_next = True; paragraph.paragraph_format.line_spacing = Pt(10.6)
        for run in paragraph.runs: run.font.size = Pt(9.2)
        paragraph = doc.add_paragraph("Exact original configuration: " + entry["exact_config"])
        paragraph.paragraph_format.keep_together = True; paragraph.paragraph_format.line_spacing = Pt(8.1)
        for run in paragraph.runs: run.font.name = "Consolas"; run.font.size = Pt(6.8)
    doc.save(path)


def _workbook(path: Path, data: dict, entries: list[dict], registry: dict) -> dict:
    definitions = pd.DataFrame([
        ("Continuous scope", "56 strategies S017–S072; each independently starts 100000 once across 1470 fixtures and 453 settlement dates, 2022/23–2025/26."),
        ("Annual reset scope", "224 rows: four separate 100000 accounts per strategy. These are diagnostics and are not deposits or resets in the continuous chart."),
        ("Full catalog scope", "All84 remain in the frozen 2025/26 graph/metrics. First-study28 have no matching earlier predictions and receive no fabricated returns."),
        ("Expected payoffs", "Probability-weighted positive/negative net payoffs at recorded allocated stakes; conditional model claims, not unconditional terminal wealth forecasts."),
        ("Selection scope", "84 strategies on original2025/26 374fixtures/113dates; two benchmarks × four circular block lengths. Each family is separate; no joint eight-family guarantee."),
        ("Selection interpretation", "Conditional on fixed recorded paths; no refit or Kelly path regeneration. Retrospective catalog, wider research search uncounted, no full White Reality Check or fresh test."),
        ("Currency", "Comparable monetary units; no specific currency or executable bookmaker funding is claimed."),
        ("Undefined ratios", "Unavailable values use a dash; no-trade totals/activity are zero. Infinity remains explicitly labelled.")], columns=["Item", "Definition"])
    sheets = {"Read me": definitions, "Continuous all metrics": data["metrics_continuous"], "Annual all metrics": data["metrics_seasonal"],
        "Season stability": data["season_stability"], "Selection all comparisons": data["selection_comparison"], "Selection global": data["selection_global"],
        "Original84 2025 metrics": data["legacy_metrics"], "Complete84 catalog": pd.DataFrame(data["catalog"]), "Final84 appendix": pd.DataFrame(entries),
        "Graph equity incl initial": data["graph"], "Continuous date equity": data["equity_continuous"], "Continuous every bet": data["bets_continuous"],
        "Annual date equity": data["equity_seasonal"], "Annual every bet": data["bets_seasonal"],
        "Input provenance": pd.DataFrame([{"path": path, **record} for path, record in registry.items()])}
    percentages = {"hit_ratio", "trading_frequency", "trade_rate_all_predictions", "price_coverage", "active_date_fraction",
        "roi", "expected_roi", "return_on_initial_capital", "max_drawdown", "avg_selected_probability",
        "avg_model_expected_return_per_unit_staked", "avg_realized_return_per_unit_staked", "empirical_break_even_hit_ratio",
        "avg_winning_return_per_unit_staked", "avg_losing_return_per_unit_staked", "max_date_capital_exposure", "mean_date_capital_exposure",
        "date_return_mean", "date_return_sample_std", "worst_season_return", "best_season_return", "return_std", "pooled_reset_roi",
        "loo_min_reset_roi", "loo_max_reset_roi", "continuous_return_on_initial_capital", "continuous_max_drawdown",
        "min_edge", "bankroll_fraction", "kelly_fraction", "max_bet_fraction", "date_cap_fraction", "date_budget_fraction", "commitment_scale",
        "probability", "edge", "full_kelly_fraction", "daily_return", "drawdown", "excess_return_on_initial_capital",
        "simultaneous_ci_low_capital_return", "simultaneous_ci_high_capital_return"}
    with pd.ExcelWriter(path, engine="xlsxwriter", engine_kwargs={"options": {"strings_to_urls": False}}) as writer:
        book = writer.book; head = book.add_format({"bold": True, "font_color": "white", "bg_color": "#12364D", "text_wrap": True})
        num = book.add_format({"num_format": "#,##0.0000"}); pct = book.add_format({"num_format": "0.00%"})
        for name, frame in sheets.items():
            safe = frame.copy()
            for column in safe:
                if isinstance(safe[column].dtype, pd.DatetimeTZDtype): safe[column] = safe[column].dt.tz_convert(None)
                if safe[column].dtype == "object": safe[column] = safe[column].map(lambda value: lab._json(value.tolist() if isinstance(value, np.ndarray) else value) if isinstance(value, (dict, list, tuple, np.ndarray)) else value)
                elif pd.api.types.is_numeric_dtype(safe[column]): safe[column] = safe[column].map(lambda value: "+∞" if value == float("inf") else "−∞" if value == float("-inf") else value)
            safe.to_excel(writer, sheet_name=name, index=False, na_rep="—")
            worksheet = writer.sheets[name]; worksheet.write_row(0, 0, list(safe.columns), head); worksheet.set_row(0, 42)
            worksheet.freeze_panes(1, 1); worksheet.autofilter(0, 0, len(safe), len(safe.columns) - 1)
            for index, column in enumerate(safe.columns):
                base = column.removesuffix("_ci_low").removesuffix("_ci_high")
                is_percent = base in percentages or column.startswith("season_return_")
                worksheet.set_column(index, index, 65 if column in ("explanation", "exact_config", "Definition", "coverage", "path", "name", "strategy_name") else min(42, max(14, len(column) * .8)), pct if is_percent else num)
    return {name: len(frame) for name, frame in sheets.items()}


def build_report(output_dir: Path) -> dict:
    output = Path(output_dir).resolve(); data = _load(output); anchor = _validate(data)
    entries = _appendix(data["catalog"])
    input_names = ("catalog.json", "compatible_catalog.json", "metrics_continuous.parquet", "metrics_seasonal.parquet", "season_stability.parquet", "bets_continuous.parquet", "equity_continuous.parquet", "bets_seasonal.parquet", "equity_seasonal.parquet", "graph_equity_continuous.csv", "selection_comparison.parquet", "selection_global.parquet", "coverage.json", "source_provenance.json")
    sources = [output / name for name in input_names if (output / name).exists()]
    sources += [Path(__file__).resolve(), Path(lab.__file__).resolve(), ROOT / "epl_strategy/report.py", LEGACY / "metrics.parquet", LEGACY / "figures/all_strategies_100k.png"]
    sources += [ROOT / name for name in ("epl_modeling/ABLATION_PROTOCOL.md", "epl_strategy/METHODOLOGY.md", "epl_strategy_lab/PROTOCOL.md", "epl_strategy_validation/PROTOCOL.md") if (ROOT / name).exists()]
    sources = sorted(set(sources))
    registry = {lab._relative(path): {"sha256": lab._sha(path), "bytes": path.stat().st_size} for path in sources}
    report = output / "report"; report.mkdir(exist_ok=True); figures = output / "figures"
    graph_info = _graphs(data["compatible_catalog"], data["graph"], figures, entries)
    pages, displays = _body(data, figures)
    for name, frame in displays.items(): frame.to_csv(report / (name.lower().replace(" ", "_") + ".csv"), index=False)
    pd.DataFrame(entries).to_csv(report / "complete84_strategy_appendix.csv", index=False)
    md = report / "validation_report.md"; pdf = md.with_suffix(".pdf"); docx = md.with_suffix(".docx")
    md.write_text(_markdown(pages, entries), encoding="utf-8"); _pdf(pages, entries, pdf); _docx(pages, entries, docx)
    sheets = _workbook(report / "validation_overview.xlsx", data, entries, registry)
    reader = PdfReader(pdf); extracted = "\n".join(page.extract_text() or "" for page in reader.pages)
    if APPENDIX_TITLE not in extracted: raise AssertionError("PDF omitted final appendix heading")
    appendix_text = " ".join(extracted.split(APPENDIX_TITLE, 1)[1].split())
    if any(" ".join(entry["title"].split()) not in appendix_text for entry in entries): raise AssertionError("PDF omitted a strategy explanation")
    if entries[-1]["strategy_id"] not in reader.pages[-1].extract_text(): raise AssertionError("Appendix is not the final report content")
    changed = [lab._relative(path) for path in sources if lab._sha(path) != registry[lab._relative(path)]["sha256"]]
    if changed: raise AssertionError("Protected report inputs changed: " + str(changed))
    produced = [path for folder in (report, figures) for path in folder.rglob("*") if path.is_file() and path.name != "report_manifest.json" and "previews" not in path.parts]
    manifest = {"status": "retrospective_exploratory_companion", "draft_date_new_york": "2026-10-15", "built_at_utc": datetime.now(timezone.utc).isoformat(),
        "continuous_strategy_count": 56, "full_catalog_count": 84, "annual_strategy_rows": 224, "continuous_matches": 1470, "continuous_dates": 453,
        "initial_capital_per_account": INITIAL, "initial_anchor_date": anchor.strftime("%Y-%m-%d"), "graph": graph_info,
        "selection_rows": len(data["selection_comparison"]), "selection_global_rows": len(data["selection_global"]),
        "appendix_ids": [entry["strategy_id"] for entry in entries], "appendix_is_last": True, "pdf_pages": len(reader.pages), "body_sections": len(pages),
        "workbook_sheet_rows": sheets, "models_refitted": False, "consumed_source_files_unchanged": not changed, "inputs": registry,
        "outputs": {path.relative_to(output).as_posix(): {"sha256": lab._sha(path), "bytes": path.stat().st_size} for path in sorted(produced)},
        "manifest_order": "Final builder envelope and independent audit are intentionally not consumed report inputs.",
        "limits": "Historical current-club cohort, retrospective catalog, archived execution unverified, conditional fixed-path inference; no fresh holdout or guaranteed selection control."}
    (report / "report_manifest.json").write_text(json.dumps(manifest, ensure_ascii=False, indent=2, allow_nan=False), encoding="utf-8")
    return {"output_directory": str(report), "continuous_strategy_count": 56, "all_visible_traces": 56, "full_catalog_count": 84, "appendix_entries": len(entries), "pdf_pages": len(reader.pages), "source_files_unchanged": not changed}


def main() -> None:
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument("--results", required=True, type=Path)
    print(json.dumps(build_report(parser.parse_args().results), indent=2))


if __name__ == "__main__": main()
