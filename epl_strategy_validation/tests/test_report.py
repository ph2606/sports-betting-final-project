"""Presentation invariants: coverage, cash paths and complete final appendix."""
import json

import numpy as np
import pandas as pd
import pytest
from docx import Document
from openpyxl import load_workbook
from pypdf import PdfReader

from epl_strategy_lab.catalog import build_catalog
from epl_strategy_validation import report


@pytest.fixture
def finished_tables():
    catalog = build_catalog()
    ids = report.COMPATIBLE_IDS
    dates = pd.date_range("2022-08-05", "2026-05-24", periods=453).normalize()
    metrics = pd.DataFrame({"strategy_id": ids, "prediction_count": 1470,
        "matched_valid_odds_count": 1470, "initial_capital_sum_units": 100000.,
        "profit": 0., "bet_count": 0})
    equity = pd.DataFrame([{"strategy_id": sid, "match_date": date, "end_bankroll": 100000.}
        for sid in ids for date in dates])
    anchor = pd.Timestamp("2022-08-04")
    graph = pd.concat([equity.assign(is_initial=False), pd.DataFrame({"strategy_id": ids,
        "match_date": anchor, "end_bankroll": 100000., "is_initial": True})], ignore_index=True)
    comparisons = pd.DataFrame([{"strategy_id": sid, "benchmark_id": bench, "block_length": block}
        for sid in report.ALL_IDS for bench in ("cash", "S042") for block in (1,3,5,10)])
    return {"catalog": catalog, "compatible_catalog": [c for c in catalog if c["strategy_id"] in ids],
        "metrics_continuous": metrics, "season_stability": metrics.copy(),
        "metrics_seasonal": pd.concat([metrics] * 4, ignore_index=True),
        "equity_continuous": equity, "graph": graph,
        "bets_continuous": pd.DataFrame(columns=["strategy_id"]),
        "selection_comparison": comparisons, "selection_global": pd.DataFrame({"row": range(8)})}


def test_all_continuous_accounts_have_one_exact_initial_cash_point(finished_tables):
    assert report._validate(finished_tables) == pd.Timestamp("2022-08-04")


def test_reject_missing_curve(finished_tables):
    finished_tables["graph"] = finished_tables["graph"].loc[lambda f: f.strategy_id.ne("S017")]
    with pytest.raises(ValueError, match="omits a compatible strategy"):
        report._validate(finished_tables)


def test_reject_a_reset_in_continuous_capital(finished_tables):
    finished_tables["metrics_continuous"].loc[0, "initial_capital_sum_units"] = 400000.
    with pytest.raises(ValueError, match="resets its 100,000"):
        report._validate(finished_tables)


def test_reject_graph_not_equal_to_cash_ledger(finished_tables):
    index = finished_tables["graph"].index[0]
    finished_tables["graph"].loc[index, "end_bankroll"] += 1
    with pytest.raises(ValueError, match="differ from the saved equity"):
        report._validate(finished_tables)


def test_reject_rewritten_temporal_rule(finished_tables):
    finished_tables["compatible_catalog"] = json.loads(json.dumps(finished_tables["compatible_catalog"]))
    finished_tables["compatible_catalog"][0]["min_edge"] = .10
    with pytest.raises(ValueError, match="changes a frozen rule"):
        report._validate(finished_tables)


def test_every_appendix_rule_keeps_exact_config_and_coverage():
    catalog = build_catalog(); entries = report._appendix(catalog)
    assert [entry["strategy_id"] for entry in entries] == report.ALL_IDS
    assert [json.loads(entry["exact_config"]) for entry in entries] == catalog
    assert sum(entry["coverage"].startswith("Four-season") for entry in entries) == 56
    assert sum(entry["coverage"].startswith("Original 2025/26") for entry in entries) == 28
    assert all(entry["coverage"] in entry["explanation"] for entry in entries)


def test_markdown_ends_with_every_strategy_explanation():
    entries = report._appendix(build_catalog())
    pages = [{"title": "Test", "orientation": "portrait", "blocks": [{"type": "paragraph", "text": "Body."}]}]
    md = report._markdown(pages, entries)
    appendix = md.split("## " + report.APPENDIX_TITLE)[1]
    assert appendix.count("### S") == 84
    assert md.rstrip().endswith("```")
    assert "### S084" in appendix and md.index("Body.") < md.index(report.APPENDIX_TITLE)


def test_pdf_and_docx_smoke_have_all84_appendix_entries_last(tmp_path):
    entries = report._appendix(build_catalog())
    pages = [{"title": "Validation smoke", "orientation": "portrait", "blocks": [{"type": "paragraph", "text": "Body before the complete appendix."}]}]
    pdf = tmp_path / "report.pdf"; docx = tmp_path / "report.docx"
    report._pdf(pages, entries, pdf); report._docx(pages, entries, docx)
    reader = PdfReader(pdf); text = " ".join(" ".join((page.extract_text() or "").split()) for page in reader.pages)
    assert report.APPENDIX_TITLE in text
    assert all(" ".join(entry["title"].split()) in text for entry in entries)
    assert "S084" in reader.pages[-1].extract_text()
    paragraphs = [p.text for p in Document(docx).paragraphs]
    assert sum(p.startswith(tuple(report.ALL_IDS)) for p in paragraphs) == 84
    assert paragraphs[-1].startswith("Exact original configuration:")
    assert json.loads(paragraphs[-1].split(": ", 1)[1])["strategy_id"] == "S084"


def test_interactive_chart_contains_every_visible_trace_offline_and_all84_appendix(tmp_path):
    catalog = [c for c in build_catalog() if c["strategy_id"] in report.COMPATIBLE_IDS]
    entries = report._appendix(build_catalog())
    graph = pd.DataFrame([{"strategy_id": sid, "match_date": date, "end_bankroll": amount, "is_initial": initial}
        for sid in report.COMPATIBLE_IDS for date, amount, initial in ((pd.Timestamp("2022-08-04"), 100000., True), (pd.Timestamp("2026-05-24"), 100001., False))])
    info = report._graphs(catalog, graph, tmp_path, entries)
    assert info["trace_ids"] == report.COMPATIBLE_IDS
    assert info["initially_visible_traces"] == 56
    text = (tmp_path / "all_strategies_4seasons_100k.html").read_text(encoding="utf-8")
    assert "plotly.js" in text and "<script src=" not in text
    assert "data:image/png;base64," in text
    assert text.count("<article id='S") == 84
    assert all((tmp_path / ("all_strategies_4seasons_100k" + suffix)).exists() for suffix in (".png", ".svg", ".html"))


def test_workbook_keeps_money_and_payoff_ratio_separate_from_percentages(tmp_path):
    data = {key: pd.DataFrame({"value": [1]}) for key in (
        "metrics_seasonal", "season_stability", "selection_comparison", "selection_global", "legacy_metrics", "graph",
        "equity_continuous", "bets_continuous", "equity_seasonal", "bets_seasonal")}
    data["metrics_continuous"] = pd.DataFrame({"payoff_ratio": [2.0], "roi": [.1], "profit": [1000.],
        "date_return_mean_to_std_unannualized": [.2], "roi_ci_low": [.05]})
    data["catalog"] = build_catalog()
    entries = report._appendix(data["catalog"])
    path = tmp_path / "report.xlsx"
    report._workbook(path, data, entries, {"source": {"sha256": "a" * 64, "bytes": 1}})
    book = load_workbook(path, read_only=True)
    sheet = book["Continuous all metrics"]
    assert [cell.value for cell in sheet[2]] == [2., .1, 1000., .2, .05]
    assert [cell.number_format for cell in sheet[2]] == ["#,##0.0000", "0.00%", "#,##0.0000", "#,##0.0000", "0.00%"]
    book.close()
